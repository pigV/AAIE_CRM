import io
import unittest

import pdfplumber
from docx import Document
from docx.oxml.ns import qn
from pypdf import PdfReader

import contract_domain
import contract_template_editor as editor


CONTRACT_TYPES = [item[0] for item in contract_domain.contract_type_options()]


class ContractMasterLayoutTests(unittest.TestCase):
    @staticmethod
    def _commercial():
        value = editor.authority_v2_commercial(
            "NewInternationalStudentContract", editor._default_commercial()
        )
        value.update({
            "commercial_schema": "authority_v2",
            "gst_treatment": "EXCLUSIVE",
            "company_name": "Australian Academy of International Education",
            "company_abn": "12 345 678 901",
            "company_address": "Level 1, 123 Collins Street, Melbourne VIC 3000",
            "company_contact": "+61 3 9000 0000 / contracts@example.test",
        })
        return value

    @staticmethod
    def _signer():
        return {
            "source": "mother",
            "name": "Kexin Parent",
            "relationship": "母亲",
            "phone": "0411000000",
            "email": "parent@example.test",
        }

    @staticmethod
    def _instance(**overrides):
        value = {
            "contract_uid": "CONTRACT-MASTER-001",
            "contract_date": "2026-08-22",
            "student_full_name": "Kexin Chen",
            "student_date_of_birth": "2011-06-14",
            "student_id_or_passport_number": "20260020",
            "student_current_year_level": "Year 8",
            "student_applying_year_level": "Year 9",
            "company_authorised_person": "Sonny Zhang",
            "parent_address": "Melbourne VIC",
            "payment_method": "Australian public account",
            "discount_amount": "100.00",
            "special_terms": "无。",
        }
        value.update(overrides)
        return value

    def _render(self, contract_type="NewInternationalStudentContract", **instance):
        snapshot = editor.build_domain_v1_snapshot(contract_type, self._commercial())
        return editor.render_published_contract_artifacts(
            snapshot,
            student_name="Kexin Chen",
            signer=self._signer(),
            selected_services=list(contract_domain.default_service_ids(contract_type)),
            contract_instance=self._instance(**instance),
        )

    def test_three_content_templates_share_one_protected_master_layout(self):
        snapshots = [
            editor.build_domain_v1_snapshot(value, self._commercial())
            for value in CONTRACT_TYPES
        ]

        policies = [snapshot["layout_policy"] for snapshot in snapshots]
        self.assertTrue(all(policy["id"] == "AAIEContractMasterLayoutV1" for policy in policies))
        self.assertEqual(policies[0], policies[1])
        self.assertEqual(policies[1], policies[2])
        self.assertEqual(
            policies[0]["section_anchors"],
            [
                "section_1_parties", "section_2_services", "section_3_fees",
                "section_4_party_a", "section_5_party_b", "section_6_refund",
                "section_7_force_majeure", "section_8_dispute",
                "section_9_variation", "section_10_special", "section_11_signature",
            ],
        )
        bodies = [tuple(block["body"] for block in snapshot["blocks"]) for snapshot in snapshots]
        self.assertEqual(len(set(bodies)), 3)

    def test_editable_title_and_optional_clause_render_inside_fixed_section(self):
        snapshot = editor.build_domain_v1_snapshot(
            "TransferApplicationContract", self._commercial()
        )
        index = next(
            i for i, block in enumerate(snapshot["blocks"])
            if block["key"] == "profile_3"
        )
        snapshot["blocks"][index]["title"] = "三、费用与付款安排"
        snapshot["blocks"].insert(index + 1, {
            "key": "profile_3::approved_extra",
            "title": "3.1 经批准的附加说明",
            "body": "本条款仅属于费用章节。",
            "enabled": True,
            "fixed_layout": False,
            "required": False,
            "legal_kind": "",
            "required_variables": [],
        })
        result = editor.render_published_contract_artifacts(
            snapshot,
            student_name="Kexin Chen",
            signer=self._signer(),
            selected_services=list(contract_domain.default_service_ids(
                "TransferApplicationContract"
            )),
            contract_instance=self._instance(special_terms=""),
        )
        text = result["authority_content"]
        self.assertIn("三、费用与付款安排", text)
        self.assertIn("3.1 经批准的附加说明", text)
        self.assertIn("本条款仅属于费用章节。", text)
        self.assertLess(text.index("三、费用与付款安排"), text.index("四、甲方权利义务"))

    def test_study_abroad_has_complete_ordered_eleven_section_corpus(self):
        result = self._render()
        text = result["authority_content"]
        headings = [
            "一、合同双方", "二、服务项目及服务范围", "三、服务费用与第三方费用",
            "四、甲方权利义务", "五、乙方权利义务", "六、退费及违约责任",
            "七、不可抗力、结果说明及免责", "八、适用法律及争议解决",
            "九、合同补充、变更和生效", "十、特别条款", "十一、签署",
        ]
        positions = [text.index(value) for value in headings]
        self.assertEqual(positions, sorted(positions))
        for sentence in (
            "本合同项下服务费不包括学校",
            "甲方不保证乙方一定获得录取、签证或其他第三方批准结果",
            "乙方应保证向甲方提交的所有文件、材料、信息和陈述合法、真实、完整、有效",
            "本合同的任何补充、变更或修改应采用书面形式",
            "包括但不限于学校申请费、学费、医疗保险费、住宿费、监护或福利安排费",
            "不对第三方是否退费、退费金额、退费时间或退费结果作出保证",
            "如乙方签证申请被拒签，且乙方明确表示不再继续申请",
            "学生最终未被任何一所学校录取",
            "自然灾害、疫情、战争、骚乱、政策变化",
            "不具备完全民事行为能力",
            "个人生活费、误工费、误学费",
        ):
            self.assertIn(sentence, text)
        compact = lambda value: "".join(value.split())
        spans = {
            "六": text[text.index("六、退费及违约责任"):text.index("七、不可抗力")],
            "七": text[text.index("七、不可抗力"):text.index("八、适用法律及争议解决")],
            "九": text[text.index("九、合同补充、变更和生效"):text.index("十、特别条款")],
        }
        self.assertGreaterEqual(len(compact(spans["六"])), 580)
        self.assertGreaterEqual(len(compact(spans["七"])), 270)
        self.assertGreaterEqual(len(compact(spans["九"])), 225)

    def test_commercial_rows_are_inside_section_three_without_diagnostics(self):
        result = self._render()
        text = result["authority_content"]
        section_three = text[text.index("三、服务费用与第三方费用"):text.index("四、甲方权利义务")]
        self.assertIn("AUD 2400.00", section_three)
        self.assertIn("优惠后服务费", section_three)
        self.assertIn("GST", section_three)
        self.assertIn("最终应付总额", section_three)
        for forbidden in ("退款配置", "拒签实际保留总额", "gst_rate", "commercial_schema"):
            self.assertNotIn(forbidden, text)

    def test_future_three_contract_types_omit_student_contact_row_across_renderers(self):
        """Catch the removed row returning in any future contract artifact."""
        for contract_type in CONTRACT_TYPES:
            with self.subTest(contract_type=contract_type):
                snapshot = editor.build_domain_v1_snapshot(
                    contract_type, self._commercial()
                )
                instance = self._instance(
                    intake_year="2027",
                    intake_term="Term 1",
                    student_phone="0400000000",
                    student_email="student@example.test",
                )
                bindings = editor.contract_display_bindings.build_frozen_display_bindings(
                    contract_type,
                    student_snapshot={
                        "student_full_name": instance["student_full_name"],
                        "date_of_birth": instance["student_date_of_birth"],
                        "student_id": instance["student_id_or_passport_number"],
                        "current_year_level": instance["student_current_year_level"],
                        "applying_year_level": instance["student_applying_year_level"],
                        "preferred_intake_year": instance["intake_year"],
                        "preferred_intake_term": instance["intake_term"],
                    },
                    signer=self._signer(),
                    commercial=self._commercial(),
                    contract_instance=instance,
                )
                result = editor.render_published_contract_artifacts(
                    snapshot,
                    student_name=instance["student_full_name"],
                    signer=self._signer(),
                    selected_services=list(
                        contract_domain.default_service_ids(contract_type)
                    ),
                    contract_instance=instance,
                    display_bindings=bindings,
                )

                model_labels = [
                    row["cells"][0]["lines"][0]
                    for node in result["model"]["content"]
                    if node["key"] == "protected_contract_identity"
                    for row in node["rows"]
                ]
                document = Document(io.BytesIO(result["docx_bytes"]))
                docx_labels = [row.cells[0].text for row in document.tables[0].rows]
                docx_text = "\n".join(
                    cell.text
                    for table in document.tables
                    for row in table.rows
                    for cell in row.cells
                )
                pdf_text = "\n".join(
                    page.extract_text() or ""
                    for page in PdfReader(io.BytesIO(result["pdf_bytes"])).pages
                )
                signing_text = result["signing_preview_bytes"].decode("utf-8")

                for visible in (
                    result["authority_content"], docx_text, pdf_text, signing_text
                ):
                    self.assertNotIn("学生联系方式", visible)
                self.assertNotIn("学生联系方式", model_labels)
                self.assertNotIn("学生联系方式", docx_labels)
                self.assertEqual(len(docx_labels), len(set(docx_labels)))
                for required in ("当前年级", "申请年级", "计划入读年份"):
                    self.assertIn(required, model_labels)
                    self.assertIn(required, docx_labels)
                self.assertEqual(instance["student_phone"], "0400000000")
                self.assertEqual(instance["student_email"], "student@example.test")

    def test_docx_keeps_signature_parties_in_one_fixed_non_splitting_table(self):
        document = Document(io.BytesIO(self._render()["docx_bytes"]))
        self.assertGreaterEqual(len(document.inline_shapes), 1)
        cover_breaks = document._element.xpath(
            "//w:p[.//w:drawing and .//w:br[@w:type='page']]"
        )
        self.assertEqual(len(cover_breaks), 1)
        signature_tables = [
            table for table in document.tables
            if "甲方（签字 / 盖章）" in "\n".join(cell.text for row in table.rows for cell in row.cells)
        ]
        self.assertEqual(len(signature_tables), 1)
        table = signature_tables[0]
        self.assertEqual(len(table.columns), 2)
        left = "\n".join(row.cells[0].text for row in table.rows)
        right = "\n".join(row.cells[1].text for row in table.rows)
        self.assertIn("AAIE", left)
        self.assertIn("委托方（乙方签字）", right)
        for row in table.rows:
            self.assertIsNotNone(row._tr.get_or_add_trPr().find(qn("w:cantSplit")))
        for table in document.tables:
            self.assertIsNotNone(table._tbl.tblGrid)

    def test_email_only_company_contact_keeps_default_phone_and_real_email(self):
        commercial = self._commercial()
        commercial["company_contact"] = "contracts@example.test"
        snapshot = editor.build_domain_v1_snapshot(
            "NewInternationalStudentContract", commercial
        )
        result = editor.render_published_contract_artifacts(
            snapshot,
            student_name="Kexin Chen",
            signer=self._signer(),
            selected_services=list(contract_domain.default_service_ids(
                "NewInternationalStudentContract"
            )),
            contract_instance=self._instance(),
        )
        document = Document(io.BytesIO(result["docx_bytes"]))
        parties = document.tables[0]
        values = {row.cells[0].text: row.cells[1].text for row in parties.rows}
        self.assertEqual(values["甲方电话"], "+61 3 8597 3046")
        self.assertEqual(values["甲方邮箱"], "contracts@example.test")

    def test_pdf_reserves_same_page_left_and_right_signature_anchors(self):
        result = self._render()
        manifest = result["anchor_manifest"]
        aaie = manifest["signature_aaie"]
        family = manifest["signature_family"]
        self.assertEqual(aaie["page"], family["page"])
        self.assertLess(aaie["x"], 250)
        self.assertGreater(family["x"], 300)
        self.assertEqual(manifest["stamp"]["page"], aaie["page"])
        reader = PdfReader(io.BytesIO(result["pdf_bytes"]))
        self.assertGreaterEqual(len(reader.pages), 4)
        final_text = reader.pages[aaie["page"] - 1].extract_text()
        self.assertIn("十一、签署", final_text)
        self.assertIn("甲方（签字 / 盖章）", final_text)
        self.assertIn("委托方（乙方签字）", final_text)
        for label in (
            "甲方法律主体名称", "甲方 ABN / ACN", "甲方授权代表",
            "甲方电话", "甲方联系邮箱", "甲方地址", "乙方签署人",
            "与学生关系", "乙方电话", "乙方邮箱", "乙方地址",
        ):
            self.assertIn(label, final_text)
        self.assertIn("请乙方/监护人自行签名并填写乙方日期", final_text)
        self.assertIn("代理方（甲方盖章）：", final_text)
        self.assertIn(
            "AustralianAcademyofInternationalEducation",
            "".join(final_text.split()),
        )
        self.assertLess(
            final_text.index("代理方（甲方盖章）："),
            final_text.index("授权代表：Sonny Zhang"),
        )
        self.assertLess(
            final_text.index("授权代表：Sonny Zhang"),
            final_text.index("甲方签署日期："),
        )
        self.assertEqual(final_text.count("Kexin Parent"), 1)
        with pdfplumber.open(io.BytesIO(result["pdf_bytes"])) as pdf:
            page = pdf.pages[aaie["page"] - 1]
            middle_dividers = sorted({
                round(float(line["x0"]), 1)
                for line in page.lines
                if abs(float(line["x0"]) - float(line["x1"])) < 0.5
                and float(line["y1"]) - float(line["y0"]) > 50
                and page.width * 0.35 < float(line["x0"]) < page.width * 0.65
            })
        self.assertEqual(middle_dividers, [297.6])

    def test_long_content_flows_before_reserved_signature_page(self):
        long_terms = "长期特别条款：" + ("本条款用于验证自动分页且不得移动签署锚点。" * 500)
        result = self._render(special_terms=long_terms)
        manifest = result["anchor_manifest"]
        reader = PdfReader(io.BytesIO(result["pdf_bytes"]))
        self.assertEqual(manifest["signature_aaie"]["page"], len(reader.pages))
        self.assertEqual(manifest["signature_family"]["page"], len(reader.pages))


if __name__ == "__main__":
    unittest.main()
