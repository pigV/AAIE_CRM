import json
import tempfile
import unittest
from pathlib import Path


class ContractDomainConfigurationTests(unittest.TestCase):
    def setUp(self):
        from contract_domain import (
            LOCAL_STUDENT_APPLICATION_CONTRACT,
            NEW_INTERNATIONAL_STUDENT_CONTRACT,
            TRANSFER_APPLICATION_CONTRACT,
        )

        self.new = NEW_INTERNATIONAL_STUDENT_CONTRACT
        self.transfer = TRANSFER_APPLICATION_CONTRACT
        self.local = LOCAL_STUDENT_APPLICATION_CONTRACT

    def test_three_stable_contract_types_have_exact_chinese_labels(self):
        from contract_domain import contract_type_config, contract_type_options

        self.assertEqual(
            contract_type_options(),
            (
                (self.new, "新留学生合同"),
                (self.transfer, "转学申请合同"),
                (self.local, "本地生申请合同"),
            ),
        )
        self.assertEqual(contract_type_config(self.new).title, "澳洲留学咨询及签证服务委托合同")
        self.assertEqual(contract_type_config(self.transfer).title, "转学申请服务委托合同")
        self.assertEqual(contract_type_config(self.local).title, "本地生申请服务委托合同")

    def test_new_international_defaults_exclude_both_590_services(self):
        from contract_domain import allowed_service_ids, default_service_ids

        defaults = default_service_ids(self.new)
        allowed = allowed_service_ids(self.new)
        self.assertIn("student_visa_lodgement", defaults)
        self.assertIn("subclass_590_guardian_visa_lodgement", allowed)
        self.assertIn("subclass_590_guardian_visa_consulting", allowed)
        self.assertNotIn("subclass_590_guardian_visa_lodgement", defaults)
        self.assertNotIn("subclass_590_guardian_visa_consulting", defaults)

    def test_transfer_defaults_include_transfer_offer_coe_oshc_but_no_visa(self):
        from contract_domain import default_service_ids

        defaults = default_service_ids(self.transfer)
        self.assertIn("school_transfer_planning_and_application", defaults)
        self.assertIn("offer_coe_oshc_process_assistance", defaults)
        self.assertFalse(any("visa" in value for value in defaults))

    def test_local_defaults_exclude_visa_oshc_and_coe(self):
        from contract_domain import default_service_ids, service_label

        defaults = default_service_ids(self.local)
        rendered = "\n".join(service_label(value) for value in defaults)
        self.assertIn("选校", rendered)
        self.assertIn("录取", rendered)
        for forbidden in ("签证", "590", "OSHC", "COE", "CAAW"):
            self.assertNotIn(forbidden, rendered)

    def test_submitted_services_are_filtered_by_selected_contract_type(self):
        from contract_domain import normalize_selected_services

        selected = normalize_selected_services(
            self.local,
            [
                "school_information_and_selection_advice",
                "student_visa_lodgement",
                "offer_and_enrolment_coordination",
                "school_information_and_selection_advice",
            ],
            submitted=True,
        )
        self.assertEqual(
            selected,
            (
                "school_information_and_selection_advice",
                "offer_and_enrolment_coordination",
            ),
        )

    def test_explicitly_cleared_services_remain_empty(self):
        from contract_domain import normalize_selected_services

        self.assertEqual(normalize_selected_services(self.new, [], submitted=True), ())

    def test_contract_sections_are_composed_by_profile_without_keyword_deletion(self):
        from contract_domain import compose_contract_sections

        international = "\n".join(compose_contract_sections(self.new).values())
        transfer = "\n".join(compose_contract_sections(self.transfer).values())
        local = "\n".join(compose_contract_sections(self.local).values())
        self.assertIn("签证", international)
        self.assertIn("COE", transfer)
        self.assertIn("OSHC", transfer)
        self.assertNotIn("签证", transfer)
        for forbidden in ("签证", "COE", "OSHC", "移民局"):
            self.assertNotIn(forbidden, local)


class ContractSignerResolutionTests(unittest.TestCase):
    def setUp(self):
        self.student_data = {
            "father_passport_first_name": "Father",
            "father_passport_last_name": "Zhang",
            "father_mobile": "0400-FATHER",
            "father_email": "father@example.com",
            "mother_passport_first_name": "Mother",
            "mother_passport_last_name": "Li",
            "mother_mobile": "0400-MOTHER",
            "mother_email": "mother@example.com",
        }
        self.parent_rows = (
            {
                "parent_type": "父亲",
                "passport_first_name": "Father",
                "passport_last_name": "Zhang",
                "mobile": "0400-FATHER",
                "email": "father@example.com",
            },
            {
                "parent_type": "母亲",
                "passport_first_name": "Mother",
                "passport_last_name": "Li",
                "mobile": "0400-MOTHER",
                "email": "mother@example.com",
            },
        )

    def test_mother_signer_is_resolved_as_one_complete_object(self):
        from contract_domain import resolve_signer

        self.assertEqual(
            resolve_signer("mother", self.student_data, self.parent_rows),
            {
                "source": "mother",
                "name": "Mother Li",
                "relationship": "母亲",
                "phone": "0400-MOTHER",
                "email": "mother@example.com",
            },
        )

    def test_father_signer_is_resolved_as_one_complete_object(self):
        from contract_domain import resolve_signer

        signer = resolve_signer("father", self.student_data, self.parent_rows)
        self.assertEqual(signer["name"], "Father Zhang")
        self.assertEqual(signer["relationship"], "父亲")
        self.assertEqual(signer["phone"], "0400-FATHER")
        self.assertEqual(signer["email"], "father@example.com")
        self.assertNotIn("MOTHER", signer["phone"])

    def test_missing_selected_parent_phone_never_borrows_other_parent(self):
        from contract_domain import ContractSignerValidationError, resolve_signer

        rows = [dict(row) for row in self.parent_rows]
        rows[1]["mobile"] = ""
        data = dict(self.student_data, mother_mobile="")
        with self.assertRaisesRegex(ContractSignerValidationError, "电话"):
            resolve_signer("mother", data, rows)

    def test_missing_selected_parent_email_accepts_snapshot_only_override(self):
        from contract_domain import resolve_signer

        rows = [dict(row) for row in self.parent_rows]
        rows[0]["email"] = ""
        data = dict(self.student_data, father_email="")
        signer = resolve_signer(
            "father",
            data,
            rows,
            overrides={"email": "father.contract@example.com"},
        )
        self.assertEqual(signer["email"], "father.contract@example.com")
        self.assertEqual(rows[0]["email"], "")
        self.assertEqual(data["father_email"], "")


class ContractSnapshotTests(unittest.TestCase):
    def test_versioned_snapshot_round_trip_freezes_signer_services_and_sections(self):
        from contract_domain import (
            NEW_INTERNATIONAL_STUDENT_CONTRACT,
            build_contract_snapshot,
            snapshot_payload,
        )

        payload = {
            "contract_type": NEW_INTERNATIONAL_STUDENT_CONTRACT,
            "service_scope_ids": ["study_planning_consulting"],
            "service_scope_items": ["澳洲留学规划咨询"],
            "contract_sections": {"二、服务项目及服务范围": "1. 澳洲留学规划咨询"},
            "signer": {
                "source": "mother",
                "name": "Mother Li",
                "relationship": "母亲",
                "phone": "0400-MOTHER",
                "email": "mother@example.com",
            },
            "student_full_name": "Student Zhang",
        }
        snapshot = build_contract_snapshot(payload)
        encoded = json.dumps({"contract_snapshot": snapshot}, ensure_ascii=False)

        restored = snapshot_payload(encoded)
        payload["service_scope_items"].append("后来新增服务")
        payload["signer"]["phone"] = "CHANGED"

        self.assertEqual(restored["snapshot_version"], 1)
        self.assertEqual(restored["service_scope_items"], ["澳洲留学规划咨询"])
        self.assertEqual(restored["signer"]["phone"], "0400-MOTHER")
        self.assertEqual(
            restored["contract_sections"]["二、服务项目及服务范围"],
            "1. 澳洲留学规划咨询",
        )


class ContractTemplateCompositionTests(unittest.TestCase):
    def test_service_scope_and_type_sections_always_override_shared_template(self):
        import app as app_module
        from contract_domain import LOCAL_STUDENT_APPLICATION_CONTRACT, compose_contract_sections
        from database import Database
        from student_doc import STUDY_CONTRACT_STANDARD_TEXT_DEFAULTS

        payload = {
            "contract_type": LOCAL_STUDENT_APPLICATION_CONTRACT,
            "service_scope_items": ["学校信息介绍与选校建议"],
            "contract_sections": compose_contract_sections(LOCAL_STUDENT_APPLICATION_CONTRACT),
            "standard_clause_texts": dict(STUDY_CONTRACT_STANDARD_TEXT_DEFAULTS),
            "refund_policy_text": "学校申请服务退费按本合同约定处理。",
            "governing_law": "澳大利亚维多利亚州法律",
            "dispute_resolution": "双方先友好协商。",
            "special_terms": "",
        }

        old_db = app_module.db
        try:
            with tempfile.TemporaryDirectory() as folder:
                app_module.db = Database(Path(folder) / "aaie_crm.db")
                app_module.db.initialize()
                overrides = app_module.contract_template_section_override_map(payload)
        finally:
            app_module.db = old_db

        self.assertEqual(
            overrides["二、服务项目及服务范围"],
            "1. 学校信息介绍与选校建议",
        )
        self.assertEqual(
            overrides["三、服务费用与第三方费用"],
            payload["contract_sections"]["三、服务费用与第三方费用"],
        )
        rendered = "\n".join(overrides.values())
        self.assertNotIn("签证", rendered)
        self.assertNotIn("移民局", rendered)

    def test_shared_engine_generates_all_three_types_and_rebuilds_service_scope(self):
        import app as app_module
        from contract_domain import (
            LOCAL_STUDENT_APPLICATION_CONTRACT,
            NEW_INTERNATIONAL_STUDENT_CONTRACT,
            TRANSFER_APPLICATION_CONTRACT,
            compose_contract_sections,
            contract_type_config,
            default_service_ids,
            service_label,
        )
        from database import Database
        from docx import Document

        old_db = app_module.db
        try:
            with tempfile.TemporaryDirectory() as folder:
                root = Path(folder)
                app_module.db = Database(root / "aaie_crm.db")
                app_module.db.initialize()
                app_module.db.set_setting("company_abn_acn", "ABN 11 222 333 444", 1)
                results = {}
                for type_id in (
                    NEW_INTERNATIONAL_STUDENT_CONTRACT,
                    TRANSFER_APPLICATION_CONTRACT,
                    LOCAL_STUDENT_APPLICATION_CONTRACT,
                ):
                    config = contract_type_config(type_id)
                    selected = list(default_service_ids(type_id)[:2])
                    standard = app_module.study_contract_standard_texts()
                    standard.update({"intro": config.intro, "service_scope_intro": config.service_scope_intro})
                    payload = {
                        **app_module.contract_company_payload(),
                        "contract_uid": f"TEST-{type_id}",
                        "contract_type": type_id,
                        "contract_title": config.title,
                        "contract_cover_title_lines": list(config.cover_title_lines),
                        "contract_cover_service_summary": config.cover_service_summary,
                        "contract_cover_title_en": config.cover_title_en,
                        "service_fee_label": config.service_fee_label,
                        "contract_date": "2026-08-12",
                        "student_full_name": "Student Test",
                        "parent_signer_name": "Mother Test",
                        "parent_signer_relationship": "母亲",
                        "parent_phone": "0400-MOTHER",
                        "parent_email": "mother@example.com",
                        "service_type": "学校申请服务",
                        "service_scope_ids": selected,
                        "service_scope_items": [service_label(value) for value in selected],
                        "service_scope_json": json.dumps(selected, ensure_ascii=False),
                        "service_fee_currency": "澳洲公户",
                        "consulting_and_visa_service_fee": "2400",
                        "final_fee_amount": "2400",
                        "refund_policy_text": "按本合同约定处理。",
                        "standard_clause_texts": standard,
                        "contract_sections": compose_contract_sections(type_id),
                    }
                    destination = root / f"{type_id}.docx"
                    app_module.generate_study_contract_document(payload, destination)
                    document = Document(destination)
                    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
                    for table in document.tables:
                        text += "\n" + "\n".join(cell.text for row in table.rows for cell in row.cells)
                    self.assertEqual(document.core_properties.title, config.title)
                    for service_id in selected:
                        self.assertIn(service_label(service_id), text)
                    results[type_id] = text

                self.assertIn("签证", results[NEW_INTERNATIONAL_STUDENT_CONTRACT])
                self.assertNotIn("签证", results[TRANSFER_APPLICATION_CONTRACT])
                for forbidden in ("签证", "COE", "OSHC", "移民局"):
                    self.assertNotIn(forbidden, results[LOCAL_STUDENT_APPLICATION_CONTRACT])

                local_config = contract_type_config(LOCAL_STUDENT_APPLICATION_CONTRACT)
                local_services = list(default_service_ids(LOCAL_STUDENT_APPLICATION_CONTRACT)[:2])
                removed_label = service_label(local_services[1])
                local_payload = {
                    **app_module.contract_company_payload(),
                    "contract_uid": "TEST-LOCAL-REBUILT",
                    "contract_type": LOCAL_STUDENT_APPLICATION_CONTRACT,
                    "contract_title": local_config.title,
                    "contract_cover_title_lines": list(local_config.cover_title_lines),
                    "contract_cover_service_summary": local_config.cover_service_summary,
                    "contract_cover_title_en": local_config.cover_title_en,
                    "service_fee_label": local_config.service_fee_label,
                    "contract_date": "2026-08-12",
                    "student_full_name": "Student Test",
                    "parent_signer_name": "Mother Test",
                    "parent_signer_relationship": "母亲",
                    "parent_phone": "0400-MOTHER",
                    "parent_email": "mother@example.com",
                    "service_type": "学校申请服务",
                    "service_scope_ids": local_services[:1],
                    "service_scope_items": [service_label(local_services[0])],
                    "service_scope_json": json.dumps(local_services[:1], ensure_ascii=False),
                    "service_fee_currency": "澳洲公户",
                    "consulting_and_visa_service_fee": "2400",
                    "final_fee_amount": "2400",
                    "refund_policy_text": "按本合同约定处理。",
                    "standard_clause_texts": {
                        **app_module.study_contract_standard_texts(),
                        "intro": local_config.intro,
                        "service_scope_intro": local_config.service_scope_intro,
                    },
                    "contract_sections": compose_contract_sections(LOCAL_STUDENT_APPLICATION_CONTRACT),
                }
                rebuilt = root / "LocalStudentApplicationContract-rebuilt.docx"
                app_module.generate_study_contract_document(local_payload, rebuilt)
                rebuilt_doc = Document(rebuilt)
                rebuilt_text = "\n".join(paragraph.text for paragraph in rebuilt_doc.paragraphs)
                for table in rebuilt_doc.tables:
                    rebuilt_text += "\n" + "\n".join(cell.text for row in table.rows for cell in row.cells)
                self.assertIn(service_label(local_services[0]), rebuilt_text)
                self.assertNotIn(removed_label, rebuilt_text)
                rebuilt_pdf = root / "LocalStudentApplicationContract-rebuilt.pdf"
                app_module.convert_docx_to_pdf_with_tables(rebuilt, rebuilt_pdf)
                pdf_text = "\n".join(
                    page.extract_text() or ""
                    for page in app_module.PdfReader(str(rebuilt_pdf)).pages
                )
                self.assertIn(service_label(local_services[0]), pdf_text)
                self.assertNotIn(removed_label, pdf_text)
        finally:
            app_module.db = old_db


class ContractAppAdapterTests(unittest.TestCase):
    def test_student_payload_uses_type_filter_and_atomic_selected_signer(self):
        import app as app_module
        from contract_domain import LOCAL_STUDENT_APPLICATION_CONTRACT
        from database import Database

        old_db = app_module.db
        try:
            with tempfile.TemporaryDirectory() as folder:
                db = Database(Path(folder) / "aaie_crm.db")
                db.initialize()
                app_module.db = db
                stamp = "2026-08-12T10:00:00+10:00"
                student_data = {
                    "father_passport_first_name": "Father",
                    "father_passport_last_name": "Zhang",
                    "father_mobile": "0400-FATHER",
                    "father_email": "father@example.com",
                    "mother_passport_first_name": "Mother",
                    "mother_passport_last_name": "Li",
                    "mother_mobile": "0400-MOTHER",
                    "mother_email": "mother@example.com",
                }
                db.execute(
                    """INSERT INTO students(
                       student_id,passport_last_name,passport_first_name,current_status,
                       created_date,updated_at,form_data_json
                    ) VALUES(?,?,?,?,?,?,?)""",
                    ("TDOMAIN001", "Student", "Test", "Active", stamp, stamp, json.dumps(student_data)),
                )
                for parent_type, first, last, phone, email in (
                    ("父亲", "Father", "Zhang", "0400-FATHER", "father@example.com"),
                    ("母亲", "Mother", "Li", "0400-MOTHER", "mother@example.com"),
                ):
                    db.execute(
                        """INSERT INTO parents(student_id,parent_type,passport_first_name,passport_last_name,mobile,email)
                           VALUES(?,?,?,?,?,?)""",
                        ("TDOMAIN001", parent_type, first, last, phone, email),
                    )
                student = db.one("SELECT * FROM students WHERE student_id='TDOMAIN001'")
                payload = app_module.contract_student_payload(
                    student,
                    {
                        "contract_type": LOCAL_STUDENT_APPLICATION_CONTRACT,
                        "parent_signer_source": "mother",
                        "service_scope_submitted": "1",
                        "service_scope": [
                            "school_information_and_selection_advice",
                            "student_visa_lodgement",
                            "offer_and_enrolment_coordination",
                        ],
                        "service_fee_currency": "支付宝",
                        "consulting_and_visa_service_fee": "2400",
                    },
                )

                self.assertEqual(payload["contract_type"], LOCAL_STUDENT_APPLICATION_CONTRACT)
                self.assertEqual(
                    payload["service_scope_ids"],
                    ["school_information_and_selection_advice", "offer_and_enrolment_coordination"],
                )
                self.assertEqual(
                    payload["signer"],
                    {
                        "source": "mother",
                        "name": "Mother Li",
                        "relationship": "母亲",
                        "phone": "0400-MOTHER",
                        "email": "mother@example.com",
                    },
                )
                self.assertEqual(payload["parent_phone"], "0400-MOTHER")
                self.assertEqual(payload["parent_email"], "mother@example.com")
                self.assertNotIn("签证", "\n".join(payload["contract_sections"].values()))
        finally:
            app_module.db = old_db

    def test_generate_route_persists_stable_type_and_versioned_snapshot_without_schema_change(self):
        import app as app_module
        from contract_domain import TRANSFER_APPLICATION_CONTRACT, snapshot_payload
        from database import Database

        class FakeNas:
            def student_folder(self, student):
                return Path(student["folder_path"]).resolve()

            def require_category_access(self, student, category, *, write=False):
                target = self.student_folder(student) / category
                target.mkdir(parents=True, exist_ok=True)
                return target

        old_db = app_module.db
        old_nas = app_module.nas
        old_generate = app_module.generate_study_contract_document
        try:
            with tempfile.TemporaryDirectory() as folder:
                root = Path(folder)
                db = Database(root / "aaie_crm.db")
                db.initialize()
                app_module.db = db
                app_module.nas = FakeNas()
                app_module.generate_study_contract_document = lambda payload, target: Path(target).write_bytes(b"isolated contract") or target
                db.set_setting("company_abn_acn", "ABN 11 222 333 444", 1)
                student_folder = root / "Student"
                student_folder.mkdir()
                stamp = "2026-08-12T10:00:00+10:00"
                student_data = {
                    "mother_passport_first_name": "Mother",
                    "mother_passport_last_name": "Li",
                    "mother_mobile": "0400-MOTHER",
                    "mother_email": "mother@example.com",
                }
                db.execute(
                    """INSERT INTO students(
                       student_id,passport_last_name,passport_first_name,current_status,
                       created_date,updated_at,folder_path,nas_status,form_data_json
                    ) VALUES(?,?,?,?,?,?,?,?,?)""",
                    (
                        "TROUTE-UPGRADE",
                        "Student",
                        "Transfer",
                        "Active",
                        stamp,
                        stamp,
                        str(student_folder),
                        "ready",
                        json.dumps(student_data),
                    ),
                )
                handler = object.__new__(app_module.CRMHandler)
                handler.form = {
                    "contract_type": TRANSFER_APPLICATION_CONTRACT,
                    "contract_date": "2026-08-12",
                    "parent_signer_source": "mother",
                    "service_scope_submitted": "1",
                    "service_scope": ["school_transfer_planning_and_application"],
                    "service_fee_currency": "支付宝",
                    "consulting_and_visa_service_fee": "2400",
                }
                handler.user = {
                    "id": 1,
                    "username": "admin",
                    "role_name": "Admin",
                    "permissions_json": json.dumps(["*"]),
                }
                handler.client_address = ("127.0.0.1", 12345)
                handler.headers = {}
                handler.sent_headers = []
                handler.send_response = lambda status: setattr(handler, "response_status", status)
                handler.send_header = lambda key, value: handler.sent_headers.append((key, value))
                handler.end_headers = lambda: None

                handler.contract_generate("POST", "TROUTE-UPGRADE")

                contract = db.one("SELECT * FROM contracts WHERE student_id='TROUTE-UPGRADE'")
                self.assertEqual(handler.response_status, 303)
                self.assertEqual(contract["contract_type"], TRANSFER_APPLICATION_CONTRACT)
                snapshot = snapshot_payload(contract["standalone_data_json"])
                self.assertEqual(snapshot["snapshot_version"], 1)
                self.assertEqual(snapshot["contract_type"], TRANSFER_APPLICATION_CONTRACT)
                self.assertEqual(snapshot["signer"]["source"], "mother")
                self.assertEqual(snapshot["service_scope_ids"], ["school_transfer_planning_and_application"])
                edit_form = app_module.contract_editor_base_form(contract, {})
                self.assertEqual(edit_form["contract_type"], TRANSFER_APPLICATION_CONTRACT)
                self.assertEqual(edit_form["parent_signer_source"], "mother")
                self.assertEqual(edit_form["parent_phone"], "0400-MOTHER")
                db.execute(
                    "UPDATE students SET form_data_json=? WHERE student_id=?",
                    (json.dumps({"mother_mobile": "CHANGED-MASTER"}), "TROUTE-UPGRADE"),
                )
                db.execute(
                    "UPDATE contracts SET service_scope_json=?,parent_phone=? WHERE id=?",
                    (json.dumps(["student_visa_lodgement"]), "CHANGED-ROW", contract["id"]),
                )
                frozen_payload = app_module.contract_payload_from_existing(
                    db.one("SELECT * FROM contracts WHERE id=?", (contract["id"],))
                )
                self.assertEqual(frozen_payload["signer"]["phone"], "0400-MOTHER")
                self.assertEqual(frozen_payload["service_scope_ids"], ["school_transfer_planning_and_application"])
        finally:
            app_module.db = old_db
            app_module.nas = old_nas
            app_module.generate_study_contract_document = old_generate

    def test_shared_contract_form_renders_three_types_and_atomic_signer_fields(self):
        import app as app_module
        from database import Database

        old_db = app_module.db
        try:
            with tempfile.TemporaryDirectory() as folder:
                db = Database(Path(folder) / "aaie_crm.db")
                db.initialize()
                app_module.db = db
                stamp = "2026-08-12T10:00:00+10:00"
                student_data = {
                    "father_passport_first_name": "Father",
                    "father_passport_last_name": "Zhang",
                    "father_mobile": "0400-FATHER",
                    "father_email": "father@example.com",
                    "mother_passport_first_name": "Mother",
                    "mother_passport_last_name": "Li",
                    "mother_mobile": "",
                    "mother_email": "mother@example.com",
                }
                db.execute(
                    """INSERT INTO students(
                       student_id,passport_last_name,passport_first_name,current_status,
                       created_date,updated_at,form_data_json
                    ) VALUES(?,?,?,?,?,?,?)""",
                    ("TFORM001", "Student", "Form", "Active", stamp, stamp, json.dumps(student_data)),
                )
                handler = object.__new__(app_module.CRMHandler)
                handler.user = {
                    "id": 1,
                    "username": "admin",
                    "role_name": "Admin",
                    "permissions_json": json.dumps(["*"]),
                }
                handler.csrf = "test-csrf"
                student = db.one("SELECT * FROM students WHERE student_id='TFORM001'")
                refunds = db.all("SELECT * FROM contract_refund_templates WHERE is_active=1 ORDER BY id")

                html = handler.contract_panel(student, [], refunds)

                self.assertEqual(html.count('name="contract_type"'), 3)
                for type_id, label in (
                    ("NewInternationalStudentContract", "新留学生合同"),
                    ("TransferApplicationContract", "转学申请合同"),
                    ("LocalStudentApplicationContract", "本地生申请合同"),
                ):
                    self.assertIn(f'value="{type_id}"', html)
                    self.assertIn(label, html)
                self.assertIn('data-contract-type-radio', html)
                self.assertIn('name="parent_signer_source" value="mother"', html)
                self.assertIn('name="parent_signer_source" value="father"', html)
                self.assertIn('name="parent_phone"', html)
                self.assertIn('name="parent_email"', html)
                self.assertIn('data-mother-phone=""', html)
                self.assertIn('data-father-phone="0400-FATHER"', html)
                self.assertIn('data-contract-service-types=', html)
                self.assertNotIn("mobile-contract-form", html)
        finally:
            app_module.db = old_db


if __name__ == "__main__":
    unittest.main()
