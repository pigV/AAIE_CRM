"""Production-derived isolated DB, local-only NAS, real CRM forms and submit handler."""
import hashlib,json,os,sqlite3,sys,secrets,signal,threading
from pathlib import Path
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT))
BOX=ROOT/'SAFEBOX/transfer-contract-workflow-20261009';RUNTIME=BOX/'runtime';PORT=19042
RUNTIME.mkdir(parents=True,exist_ok=True)
# No startup loop, provider, mail, sync, backup, or shared service is started.
import app
from database import Database
from security import password_hash
from tests.transfer_contract_final_signed_browser_harness import IsolatedNas
snapshot=json.loads((BOX/'evidence/snapshot.json').read_text())
path=Path(snapshot['isolated_database']).resolve()
if RUNTIME.resolve() not in path.parents or not snapshot.get('source_read_only'):
 raise RuntimeError('Isolated coherent read-only snapshot required')

class WorkflowNas(IsolatedNas):
 def require_category_access(self, student, category, *, write=False):
  folder=(self.student_folder(student)/str(category)).resolve()
  if self.root not in folder.parents:raise RuntimeError('isolated category escaped local NAS')
  if write:folder.mkdir(parents=True,exist_ok=True)
  return folder

app.db=Database(path);app.DB_PATH=path;app.nas=WorkflowNas(RUNTIME/'nas');app.SAFE_PRODUCTION_ACTIVE=True
for name,sub in [('PREVIEW_DIR','previews'),('UPLOAD_REVIEW_DIR','upload-staging'),('SCHOOL_PLAN_REPORT_DIR','school-plans'),('CONTRACT_SIGNING_DIR','signing'),('LOG_DIR','logs'),('BACKUP_DIR','backups'),('SYSTEM_BACKUP_DIR','backups')]:setattr(app,name,RUNTIME/sub)
prepared='--serve-prepared' in sys.argv
password=secrets.token_urlsafe(24)
(RUNTIME/'browser-login.json').write_text(json.dumps({'username':'admin','password':password}));os.chmod(RUNTIME/'browser-login.json',0o600)
with app.db.transaction() as c:
 admin=c.execute("select id from users where username='admin'").fetchone()[0]
 c.execute('update users set password_hash=?,must_change_password=0 where id=?',(password_hash(password),admin))

def fingerprint(c):
 out={}
 for table in ('contracts','contract_versions','contract_signatures','contract_signing_links','contract_signing_logs'):
  ids=[r[0] for r in c.execute("select id from contracts where contract_type='TransferApplicationContract' and id<211")]
  key='id' if table=='contracts' else 'contract_id'
  rows=list(c.execute(f'SELECT * FROM {table} WHERE {key} IN ({",".join("?" for _ in ids)}) ORDER BY id',ids))
  out[table]=hashlib.sha256(json.dumps([list(r) for r in rows],ensure_ascii=False,default=str).encode()).hexdigest()
 return {'count':len(ids),'hashes':out}

if not prepared:
 with app.db.transaction() as c:
  before=fingerprint(c);(BOX/'evidence/historical-before.json').write_text(json.dumps(before,indent=2))
  orig=dict(c.execute("select * from students where student_id='20260163'").fetchone())
  for student in ('99961001','99961002','99961003','99961004'):
   row=dict(orig,student_id=student,passport_first_name='Workflow',passport_last_name=student,folder_path=str(RUNTIME/'nas'/student),form_data_json=json.dumps({**json.loads(orig['form_data_json'] or '{}'), 'father_passport_first_name':'Synthetic','father_passport_last_name':'Father','mother_passport_first_name':'Synthetic','mother_passport_last_name':'Mother','father_mobile':'0411111111','mother_mobile':'0422222222','father_email':'father@example.test','mother_email':'mother@example.test'}),nas_status='ready')
   for key in ('source_parent_form_token','converted_from_parent_token'):
    if key in row:row[key]=None
   keys=list(row);c.execute(f'INSERT INTO students({",".join(keys)}) VALUES({",".join("?" for _ in keys)})',[row[k] for k in keys])
  # Relocate only isolated student storage. Historical contract rows remain byte-identical.
  c.execute('update students set folder_path=?,nas_status=? where student_id=?',(str(RUNTIME/'nas'/'20260163'),'ready','20260163'))
  for student in ('99961001','99961003','99961004'):
   c.execute("insert into parent_form_tokens(token,student_id,intake_type,submitted_at,created_at,expires_at,created_by) values(?,?,'new_student_transfer',?,?,?,?)",('workflow-'+student,student,app.now_iso(),app.now_iso(),'2030-01-01',admin))
  c.execute("insert into parent_form_tokens(token,student_id,intake_type,submitted_at,created_at,expires_at,created_by) values('workflow-second-numeric','99961004','new_student_transfer',?,?,?,?)",(app.now_iso(),app.now_iso(),'2030-01-01',admin))
  for student in ('99961002','99961003'):
   cycle=c.execute("insert into application_cycles(student_id,cycle_uid,sequence_no,cycle_type,education_stream,status,opened_at,created_at,updated_at) values(?,?,1,'initial','school_education','active',?,?,?)",(student,'workflow-cycle-'+student,app.now_iso(),app.now_iso(),app.now_iso())).lastrowid
   c.execute("insert into contracts(student_id,contract_uid,contract_type,contract_status,application_cycle_id,generated_at,contract_category,contract_title,generated_by,updated_at) values(?,?,'LocalStudentApplicationContract','FullyExecuted',?,?,'普通合同','Synthetic Ordinary',?,?)",(student,'ORDINARY-'+student,cycle,app.now_iso(),admin,app.now_iso()))
   c.execute("insert into student_school_enrolments(student_id,enrolment_role,status,confirmed_at,school_name_snapshot,created_at,updated_at,created_by,updated_by,confirmed_by,confirmation_reason) values(?,'current','active',?,'Synthetic School',?,?,?,?,?,'Synthetic confirmed enrolment')",(student,app.now_iso(),app.now_iso(),app.now_iso(),admin,admin,admin))
  for key,value in [('STUDENT_CASES_READ_ENABLED','true'),('TRANSFER_CASES_WRITE_ENABLED','true'),('auto_backup_enabled','false'),('CRM_PUBLIC_BASE_URL',f'http://127.0.0.1:{PORT}'),('PARENT_FORM_PUBLIC_BASE_URL',f'http://127.0.0.1:{PORT}'),('CONTRACT_SIGNING_BASE_URL',f'http://127.0.0.1:{PORT}')]:
   c.execute('insert into system_settings(key,value,updated_at) values(?,?,?) on conflict(key) do update set value=excluded.value',(key,value,app.now_iso()))

class Handler(app.CRMHandler):
 def do_GET(self):
  from urllib.parse import urlparse,parse_qs
  parsed=urlparse(self.path)
  if parsed.path=='/qa/result':
   student=parse_qs(parsed.query).get('student',[''])[0]
   if student not in ('99961001','99961002','99961003','99961004','20260163'):return self.send_error(404)
   rows=app.db.all("select * from contracts where student_id=? and contract_type='TransferApplicationContract' ORDER BY id",(student,));out={'count':len(rows)}
   if rows:
    row=rows[-1];snapshot=app.snapshot_payload(row['standalone_data_json'] or '');out.update({k:row[k] for k in ('id','consulting_and_visa_service_fee','discount_amount','final_fee_amount','service_fee_currency')});out['payment_plan']=snapshot.get('commercial',{}).get('payment_plan');out['gst']=snapshot.get('commercial',{}).get('gst_amount');out['source_count']=app.db.one('select count(*) n from transfer_source_authorities where transfer_contract_id=?',(row['id'],))['n']
    if row['contract_file_path']:
     from docx import Document
     doc=Document(row['contract_file_path']);text='\n'.join([p.text for p in doc.paragraphs]+[cell.text for t in doc.tables for r in t.rows for cell in r.cells]);out['file_has_1320']='1320' in text;out['file_has_installments']='特殊分期付款计划' in text;out['file_has_single_payment']='一次性' in text;out['file_sha256']=hashlib.sha256(Path(row['contract_file_path']).read_bytes()).hexdigest()
   payload=json.dumps(out).encode();self.send_response(200);self.send_header('Content-Type','application/json');self.send_header('Content-Length',str(len(payload)));self.end_headers();self.wfile.write(payload);return
  return super().do_GET()

server=app.ReusableThreadingHTTPServer(('127.0.0.1',PORT),Handler);server.service_mode='admin'
def stop(*_):threading.Thread(target=server.shutdown,daemon=True).start()
signal.signal(signal.SIGTERM,stop);signal.signal(signal.SIGINT,stop)
print('ISOLATED_BROWSER_READY http://127.0.0.1:19042',flush=True)
try:server.serve_forever()
finally:
 server.server_close()
 (RUNTIME/'browser-login.json').write_text(json.dumps({'username':'admin','password':None,'status':'ephemeral credential cleared'}))
 c=sqlite3.connect(path);(BOX/'evidence/historical-after.json').write_text(json.dumps(fingerprint(c),indent=2));c.close()
