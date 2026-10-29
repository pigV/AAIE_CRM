import hashlib
import json
import sqlite3
import tempfile
import unittest
from contextlib import contextmanager
from pathlib import Path
from unittest import mock

import contract_domain

try:
    import contract_template_editor as editor
except ModuleNotFoundError:
    editor = None


ROOT = Path(__file__).resolve().parents[1]
MIGRATION = ROOT / "migrations" / "20260818_025_contract_template_editor_v1.sql"


class ContractTemplateEditorV1Tests(unittest.TestCase):
    def setUp(self):
        self.assertIsNotNone(editor, "Contract Template Editor V1 module is not implemented")
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.conn = sqlite3.connect(self.root / "editor.sqlite3")
        self.conn.row_factory = sqlite3.Row
        self.addCleanup(self.conn.close)
        self.conn.execute("PRAGMA foreign_keys=ON")
        self.conn.executescript(MIGRATION.read_text(encoding="utf-8"))
        self.service = editor.ContractTemplateEditorService(self.conn)

    def _import_v1(self):
        return self.service.import_domain_v1(actor_user_id=1)

    def _authority_v2_published(self, template_type):
        current = self.service.current_published(template_type)
        snapshot = self.service._snapshot(current["id"])
        snapshot["commercial"] = editor.authority_v2_commercial(
            template_type, snapshot["commercial"]
        )
        version_id = self.service._insert_snapshot(
            int(current["identity_id"]),
            int(current["version_number"]) + 1,
            "Published",
            snapshot,
            actor=1,
            source_version_id=current["id"],
            immutable=True,
        )
        self.conn.execute(
            "UPDATE contract_template_identities "
            "SET current_published_version_id=? WHERE id=?",
            (version_id, current["identity_id"]),
        )
        self.conn.execute(
            "UPDATE contract_template_versions SET status='Historical' WHERE id=?",
            (current["id"],),
        )
        self.conn.commit()
        return self.service._version(version_id)

    def _valid_commercial(self):
        return {
            "currency": "AUD",
            "service_fee": "3000.00",
            "payment_terms": "签署后七日内支付。",
            "company_name": "Australian Academy of International Education",
            "company_abn": "12 345 678 901",
            "company_address": "Melbourne VIC",
            "company_contact": "contracts@example.test",
            "refund_policy": "按已完成服务比例及第三方政策处理退款。",
        }

    @staticmethod
    def _simulation_student():
        return {
            "student_name": "Boundary Student",
            "student_id": "T-EDITOR-001",
            "date_of_birth": "2011-02-03",
            "current_year_level": "Year 7",
            "applying_year_level": "Year 8",
            "preferred_intake_year": "2027",
            "preferred_intake_term": "02",
            "father_passport_first_name": "Father",
            "father_passport_last_name": "Only",
            "father_mobile": "",
            "father_email": "father@example.test",
            "mother_passport_first_name": "Mother",
            "mother_passport_last_name": "Other",
            "mother_mobile": "0400000000",
            "mother_email": "mother@example.test",
        }

    def _simulate_ready_draft(self, *, block_body=None, service=None):
        self._import_v1()
        service = service or self.service
        draft = service.create_draft("NewInternationalStudentContract", actor_user_id=1)
        lock = service.acquire_lock(
            "NewInternationalStudentContract", draft["id"], owner_user_id=1
        )
        draft = service.autosave_settings(
            draft["id"],
            commercial=self._valid_commercial(),
            services=[
                {
                    "service_id": "study_planning_consulting",
                    "enabled": True,
                    "default_selected": True,
                }
            ],
            expected_revision=draft["revision"],
            lock_token=lock["token"],
            actor_user_id=1,
        )
        if block_body is not None:
            draft = service.autosave_block(
                draft["id"],
                "profile_3",
                title="三、服务费用与第三方费用",
                body=block_body,
                enabled=True,
                expected_revision=draft["revision"],
                lock_token=lock["token"],
                actor_user_id=1,
            )
        result = service.simulate(
            draft["id"],
            source_type="system",
            student_snapshot=self._simulation_student(),
            signer_source="father",
            signer_overrides={"phone": "0411111111"},
            selected_services=["study_planning_consulting"],
            output_root=self.root / "previews",
            actor_user_id=1,
            expected_revision=draft["revision"],
            lock_token=lock["token"],
        )
        return draft, lock, result

    def _latest_simulation(self, version_id):
        return self.conn.execute(
            "SELECT * FROM contract_template_simulations WHERE version_id=? ORDER BY id DESC LIMIT 1",
            (version_id,),
        ).fetchone()

    def _insert_simulation_copy(self, row, **overrides):
        values = dict(row)
        values.pop("id", None)
        values.update(overrides)
        columns = list(values)
        cursor = self.conn.execute(
            f"INSERT INTO contract_template_simulations({','.join(columns)}) "
            f"VALUES({','.join('?' for _ in columns)})",
            [values[column] for column in columns],
        )
        return cursor.lastrowid

    def _insert_rebound_artifact_simulation(self, row, filename, artifact_path):
        values = dict(row)
        artifact_hash = hashlib.sha256(artifact_path.read_bytes()).hexdigest()
        manifest = json.loads(values["artifact_manifest_json"])
        manifest[filename] = {
            "sha256": artifact_hash,
            "bytes": artifact_path.stat().st_size,
        }
        hash_column = {
            "preview.docx": "docx_artifact_hash",
            "preview.pdf": "pdf_artifact_hash",
            "signing-preview.html": "signing_artifact_hash",
        }[filename]
        values[hash_column] = artifact_hash
        values["artifact_manifest_json"] = json.dumps(
            manifest, ensure_ascii=False, sort_keys=True, separators=(",", ":")
        )
        values["simulation_result_hash"] = editor._hash(
            editor._simulation_result_binding(values)
        )
        return self._insert_simulation_copy(values)

    def test_v1_converter_preserves_three_contract_domain_authorities(self):
        snapshots = {
            type_id: editor.build_domain_v1_snapshot(type_id, self._valid_commercial())
            for type_id, _label in contract_domain.contract_type_options()
        }

        self.assertEqual(
            list(snapshots),
            [
                "NewInternationalStudentContract",
                "TransferApplicationContract",
                "LocalStudentApplicationContract",
            ],
        )
        for type_id, snapshot in snapshots.items():
            config = contract_domain.contract_type_config(type_id)
            self.assertEqual(snapshot["template_type"], type_id)
            self.assertEqual(snapshot["display_name"], config.label)
            self.assertEqual(
                [item["service_id"] for item in snapshot["services"]],
                list(config.allowed_services),
            )
            self.assertEqual(
                [item["service_id"] for item in snapshot["services"] if item["default_selected"]],
                list(config.default_services),
            )
            self.assertEqual(
                snapshot["mandatory_display_bindings"]["contract_type"], type_id
            )
            self.assertIn(
                "application_grade",
                snapshot["mandatory_display_bindings"]["required_fields"],
            )
            self.assertIn(
                "planned_entry_year",
                snapshot["mandatory_display_bindings"]["required_fields"],
            )
            expected_sections = contract_domain.compose_contract_sections(type_id)
            self.assertEqual(
                {block["title"]: block["body"] for block in snapshot["blocks"] if block["key"].startswith("profile_")},
                expected_sections,
            )

        international = snapshots["NewInternationalStudentContract"]
        self.assertFalse(next(item for item in international["services"] if item["service_id"] == "subclass_590_guardian_visa_lodgement")["default_selected"])
        self.assertFalse(next(item for item in international["services"] if item["service_id"] == "subclass_590_guardian_visa_consulting")["default_selected"])
        self.assertNotIn("student_visa_lodgement", {item["service_id"] for item in snapshots["TransferApplicationContract"]["services"]})
        self.assertNotIn("offer_coe_oshc_process_assistance", {item["service_id"] for item in snapshots["LocalStudentApplicationContract"]["services"]})

    def test_publish_gate_rejects_missing_mandatory_display_binding_profile(self):
        self._import_v1()
        draft = self.service.create_draft(
            "NewInternationalStudentContract", actor_user_id=1
        )
        version = self.service._version(draft["id"])
        snapshot = json.loads(version["snapshot_json"])
        snapshot.pop("mandatory_display_bindings", None)
        issues = self.service._publish_issues(version, snapshot)

        self.assertIn(
            "mandatory_display_binding_profile_missing",
            {issue["code"] for issue in issues},
        )

    def test_publish_gate_rejects_removing_each_p0_mandatory_binding(self):
        self._import_v1()
        draft = self.service.create_draft(
            "NewInternationalStudentContract", actor_user_id=1
        )
        version = self.service._version(draft["id"])
        original = json.loads(version["snapshot_json"])

        for binding in ("application_grade", "planned_entry_year"):
            with self.subTest(binding=binding):
                snapshot = json.loads(json.dumps(original))
                snapshot["mandatory_display_bindings"]["required_fields"].remove(binding)
                issues = self.service._publish_issues(version, snapshot)
                self.assertIn(
                    "mandatory_display_binding_profile_mismatch",
                    {issue["code"] for issue in issues},
                )

    def test_master_simulation_rejects_blank_mandatory_grade_and_entry_year(self):
        snapshot = editor.build_domain_v1_snapshot(
            "NewInternationalStudentContract", self._valid_commercial()
        )
        base_input = {
            "source_type": "system",
            "student_snapshot": self._simulation_student(),
            "signer_source": "father",
            "signer_overrides": {"phone": "0411111111"},
            "selected_services": ["study_planning_consulting"],
        }

        for field, message in (
            ("applying_year_level", "缺少申请年级"),
            ("preferred_intake_year", "缺少计划入读年份"),
        ):
            with self.subTest(field=field):
                input_payload = json.loads(json.dumps(base_input))
                input_payload["student_snapshot"][field] = ""
                with self.assertRaisesRegex(editor.TemplateRuleError, message):
                    editor._derive_render_expectation(snapshot, input_payload)

    def test_import_creates_independent_published_v1_versions(self):
        self.conn.execute("CREATE TABLE system_settings(key TEXT PRIMARY KEY,value TEXT NOT NULL)")
        self.conn.executemany(
            "INSERT INTO system_settings(key,value) VALUES(?,?)",
            [
                ("company_legal_name", "Authoritative Company Pty Ltd"),
                ("company_abn_acn", "50169160838"),
                ("company_address", "Authoritative Address"),
                ("company_phone", "+61 3 0000 0000"),
                ("company_email", "contracts@example.test"),
            ],
        )
        imported = self._import_v1()

        self.assertEqual(len(imported), 3)
        identities = self.service.list_templates()
        self.assertEqual([row["template_type"] for row in identities], [
            "NewInternationalStudentContract",
            "TransferApplicationContract",
            "LocalStudentApplicationContract",
        ])
        for identity in identities:
            published = self.service.current_published(identity["template_type"])
            self.assertEqual(published["status"], "Published")
            self.assertEqual(published["version_number"], 1)
            self.assertTrue(published["immutable"])
            commercial = json.loads(published["snapshot_json"])["commercial"]
            self.assertEqual(commercial["company_name"], "Authoritative Company Pty Ltd")
            self.assertEqual(commercial["company_abn"], "50169160838")

    def test_draft_and_autosave_do_not_change_current_published_version(self):
        self._import_v1()
        before = self.service.current_published("TransferApplicationContract")
        draft = self.service.create_draft("TransferApplicationContract", actor_user_id=1)
        lock = self.service.acquire_lock("TransferApplicationContract", draft["id"], owner_user_id=1)

        changed = self.service.autosave_block(
            draft["id"],
            "profile_3",
            title="三、服务费用与第三方费用",
            body="转学服务费及第三方费用的新草稿文字。",
            enabled=True,
            expected_revision=draft["revision"],
            lock_token=lock["token"],
            actor_user_id=1,
        )

        after = self.service.current_published("TransferApplicationContract")
        self.assertEqual(after["id"], before["id"])
        self.assertEqual(after["content_hash"], before["content_hash"])
        self.assertEqual(changed["revision"], draft["revision"] + 1)

    def test_all_three_draft_autosaves_ignore_publication_gates_but_apply_blocks(self):
        self._import_v1()
        template_types = (
            "NewInternationalStudentContract",
            "TransferApplicationContract",
            "LocalStudentApplicationContract",
        )

        for template_type in template_types:
            with self.subTest(template_type=template_type):
                published_before = self.service.current_published(template_type)
                draft = self.service.create_draft(template_type, actor_user_id=1)
                lock = self.service.acquire_lock(
                    template_type, draft["id"], owner_user_id=1
                )
                snapshot = self.service._snapshot(draft["id"])
                snapshot.pop("mandatory_display_bindings", None)
                self.conn.execute(
                    "UPDATE contract_template_versions "
                    "SET status='Preview',snapshot_json=?,content_hash=? WHERE id=?",
                    (
                        editor._canonical_json(snapshot),
                        editor._hash(snapshot),
                        draft["id"],
                    ),
                )
                self.conn.commit()
                draft = self.service._version(draft["id"])
                issue_codes = {
                    issue["code"]
                    for issue in self.service._publish_issues(
                        draft, self.service._snapshot(draft["id"])
                    )
                }
                self.assertIn("mandatory_display_binding_profile_missing", issue_codes)
                self.assertIn("simulation_required", issue_codes)

                block = self.conn.execute(
                    "SELECT * FROM contract_template_clause_blocks "
                    "WHERE version_id=? AND block_key='profile_5'",
                    (draft["id"],),
                ).fetchone()
                marker = f"\nDRAFT-SAVE-GATE-{template_type}"
                changed = self.service.autosave_block(
                    draft["id"],
                    block["block_key"],
                    title=block["title"],
                    body=f"{block['body']}{marker}",
                    enabled=bool(block["enabled"]),
                    expected_revision=draft["revision"],
                    lock_token=lock["token"],
                    actor_user_id=1,
                )
                self.assertEqual(changed["revision"], draft["revision"] + 1)
                self.assertIn(
                    marker,
                    self.conn.execute(
                        "SELECT body FROM contract_template_clause_blocks "
                        "WHERE version_id=? AND block_key='profile_5'",
                        (draft["id"],),
                    ).fetchone()[0],
                )

                with self.assertRaises(editor.PublishBlocked):
                    self.service.apply_template(
                        draft["id"],
                        actor_user_id=1,
                        lock_token=lock["token"],
                        expected_revision=changed["revision"],
                        output_root=self.root / "apply-gate" / template_type,
                        actor_is_admin=True,
                    )
                self.assertEqual(
                    self.service.current_published(template_type)["id"],
                    published_before["id"],
                )
                self.assertIn(
                    marker,
                    self.conn.execute(
                        "SELECT body FROM contract_template_clause_blocks "
                        "WHERE version_id=? AND block_key='profile_5'",
                        (draft["id"],),
                    ).fetchone()[0],
                )

                with self.assertRaises(editor.StaleRevision):
                    self.service.autosave_block(
                        draft["id"], block["block_key"], title=block["title"],
                        body=block["body"], enabled=bool(block["enabled"]),
                        expected_revision=draft["revision"],
                        lock_token=lock["token"], actor_user_id=1,
                    )
                self.conn.execute(
                    "UPDATE contract_template_edit_locks SET owner_user_id=2 "
                    "WHERE version_id=?",
                    (draft["id"],),
                )
                self.conn.commit()
                with self.assertRaises(editor.EditLockConflict):
                    self.service.autosave_block(
                        draft["id"], block["block_key"], title=block["title"],
                        body=block["body"], enabled=bool(block["enabled"]),
                        expected_revision=changed["revision"],
                        lock_token=lock["token"], actor_user_id=1,
                    )

    def test_repeated_normal_draft_creation_resumes_one_active_draft(self):
        """Catch repeated clicks creating multiple simultaneously active normal drafts."""
        self._import_v1()

        first = self.service.create_draft(
            "TransferApplicationContract", actor_user_id=9
        )
        second = self.service.create_draft(
            "TransferApplicationContract", actor_user_id=9
        )

        self.assertEqual(second["id"], first["id"])
        active = self.conn.execute(
            """SELECT id FROM contract_template_versions
               WHERE identity_id=? AND status IN ('Draft','Simulation','Validation','Preview')""",
            (first["identity_id"],),
        ).fetchall()
        self.assertEqual([row["id"] for row in active], [first["id"]])

    def test_reusing_legacy_draft_repairs_only_missing_mandatory_profile_for_all_types(self):
        """Legacy active drafts must gain today's canonical profile on reuse."""
        self._import_v1()

        for template_type, _label in contract_domain.contract_type_options():
            with self.subTest(template_type=template_type):
                draft = self.service.create_draft(template_type, actor_user_id=9)
                snapshot = self.service._snapshot(draft["id"])
                original_blocks = snapshot["blocks"]
                original_services = snapshot["services"]
                snapshot.pop("mandatory_display_bindings", None)
                self.conn.execute(
                    "UPDATE contract_template_versions SET snapshot_json=?,content_hash=? WHERE id=?",
                    (
                        editor._canonical_json(snapshot),
                        editor._hash(snapshot),
                        draft["id"],
                    ),
                )
                self.conn.commit()
                before = dict(self.service._version(draft["id"]))

                repaired = self.service.create_draft(template_type, actor_user_id=9)
                repaired_snapshot = self.service._snapshot(draft["id"])

                self.assertEqual(repaired["id"], draft["id"])
                self.assertEqual(repaired["revision"], before["revision"] + 1)
                self.assertEqual(
                    repaired_snapshot["mandatory_display_bindings"],
                    editor.contract_display_bindings.mandatory_profile(template_type),
                )
                self.assertEqual(repaired["content_hash"], editor._hash(repaired_snapshot))
                self.assertEqual(repaired_snapshot["blocks"], original_blocks)
                self.assertEqual(repaired_snapshot["services"], original_services)
                audits = self.conn.execute(
                    "SELECT action FROM contract_template_audit_events "
                    "WHERE version_id=? AND action='template_mandatory_display_binding_profile_repaired'",
                    (draft["id"],),
                ).fetchall()
                self.assertEqual(len(audits), 1)

                unchanged = self.service.create_draft(template_type, actor_user_id=9)
                self.assertEqual(unchanged["revision"], repaired["revision"])
                audits = self.conn.execute(
                    "SELECT action FROM contract_template_audit_events "
                    "WHERE version_id=? AND action='template_mandatory_display_binding_profile_repaired'",
                    (draft["id"],),
                ).fetchall()
                self.assertEqual(len(audits), 1)

    def test_reusing_legacy_draft_repairs_explicit_null_profile(self):
        self._import_v1()
        template_type = "TransferApplicationContract"
        draft = self.service.create_draft(template_type, actor_user_id=9)
        snapshot = self.service._snapshot(draft["id"])
        snapshot["mandatory_display_bindings"] = None
        self.conn.execute(
            "UPDATE contract_template_versions SET snapshot_json=?,content_hash=? WHERE id=?",
            (editor._canonical_json(snapshot), editor._hash(snapshot), draft["id"]),
        )
        self.conn.commit()

        repaired = self.service.create_draft(template_type, actor_user_id=9)

        repaired_snapshot = self.service._snapshot(draft["id"])
        self.assertEqual(
            repaired_snapshot["mandatory_display_bindings"],
            editor.contract_display_bindings.mandatory_profile(template_type),
        )
        self.assertEqual(repaired["content_hash"], editor._hash(repaired_snapshot))

    def test_legacy_profile_repair_rolls_back_when_audit_write_fails(self):
        self._import_v1()
        template_type = "NewInternationalStudentContract"
        draft = self.service.create_draft(template_type, actor_user_id=9)
        snapshot = self.service._snapshot(draft["id"])
        snapshot.pop("mandatory_display_bindings", None)
        self.conn.execute(
            "UPDATE contract_template_versions SET snapshot_json=?,content_hash=? WHERE id=?",
            (editor._canonical_json(snapshot), editor._hash(snapshot), draft["id"]),
        )
        self.conn.commit()
        before = dict(self.service._version(draft["id"]))

        with mock.patch.object(
            self.service, "_audit", side_effect=RuntimeError("audit failed")
        ):
            with self.assertRaisesRegex(RuntimeError, "audit failed"):
                self.service.create_draft(template_type, actor_user_id=9)

        after = dict(self.service._version(draft["id"]))
        self.assertEqual(after["revision"], before["revision"])
        self.assertEqual(after["content_hash"], before["content_hash"])
        self.assertNotIn(
            "mandatory_display_bindings", self.service._snapshot(draft["id"])
        )

    def test_reusing_draft_with_mismatched_mandatory_profile_fails_closed(self):
        self._import_v1()
        template_type = "NewInternationalStudentContract"
        draft = self.service.create_draft(template_type, actor_user_id=9)
        snapshot = self.service._snapshot(draft["id"])
        snapshot["mandatory_display_bindings"]["required_fields"].remove(
            "planned_entry_year"
        )
        self.conn.execute(
            "UPDATE contract_template_versions SET snapshot_json=?,content_hash=? WHERE id=?",
            (editor._canonical_json(snapshot), editor._hash(snapshot), draft["id"]),
        )
        self.conn.commit()
        before = dict(self.service._version(draft["id"]))

        with self.assertRaisesRegex(
            editor.TemplateRuleError, "mandatory display binding profile"
        ):
            self.service.create_draft(template_type, actor_user_id=9)

        after = dict(self.service._version(draft["id"]))
        self.assertEqual(after["revision"], before["revision"])
        self.assertEqual(after["content_hash"], before["content_hash"])

    def test_all_existing_draft_entrypoints_repair_the_same_legacy_draft(self):
        self._import_v1()
        template_type = "LocalStudentApplicationContract"
        published = self.service.current_published(template_type)
        draft = self.service.create_draft(template_type, actor_user_id=9)

        for entrypoint in (
            lambda: self.service.create_draft_from_version(
                published["id"], actor_user_id=9
            ),
            lambda: self.service.create_master_layout_draft(
                template_type, actor_user_id=9
            ),
        ):
            snapshot = self.service._snapshot(draft["id"])
            snapshot.pop("mandatory_display_bindings", None)
            self.conn.execute(
                "UPDATE contract_template_versions SET snapshot_json=?,content_hash=? WHERE id=?",
                (
                    editor._canonical_json(snapshot),
                    editor._hash(snapshot),
                    draft["id"],
                ),
            )
            self.conn.commit()
            before_revision = int(self.service._version(draft["id"])["revision"])

            resumed = entrypoint()

            self.assertEqual(resumed["id"], draft["id"])
            self.assertEqual(resumed["revision"], before_revision + 1)
            self.assertEqual(
                self.service._snapshot(draft["id"])["mandatory_display_bindings"],
                editor.contract_display_bindings.mandatory_profile(template_type),
            )

    def test_reusing_legacy_draft_restores_only_missing_authority_v2_metadata_for_all_types(self):
        self._import_v1()
        required_metadata = {
            "commercial_schema",
            "standard_service_fee",
            "standard_service_fee_ex_gst",
            "gst_treatment",
            "gst_rate",
            "default_gst_amount",
            "default_final_total",
            "default_payment_plan",
            "payment_methods_json",
            "installment_approval_role",
            "discount_approval_role",
            "discount_reason_required",
            "base_fee_override_allowed",
        }

        for template_type, _label in contract_domain.contract_type_options():
            with self.subTest(template_type=template_type):
                published = self._authority_v2_published(template_type)
                published_snapshot = self.service._snapshot(published["id"])
                draft = self.service.create_master_layout_draft(
                    template_type, actor_user_id=9
                )
                draft_snapshot = self.service._snapshot(draft["id"])
                original_blocks = draft_snapshot["blocks"]
                original_services = draft_snapshot["services"]
                original_refund = "DRAFT REFUND BODY MUST NOT CHANGE"
                draft_snapshot["commercial"]["refund_policy"] = original_refund
                source_commercial = published_snapshot["commercial"]
                expected_keys = set(required_metadata)
                if template_type == "NewInternationalStudentContract":
                    expected_keys.update(
                        {"visa_refusal_retained_ex_gst", "visa_refusal_rule_source"}
                    )
                for key in expected_keys:
                    draft_snapshot["commercial"].pop(key, None)
                self.conn.execute(
                    "UPDATE contract_template_versions SET snapshot_json=?,content_hash=? WHERE id=?",
                    (
                        editor._canonical_json(draft_snapshot),
                        editor._hash(draft_snapshot),
                        draft["id"],
                    ),
                )
                self.conn.commit()
                before = dict(self.service._version(draft["id"]))

                repaired = self.service.create_draft(template_type, actor_user_id=9)
                repaired_snapshot = self.service._snapshot(draft["id"])

                self.assertEqual(repaired["id"], draft["id"])
                self.assertEqual(repaired["revision"], before["revision"] + 1)
                self.assertEqual(repaired_snapshot["blocks"], original_blocks)
                self.assertEqual(repaired_snapshot["services"], original_services)
                self.assertEqual(
                    repaired_snapshot["commercial"]["refund_policy"], original_refund
                )
                for key in expected_keys:
                    self.assertEqual(
                        repaired_snapshot["commercial"][key], source_commercial[key]
                    )
                self.assertEqual(repaired["content_hash"], editor._hash(repaired_snapshot))
                audit = self.conn.execute(
                    "SELECT after_json FROM contract_template_audit_events "
                    "WHERE version_id=? AND action="
                    "'template_authority_v2_commercial_metadata_repaired' "
                    "ORDER BY id DESC LIMIT 1",
                    (draft["id"],),
                ).fetchone()
                self.assertIsNotNone(audit)
                self.assertEqual(
                    set(json.loads(audit["after_json"])["restored_fields"]),
                    expected_keys,
                )

    def test_reusing_draft_never_overwrites_present_commercial_metadata(self):
        self._import_v1()
        template_type = "NewInternationalStudentContract"
        published = self._authority_v2_published(template_type)
        published_snapshot = self.service._snapshot(published["id"])
        draft = self.service.create_master_layout_draft(template_type, actor_user_id=9)
        draft_snapshot = self.service._snapshot(draft["id"])
        draft_snapshot["commercial"]["refund_policy"] = "KEEP DRAFT REFUND"
        draft_snapshot["commercial"]["gst_rate"] = "0.10"
        draft_snapshot["commercial"].pop("commercial_schema")
        self.conn.execute(
            "UPDATE contract_template_versions SET snapshot_json=?,content_hash=? WHERE id=?",
            (
                editor._canonical_json(draft_snapshot),
                editor._hash(draft_snapshot),
                draft["id"],
            ),
        )
        self.conn.commit()

        self.service.create_draft(template_type, actor_user_id=9)

        repaired = self.service._snapshot(draft["id"])
        self.assertEqual(repaired["commercial"]["refund_policy"], "KEEP DRAFT REFUND")
        self.assertEqual(repaired["commercial"]["gst_rate"], "0.10")
        self.assertEqual(repaired["commercial"]["commercial_schema"], "authority_v2")

    def test_new_draft_entrypoints_inherit_authority_v2_metadata(self):
        self._import_v1()
        creators = {
            "NewInternationalStudentContract": lambda published: (
                self.service.create_draft(
                    "NewInternationalStudentContract", actor_user_id=9
                )
            ),
            "TransferApplicationContract": lambda published: (
                self.service.create_master_layout_draft(
                    "TransferApplicationContract", actor_user_id=9
                )
            ),
            "LocalStudentApplicationContract": lambda published: (
                self.service.create_draft_from_version(
                    published["id"], actor_user_id=9
                )
            ),
        }

        for template_type, creator in creators.items():
            with self.subTest(template_type=template_type):
                published = self._authority_v2_published(template_type)
                source = self.service._snapshot(published["id"])
                draft = creator(published)
                created = self.service._snapshot(draft["id"])
                self.assertEqual(
                    created["mandatory_display_bindings"],
                    editor.contract_display_bindings.mandatory_profile(template_type),
                )
                for key in editor.AUTHORITY_V2_INHERITED_METADATA_FIELDS:
                    if key in source["commercial"]:
                        self.assertEqual(
                            created["commercial"][key], source["commercial"][key]
                        )

    def test_repeated_master_draft_creation_resumes_one_active_draft(self):
        """Catch repeated master-layout clicks orphaning an earlier active draft."""
        self._import_v1()

        first = self.service.create_master_layout_draft(
            "LocalStudentApplicationContract", actor_user_id=9
        )
        second = self.service.create_master_layout_draft(
            "LocalStudentApplicationContract", actor_user_id=9
        )

        self.assertEqual(second["id"], first["id"])
        active = self.conn.execute(
            """SELECT id FROM contract_template_versions
               WHERE identity_id=? AND status IN ('Draft','Simulation','Validation','Preview')""",
            (first["identity_id"],),
        ).fetchall()
        self.assertEqual([row["id"] for row in active], [first["id"]])

    def test_repeated_history_republish_resumes_one_active_draft(self):
        """Catch repeated history republish creating parallel active drafts."""
        self._import_v1()
        published = self.service.current_published(
            "TransferApplicationContract"
        )

        first = self.service.create_draft_from_version(
            published["id"], actor_user_id=9
        )
        second = self.service.create_draft_from_version(
            published["id"], actor_user_id=9
        )

        self.assertEqual(second["id"], first["id"])
        active = self.conn.execute(
            """SELECT id FROM contract_template_versions
               WHERE identity_id=? AND status IN ('Draft','Simulation','Validation','Preview')""",
            (first["identity_id"],),
        ).fetchall()
        self.assertEqual([row["id"] for row in active], [first["id"]])

    def test_concurrent_history_republish_resumes_competing_active_draft(self):
        """Catch two republish requests racing before either sees an active draft."""
        self._import_v1()
        published = self.service.current_published(
            "TransferApplicationContract"
        )
        competing_connection = sqlite3.connect(self.root / "editor.sqlite3")
        competing_connection.row_factory = sqlite3.Row
        competing_connection.execute("PRAGMA foreign_keys=ON")
        self.addCleanup(competing_connection.close)
        competing_service = editor.ContractTemplateEditorService(
            competing_connection
        )
        interposed = []

        def interpose_before_immediate(statement):
            if statement.strip().upper() != "BEGIN IMMEDIATE" or interposed:
                return
            interposed.append(
                competing_service.create_draft_from_version(
                    published["id"], actor_user_id=10
                )
            )

        self.conn.set_trace_callback(interpose_before_immediate)
        self.addCleanup(self.conn.set_trace_callback, None)
        resumed = self.service.create_draft_from_version(
            published["id"], actor_user_id=9
        )

        self.assertEqual(resumed["id"], interposed[0]["id"])
        active_count = self.conn.execute(
            """SELECT COUNT(*) FROM contract_template_versions
               WHERE identity_id=? AND status IN ('Draft','Simulation','Validation','Preview')""",
            (published["identity_id"],),
        ).fetchone()[0]
        self.assertEqual(active_count, 1)

    def test_lock_acquisition_rejects_non_mutable_version_without_replacing_lock(self):
        """Catch a history click moving an active identity lock to an immutable version."""
        self._import_v1()
        draft = self.service.create_draft(
            "LocalStudentApplicationContract", actor_user_id=9
        )
        lock = self.service.acquire_lock(
            "LocalStudentApplicationContract", draft["id"], owner_user_id=9
        )
        published = self.service.current_published(
            "LocalStudentApplicationContract"
        )

        with self.assertRaises(editor.ImmutableVersionError):
            self.service.acquire_lock(
                "LocalStudentApplicationContract",
                published["id"],
                owner_user_id=9,
            )

        retained = self.conn.execute(
            "SELECT version_id,token FROM contract_template_edit_locks WHERE identity_id=?",
            (draft["identity_id"],),
        ).fetchone()
        self.assertEqual(retained["version_id"], draft["id"])
        self.assertEqual(retained["token"], lock["token"])

    def test_edit_lock_and_stale_revision_fail_closed(self):
        self._import_v1()
        draft = self.service.create_draft("LocalStudentApplicationContract", actor_user_id=1)
        lock = self.service.acquire_lock("LocalStudentApplicationContract", draft["id"], owner_user_id=1)

        with self.assertRaisesRegex(editor.EditLockConflict, "另一位用户"):
            self.service.acquire_lock("LocalStudentApplicationContract", draft["id"], owner_user_id=2)
        with self.assertRaisesRegex(editor.StaleRevision, "版本已变化"):
            self.service.reorder_blocks(
                draft["id"],
                ["profile_3", "profile_4"],
                expected_revision=draft["revision"] + 1,
                lock_token=lock["token"],
                actor_user_id=1,
            )

    def test_same_owner_heartbeat_and_reopen_preserve_one_authoritative_token(self):
        """Catch a same-user heartbeat rotating the token and locking out another tab."""
        self._import_v1()
        draft = self.service.create_draft(
            "TransferApplicationContract", actor_user_id=2
        )
        first = self.service.acquire_lock(
            "TransferApplicationContract", draft["id"], owner_user_id=2
        )
        second = self.service.acquire_lock(
            "TransferApplicationContract", draft["id"], owner_user_id=2
        )

        self.assertEqual(second["token"], first["token"])
        self.assertEqual(
            self.service.live_lock(draft["id"], owner_user_id=2)["token"],
            first["token"],
        )
        self.assertEqual(
            self.conn.execute(
                "SELECT COUNT(*) FROM contract_template_edit_locks WHERE identity_id=?",
                (draft["identity_id"],),
            ).fetchone()[0],
            1,
        )
        self.conn.execute(
            "UPDATE contract_template_edit_locks SET expires_at=? WHERE version_id=?",
            ("2000-01-01T00:00:00+00:00", draft["id"]),
        )
        self.conn.commit()
        self.assertIsNone(self.service.live_lock(draft["id"], owner_user_id=2))
        reopened = self.service.acquire_lock(
            "TransferApplicationContract", draft["id"], owner_user_id=2
        )
        self.assertEqual(reopened["token"], first["token"])

    def test_expired_exact_owner_token_reclaims_and_autosaves_without_reload(self):
        """Catch a suspended pre-deploy page losing its unsaved body after lock expiry."""
        self._import_v1()
        draft = self.service.create_draft(
            "LocalStudentApplicationContract", actor_user_id=2
        )
        lock = self.service.acquire_lock(
            "LocalStudentApplicationContract", draft["id"], owner_user_id=2
        )
        self.conn.execute(
            "UPDATE contract_template_edit_locks SET expires_at=? WHERE version_id=?",
            ("2000-01-01T00:00:00+00:00", draft["id"]),
        )
        self.conn.commit()

        changed = self.service.autosave_block(
            draft["id"],
            "profile_3",
            title="三、服务费用与第三方费用",
            body="UNSAVED CONTENT SURVIVES EXPIRED LOCK",
            enabled=True,
            expected_revision=draft["revision"],
            lock_token=lock["token"],
            actor_user_id=2,
        )

        self.assertEqual(changed["revision"], draft["revision"] + 1)
        live = self.service.live_lock(draft["id"], owner_user_id=2)
        self.assertIsNotNone(live)
        self.assertEqual(live["token"], lock["token"])
        saved = self.conn.execute(
            "SELECT body FROM contract_template_clause_blocks "
            "WHERE version_id=? AND block_key='profile_3'",
            (draft["id"],),
        ).fetchone()[0]
        self.assertEqual(saved, "UNSAVED CONTENT SURVIVES EXPIRED LOCK")

    def test_restart_recovery_uses_persisted_owner_token_without_input_loss(self):
        """Catch a new server process rejecting the old page's persisted DB token."""
        self._import_v1()
        draft = self.service.create_draft(
            "TransferApplicationContract", actor_user_id=2
        )
        lock = self.service.acquire_lock(
            "TransferApplicationContract", draft["id"], owner_user_id=2
        )
        self.conn.execute(
            "UPDATE contract_template_edit_locks SET expires_at=? WHERE version_id=?",
            ("2000-01-01T00:00:00+00:00", draft["id"]),
        )
        self.conn.commit()

        restarted_conn = sqlite3.connect(self.root / "editor.sqlite3")
        restarted_conn.row_factory = sqlite3.Row
        restarted_conn.execute("PRAGMA foreign_keys=ON")
        self.addCleanup(restarted_conn.close)
        restarted_service = editor.ContractTemplateEditorService(restarted_conn)
        changed = restarted_service.autosave_block(
            draft["id"],
            "profile_4",
            title="四、甲方权利义务",
            body="RESTART RECOVERY CONTENT",
            enabled=True,
            expected_revision=draft["revision"],
            lock_token=lock["token"],
            actor_user_id=2,
        )

        self.assertEqual(changed["revision"], draft["revision"] + 1)
        self.assertEqual(
            restarted_service.live_lock(draft["id"], owner_user_id=2)["token"],
            lock["token"],
        )

    def test_same_user_two_tabs_share_lock_but_stale_revision_cannot_overwrite(self):
        """Catch multi-tab recovery weakening optimistic concurrency protection."""
        self._import_v1()
        draft = self.service.create_draft(
            "TransferApplicationContract", actor_user_id=2
        )
        tab_one = self.service.acquire_lock(
            "TransferApplicationContract", draft["id"], owner_user_id=2
        )
        tab_two = self.service.acquire_lock(
            "TransferApplicationContract", draft["id"], owner_user_id=2
        )
        self.assertEqual(tab_one["token"], tab_two["token"])

        changed = self.service.autosave_block(
            draft["id"], "profile_3", title="三、服务费用与第三方费用", body="TAB ONE",
            enabled=True, expected_revision=draft["revision"],
            lock_token=tab_one["token"], actor_user_id=2,
        )
        with self.assertRaises(editor.StaleRevision):
            self.service.autosave_block(
                draft["id"], "profile_3", title="三、服务费用与第三方费用", body="TAB TWO STALE",
                enabled=True, expected_revision=draft["revision"],
                lock_token=tab_two["token"], actor_user_id=2,
            )
        saved = self.conn.execute(
            "SELECT body FROM contract_template_clause_blocks "
            "WHERE version_id=? AND block_key='profile_3'",
            (draft["id"],),
        ).fetchone()[0]
        self.assertEqual(saved, "TAB ONE")
        self.assertEqual(changed["revision"], draft["revision"] + 1)

    def test_fixed_blocks_and_required_variables_are_protected(self):
        self._import_v1()
        draft = self.service.create_draft("NewInternationalStudentContract", actor_user_id=1)
        lock = self.service.acquire_lock(
            "NewInternationalStudentContract", draft["id"], owner_user_id=1
        )

        with self.assertRaisesRegex(editor.TemplateRuleError, "版式"):
            self.service.delete_block(draft["id"], "fixed_parties", draft["revision"], lock["token"], 1)
        with self.assertRaisesRegex(editor.TemplateRuleError, "版式"):
            self.service.autosave_block(
                draft["id"], "fixed_parties", title="Changed", body="学生：{{student_name}}\n签署人：{{signer_name}}\n电话：{{signer_phone}}\n邮箱：{{signer_email}}",
                enabled=True, expected_revision=draft["revision"], lock_token=lock["token"], actor_user_id=1,
            )
        with self.assertRaisesRegex(editor.TemplateRuleError, "必要变量"):
            self.service.autosave_block(
                draft["id"], "fixed_parties", title="合同双方", body="学生资料已删除",
                enabled=True, expected_revision=draft["revision"], lock_token=lock["token"], actor_user_id=1,
            )
        with self.assertRaisesRegex(editor.TemplateRuleError, "位置"):
            self.service.autosave_block(
                draft["id"], "fixed_parties", title="合同双方",
                body="签署人：{{signer_name}}\n学生：{{student_name}}\n电话：{{signer_phone}}\n邮箱：{{signer_email}}",
                enabled=True, expected_revision=draft["revision"], lock_token=lock["token"], actor_user_id=1,
            )

    def test_master_layout_allows_content_operations_without_moving_anchors(self):
        self._import_v1()
        draft = self.service.create_master_layout_draft(
            "NewInternationalStudentContract", actor_user_id=1
        )
        snapshot = json.loads(draft["snapshot_json"])
        self.assertEqual(snapshot["layout_policy"]["id"], "AAIEContractMasterLayoutV1")
        lock = self.service.acquire_lock(
            "NewInternationalStudentContract", draft["id"], owner_user_id=1
        )
        profile = next(block for block in snapshot["blocks"] if block["key"] == "profile_3")
        saved = self.service.autosave_block(
            draft["id"], "profile_3", title="三、费用与付款安排",
            body=profile["body"] + "\n经批准的内容编辑。", enabled=True,
            expected_revision=draft["revision"], lock_token=lock["token"],
            actor_user_id=1,
        )
        first = self.service.create_block(
            draft["id"], "profile_3::custom_one", "3.1 附加费用说明", "附加条款一。",
            saved["revision"], lock["token"], 1,
        )
        second = self.service.create_block(
            draft["id"], "profile_3::custom_two", "3.2 第二项说明", "附加条款二。",
            first["revision"], lock["token"], 1,
        )
        ordered = [block["key"] for block in json.loads(second["snapshot_json"])["blocks"]]
        left = ordered.index("profile_3::custom_one")
        right = ordered.index("profile_3::custom_two")
        ordered[left], ordered[right] = ordered[right], ordered[left]
        reordered = self.service.reorder_blocks(
            draft["id"], ordered, expected_revision=second["revision"],
            lock_token=lock["token"], actor_user_id=1,
        )
        deleted = self.service.delete_block(
            draft["id"], "profile_3::custom_one", reordered["revision"], lock["token"], 1,
        )
        final_snapshot = json.loads(deleted["snapshot_json"])
        self.assertEqual(
            next(item for item in final_snapshot["blocks"] if item["key"] == "profile_3")["title"],
            "三、费用与付款安排",
        )
        self.assertIn("profile_3::custom_two", [item["key"] for item in final_snapshot["blocks"]])
        self.assertNotIn("profile_3::custom_one", [item["key"] for item in final_snapshot["blocks"]])

        with self.assertRaisesRegex(editor.TemplateRuleError, "序号|标题"):
            self.service.autosave_block(
                draft["id"], "profile_3", title="费用改名", body=profile["body"], enabled=True,
                expected_revision=deleted["revision"], lock_token=lock["token"], actor_user_id=1,
            )
        with self.assertRaisesRegex(editor.TemplateRuleError, "锚点|顺序"):
            invalid = [key for key in [item["key"] for item in final_snapshot["blocks"]] if key != "profile_3"]
            invalid.insert(0, "profile_3")
            self.service.reorder_blocks(
                draft["id"], invalid, expected_revision=deleted["revision"],
                lock_token=lock["token"], actor_user_id=1,
            )
        with self.assertRaisesRegex(editor.TemplateRuleError, "固定|主版式"):
            self.service.create_block(
                draft["id"], "fixed_parties::free", "自由区块", "内容",
                deleted["revision"], lock["token"], 1,
            )
        with self.assertRaisesRegex(editor.TemplateRuleError, "锚点|必要"):
            self.service.delete_block(
                draft["id"], "profile_3", deleted["revision"], lock["token"], 1,
            )

    def test_simulation_binds_one_parent_and_writes_only_under_temporary_root(self):
        self._import_v1()
        draft = self.service.create_draft("NewInternationalStudentContract", actor_user_id=1)
        lock = self.service.acquire_lock(
            "NewInternationalStudentContract", draft["id"], owner_user_id=1
        )
        student = {
            "student_name": "Test Student",
            "student_id": "T-EDITOR-002",
            "date_of_birth": "2011-02-03",
            "current_year_level": "Year 7",
            "applying_year_level": "Year 8",
            "preferred_intake_year": "2027",
            "company_authorised_person": "Sonny Zhang",
            "father_passport_first_name": "Father",
            "father_passport_last_name": "Only",
            "father_mobile": "",
            "father_email": "father@example.test",
            "mother_passport_first_name": "Mother",
            "mother_passport_last_name": "Other",
            "mother_mobile": "0400000000",
            "mother_email": "mother@example.test",
        }

        result = self.service.simulate(
            draft["id"], source_type="system", student_snapshot=student,
            signer_source="father", signer_overrides={"phone": "0411111111"},
            selected_services=["study_planning_consulting"], output_root=self.root / "previews", actor_user_id=1,
            expected_revision=draft["revision"], lock_token=lock["token"],
        )

        self.assertEqual(result["signer"], {
            "source": "father", "name": "Father Only", "relationship": "父亲",
            "phone": "0411111111", "email": "father@example.test",
        })
        self.assertNotIn("Mother Other", result["authority_content"])
        self.assertIn("费用明细", result["authority_content"])
        self.assertIn("标准服务费（不含 GST）：AUD 0.00", result["authority_content"])
        self.assertEqual(result["docx_content_hash"], result["pdf_content_hash"])
        self.assertEqual(result["docx_content_hash"], result["signing_content_hash"])
        self.assertTrue(Path(result["artifact_dir"]).is_relative_to((self.root / "previews").resolve()))
        self.assertTrue((Path(result["artifact_dir"]) / "preview.docx").is_file())
        self.assertTrue((Path(result["artifact_dir"]) / "preview.pdf").is_file())
        self.assertTrue((Path(result["artifact_dir"]) / "signing-preview.html").is_file())

    def test_simulation_schema_and_row_persist_exact_provenance(self):
        draft, _lock, result = self._simulate_ready_draft()
        required_columns = {
            "template_identity_id",
            "template_revision",
            "template_content_hash",
            "template_snapshot_hash",
            "simulation_input_hash",
            "simulation_result_hash",
            "docx_semantic_hash",
            "pdf_semantic_hash",
            "signing_semantic_hash",
            "docx_semantic_trace_json",
            "pdf_semantic_trace_json",
            "signing_semantic_trace_json",
            "docx_artifact_hash",
            "pdf_artifact_hash",
            "signing_artifact_hash",
            "artifact_manifest_json",
            "artifact_root",
            "simulated_at",
        }
        columns = {
            row["name"]
            for row in self.conn.execute("PRAGMA table_info(contract_template_simulations)")
        }
        self.assertTrue(
            required_columns.issubset(columns),
            f"missing simulation provenance columns: {sorted(required_columns - columns)}",
        )

        row = self.conn.execute(
            "SELECT * FROM contract_template_simulations WHERE version_id=? ORDER BY id DESC LIMIT 1",
            (draft["id"],),
        ).fetchone()
        self.assertEqual(row["template_identity_id"], draft["identity_id"])
        self.assertEqual(row["template_revision"], draft["revision"])
        self.assertEqual(row["template_content_hash"], draft["content_hash"])
        snapshot = json.loads(draft["snapshot_json"])
        expected_snapshot_hash = hashlib.sha256(
            json.dumps(
                snapshot, ensure_ascii=False, sort_keys=True, separators=(",", ":")
            ).encode("utf-8")
        ).hexdigest()
        self.assertEqual(row["template_snapshot_hash"], expected_snapshot_hash)
        expected_input = {
            "selected_services": ["study_planning_consulting"],
            "signer_overrides": {"phone": "0411111111"},
            "signer_source": "father",
            "source_type": "system",
            "student_snapshot": self._simulation_student(),
        }
        expected_input_json = json.dumps(
            expected_input, ensure_ascii=False, sort_keys=True, separators=(",", ":")
        )
        self.assertEqual(row["input_json"], expected_input_json)
        self.assertEqual(
            row["simulation_input_hash"],
            hashlib.sha256(expected_input_json.encode("utf-8")).hexdigest(),
        )
        self.assertEqual(row["simulated_at"], row["created_at"])
        self.assertEqual(row["artifact_root"], str((self.root / "previews").resolve()))
        self.assertEqual(result["artifact_root"], row["artifact_root"])
        for name in required_columns - {"simulated_at"}:
            self.assertIsNotNone(row[name])
        self.assertEqual(result["simulation_result_hash"], row["simulation_result_hash"])

    def test_validation_rejects_stale_simulation_revision(self):
        draft, lock, _result = self._simulate_ready_draft()
        self.conn.execute(
            "UPDATE contract_template_versions SET revision=revision+1 WHERE id=?",
            (draft["id"],),
        )
        validation = self.service.validate_for_publish(
            draft["id"], actor_user_id=1, lock_token=lock["token"], actor_is_admin=True
        )
        self.assertIn(
            "simulation_stale_revision", {issue["code"] for issue in validation["issues"]}
        )

    def test_validation_rejects_stale_simulation_content_hash(self):
        draft, lock, _result = self._simulate_ready_draft()
        self.conn.execute(
            "UPDATE contract_template_versions SET content_hash=? WHERE id=?",
            ("f" * 64, draft["id"]),
        )
        validation = self.service.validate_for_publish(
            draft["id"], actor_user_id=1, lock_token=lock["token"], actor_is_admin=True
        )
        self.assertIn(
            "simulation_stale_content_hash",
            {issue["code"] for issue in validation["issues"]},
        )

    def test_validation_rejects_stale_simulation_snapshot_hash(self):
        draft, lock, _result = self._simulate_ready_draft()
        snapshot = json.loads(draft["snapshot_json"])
        snapshot["unauthorised_after_simulation"] = True
        self.conn.execute(
            "UPDATE contract_template_versions SET snapshot_json=? WHERE id=?",
            (
                json.dumps(
                    snapshot,
                    ensure_ascii=False,
                    sort_keys=True,
                    separators=(",", ":"),
                ),
                draft["id"],
            ),
        )
        validation = self.service.validate_for_publish(
            draft["id"], actor_user_id=1, lock_token=lock["token"], actor_is_admin=True
        )
        self.assertIn(
            "simulation_stale_snapshot_hash",
            {issue["code"] for issue in validation["issues"]},
        )

    def test_validation_recomputes_simulation_input_and_result_bindings(self):
        draft, lock, _result = self._simulate_ready_draft()
        simulation = self._latest_simulation(draft["id"])
        self._insert_simulation_copy(simulation, input_json="{}")
        validation = self.service.validate_for_publish(
            draft["id"], actor_user_id=1, lock_token=lock["token"], actor_is_admin=True
        )
        self.assertIn(
            "simulation_input_mismatch", {issue["code"] for issue in validation["issues"]}
        )

        self._insert_simulation_copy(
            simulation,
            docx_content_hash="0" * 64,
            pdf_content_hash="0" * 64,
            signing_content_hash="0" * 64,
        )
        validation = self.service.validate_for_publish(
            draft["id"], actor_user_id=1, lock_token=lock["token"], actor_is_admin=True
        )
        self.assertIn(
            "simulation_result_mismatch", {issue["code"] for issue in validation["issues"]}
        )

    def test_validation_rejects_missing_or_tampered_actual_artifacts(self):
        draft, lock, result = self._simulate_ready_draft()
        artifact_dir = Path(result["artifact_dir"])
        pdf_path = artifact_dir / "preview.pdf"
        pdf_path.write_bytes(pdf_path.read_bytes() + b"\nTAMPERED")
        validation = self.service.validate_for_publish(
            draft["id"], actor_user_id=1, lock_token=lock["token"], actor_is_admin=True
        )
        self.assertIn(
            "simulation_artifact_tampered",
            {issue["code"] for issue in validation["issues"]},
        )

        pdf_path.write_bytes(pdf_path.read_bytes()[: -len(b"\nTAMPERED")])
        signing_path = artifact_dir / "signing-preview.html"
        signing_path.rename(artifact_dir / "signing-preview.html.missing")
        validation = self.service.validate_for_publish(
            draft["id"], actor_user_id=1, lock_token=lock["token"], actor_is_admin=True
        )
        self.assertIn(
            "simulation_artifact_missing",
            {issue["code"] for issue in validation["issues"]},
        )

    def test_validation_reextracts_artifact_semantics_after_hash_manifest_rebinding(self):
        draft, lock, result = self._simulate_ready_draft()
        simulation = dict(self._latest_simulation(draft["id"]))
        signing_path = Path(result["artifact_dir"]) / "signing-preview.html"
        signing_path.write_text(
            signing_path.read_text(encoding="utf-8").replace(
                "Boundary Student", "FORGED VISIBLE STUDENT"
            ),
            encoding="utf-8",
        )
        forged_hash = hashlib.sha256(signing_path.read_bytes()).hexdigest()
        manifest = json.loads(simulation["artifact_manifest_json"])
        manifest["signing-preview.html"] = {
            "sha256": forged_hash,
            "bytes": signing_path.stat().st_size,
        }
        simulation.update(
            signing_artifact_hash=forged_hash,
            artifact_manifest_json=json.dumps(
                manifest, ensure_ascii=False, sort_keys=True, separators=(",", ":")
            ),
        )
        simulation["simulation_result_hash"] = editor._hash(
            editor._simulation_result_binding(simulation)
        )
        self._insert_simulation_copy(simulation)

        validation = self.service.validate_for_publish(
            draft["id"], actor_user_id=1, lock_token=lock["token"], actor_is_admin=True
        )
        self.assertFalse(validation["passed"])
        self.assertIn(
            "simulation_semantic_mismatch",
            {issue["code"] for issue in validation["issues"]},
        )

    def test_validation_rechecks_pdf_coordinates_after_hash_manifest_rebinding(self):
        from pypdf import PdfReader, PdfWriter
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfgen import canvas

        draft, lock, result = self._simulate_ready_draft()
        simulation = dict(self._latest_simulation(draft["id"]))
        artifact_dir = Path(result["artifact_dir"])
        pdf_path = artifact_dir / "preview.pdf"
        outside_page = artifact_dir / "outside-page.pdf"
        outside = canvas.Canvas(str(outside_page), pagesize=A4)
        outside.drawString(A4[0] + 25, 100, "OUTSIDE-MEDIABOX")
        outside.save()
        writer = PdfWriter()
        for page in PdfReader(pdf_path).pages:
            writer.add_page(page)
        writer.add_page(PdfReader(outside_page).pages[0])
        rebound_pdf = artifact_dir / "rebound-preview.pdf"
        with rebound_pdf.open("wb") as stream:
            writer.write(stream)
        rebound_pdf.replace(pdf_path)

        forged_hash = hashlib.sha256(pdf_path.read_bytes()).hexdigest()
        manifest = json.loads(simulation["artifact_manifest_json"])
        manifest["preview.pdf"] = {
            "sha256": forged_hash,
            "bytes": pdf_path.stat().st_size,
        }
        simulation.update(
            pdf_artifact_hash=forged_hash,
            artifact_manifest_json=json.dumps(
                manifest, ensure_ascii=False, sort_keys=True, separators=(",", ":")
            ),
        )
        simulation["simulation_result_hash"] = editor._hash(
            editor._simulation_result_binding(simulation)
        )
        self._insert_simulation_copy(simulation)

        validation = self.service.validate_for_publish(
            draft["id"], actor_user_id=1, lock_token=lock["token"], actor_is_admin=True
        )
        self.assertFalse(validation["passed"])
        self.assertIn(
            "simulation_semantic_mismatch",
            {issue["code"] for issue in validation["issues"]},
        )

    def test_validation_rejects_rebound_pdf_with_extra_in_bounds_visible_page(self):
        from pypdf import PdfReader, PdfWriter
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfgen import canvas

        draft, lock, result = self._simulate_ready_draft()
        simulation = dict(self._latest_simulation(draft["id"]))
        artifact_dir = Path(result["artifact_dir"])
        pdf_path = artifact_dir / "preview.pdf"
        extra_path = artifact_dir / "forged-extra-page.pdf"
        extra = canvas.Canvas(str(extra_path), pagesize=A4)
        extra.drawString(40, 100, "FORGED-EXTRA-PAYMENT-TERM-INSIDE-PAGE")
        extra.save()
        writer = PdfWriter()
        for page in PdfReader(pdf_path).pages:
            writer.add_page(page)
        writer.add_page(PdfReader(extra_path).pages[0])
        rebound_path = artifact_dir / "rebound-extra-page.pdf"
        with rebound_path.open("wb") as stream:
            writer.write(stream)
        rebound_path.replace(pdf_path)
        self._insert_rebound_artifact_simulation(simulation, "preview.pdf", pdf_path)

        validation = self.service.validate_for_publish(
            draft["id"], actor_user_id=1, lock_token=lock["token"], actor_is_admin=True
        )

        self.assertFalse(validation["passed"])
        self.assertIn(
            "simulation_semantic_mismatch",
            {issue["code"] for issue in validation["issues"]},
        )

    def test_validation_rejects_rebound_pdf_with_extra_in_bounds_visible_run(self):
        from pypdf import PdfReader, PdfWriter
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfgen import canvas

        draft, lock, result = self._simulate_ready_draft()
        simulation = dict(self._latest_simulation(draft["id"]))
        artifact_dir = Path(result["artifact_dir"])
        pdf_path = artifact_dir / "preview.pdf"
        overlay_path = artifact_dir / "forged-overlay.pdf"
        overlay = canvas.Canvas(str(overlay_path), pagesize=A4)
        overlay.drawString(40, 55, "FORGED-EXTRA-IN-BOUNDS-RUN")
        overlay.save()
        reader = PdfReader(pdf_path)
        writer = PdfWriter()
        for page in reader.pages:
            writer.add_page(page)
        writer.pages[0].merge_page(PdfReader(overlay_path).pages[0])
        rebound_path = artifact_dir / "rebound-overlay.pdf"
        with rebound_path.open("wb") as stream:
            writer.write(stream)
        rebound_path.replace(pdf_path)
        self._insert_rebound_artifact_simulation(simulation, "preview.pdf", pdf_path)

        validation = self.service.validate_for_publish(
            draft["id"], actor_user_id=1, lock_token=lock["token"], actor_is_admin=True
        )

        self.assertFalse(validation["passed"])
        self.assertIn(
            "simulation_semantic_mismatch",
            {issue["code"] for issue in validation["issues"]},
        )

    def test_validation_rejects_rebound_signing_artifact_with_hidden_article(self):
        draft, lock, result = self._simulate_ready_draft()
        simulation = dict(self._latest_simulation(draft["id"]))
        signing_path = Path(result["artifact_dir"]) / "signing-preview.html"
        signing_path.write_text(
            signing_path.read_text(encoding="utf-8").replace(
                "<article>", "<article hidden>", 1
            ),
            encoding="utf-8",
        )
        self._insert_rebound_artifact_simulation(
            simulation, "signing-preview.html", signing_path
        )

        validation = self.service.validate_for_publish(
            draft["id"], actor_user_id=1, lock_token=lock["token"], actor_is_admin=True
        )

        self.assertFalse(validation["passed"])
        self.assertIn(
            "simulation_semantic_mismatch",
            {issue["code"] for issue in validation["issues"]},
        )

    def test_signing_extractor_rejects_visibility_markup_and_unexpected_structure(self):
        _draft, _lock, result = self._simulate_ready_draft()
        signing_path = Path(result["artifact_dir"]) / "signing-preview.html"
        original = signing_path.read_text(encoding="utf-8")
        mutations = {
            "aria-hidden": original.replace("<article>", '<article aria-hidden="true">', 1),
            "display-none": original.replace("<article>", '<article style="display:none">', 1),
            "visibility-hidden": original.replace(
                "<article>", '<article style="visibility:hidden">', 1
            ),
            "unexpected-class": original.replace("<article>", '<article class="hidden">', 1),
            "style-element": original.replace("<article>", "<style>article{display:none}</style><article>", 1),
            "script-element": original.replace("<article>", "<script>void 0</script><article>", 1),
            "template-element": original.replace("<article>", "<template><article>", 1),
        }
        for name, forged in mutations.items():
            with self.subTest(name=name):
                signing_path.write_text(forged, encoding="utf-8")
                with self.assertRaisesRegex(editor.TemplateRuleError, "成品"):
                    editor._extract_signing_semantic_trace(signing_path)

    def test_validation_rejects_rebound_docx_with_hidden_visible_run(self):
        from docx import Document

        draft, lock, result = self._simulate_ready_draft()
        simulation = dict(self._latest_simulation(draft["id"]))
        docx_path = Path(result["artifact_dir"]) / "preview.docx"
        document = Document(docx_path)
        self.assertTrue(document.paragraphs[0].runs)
        document.paragraphs[0].runs[0].font.hidden = True
        document.save(docx_path)
        self._insert_rebound_artifact_simulation(simulation, "preview.docx", docx_path)

        validation = self.service.validate_for_publish(
            draft["id"], actor_user_id=1, lock_token=lock["token"], actor_is_admin=True
        )

        self.assertFalse(validation["passed"])
        self.assertIn(
            "simulation_semantic_mismatch",
            {issue["code"] for issue in validation["issues"]},
        )

    def test_docx_extractor_rejects_tiny_white_hidden_and_off_page_properties(self):
        from docx import Document
        from docx.dml.color import RGBColor
        from docx.oxml import OxmlElement
        from docx.shared import Inches, Pt

        _draft, _lock, result = self._simulate_ready_draft()
        source_path = Path(result["artifact_dir"]) / "preview.docx"
        original = source_path.read_bytes()

        def tiny_run(document):
            document.paragraphs[0].runs[0].font.size = Pt(1)

        def white_run(document):
            document.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)

        def hidden_spec_vanish(document):
            properties = document.paragraphs[0].runs[0]._r.get_or_add_rPr()
            properties.append(OxmlElement("w:specVanish"))

        def off_page_paragraph(document):
            document.paragraphs[0].paragraph_format.left_indent = Inches(20)

        def tiny_default_style(document):
            document.styles["Normal"].font.size = Pt(1)

        for name, mutation in {
            "tiny-run": tiny_run,
            "white-run": white_run,
            "spec-vanish": hidden_spec_vanish,
            "off-page-paragraph": off_page_paragraph,
            "tiny-default-style": tiny_default_style,
        }.items():
            with self.subTest(name=name):
                source_path.write_bytes(original)
                document = Document(source_path)
                mutation(document)
                document.save(source_path)
                with self.assertRaisesRegex(editor.TemplateRuleError, "DOCX成品"):
                    editor._extract_docx_semantic_trace(source_path)

    def test_validation_rejects_rebound_pdf_with_tiny_flattened_visible_text(self):
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.pdfgen import canvas

        draft, lock, result = self._simulate_ready_draft()
        simulation = dict(self._latest_simulation(draft["id"]))
        trace = json.loads(simulation["pdf_semantic_trace_json"])
        visible = "".join(
            line
            for lines in editor._trace_line_lists(trace)
            for line in lines
        )
        font_name = "AAIETinyForgedNoto"
        font_path = (
            ROOT / "assets" / "fonts" / "noto-sans-cjk-sc" / "NotoSansSC-Regular.ttf"
        )
        pdfmetrics.registerFont(TTFont(font_name, str(font_path)))
        pdf_path = Path(result["artifact_dir"]) / "preview.pdf"
        forged = canvas.Canvas(str(pdf_path), pagesize=A4)
        forged.setFont(font_name, 1.5)
        x = 20.0
        y = A4[1] - 20.0
        line = ""
        for character in visible:
            candidate = line + character
            if pdfmetrics.stringWidth(candidate, font_name, 1.5) > A4[0] - 40:
                forged.drawString(x, y, line)
                y -= 2.0
                line = character
            else:
                line = candidate
        if line:
            forged.drawString(x, y, line)
        forged.save()
        self._insert_rebound_artifact_simulation(simulation, "preview.pdf", pdf_path)

        validation = self.service.validate_for_publish(
            draft["id"], actor_user_id=1, lock_token=lock["token"], actor_is_admin=True
        )

        self.assertFalse(validation["passed"])
        self.assertIn(
            "simulation_semantic_mismatch",
            {issue["code"] for issue in validation["issues"]},
        )

    def test_validation_rejects_rebound_pdf_with_invisible_exact_plan_text(self):
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfgen import canvas

        draft, lock, result = self._simulate_ready_draft()
        simulation = dict(self._latest_simulation(draft["id"]))
        model = result["visible_render_model"]
        plan = editor._build_pdf_layout_plan(model)
        pdf_path = Path(result["artifact_dir"]) / "preview.pdf"
        forged = canvas.Canvas(str(pdf_path), pagesize=A4)
        runs_by_page = {}
        for run in plan["runs"]:
            runs_by_page.setdefault(int(run["page"]), []).append(run)
        for page_number in range(1, int(plan["page_count"]) + 1):
            if page_number > 1:
                forged.showPage()
            forged._code.append("3 Tr")
            for run in runs_by_page.get(page_number, []):
                forged.setFont(str(plan["font_name"]), float(run["font_size"]))
                forged.setFillColorRGB(0, 0, 0)
                forged.drawString(
                    float(run["x"]), float(run["y"]), str(run["text"])
                )
        forged.save()
        with self.assertRaisesRegex(editor.TemplateRuleError, "物理结构"):
            editor._extract_verified_pdf_semantic_trace(pdf_path, plan)
        self._insert_rebound_artifact_simulation(simulation, "preview.pdf", pdf_path)

        validation = self.service.validate_for_publish(
            draft["id"], actor_user_id=1, lock_token=lock["token"], actor_is_admin=True
        )

        self.assertFalse(validation["passed"])
        self.assertIn(
            "simulation_semantic_mismatch",
            {issue["code"] for issue in validation["issues"]},
        )

    def test_validation_rejects_rebound_pdf_with_non_text_whiteout_overlay(self):
        from pypdf import PdfReader, PdfWriter
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfgen import canvas

        draft, lock, result = self._simulate_ready_draft()
        simulation = dict(self._latest_simulation(draft["id"]))
        artifact_dir = Path(result["artifact_dir"])
        pdf_path = artifact_dir / "preview.pdf"
        overlay_path = artifact_dir / "whiteout-overlay.pdf"
        overlay = canvas.Canvas(str(overlay_path), pagesize=A4)
        overlay.setFillColorRGB(1, 1, 1)
        overlay.rect(0, 0, A4[0], A4[1], stroke=0, fill=1)
        overlay.save()
        writer = PdfWriter()
        for page in PdfReader(pdf_path).pages:
            writer.add_page(page)
        writer.pages[0].merge_page(PdfReader(overlay_path).pages[0])
        rebound_path = artifact_dir / "rebound-whiteout.pdf"
        with rebound_path.open("wb") as stream:
            writer.write(stream)
        rebound_path.replace(pdf_path)
        self._insert_rebound_artifact_simulation(simulation, "preview.pdf", pdf_path)

        validation = self.service.validate_for_publish(
            draft["id"], actor_user_id=1, lock_token=lock["token"], actor_is_admin=True
        )

        self.assertFalse(validation["passed"])
        self.assertIn(
            "simulation_semantic_mismatch",
            {issue["code"] for issue in validation["issues"]},
        )

    def test_validation_rejects_rebound_pdf_with_catalog_names_action(self):
        from pypdf import PdfReader, PdfWriter
        from pypdf.generic import (
            ArrayObject,
            DictionaryObject,
            NameObject,
            TextStringObject,
        )

        draft, lock, result = self._simulate_ready_draft()
        simulation = dict(self._latest_simulation(draft["id"]))
        pdf_path = Path(result["artifact_dir"]) / "preview.pdf"
        writer = PdfWriter()
        writer.clone_document_from_reader(PdfReader(pdf_path))
        action = DictionaryObject({
            NameObject("/S"): NameObject("/JavaScript"),
            NameObject("/JS"): TextStringObject("void(0)"),
        })
        javascript_names = DictionaryObject({
            NameObject("/Names"): ArrayObject([
                TextStringObject("non-executed-regression-fixture"),
                writer._add_object(action),
            ]),
        })
        catalog_names = DictionaryObject({
            NameObject("/JavaScript"): writer._add_object(javascript_names),
        })
        writer._root_object[NameObject("/Names")] = writer._add_object(catalog_names)
        rebound_path = pdf_path.with_name("rebound-catalog-names.pdf")
        with rebound_path.open("wb") as stream:
            writer.write(stream)
        rebound_path.replace(pdf_path)
        self._insert_rebound_artifact_simulation(simulation, "preview.pdf", pdf_path)

        validation = self.service.validate_for_publish(
            draft["id"], actor_user_id=1, lock_token=lock["token"], actor_is_admin=True
        )

        self.assertFalse(validation["passed"])
        self.assertIn(
            "simulation_semantic_mismatch",
            {issue["code"] for issue in validation["issues"]},
        )

    def test_validation_rejects_rebound_pdf_with_attachment_and_open_action(self):
        from pypdf import PdfReader, PdfWriter
        from pypdf.generic import (
            ArrayObject,
            DecodedStreamObject,
            DictionaryObject,
            NameObject,
            TextStringObject,
        )

        draft, lock, result = self._simulate_ready_draft()
        simulation = dict(self._latest_simulation(draft["id"]))
        pdf_path = Path(result["artifact_dir"]) / "preview.pdf"
        writer = PdfWriter()
        writer.clone_document_from_reader(PdfReader(pdf_path))
        attachment = DecodedStreamObject()
        attachment.set_data(b"non-executed attachment regression fixture")
        attachment.update({NameObject("/Type"): NameObject("/EmbeddedFile")})
        file_specification = DictionaryObject({
            NameObject("/Type"): NameObject("/Filespec"),
            NameObject("/F"): TextStringObject("fixture.txt"),
            NameObject("/EF"): DictionaryObject({
                NameObject("/F"): writer._add_object(attachment),
            }),
        })
        embedded_files = DictionaryObject({
            NameObject("/Names"): ArrayObject([
                TextStringObject("fixture.txt"),
                writer._add_object(file_specification),
            ]),
        })
        writer._root_object[NameObject("/Names")] = writer._add_object(
            DictionaryObject({
                NameObject("/EmbeddedFiles"): writer._add_object(embedded_files),
            })
        )
        writer._root_object[NameObject("/OpenAction")] = writer._add_object(
            DictionaryObject({
                NameObject("/S"): NameObject("/JavaScript"),
                NameObject("/JS"): TextStringObject("void(0)"),
            })
        )
        rebound_path = pdf_path.with_name("rebound-attachment-action.pdf")
        with rebound_path.open("wb") as stream:
            writer.write(stream)
        rebound_path.replace(pdf_path)
        self._insert_rebound_artifact_simulation(simulation, "preview.pdf", pdf_path)

        validation = self.service.validate_for_publish(
            draft["id"], actor_user_id=1, lock_token=lock["token"], actor_is_admin=True
        )

        self.assertFalse(validation["passed"])
        self.assertIn(
            "simulation_semantic_mismatch",
            {issue["code"] for issue in validation["issues"]},
        )

    def test_pdf_package_fingerprint_explicitly_rejects_encryption(self):
        from pypdf import PdfReader, PdfWriter

        _draft, _lock, result = self._simulate_ready_draft()
        pdf_path = Path(result["artifact_dir"]) / "preview.pdf"
        writer = PdfWriter()
        for page in PdfReader(pdf_path).pages:
            writer.add_page(page)
        writer.encrypt("non-secret-regression-fixture")
        encrypted_path = pdf_path.with_name("encrypted-regression-fixture.pdf")
        with encrypted_path.open("wb") as stream:
            writer.write(stream)

        with self.assertRaisesRegex(editor.TemplateRuleError, "加密"):
            editor._pdf_layout_fingerprint(encrypted_path)

    def test_validation_rejects_rebound_pdf_with_page_additional_action(self):
        from pypdf import PdfReader, PdfWriter
        from pypdf.generic import DictionaryObject, NameObject, TextStringObject

        draft, lock, result = self._simulate_ready_draft()
        simulation = dict(self._latest_simulation(draft["id"]))
        pdf_path = Path(result["artifact_dir"]) / "preview.pdf"
        writer = PdfWriter()
        writer.clone_document_from_reader(PdfReader(pdf_path))
        action = writer._add_object(DictionaryObject({
            NameObject("/S"): NameObject("/JavaScript"),
            NameObject("/JS"): TextStringObject("void(0)"),
        }))
        writer.pages[0][NameObject("/AA")] = DictionaryObject({
            NameObject("/O"): action,
        })
        rebound_path = pdf_path.with_name("rebound-page-action.pdf")
        with rebound_path.open("wb") as stream:
            writer.write(stream)
        rebound_path.replace(pdf_path)
        self._insert_rebound_artifact_simulation(simulation, "preview.pdf", pdf_path)

        validation = self.service.validate_for_publish(
            draft["id"], actor_user_id=1, lock_token=lock["token"], actor_is_admin=True
        )

        self.assertFalse(validation["passed"])
        self.assertIn(
            "simulation_semantic_mismatch",
            {issue["code"] for issue in validation["issues"]},
        )

    def test_validation_rejects_rebound_pdf_with_catalog_and_page_xmp_metadata(self):
        from pypdf import PdfReader, PdfWriter
        from pypdf.generic import DecodedStreamObject, NameObject

        draft, lock, result = self._simulate_ready_draft()
        simulation = dict(self._latest_simulation(draft["id"]))
        pdf_path = Path(result["artifact_dir"]) / "preview.pdf"
        writer = PdfWriter()
        writer.clone_document_from_reader(PdfReader(pdf_path))
        xmp = DecodedStreamObject()
        xmp.set_data(
            b'<?xpacket begin=""?><x:xmpmeta xmlns:x="adobe:ns:meta/">'
            b'<fixture>metadata-only-regression</fixture></x:xmpmeta><?xpacket end="w"?>'
        )
        xmp.update({
            NameObject("/Type"): NameObject("/Metadata"),
            NameObject("/Subtype"): NameObject("/XML"),
        })
        xmp_reference = writer._add_object(xmp)
        writer._root_object[NameObject("/Metadata")] = xmp_reference
        writer.pages[0][NameObject("/Metadata")] = xmp_reference
        rebound_path = pdf_path.with_name("rebound-xmp-metadata.pdf")
        with rebound_path.open("wb") as stream:
            writer.write(stream)
        rebound_path.replace(pdf_path)
        self._insert_rebound_artifact_simulation(simulation, "preview.pdf", pdf_path)

        validation = self.service.validate_for_publish(
            draft["id"], actor_user_id=1, lock_token=lock["token"], actor_is_admin=True
        )

        self.assertFalse(validation["passed"])
        self.assertIn(
            "simulation_semantic_mismatch",
            {issue["code"] for issue in validation["issues"]},
        )

    def test_validation_rejects_rebound_pdf_with_trailer_info_metadata(self):
        from pypdf import PdfReader, PdfWriter

        draft, lock, result = self._simulate_ready_draft()
        simulation = dict(self._latest_simulation(draft["id"]))
        pdf_path = Path(result["artifact_dir"]) / "preview.pdf"
        writer = PdfWriter()
        writer.clone_document_from_reader(PdfReader(pdf_path))
        writer.add_metadata({
            "/Title": "metadata-only-regression-fixture",
            "/Author": "not-executed",
        })
        rebound_path = pdf_path.with_name("rebound-info-metadata.pdf")
        with rebound_path.open("wb") as stream:
            writer.write(stream)
        rebound_path.replace(pdf_path)
        self._insert_rebound_artifact_simulation(simulation, "preview.pdf", pdf_path)

        validation = self.service.validate_for_publish(
            draft["id"], actor_user_id=1, lock_token=lock["token"], actor_is_admin=True
        )

        self.assertFalse(validation["passed"])
        self.assertIn(
            "simulation_semantic_mismatch",
            {issue["code"] for issue in validation["issues"]},
        )

    def test_all_artifact_builders_are_exactly_deterministic_across_processes(self):
        import subprocess
        import sys

        _draft, _lock, result = self._simulate_ready_draft()
        model_json = editor._canonical_json(result["visible_render_model"])
        script = (
            "import hashlib,json,sys; import contract_template_editor as e; "
            "model=json.loads(sys.stdin.read()); "
            "plan=e._build_pdf_layout_plan(model); "
            "artifacts={"
            "'docx':e._build_docx_artifact_bytes(model),"
            "'pdf':e._render_pdf_plan_bytes(plan),"
            "'signing':e._build_signing_artifact_bytes(model)}; "
            "print(json.dumps({name:{'bytes':len(data),'sha256':hashlib.sha256(data).hexdigest()} "
            "for name,data in artifacts.items()},sort_keys=True))"
        )

        def build_in_new_process():
            completed = subprocess.run(
                [sys.executable, "-c", script],
                cwd=ROOT,
                input=model_json,
                text=True,
                capture_output=True,
                check=True,
            )
            return json.loads(completed.stdout)

        first = build_in_new_process()
        second = build_in_new_process()
        local = {
            "docx": editor._build_docx_artifact_bytes(result["visible_render_model"]),
            "pdf": editor._render_pdf_plan_bytes(
                editor._build_pdf_layout_plan(result["visible_render_model"])
            ),
            "signing": editor._build_signing_artifact_bytes(
                result["visible_render_model"]
            ),
        }

        self.assertEqual(first, second)
        self.assertEqual(
            first,
            {
                name: {
                    "bytes": len(data),
                    "sha256": hashlib.sha256(data).hexdigest(),
                }
                for name, data in local.items()
            },
        )

    def test_validation_rejects_rebound_docx_with_off_page_header_footer_distance(self):
        from docx import Document
        from docx.shared import Inches

        draft, lock, result = self._simulate_ready_draft()
        simulation = dict(self._latest_simulation(draft["id"]))
        docx_path = Path(result["artifact_dir"]) / "preview.docx"
        document = Document(docx_path)
        document.sections[0].header_distance = Inches(20)
        document.sections[0].footer_distance = Inches(20)
        document.save(docx_path)
        self.assertEqual(
            editor._canonical_json(editor._extract_docx_semantic_trace(docx_path)),
            simulation["docx_semantic_trace_json"],
        )
        self._insert_rebound_artifact_simulation(simulation, "preview.docx", docx_path)

        validation = self.service.validate_for_publish(
            draft["id"], actor_user_id=1, lock_token=lock["token"], actor_is_admin=True
        )

        self.assertFalse(validation["passed"])
        self.assertIn(
            "simulation_semantic_mismatch",
            {issue["code"] for issue in validation["issues"]},
        )

    def test_validation_rejects_rebound_docx_with_table_cell_layout_mutation(self):
        from docx import Document
        from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
        from docx.shared import Inches

        draft, lock, result = self._simulate_ready_draft()
        simulation = dict(self._latest_simulation(draft["id"]))
        docx_path = Path(result["artifact_dir"]) / "preview.docx"
        document = Document(docx_path)
        table = document.tables[0]
        table.alignment = WD_TABLE_ALIGNMENT.RIGHT
        table.autofit = False
        table.rows[0].cells[0].width = Inches(20)
        table.rows[0].cells[0].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.BOTTOM
        document.save(docx_path)
        self.assertEqual(
            editor._canonical_json(editor._extract_docx_semantic_trace(docx_path)),
            simulation["docx_semantic_trace_json"],
        )
        self._insert_rebound_artifact_simulation(simulation, "preview.docx", docx_path)

        validation = self.service.validate_for_publish(
            draft["id"], actor_user_id=1, lock_token=lock["token"], actor_is_admin=True
        )

        self.assertFalse(validation["passed"])
        self.assertIn(
            "simulation_semantic_mismatch",
            {issue["code"] for issue in validation["issues"]},
        )

    def test_docx_embeds_tracked_noto_font_with_portable_ooxml_relationship(self):
        import zipfile
        from xml.etree import ElementTree

        _draft, _lock, result = self._simulate_ready_draft()
        docx_path = Path(result["artifact_dir"]) / "preview.docx"
        expected_font = (
            ROOT / "assets" / "fonts" / "noto-sans-cjk-sc" / "NotoSansSC-Regular.ttf"
        ).read_bytes()
        with zipfile.ZipFile(docx_path) as package:
            embedded_name = "word/fonts/NotoSansSC-Regular.odttf"
            self.assertIn(embedded_name, package.namelist())
            self.assertIn(b"odttf", package.read("[Content_Types].xml"))
            relationships = package.read("word/_rels/fontTable.xml.rels")
            self.assertIn(
                b"http://schemas.openxmlformats.org/officeDocument/2006/relationships/font",
                relationships,
            )
            self.assertIn(b"fonts/NotoSansSC-Regular.odttf", relationships)
            self.assertIn(editor.DOCX_FONT_RELATIONSHIP_ID.encode(), relationships)
            font_table = ElementTree.fromstring(package.read("word/fontTable.xml"))
            namespace = {
                "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
            }
            embedded = font_table.find(
                ".//w:font[@w:name='Noto Sans SC']/w:embedRegular", namespace
            )
            self.assertIsNotNone(embedded)
            self.assertEqual(
                embedded.attrib[
                    "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
                ],
                editor.DOCX_FONT_RELATIONSHIP_ID,
            )
            font_key = embedded.attrib[
                "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}fontKey"
            ]
            key = bytes.fromhex(font_key.strip("{}").replace("-", ""))[::-1]
            obfuscated = bytearray(package.read(embedded_name))
            for index in range(min(32, len(obfuscated))):
                obfuscated[index] ^= key[index % len(key)]
            self.assertEqual(bytes(obfuscated), expected_font)
            settings = package.read("word/settings.xml")
            self.assertIn(b"embedTrueTypeFonts", settings)
            self.assertIn(b"saveSubsetFonts", settings)

    def test_invalid_artifact_directory_is_never_opened_or_extracted(self):
        draft, lock, result = self._simulate_ready_draft()
        simulation = dict(self._latest_simulation(draft["id"]))
        source_dir = Path(result["artifact_dir"])
        invalid_dir = self.root / "unowned" / source_dir.name
        invalid_dir.mkdir(parents=True)
        for filename in ("preview.docx", "preview.pdf", "signing-preview.html"):
            (invalid_dir / filename).write_bytes((source_dir / filename).read_bytes())
        simulation["artifact_dir"] = str(invalid_dir)
        simulation["simulation_result_hash"] = editor._hash(
            editor._simulation_result_binding(simulation)
        )
        self._insert_simulation_copy(simulation)

        with (
            mock.patch.object(editor, "_extract_docx_semantic_trace") as docx_extract,
            mock.patch.object(editor, "_extract_verified_pdf_semantic_trace") as pdf_extract,
            mock.patch.object(editor, "_extract_signing_semantic_trace") as signing_extract,
            mock.patch.object(editor, "_file_hash") as file_hash,
        ):
            validation = self.service.validate_for_publish(
                draft["id"], actor_user_id=1, lock_token=lock["token"],
                actor_is_admin=True,
            )

        self.assertFalse(validation["passed"])
        self.assertIn(
            "simulation_artifact_path_invalid",
            {issue["code"] for issue in validation["issues"]},
        )
        docx_extract.assert_not_called()
        pdf_extract.assert_not_called()
        signing_extract.assert_not_called()
        file_hash.assert_not_called()

    def test_validation_rebuilds_current_render_and_rejects_rebound_old_artifacts(self):
        draft, lock, result = self._simulate_ready_draft()
        old_simulation = dict(self._latest_simulation(draft["id"]))
        old_signing = Path(result["artifact_dir"]) / "signing-preview.html"
        changed_marker = "POST-SIMULATION-EDIT-MUST-NOT-PASS"
        changed = self.service.autosave_block(
            draft["id"],
            "profile_3",
            title="三、服务费用与第三方费用",
            body=changed_marker,
            enabled=True,
            expected_revision=draft["revision"],
            lock_token=lock["token"],
            actor_user_id=1,
        )
        forged = dict(old_simulation)
        forged.update(
            template_identity_id=changed["identity_id"],
            template_revision=changed["revision"],
            template_content_hash=changed["content_hash"],
            template_snapshot_hash=hashlib.sha256(
                json.dumps(
                    json.loads(changed["snapshot_json"]),
                    ensure_ascii=False,
                    sort_keys=True,
                    separators=(",", ":"),
                ).encode("utf-8")
            ).hexdigest(),
        )
        forged["simulation_result_hash"] = editor._hash(
            editor._simulation_result_binding(forged)
        )
        self._insert_simulation_copy(forged)

        validation = self.service.validate_for_publish(
            draft["id"], actor_user_id=1, lock_token=lock["token"], actor_is_admin=True
        )
        self.assertIn(changed_marker, changed["snapshot_json"])
        self.assertNotIn(changed_marker, old_signing.read_text(encoding="utf-8"))
        self.assertFalse(validation["passed"])
        self.assertIn(
            "simulation_semantic_mismatch",
            {issue["code"] for issue in validation["issues"]},
        )

    def test_simulation_and_validation_records_are_permanent_and_relationship_bound(self):
        draft, lock, _result = self._simulate_ready_draft()
        validation = self.service.validate_for_publish(
            draft["id"], actor_user_id=1, lock_token=lock["token"], actor_is_admin=True
        )
        self.assertTrue(validation["passed"])
        simulation = self._latest_simulation(draft["id"])
        validation_row = self.conn.execute(
            "SELECT * FROM contract_template_validations WHERE version_id=? ORDER BY id DESC LIMIT 1",
            (draft["id"],),
        ).fetchone()

        for table, row in (
            ("contract_template_simulations", simulation),
            ("contract_template_validations", validation_row),
        ):
            with self.assertRaisesRegex(sqlite3.IntegrityError, "permanent"):
                self.conn.execute(f"UPDATE {table} SET id=id WHERE id=?", (row["id"],))
            with self.assertRaisesRegex(sqlite3.IntegrityError, "permanent"):
                self.conn.execute(f"UPDATE OR REPLACE {table} SET id=id WHERE id=?", (row["id"],))
            with self.assertRaisesRegex(sqlite3.IntegrityError, "permanent"):
                self.conn.execute(f"DELETE FROM {table} WHERE id=?", (row["id"],))
            columns = list(row.keys())
            with self.assertRaisesRegex(sqlite3.IntegrityError, "permanent"):
                self.conn.execute(
                    f"INSERT OR REPLACE INTO {table}({','.join(columns)}) "
                    f"VALUES({','.join('?' for _ in columns)})",
                    [row[column] for column in columns],
                )

        other_identity_id = self.conn.execute(
            "SELECT id FROM contract_template_identities WHERE id<>? ORDER BY id LIMIT 1",
            (draft["identity_id"],),
        ).fetchone()["id"]
        mismatched = dict(simulation)
        mismatched["template_identity_id"] = other_identity_id
        mismatched["simulation_result_hash"] = editor._hash(
            editor._simulation_result_binding(mismatched)
        )
        with self.assertRaisesRegex(sqlite3.IntegrityError, "identity|relationship|FOREIGN KEY"):
            self._insert_simulation_copy(mismatched)

    def test_malformed_simulation_provenance_returns_stable_fail_closed_issues(self):
        draft, lock, _result = self._simulate_ready_draft()
        simulation = self._latest_simulation(draft["id"])

        self._insert_simulation_copy(simulation, template_revision="not-an-integer")
        validation = self.service.validate_for_publish(
            draft["id"], actor_user_id=1, lock_token=lock["token"], actor_is_admin=True
        )
        self.assertFalse(validation["passed"])
        self.assertIn(
            "simulation_binding_invalid", {issue["code"] for issue in validation["issues"]}
        )

        self._insert_simulation_copy(simulation, input_json="{")
        validation = self.service.validate_for_publish(
            draft["id"], actor_user_id=1, lock_token=lock["token"], actor_is_admin=True
        )
        self.assertFalse(validation["passed"])
        self.assertIn(
            "simulation_input_mismatch", {issue["code"] for issue in validation["issues"]}
        )

        invalid_path = "\0invalid-artifact-path"
        path_row = dict(simulation)
        path_row["artifact_dir"] = invalid_path
        path_row["simulation_result_hash"] = editor._hash(
            editor._simulation_result_binding(path_row)
        )
        self._insert_simulation_copy(path_row)
        validation = self.service.validate_for_publish(
            draft["id"], actor_user_id=1, lock_token=lock["token"], actor_is_admin=True
        )
        self.assertFalse(validation["passed"])
        self.assertIn(
            "simulation_artifact_path_invalid",
            {issue["code"] for issue in validation["issues"]},
        )

        self._insert_simulation_copy(simulation)
        with mock.patch.object(editor, "_file_hash", side_effect=OSError("unreadable")):
            validation = self.service.validate_for_publish(
                draft["id"], actor_user_id=1, lock_token=lock["token"], actor_is_admin=True
            )
        self.assertFalse(validation["passed"])
        self.assertIn(
            "simulation_artifact_unreadable",
            {issue["code"] for issue in validation["issues"]},
        )

    def test_malformed_simulation_identity_is_fail_closed_even_for_preexisting_corruption(self):
        draft, lock, _result = self._simulate_ready_draft()
        simulation = self._latest_simulation(draft["id"])
        self.conn.commit()
        self.conn.execute("PRAGMA foreign_keys=OFF")
        self.conn.execute("DROP TRIGGER IF EXISTS contract_template_simulations_relationship_insert")
        self._insert_simulation_copy(simulation, template_identity_id="not-an-integer")
        self.conn.commit()
        self.conn.execute("PRAGMA foreign_keys=ON")

        validation = self.service.validate_for_publish(
            draft["id"], actor_user_id=1, lock_token=lock["token"], actor_is_admin=True
        )
        self.assertFalse(validation["passed"])
        self.assertIn(
            "simulation_binding_invalid", {issue["code"] for issue in validation["issues"]}
        )

    def test_pdf_wraps_long_visible_lines_across_pages_without_truncation(self):
        from pypdf import PdfReader

        body = "LONG-BEGIN " + ("wrap-segment " * 1600) + "LONG-END-UNTRUNCATED"
        _draft, _lock, result = self._simulate_ready_draft(block_body=body)
        pdf = PdfReader(Path(result["artifact_dir"]) / "preview.pdf")
        extracted = "\n".join(page.extract_text() or "" for page in pdf.pages)
        self.assertGreater(len(pdf.pages), 1)
        self.assertIn("LONG-BEGIN", extracted)
        self.assertIn("LONG-END-UNTRUNCATED", extracted)

    def test_pdf_wraps_long_header_and_footer_inside_every_page_boundary(self):
        import pdfplumber

        self._import_v1()
        draft = self.service.create_draft("NewInternationalStudentContract", actor_user_id=1)
        snapshot = json.loads(draft["snapshot_json"])
        snapshot["contract_title"] = "HEADER-BEGIN " + ("header-segment " * 700) + "HEADER-END"
        snapshot["commercial"]["company_name"] = (
            "FOOTER-BEGIN " + ("footer-segment " * 700) + "FOOTER-END"
        )
        snapshot["commercial"]["company_contact"] = "footer-contact@example.test"
        model = editor.build_visible_render_model(
            snapshot,
            {
                "{{student_name}}": "Boundary Student",
                "{{signer_name}}": "Father Only",
                "{{signer_phone}}": "0411111111",
                "{{signer_email}}": "father@example.test",
                "{{service_scope}}": "Study planning",
            },
            {
                "name": "Father Only",
                "relationship": "父亲",
                "phone": "0411111111",
                "email": "father@example.test",
            },
        )
        path = self.root / "bounded-header-footer.pdf"
        self.service._render_pdf_artifact(model, path)

        with pdfplumber.open(path) as pdf:
            extracted = "\n".join(page.extract_text() or "" for page in pdf.pages)
            self.assertIn("HEADER-END", extracted)
            self.assertIn("FOOTER-END", extracted)
            for page in pdf.pages:
                for character in page.chars:
                    self.assertGreaterEqual(character["x0"], -0.01)
                    self.assertLessEqual(character["x1"], page.width + 0.01)
                    self.assertGreaterEqual(character["top"], -0.01)
                    self.assertLessEqual(character["bottom"], page.height + 0.01)

    def test_pdf_renders_multiline_table_cells_as_distinct_visible_lines(self):
        import pdfplumber

        self._import_v1()
        draft = self.service.create_draft("NewInternationalStudentContract", actor_user_id=1)
        snapshot = json.loads(draft["snapshot_json"])
        snapshot["commercial"]["company_address"] = "ADDRESS-LINE-ONE\nADDRESS-LINE-TWO-END"
        snapshot["commercial"]["company_contact"] = "CONTACT-LINE-ONE\nCONTACT-LINE-TWO-END"
        model = editor.build_visible_render_model(
            snapshot,
            {
                "{{student_name}}": "Boundary Student",
                "{{signer_name}}": "Father Only",
                "{{signer_phone}}": "0411111111",
                "{{signer_email}}": "father@example.test",
                "{{service_scope}}": "Study planning",
            },
            {
                "name": "Father Only",
                "relationship": "父亲",
                "phone": "0411111111",
                "email": "father@example.test",
            },
        )
        path = self.root / "multiline-table.pdf"
        self.service._render_pdf_artifact(model, path)

        with pdfplumber.open(path) as pdf:
            pages = pdf.pages
            extracted = "\n".join(page.extract_text() or "" for page in pages)
            self.assertNotIn("(cid:0)", extracted)
            first = [match for page in pages for match in page.search("ADDRESS-LINE-ONE")]
            second = [match for page in pages for match in page.search("ADDRESS-LINE-TWO-END")]
            self.assertEqual(len(first), 1)
            self.assertEqual(len(second), 1)
            self.assertNotEqual(first[0]["top"], second[0]["top"])

    def test_pdf_embeds_local_cjk_font_for_portable_visible_output(self):
        from pypdf import PdfReader

        _draft, _lock, result = self._simulate_ready_draft()
        pdf = PdfReader(Path(result["artifact_dir"]) / "preview.pdf")
        embedded_fonts = []
        for page in pdf.pages:
            fonts = page["/Resources"].get("/Font", {}).get_object()
            for font_ref in fonts.values():
                font = font_ref.get_object()
                descriptor_ref = font.get("/FontDescriptor")
                if descriptor_ref is None:
                    continue
                descriptor = descriptor_ref.get_object()
                if any(key in descriptor for key in ("/FontFile", "/FontFile2", "/FontFile3")):
                    embedded_fonts.append(str(font.get("/BaseFont") or ""))
        self.assertTrue(embedded_fonts, "PDF must embed the local CJK font used for visible text")

    def test_docx_declares_cjk_font_for_body_headings_header_and_footer(self):
        from docx import Document
        from docx.oxml.ns import qn

        _draft, _lock, result = self._simulate_ready_draft()
        document = Document(Path(result["artifact_dir"]) / "preview.docx")
        for style_name in ("Normal", "Heading 2", "Header", "Footer"):
            style = document.styles[style_name]
            self.assertEqual(style.font.name, "Noto Sans SC")
            self.assertEqual(
                style.element.rPr.rFonts.get(qn("w:eastAsia")), "Noto Sans SC"
            )

    def test_simulation_rejects_divergent_renderer_semantic_trace(self):
        class DivergentPdfService(editor.ContractTemplateEditorService):
            def _render_pdf_artifact(self, model, path):
                rendered = super()._render_pdf_artifact(model, path)
                rendered["semantic_trace"]["header"][0]["lines"] = ["DIVERGENT"]
                return rendered

        service = DivergentPdfService(self.conn)
        self._import_v1()
        draft = service.create_draft("NewInternationalStudentContract", actor_user_id=1)
        lock = service.acquire_lock(
            "NewInternationalStudentContract", draft["id"], owner_user_id=1
        )
        with self.assertRaisesRegex(editor.TemplateRuleError, "语义"):
            service.simulate(
                draft["id"],
                source_type="system",
                student_snapshot=self._simulation_student(),
                signer_source="father",
                signer_overrides={"phone": "0411111111"},
                selected_services=["study_planning_consulting"],
                output_root=self.root / "previews",
                actor_user_id=1,
                expected_revision=draft["revision"],
                lock_token=lock["token"],
            )

    def test_simulation_rejects_post_save_signing_artifact_omission(self):
        class OmittingSigningService(editor.ContractTemplateEditorService):
            @staticmethod
            def _write_signing_artifact(model, path):
                path.write_text(
                    "<!doctype html><article><p>visible output omitted</p></article>",
                    encoding="utf-8",
                )

        service = OmittingSigningService(self.conn)
        self._import_v1()
        draft = service.create_draft("NewInternationalStudentContract", actor_user_id=1)
        lock = service.acquire_lock(
            "NewInternationalStudentContract", draft["id"], owner_user_id=1
        )
        with self.assertRaisesRegex(editor.TemplateRuleError, "成品|语义"):
            service.simulate(
                draft["id"],
                source_type="system",
                student_snapshot=self._simulation_student(),
                signer_source="father",
                signer_overrides={"phone": "0411111111"},
                selected_services=["study_planning_consulting"],
                output_root=self.root / "previews",
                actor_user_id=1,
                expected_revision=draft["revision"],
                lock_token=lock["token"],
            )

    def test_visible_render_model_and_artifacts_preserve_all_semantic_regions(self):
        from docx import Document

        _draft, _lock, result = self._simulate_ready_draft(
            block_body="First paragraph line one\nline two\n\nSecond paragraph END-LINE-MARKER"
        )
        self.assertIn("visible_render_model", result)
        model = result["visible_render_model"]
        self.assertEqual(
            model["header"], [{"lines": ["澳洲留学咨询及签证服务委托合同"]}]
        )
        parties = next(node for node in model["content"] if node.get("key") == "protected_contract_identity")
        party_rows = {
            row["cells"][0]["lines"][0]: row["cells"][1]["lines"][0]
            for row in parties["rows"]
        }
        self.assertEqual(party_rows["学生姓名"], "Boundary Student")
        self.assertEqual(party_rows["乙方签署人"], "Father Only")
        changed = next(node for node in model["content"] if node.get("key") == "profile_3")
        self.assertEqual(
            changed["paragraphs"],
            [
                {"lines": ["First paragraph line one", "line two"]},
                {"lines": ["Second paragraph END-LINE-MARKER"]},
            ],
        )
        commercial = next(node for node in model["content"] if node.get("key") == "protected_pricing_table")
        self.assertEqual(
            commercial["rows"][:2],
            [
                {"cells": [{"lines": ["费用名称"]}, {"lines": ["留学咨询及签证服务费"]}]},
                {"cells": [{"lines": ["服务类型"]}, {"lines": ["留学咨询 / 学校申请 / 签证服务"]}]},
            ],
        )
        self.assertEqual(
            model["footer"],
            [
                {
                    "lines": [
                        "Australian Academy of International Education",
                        "contracts@example.test",
                    ]
                }
            ],
        )
        self.assertEqual(
            model["signature_region"],
            {
                "title": "十一、签署",
                "paragraphs": [
                    {
                        "lines": [
                            "AAIE / 甲方（签字 / 盖章）",
                            "代理方（甲方盖章）：Australian Academy of International Education",
                            "授权代表：",
                            "甲方签署日期：",
                            "委托方（乙方签字）",
                            "法定代理人或者监护人：Father Only",
                            "关系：父亲",
                            "家长/乙方签署日期：",
                        ]
                    }
                ],
            },
        )

        artifact_dir = Path(result["artifact_dir"])
        document = Document(artifact_dir / "preview.docx")
        self.assertGreaterEqual(len(document.inline_shapes), 1)
        docx_rows = [[cell.text for cell in row.cells] for table in document.tables for row in table.rows]
        self.assertIn(
            ["费用名称", "留学咨询及签证服务费"],
            docx_rows,
        )
        docx_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        self.assertIn("First paragraph line one\nline two", docx_text)
        self.assertIn("Second paragraph END-LINE-MARKER", docx_text)
        self.assertIn("十一、签署", docx_text)

        signing_html = (artifact_dir / "signing-preview.html").read_text(encoding="utf-8")
        self.assertIn('data-master-anchor="section_11_signature"', signing_html)
        self.assertIn('data-stamp-anchor="aaie-only"', signing_html)
        self.assertIn("First paragraph line one<br>line two", signing_html)
        self.assertIn('data-signature-anchor="signature_aaie"', signing_html)
        self.assertIn('data-signature-anchor="signature_family"', signing_html)

        manifest = result["artifact_manifest"]
        for filename, result_key in {
            "preview.docx": "docx_artifact_hash",
            "preview.pdf": "pdf_artifact_hash",
            "signing-preview.html": "signing_artifact_hash",
        }.items():
            digest = hashlib.sha256((artifact_dir / filename).read_bytes()).hexdigest()
            self.assertEqual(result[result_key], digest)
            self.assertEqual(manifest[filename]["sha256"], digest)
            self.assertEqual(manifest[filename]["bytes"], (artifact_dir / filename).stat().st_size)
        self.assertEqual(
            len(
                {
                    result["docx_semantic_hash"],
                    result["pdf_semantic_hash"],
                    result["signing_semantic_hash"],
                }
            ),
            1,
        )

    def test_publish_validation_blocks_invalid_draft_without_switching_authority(self):
        self._import_v1()
        before = self.service.current_published("TransferApplicationContract")
        draft = self.service.create_draft("TransferApplicationContract", actor_user_id=1)
        lock = self.service.acquire_lock("TransferApplicationContract", draft["id"], owner_user_id=1)

        validation = self.service.validate_for_publish(
            draft["id"], actor_user_id=1, lock_token=lock["token"], actor_is_admin=True,
        )

        self.assertFalse(validation["passed"])
        self.assertIn("simulation_required", {issue["code"] for issue in validation["issues"]})
        with self.assertRaises(editor.PublishBlocked):
            self.service.publish(draft["id"], actor_user_id=1, lock_token=lock["token"], actor_is_admin=True)
        self.assertEqual(self.service.current_published("TransferApplicationContract")["id"], before["id"])

    def test_publish_rechecks_after_begin_immediate_and_rejects_interposed_legal_autosave(self):
        draft, lock, _result = self._simulate_ready_draft()
        self.service.prepare_preview(
            draft["id"], actor_user_id=1, lock_token=lock["token"], actor_is_admin=True,
            expected_revision=draft["revision"],
        )
        previous_id = self.service.current_published(
            "NewInternationalStudentContract"
        )["id"]
        competing_connection = sqlite3.connect(self.root / "editor.sqlite3")
        competing_connection.row_factory = sqlite3.Row
        competing_connection.execute("PRAGMA foreign_keys=ON")
        self.addCleanup(competing_connection.close)
        competing_service = editor.ContractTemplateEditorService(competing_connection)
        interposed = []

        def interpose_before_immediate(statement):
            if statement.strip().upper() != "BEGIN IMMEDIATE" or interposed:
                return
            interposed.append(True)
            competing_service.autosave_settings(
                draft["id"],
                commercial={**self._valid_commercial(), "service_fee": "3999.00"},
                services=[{
                    "service_id": "study_planning_consulting",
                    "enabled": True,
                    "default_selected": True,
                }],
                expected_revision=draft["revision"],
                lock_token=lock["token"],
                actor_user_id=1,
            )

        self.conn.set_trace_callback(interpose_before_immediate)
        self.addCleanup(self.conn.set_trace_callback, None)

        with self.assertRaises(editor.StaleRevision):
            self.service.publish(
                draft["id"], actor_user_id=1, lock_token=lock["token"],
                expected_revision=draft["revision"], actor_is_admin=True,
            )

        self.assertEqual(interposed, [True])
        current = self.service.current_published("NewInternationalStudentContract")
        rejected = self.service._version(draft["id"])
        self.assertEqual(current["id"], previous_id)
        self.assertFalse(rejected["immutable"])
        self.assertNotEqual(rejected["status"], "Published")

    def test_publish_rolls_back_validation_preview_pointer_history_and_audit_together(self):
        class FailingPublishAuditService(editor.ContractTemplateEditorService):
            def __init__(self, connection):
                super().__init__(connection)
                self.publish_audit_attempted = False

            def _audit(self, action, actor, identity_id=None, version_id=None, before=None, after=None):
                if action == "template_published":
                    self.publish_audit_attempted = True
                    raise RuntimeError("publish audit failure")
                return super()._audit(
                    action, actor, identity_id, version_id, before, after
                )

        service = FailingPublishAuditService(self.conn)
        draft, lock, _result = self._simulate_ready_draft(service=service)
        previous_id = service.current_published("NewInternationalStudentContract")["id"]
        before_status = service._version(draft["id"])["status"]
        before_validations = self.conn.execute(
            "SELECT COUNT(*) FROM contract_template_validations WHERE version_id=?",
            (draft["id"],),
        ).fetchone()[0]
        before_audits = self.conn.execute(
            "SELECT COUNT(*) FROM contract_template_audit_events WHERE version_id=?",
            (draft["id"],),
        ).fetchone()[0]

        with self.assertRaisesRegex(RuntimeError, "publish audit failure"):
            service.publish(
                draft["id"], actor_user_id=1, lock_token=lock["token"],
                expected_revision=draft["revision"], actor_is_admin=True,
            )
        self.conn.rollback()

        self.assertTrue(service.publish_audit_attempted)
        self.assertEqual(service._version(draft["id"])["status"], before_status)
        self.assertFalse(service._version(draft["id"])["immutable"])
        self.assertEqual(
            service.current_published("NewInternationalStudentContract")["id"],
            previous_id,
        )
        self.assertEqual(
            self.conn.execute(
                "SELECT COUNT(*) FROM contract_template_validations WHERE version_id=?",
                (draft["id"],),
            ).fetchone()[0],
            before_validations,
        )
        self.assertEqual(
            self.conn.execute(
                "SELECT COUNT(*) FROM contract_template_audit_events WHERE version_id=?",
                (draft["id"],),
            ).fetchone()[0],
            before_audits,
        )

    def test_simulate_requires_current_revision_and_live_owner_lock(self):
        self._import_v1()
        draft = self.service.create_draft(
            "NewInternationalStudentContract", actor_user_id=1
        )
        lock = self.service.acquire_lock(
            "NewInternationalStudentContract", draft["id"], owner_user_id=1
        )
        changed = self.service.autosave_settings(
            draft["id"], commercial=self._valid_commercial(), services=[],
            expected_revision=draft["revision"], lock_token=lock["token"],
            actor_user_id=1,
        )
        attempts = [
            (editor.StaleRevision, draft["revision"], lock["token"], 1),
            (editor.EditLockConflict, changed["revision"], lock["token"], 2),
            (editor.EditLockConflict, changed["revision"], "stale-token", 1),
        ]
        for expected_error, revision, token, actor in attempts:
            caught = None
            try:
                self.service.simulate(
                    draft["id"], source_type="system",
                    student_snapshot=self._simulation_student(),
                    signer_source="father", signer_overrides={},
                    selected_services=[], output_root=self.root / "previews",
                    actor_user_id=actor, expected_revision=revision,
                    lock_token=token,
                )
            except Exception as exc:
                caught = exc
            self.assertIsInstance(caught, expected_error)

    def test_different_user_lock_blocks_simulate_validate_preview_and_publish_without_writes(self):
        draft, lock, _result = self._simulate_ready_draft()
        self.conn.execute(
            "UPDATE contract_template_edit_locks SET owner_user_id=?,expires_at=? WHERE version_id=?",
            (2, "2999-01-01T00:00:00+00:00", draft["id"]),
        )
        self.conn.commit()
        before_status = self.service._version(draft["id"])["status"]
        before_validations = self.conn.execute(
            "SELECT COUNT(*) FROM contract_template_validations WHERE version_id=?",
            (draft["id"],),
        ).fetchone()[0]
        before_audits = self.conn.execute(
            "SELECT COUNT(*) FROM contract_template_audit_events WHERE version_id=?",
            (draft["id"],),
        ).fetchone()[0]

        operations = [
            lambda: self.service.simulate(
                draft["id"], source_type="system",
                student_snapshot=self._simulation_student(), signer_source="father",
                signer_overrides={}, selected_services=[],
                output_root=self.root / "expired", actor_user_id=1,
                expected_revision=draft["revision"], lock_token=lock["token"],
            ),
            lambda: self.service.validate_for_publish(
                draft["id"], actor_user_id=1, lock_token=lock["token"],
                expected_revision=draft["revision"], actor_is_admin=True,
            ),
            lambda: self.service.prepare_preview(
                draft["id"], actor_user_id=1, lock_token=lock["token"],
                expected_revision=draft["revision"], actor_is_admin=True,
            ),
            lambda: self.service.publish(
                draft["id"], actor_user_id=1, lock_token=lock["token"],
                expected_revision=draft["revision"], actor_is_admin=True,
            ),
        ]
        for operation in operations:
            with self.assertRaises(editor.EditLockConflict):
                operation()

        self.assertEqual(self.service._version(draft["id"])["status"], before_status)
        self.assertEqual(
            self.conn.execute(
                "SELECT COUNT(*) FROM contract_template_validations WHERE version_id=?",
                (draft["id"],),
            ).fetchone()[0],
            before_validations,
        )
        self.assertEqual(
            self.conn.execute(
                "SELECT COUNT(*) FROM contract_template_audit_events WHERE version_id=?",
                (draft["id"],),
            ).fetchone()[0],
            before_audits,
        )

    def test_validate_and_publish_reject_stale_revision_and_other_lock_owner_without_writes(self):
        draft, lock, _result = self._simulate_ready_draft()
        before_status = self.service._version(draft["id"])["status"]
        before_validations = self.conn.execute(
            "SELECT COUNT(*) FROM contract_template_validations WHERE version_id=?",
            (draft["id"],),
        ).fetchone()[0]
        before_audits = self.conn.execute(
            "SELECT COUNT(*) FROM contract_template_audit_events WHERE version_id=?",
            (draft["id"],),
        ).fetchone()[0]

        for operation in (
            lambda: self.service.validate_for_publish(
                draft["id"], actor_user_id=1, lock_token=lock["token"],
                expected_revision=draft["revision"] - 1, actor_is_admin=True,
            ),
            lambda: self.service.publish(
                draft["id"], actor_user_id=1, lock_token=lock["token"],
                expected_revision=draft["revision"] - 1, actor_is_admin=True,
            ),
        ):
            with self.assertRaises(editor.StaleRevision):
                operation()
        for operation in (
            lambda: self.service.validate_for_publish(
                draft["id"], actor_user_id=2, lock_token=lock["token"],
                expected_revision=draft["revision"], actor_is_admin=True,
            ),
            lambda: self.service.publish(
                draft["id"], actor_user_id=2, lock_token=lock["token"],
                expected_revision=draft["revision"], actor_is_admin=True,
            ),
        ):
            with self.assertRaises(editor.EditLockConflict):
                operation()

        self.assertEqual(self.service._version(draft["id"])["status"], before_status)
        self.assertEqual(
            self.conn.execute(
                "SELECT COUNT(*) FROM contract_template_validations WHERE version_id=?",
                (draft["id"],),
            ).fetchone()[0],
            before_validations,
        )
        self.assertEqual(
            self.conn.execute(
                "SELECT COUNT(*) FROM contract_template_audit_events WHERE version_id=?",
                (draft["id"],),
            ).fetchone()[0],
            before_audits,
        )

    def test_unauthorized_validation_is_rejected_before_any_service_write(self):
        self._import_v1()
        draft = self.service.create_draft(
            "TransferApplicationContract", actor_user_id=1
        )
        lock = self.service.acquire_lock(
            "TransferApplicationContract", draft["id"], owner_user_id=9
        )
        before_status = self.service._version(draft["id"])["status"]
        before_audits = len(self.service.audit_events())

        with self.assertRaises(editor.PermissionDenied):
            self.service.validate_for_publish(
                draft["id"], actor_user_id=9, lock_token=lock["token"],
                expected_revision=draft["revision"], actor_is_admin=False,
            )

        self.assertEqual(self.service._version(draft["id"])["status"], before_status)
        self.assertEqual(
            self.conn.execute(
                "SELECT COUNT(*) FROM contract_template_validations WHERE version_id=?",
                (draft["id"],),
            ).fetchone()[0],
            0,
        )
        self.assertEqual(len(self.service.audit_events()), before_audits)

    def test_service_does_not_trust_forged_admin_flag_when_live_role_is_staff(self):
        self.conn.executescript(
            """
            CREATE TABLE roles(id INTEGER PRIMARY KEY,name TEXT NOT NULL);
            CREATE TABLE users(id INTEGER PRIMARY KEY,role_id INTEGER NOT NULL);
            INSERT INTO roles(id,name) VALUES(1,'Staff');
            INSERT INTO users(id,role_id) VALUES(9,1);
            """
        )
        self._import_v1()
        draft = self.service.create_draft(
            "TransferApplicationContract", actor_user_id=1
        )
        lock = self.service.acquire_lock(
            "TransferApplicationContract", draft["id"], owner_user_id=9
        )

        with self.assertRaises(editor.PermissionDenied):
            self.service.validate_for_publish(
                draft["id"], actor_user_id=9, lock_token=lock["token"],
                expected_revision=draft["revision"], actor_is_admin=True,
            )
        self.assertEqual(
            self.conn.execute(
                "SELECT COUNT(*) FROM contract_template_validations WHERE version_id=?",
                (draft["id"],),
            ).fetchone()[0],
            0,
        )

    def test_successful_publish_is_immutable_and_historical_republish_creates_new_version(self):
        self._import_v1()
        draft = self.service.create_draft("LocalStudentApplicationContract", actor_user_id=1)
        lock = self.service.acquire_lock("LocalStudentApplicationContract", draft["id"], owner_user_id=1)
        draft = self.service.autosave_settings(
            draft["id"], commercial=self._valid_commercial(),
            services=[
                {"service_id": service_id, "enabled": True, "default_selected": True}
                for service_id in contract_domain.default_service_ids("LocalStudentApplicationContract")
            ], expected_revision=draft["revision"], lock_token=lock["token"], actor_user_id=1,
        )
        self.service.simulate(
            draft["id"], source_type="system",
            student_snapshot={
                "student_name": "Local Student", "student_id": "T-LOCAL-001",
                "date_of_birth": "2011-02-03", "current_year_level": "Year 7",
                "applying_year_level": "Year 8", "preferred_intake_year": "2027",
                "company_authorised_person": "Sonny Zhang",
                "mother_passport_first_name": "Mum",
                "mother_passport_last_name": "Local", "mother_mobile": "0400000000",
                "mother_email": "mum@example.test",
            }, signer_source="mother", signer_overrides={},
            selected_services=list(contract_domain.default_service_ids("LocalStudentApplicationContract")),
            output_root=self.root / "previews", actor_user_id=1,
            expected_revision=draft["revision"], lock_token=lock["token"],
        )
        preview = self.service.prepare_preview(
            draft["id"], actor_user_id=1, lock_token=lock["token"], actor_is_admin=True,
        )
        self.assertEqual(preview["status"], "Preview")
        published = self.service.publish(draft["id"], actor_user_id=1, lock_token=lock["token"], actor_is_admin=True)

        self.assertEqual(published["status"], "Published")
        self.assertTrue(published["immutable"])
        history = self.service.version_history("LocalStudentApplicationContract")
        self.assertEqual([row["status"] for row in history[:2]], ["Published", "Historical"])
        with self.assertRaises(editor.ImmutableVersionError):
            self.service.create_block(published["id"], "late", "Late", "No", 1, "invalid", 1)

        republish_draft = self.service.create_draft_from_version(published["id"], actor_user_id=1)
        self.assertGreater(republish_draft["version_number"], published["version_number"])
        self.assertNotEqual(republish_draft["id"], published["id"])

        with self.assertRaises(sqlite3.IntegrityError):
            self.conn.execute(
                "INSERT INTO contract_template_clause_blocks(version_id,block_key,title,body,sort_order,enabled,fixed_layout,required,required_variables_json) VALUES(?,?,?,?,?,?,?,?,?)",
                (published["id"], "late_insert", "Late", "No", 999, 1, 0, 0, "[]"),
            )
        self.conn.rollback()
        with self.assertRaises(sqlite3.IntegrityError):
            self.conn.execute(
                "INSERT INTO contract_template_service_items(version_id,service_id,label,sort_order,enabled,default_selected) VALUES(?,?,?,?,?,?)",
                (published["id"], "late_insert", "Late", 999, 1, 0),
            )
        self.conn.rollback()

    def test_direct_sql_cannot_unfreeze_an_immutable_version_before_mutating_it(self):
        """Catch an attacker that downgrades `immutable` then changes published content."""
        self._import_v1()
        published = self.service.current_published("NewInternationalStudentContract")

        with self.assertRaisesRegex(sqlite3.IntegrityError, "immutable"):
            self.conn.execute(
                "UPDATE contract_template_versions SET immutable=0 WHERE id=?",
                (published["id"],),
            )
        self.conn.rollback()
        self.assertEqual(
            1,
            self.conn.execute(
                "SELECT immutable FROM contract_template_versions WHERE id=?",
                (published["id"],),
            ).fetchone()[0],
        )
        with self.assertRaisesRegex(sqlite3.IntegrityError, "immutable"):
            self.conn.execute(
                "UPDATE contract_template_versions SET snapshot_json='{}' WHERE id=?",
                (published["id"],),
            )
        self.conn.rollback()

    def test_direct_sql_cannot_downgrade_or_rewrite_an_immutable_version(self):
        """Catch direct published-status, snapshot, revision, lineage and publication attacks."""
        self._import_v1()
        published = self.service.current_published("TransferApplicationContract")
        other_published = self.service.current_published("NewInternationalStudentContract")
        attacks = (
            ("status='Draft'", "status downgrade"),
            ("identity_id=identity_id+1", "identity rewrite"),
            ("version_number=version_number+100", "version number rewrite"),
            ("snapshot_json='{}'", "snapshot rewrite"),
            ("content_hash='forged'", "content hash rewrite"),
            ("revision=revision+1", "revision rewrite"),
            (f"source_version_id={other_published['id']}", "lineage rewrite"),
            ("created_by=999", "creator rewrite"),
            ("created_at='2099-01-01T00:00:00+00:00'", "creation timestamp rewrite"),
            ("published_by=999", "publisher rewrite"),
            ("published_at='2099-01-01T00:00:00+00:00'", "publication timestamp rewrite"),
        )

        for assignment, attack_name in attacks:
            with self.subTest(attack_name=attack_name), self.assertRaisesRegex(
                sqlite3.IntegrityError, "immutable"
            ):
                self.conn.execute(
                    f"UPDATE contract_template_versions SET {assignment} WHERE id=?",
                    (published["id"],),
                )
            self.conn.rollback()

    def test_legal_clone_and_publish_to_historical_transition_remain_possible(self):
        """Catch an over-broad trigger that blocks clone or legal immutable transitions."""
        self._import_v1()
        published = self.service.current_published("LocalStudentApplicationContract")

        draft = self.service.create_draft_from_version(published["id"], actor_user_id=1)
        self.assertEqual("Draft", draft["status"])
        self.assertEqual(0, draft["immutable"])
        self.assertEqual(published["id"], draft["source_version_id"])

        self.conn.execute(
            """UPDATE contract_template_versions
               SET status='Published', immutable=1, published_by=1,
                   published_at='2026-08-19T00:00:00+00:00'
               WHERE id=?""",
            (draft["id"],),
        )
        self.conn.execute(
            "UPDATE contract_template_versions SET status='Historical' WHERE id=?",
            (draft["id"],),
        )
        self.conn.commit()
        historical = self.service._version(draft["id"])
        self.assertEqual("Historical", historical["status"])
        self.assertEqual(1, historical["immutable"])

    def test_direct_sql_cannot_reparent_mutable_children_into_an_immutable_version(self):
        """Catch moving new clause/service rows into the effective published content."""
        self._import_v1()
        published = self.service.current_published("NewInternationalStudentContract")
        draft = self.service.create_draft("NewInternationalStudentContract", actor_user_id=1)

        self.conn.execute(
            """INSERT INTO contract_template_clause_blocks(
               version_id,block_key,title,body,sort_order,enabled,fixed_layout,required,required_variables_json
               ) VALUES(?,?,?,?,?,?,?,?,?)""",
            (draft["id"], "rogue_clause", "Rogue", "unapproved", 99, 1, 0, 0, "[]"),
        )
        clause_id = self.conn.execute(
            "SELECT id FROM contract_template_clause_blocks WHERE version_id=? AND block_key='rogue_clause'",
            (draft["id"],),
        ).fetchone()[0]
        with self.assertRaisesRegex(sqlite3.IntegrityError, "immutable"):
            self.conn.execute(
                "UPDATE contract_template_clause_blocks SET version_id=? WHERE id=?",
                (published["id"], clause_id),
            )
        self.conn.rollback()

        self.conn.execute(
            """INSERT INTO contract_template_service_items(
               version_id,service_id,label,sort_order,enabled,default_selected
               ) VALUES(?,?,?,?,?,?)""",
            (draft["id"], "rogue_service", "Rogue", 99, 1, 0),
        )
        service_id = self.conn.execute(
            "SELECT id FROM contract_template_service_items WHERE version_id=? AND service_id='rogue_service'",
            (draft["id"],),
        ).fetchone()[0]
        with self.assertRaisesRegex(sqlite3.IntegrityError, "immutable"):
            self.conn.execute(
                "UPDATE contract_template_service_items SET version_id=? WHERE id=?",
                (published["id"], service_id),
            )
        self.conn.rollback()

    def test_direct_sql_cannot_create_or_point_to_a_mutable_published_version(self):
        """Catch Published+immutable=0 inserts, transitions, and authority-pointer bypasses."""
        self._import_v1()
        identity = self.service._identity("TransferApplicationContract")
        published = self.service.current_published("TransferApplicationContract")
        draft = self.service.create_draft("TransferApplicationContract", actor_user_id=1)

        with self.assertRaisesRegex(sqlite3.IntegrityError, "published"):
            self.conn.execute(
                """INSERT INTO contract_template_versions(
                   identity_id,version_number,status,revision,snapshot_json,content_hash,immutable,created_by,created_at
                   ) VALUES(?,?, 'Published', 1, '{}', 'forged', 0, 1, '2026-08-19T00:00:00+00:00')""",
                (identity["id"], 99),
            )
        self.conn.rollback()
        with self.assertRaisesRegex(sqlite3.IntegrityError, "published"):
            self.conn.execute(
                "UPDATE contract_template_versions SET status='Published',immutable=0 WHERE id=?",
                (draft["id"],),
            )
        self.conn.rollback()
        with self.assertRaisesRegex(sqlite3.IntegrityError, "published"):
            self.conn.execute(
                "UPDATE contract_template_identities SET current_published_version_id=? WHERE id=?",
                (draft["id"], identity["id"]),
            )
        self.conn.rollback()
        self.assertEqual(published["id"], self.service.current_published("TransferApplicationContract")["id"])

    def test_direct_sql_cannot_insert_a_forged_immutable_published_version(self):
        """Catch a fully-provenanced Published insert that skips Draft-to-Published."""
        self._import_v1()
        identity = self.service._identity("TransferApplicationContract")

        with self.assertRaisesRegex(sqlite3.IntegrityError, "published"):
            self.conn.execute(
                """INSERT INTO contract_template_versions(
                   identity_id,version_number,status,revision,snapshot_json,content_hash,immutable,
                   created_by,created_at,published_by,published_at
                   ) VALUES(?,?, 'Published', 1, '{\"forged\":true}', 'forged', 1,
                             999, '2026-08-19T00:00:00+00:00', 999, '2026-08-19T00:00:00+00:00')""",
                (identity["id"], 99),
            )
        self.conn.rollback()

    def test_direct_sql_cannot_create_an_identity_with_another_identitys_current_version(self):
        """Catch cross-identity current-published assignment during identity creation."""
        self._import_v1()
        published = self.service.current_published("NewInternationalStudentContract")

        with self.assertRaisesRegex(sqlite3.IntegrityError, "published"):
            self.conn.execute(
                """INSERT INTO contract_template_identities(
                   template_type,display_name,current_published_version_id,created_at,updated_at
                   ) VALUES(?,?,?,?,?)""",
                (
                    "ForgedContractTemplate",
                    "Forged template",
                    published["id"],
                    "2026-08-19T00:00:00+00:00",
                    "2026-08-19T00:00:00+00:00",
                ),
            )
        self.conn.rollback()

    def test_insert_or_replace_cannot_replace_an_immutable_published_version(self):
        """Catch REPLACE deleting a published row before inserting a mutable forgery."""
        self._import_v1()
        published = self.service.current_published("NewInternationalStudentContract")
        self.assertEqual(0, self.conn.execute("PRAGMA recursive_triggers").fetchone()[0])
        self.conn.execute("PRAGMA defer_foreign_keys=ON")

        with self.assertRaisesRegex(sqlite3.IntegrityError, "permanent"):
            self.conn.execute(
                """INSERT OR REPLACE INTO contract_template_versions(
                   id,identity_id,version_number,status,revision,snapshot_json,content_hash,immutable,
                   created_by,created_at
                   ) VALUES(?,?,?,'Draft',1,'{\"forged\":true}','forged',0,999,'2026-08-19T00:00:00+00:00')""",
                (published["id"], published["identity_id"], published["version_number"]),
            )
        self.conn.rollback()

    def test_insert_or_replace_cannot_replace_an_append_only_audit_event(self):
        """Catch REPLACE bypassing append-only audit UPDATE/DELETE triggers."""
        self._import_v1()
        event = self.conn.execute(
            "SELECT id FROM contract_template_audit_events ORDER BY id LIMIT 1"
        ).fetchone()
        self.assertIsNotNone(event)
        self.assertEqual(0, self.conn.execute("PRAGMA recursive_triggers").fetchone()[0])

        with self.assertRaisesRegex(sqlite3.IntegrityError, "append-only"):
            self.conn.execute(
                """INSERT OR REPLACE INTO contract_template_audit_events(id,action,created_at)
                   VALUES(?, 'forged_audit', '2026-08-19T00:00:00+00:00')""",
                (event["id"],),
            )
        self.conn.rollback()

    def test_insert_or_replace_cannot_remove_an_immutable_clause_or_service_row(self):
        """Catch child-primary-key REPLACE that moves an immutable child into a draft."""
        self._import_v1()
        published = self.service.current_published("TransferApplicationContract")
        draft = self.service.create_draft("TransferApplicationContract", actor_user_id=1)
        block = self.conn.execute(
            "SELECT id FROM contract_template_clause_blocks WHERE version_id=? ORDER BY id LIMIT 1",
            (published["id"],),
        ).fetchone()
        service = self.conn.execute(
            "SELECT id FROM contract_template_service_items WHERE version_id=? ORDER BY id LIMIT 1",
            (published["id"],),
        ).fetchone()
        self.assertEqual(0, self.conn.execute("PRAGMA recursive_triggers").fetchone()[0])

        with self.assertRaisesRegex(sqlite3.IntegrityError, "immutable"):
            self.conn.execute(
                """INSERT OR REPLACE INTO contract_template_clause_blocks(
                   id,version_id,block_key,title,body,sort_order,enabled,fixed_layout,required,required_variables_json
                   ) VALUES(?,?, 'forged_clause', 'Forged', 'forged', 99,1,0,0,'[]')""",
                (block["id"], draft["id"]),
            )
        self.conn.rollback()
        with self.assertRaisesRegex(sqlite3.IntegrityError, "immutable"):
            self.conn.execute(
                """INSERT OR REPLACE INTO contract_template_service_items(
                   id,version_id,service_id,label,sort_order,enabled,default_selected
                   ) VALUES(?,?, 'forged_service', 'Forged', 99,1,0)""",
                (service["id"], draft["id"]),
            )
        self.conn.rollback()

    def test_insert_or_replace_cannot_replace_an_existing_template_identity(self):
        """Catch identity-primary-key REPLACE clearing current authority under deferred FKs."""
        self._import_v1()
        identity = self.service._identity("LocalStudentApplicationContract")
        self.assertEqual(0, self.conn.execute("PRAGMA recursive_triggers").fetchone()[0])
        self.conn.execute("PRAGMA defer_foreign_keys=ON")

        with self.assertRaisesRegex(sqlite3.IntegrityError, "permanent"):
            self.conn.execute(
                """INSERT OR REPLACE INTO contract_template_identities(
                   id,template_type,display_name,current_published_version_id,created_at,updated_at
                   ) VALUES(?,?,?,NULL,?,?)""",
                (
                    identity["id"],
                    "ForgedContractTemplate",
                    "Forged template",
                    "2026-08-19T00:00:00+00:00",
                    "2026-08-19T00:00:00+00:00",
                ),
            )
        self.conn.rollback()

    def test_direct_sql_cannot_clear_an_established_current_published_pointer(self):
        """Catch clearing canonical current authority after first publication."""
        self._import_v1()
        identity = self.service._identity("NewInternationalStudentContract")

        with self.assertRaisesRegex(sqlite3.IntegrityError, "published"):
            self.conn.execute(
                "UPDATE contract_template_identities SET current_published_version_id=NULL WHERE id=?",
                (identity["id"],),
            )
        self.conn.rollback()

    def test_direct_sql_cannot_historicalize_the_current_published_target(self):
        """Catch leaving an identity pointing at a Historical version."""
        self._import_v1()
        published = self.service.current_published("NewInternationalStudentContract")

        with self.assertRaisesRegex(sqlite3.IntegrityError, "published"):
            self.conn.execute(
                "UPDATE contract_template_versions SET status='Historical' WHERE id=?",
                (published["id"],),
            )
        self.conn.rollback()

    def test_direct_sql_cannot_rewrite_an_immutable_version_primary_key(self):
        """Catch primary-key mutation even when no child row happens to block it."""
        self._import_v1()
        identity = self.service._identity("NewInternationalStudentContract")
        self.conn.execute(
            """INSERT INTO contract_template_versions(
               identity_id,version_number,status,revision,snapshot_json,content_hash,immutable,created_by,created_at
               ) VALUES(?,99,'Draft',1,'{}','draft',0,1,'2026-08-19T00:00:00+00:00')""",
            (identity["id"],),
        )
        version_id = self.conn.execute("SELECT last_insert_rowid()").fetchone()[0]
        self.conn.execute(
            """UPDATE contract_template_versions
               SET status='Published',immutable=1,published_by=1,published_at='2026-08-19T00:00:00+00:00'
               WHERE id=?""",
            (version_id,),
        )

        with self.assertRaisesRegex(sqlite3.IntegrityError, "immutable"):
            self.conn.execute(
                "UPDATE contract_template_versions SET id=id+1000 WHERE id=?", (version_id,)
            )
        self.conn.rollback()

    def test_direct_sql_cannot_transition_a_draft_directly_to_historical(self):
        """Catch forged history that skips the immutable Published state."""
        self._import_v1()
        draft = self.service.create_draft("NewInternationalStudentContract", actor_user_id=1)

        with self.assertRaisesRegex(sqlite3.IntegrityError, "historical"):
            self.conn.execute(
                """UPDATE contract_template_versions
                   SET status='Historical',immutable=1,published_by=999,
                       published_at='2026-08-19T00:00:00+00:00' WHERE id=?""",
                (draft["id"],),
            )
        self.conn.rollback()

    def test_update_or_replace_cannot_replace_an_immutable_version_with_a_mutable_draft(self):
        """Catch conflict UPDATE deleting the immutable version rather than the draft source."""
        self._import_v1()
        published = self.service.current_published("NewInternationalStudentContract")
        self.conn.execute(
            """INSERT INTO contract_template_versions(
               identity_id,version_number,status,revision,snapshot_json,content_hash,immutable,created_by,created_at
               ) VALUES(?,99,'Draft',1,'{}','draft',0,1,'2026-08-19T00:00:00+00:00')""",
            (published["identity_id"],),
        )
        draft_id = self.conn.execute("SELECT last_insert_rowid()").fetchone()[0]
        self.assertEqual(0, self.conn.execute("PRAGMA recursive_triggers").fetchone()[0])
        self.conn.execute("PRAGMA defer_foreign_keys=ON")

        with self.assertRaisesRegex(sqlite3.IntegrityError, "permanent"):
            self.conn.execute(
                """UPDATE OR REPLACE contract_template_versions
                   SET id=?,identity_id=?,version_number=? WHERE id=?""",
                (published["id"], published["identity_id"], published["version_number"], draft_id),
            )
        self.conn.rollback()

    def test_update_or_replace_cannot_delete_an_immutable_clause_or_service_conflict_target(self):
        """Catch mutable child updates whose PK conflict would delete published child content."""
        self._import_v1()
        published = self.service.current_published("TransferApplicationContract")
        draft = self.service.create_draft("TransferApplicationContract", actor_user_id=1)
        published_block = self.conn.execute(
            "SELECT id FROM contract_template_clause_blocks WHERE version_id=? ORDER BY id LIMIT 1",
            (published["id"],),
        ).fetchone()
        draft_block = self.conn.execute(
            "SELECT id FROM contract_template_clause_blocks WHERE version_id=? ORDER BY id LIMIT 1",
            (draft["id"],),
        ).fetchone()
        published_service = self.conn.execute(
            "SELECT id FROM contract_template_service_items WHERE version_id=? ORDER BY id LIMIT 1",
            (published["id"],),
        ).fetchone()
        draft_service = self.conn.execute(
            "SELECT id FROM contract_template_service_items WHERE version_id=? ORDER BY id LIMIT 1",
            (draft["id"],),
        ).fetchone()
        self.assertEqual(0, self.conn.execute("PRAGMA recursive_triggers").fetchone()[0])

        with self.assertRaisesRegex(sqlite3.IntegrityError, "immutable"):
            self.conn.execute(
                "UPDATE OR REPLACE contract_template_clause_blocks SET id=? WHERE id=?",
                (published_block["id"], draft_block["id"]),
            )
        self.conn.rollback()
        with self.assertRaisesRegex(sqlite3.IntegrityError, "immutable"):
            self.conn.execute(
                "UPDATE OR REPLACE contract_template_service_items SET id=? WHERE id=?",
                (published_service["id"], draft_service["id"]),
            )
        self.conn.rollback()

    def test_update_or_replace_cannot_replace_a_canonical_identity(self):
        """Catch a blank identity update that deletes a populated canonical identity."""
        self._import_v1()
        target = self.service._identity("LocalStudentApplicationContract")
        self.conn.execute(
            """INSERT INTO contract_template_identities(
               template_type,display_name,current_published_version_id,created_at,updated_at
               ) VALUES('MutableIdentity','Mutable',NULL,'2026-08-19T00:00:00+00:00','2026-08-19T00:00:00+00:00')"""
        )
        source_id = self.conn.execute("SELECT last_insert_rowid()").fetchone()[0]
        self.assertEqual(0, self.conn.execute("PRAGMA recursive_triggers").fetchone()[0])
        self.conn.execute("PRAGMA defer_foreign_keys=ON")

        with self.assertRaisesRegex(sqlite3.IntegrityError, "(permanent|immutable)"):
            self.conn.execute(
                """UPDATE OR REPLACE contract_template_identities
                   SET id=?,template_type=? WHERE id=?""",
                (target["id"], target["template_type"], source_id),
            )
        self.conn.rollback()

    def test_insert_or_replace_cannot_replace_an_identity_by_template_type(self):
        """Catch the unique template-type conflict not covered by an id-only guard."""
        self.conn.execute(
            """INSERT INTO contract_template_identities(
               template_type,display_name,current_published_version_id,created_at,updated_at
               ) VALUES('VersionlessIdentity','Versionless',NULL,'2026-08-19T00:00:00+00:00','2026-08-19T00:00:00+00:00')"""
        )
        identity = self.conn.execute(
            "SELECT * FROM contract_template_identities WHERE template_type='VersionlessIdentity'"
        ).fetchone()
        self.assertEqual(0, self.conn.execute("PRAGMA recursive_triggers").fetchone()[0])

        with self.assertRaisesRegex(sqlite3.IntegrityError, "permanent"):
            self.conn.execute(
                """INSERT OR REPLACE INTO contract_template_identities(
                   id,template_type,display_name,current_published_version_id,created_at,updated_at
                   ) VALUES(?,?,?,NULL,?,?)""",
                (
                    identity["id"] + 1000,
                    identity["template_type"],
                    "Forged replacement",
                    "2026-08-19T00:00:00+00:00",
                    "2026-08-19T00:00:00+00:00",
                ),
            )
        self.conn.rollback()

    def test_deferred_foreign_keys_cannot_delete_and_reinsert_a_canonical_identity(self):
        """Catch DELETE+INSERT replacing canonical identity metadata under deferred FKs."""
        self._import_v1()
        identity = self.service._identity("NewInternationalStudentContract")
        self.conn.execute("PRAGMA defer_foreign_keys=ON")

        with self.assertRaisesRegex(sqlite3.IntegrityError, "permanent"):
            self.conn.execute("DELETE FROM contract_template_identities WHERE id=?", (identity["id"],))
            self.conn.execute(
                """INSERT INTO contract_template_identities(
                   id,template_type,display_name,current_published_version_id,created_at,updated_at
                   ) VALUES(?,?,?,NULL,?,?)""",
                (
                    identity["id"],
                    identity["template_type"],
                    "Forged identity",
                    "2026-08-19T00:00:00+00:00",
                    "2026-08-19T00:00:00+00:00",
                ),
            )
            self.conn.commit()
        self.conn.rollback()

    def test_direct_sql_cannot_swap_canonical_contract_template_types(self):
        """Catch the three-step unique-key swap that redirects type lookup to another snapshot."""
        self._import_v1()
        left = self.service._identity("NewInternationalStudentContract")
        right = self.service._identity("TransferApplicationContract")

        with self.assertRaisesRegex(sqlite3.IntegrityError, "immutable"):
            self.conn.execute(
                "UPDATE contract_template_identities SET template_type='TemporaryType' WHERE id=?",
                (left["id"],),
            )
            self.conn.execute(
                "UPDATE contract_template_identities SET template_type=? WHERE id=?",
                (left["template_type"], right["id"]),
            )
            self.conn.execute(
                "UPDATE contract_template_identities SET template_type=? WHERE id=?",
                (right["template_type"], left["id"]),
            )
        self.conn.rollback()

    def test_direct_sql_cannot_rewrite_canonical_identity_display_or_creation_provenance(self):
        """Catch identity metadata rewrite while preserving normal pointer/updated-at changes."""
        self._import_v1()
        identity = self.service._identity("LocalStudentApplicationContract")

        for assignment, attack_name in (
            ("display_name='Forged label'", "display name"),
            ("created_at='2099-01-01T00:00:00+00:00'", "creation timestamp"),
        ):
            with self.subTest(attack_name=attack_name), self.assertRaisesRegex(
                sqlite3.IntegrityError, "immutable"
            ):
                self.conn.execute(
                    f"UPDATE contract_template_identities SET {assignment} WHERE id=?",
                    (identity["id"],),
                )
            self.conn.rollback()

    def test_staff_permission_grants_are_admin_only_and_audited(self):
        self._import_v1()
        with self.assertRaises(editor.PermissionDenied):
            self.service.set_user_permission(9, can_edit=True, can_publish=True, actor_user_id=2, actor_is_admin=False)
        self.service.set_user_permission(9, can_edit=True, can_publish=True, actor_user_id=1, actor_is_admin=True)

        permission = self.service.permission_for(9, actor_is_admin=False)
        self.assertTrue(permission["can_edit"])
        self.assertTrue(permission["can_publish"])
        actions = [row["action"] for row in self.service.audit_events()]
        self.assertIn("template_permission_changed", actions)

    def test_permission_grant_uses_live_admin_role_and_rejects_forged_flag_without_writes(self):
        self.conn.executescript(
            """
            CREATE TABLE roles(id INTEGER PRIMARY KEY,name TEXT NOT NULL);
            CREATE TABLE users(id INTEGER PRIMARY KEY,role_id INTEGER NOT NULL);
            INSERT INTO roles(id,name) VALUES(1,'Admin'),(2,'Staff');
            INSERT INTO users(id,role_id) VALUES(1,1),(2,2);
            """
        )
        before_audits = len(self.service.audit_events())

        with self.assertRaises(editor.PermissionDenied):
            self.service.set_user_permission(
                9, can_edit=True, can_publish=True,
                actor_user_id=2, actor_is_admin=True,
            )

        self.assertIsNone(self.conn.execute(
            "SELECT * FROM contract_template_permissions WHERE user_id=9"
        ).fetchone())
        self.assertEqual(len(self.service.audit_events()), before_audits)

        self.service.set_user_permission(
            9, can_edit=True, can_publish=True,
            actor_user_id=1, actor_is_admin=False,
        )
        granted = self.conn.execute(
            "SELECT * FROM contract_template_permissions WHERE user_id=9"
        ).fetchone()
        self.assertEqual(granted["granted_by"], 1)
        self.assertTrue(granted["can_edit"])
        self.assertTrue(granted["can_publish"])
        self.assertEqual(len(self.service.audit_events()), before_audits + 1)

    def test_permission_grant_begins_immediate_before_live_role_check_and_prebegin_downgrade_wins(self):
        self.conn.executescript(
            """
            CREATE TABLE roles(id INTEGER PRIMARY KEY,name TEXT NOT NULL);
            CREATE TABLE users(id INTEGER PRIMARY KEY,role_id INTEGER NOT NULL);
            INSERT INTO roles(id,name) VALUES(1,'Admin'),(2,'Staff');
            INSERT INTO users(id,role_id) VALUES(1,1);
            """
        )
        original_transaction = self.service._immediate_transaction
        statements = []
        self.conn.set_trace_callback(statements.append)

        @contextmanager
        def downgrade_before_begin():
            self.conn.execute("UPDATE users SET role_id=2 WHERE id=1")
            self.conn.commit()
            with original_transaction():
                yield

        with mock.patch.object(
            self.service, "_immediate_transaction", downgrade_before_begin
        ):
            with self.assertRaises(editor.PermissionDenied):
                self.service.set_user_permission(
                    9, can_edit=True, can_publish=True,
                    actor_user_id=1, actor_is_admin=True,
                )
        self.conn.set_trace_callback(None)

        begin_index = next(
            index for index, statement in enumerate(statements)
            if statement.strip().upper() == "BEGIN IMMEDIATE"
        )
        role_index = next(
            index for index, statement in enumerate(statements)
            if "SELECT r.name FROM users" in statement
        )
        self.assertLess(begin_index, role_index)
        self.assertIsNone(self.conn.execute(
            "SELECT * FROM contract_template_permissions WHERE user_id=9"
        ).fetchone())
        self.assertEqual(self.service.audit_events(), [])

    def test_permission_grant_rolls_back_permission_when_audit_fails(self):
        self.conn.executescript(
            """
            CREATE TABLE roles(id INTEGER PRIMARY KEY,name TEXT NOT NULL);
            CREATE TABLE users(id INTEGER PRIMARY KEY,role_id INTEGER NOT NULL);
            INSERT INTO roles(id,name) VALUES(1,'Admin');
            INSERT INTO users(id,role_id) VALUES(1,1);
            """
        )

        with mock.patch.object(
            self.service, "_audit", side_effect=RuntimeError("audit write failed")
        ):
            with self.assertRaisesRegex(RuntimeError, "audit write failed"):
                self.service.set_user_permission(
                    9, can_edit=True, can_publish=True,
                    actor_user_id=1, actor_is_admin=False,
                )

        self.assertFalse(self.conn.in_transaction)
        self.assertIsNone(self.conn.execute(
            "SELECT * FROM contract_template_permissions WHERE user_id=9"
        ).fetchone())
        self.assertEqual(self.service.audit_events(), [])

    def test_version_diff_reports_block_service_and_commercial_changes(self):
        self._import_v1()
        initial = self.service.current_published("NewInternationalStudentContract")
        draft = self.service.create_draft("NewInternationalStudentContract", actor_user_id=1)
        lock = self.service.acquire_lock("NewInternationalStudentContract", draft["id"], owner_user_id=1)
        changed = self.service.autosave_block(
            draft["id"], "profile_3", title="三、服务费用与第三方费用", body="Changed text",
            enabled=True, expected_revision=draft["revision"], lock_token=lock["token"], actor_user_id=1,
        )

        diff = self.service.compare_versions(initial["id"], changed["id"])
        self.assertEqual(diff["blocks"][0]["key"], "profile_3")
        self.assertEqual(diff["blocks"][0]["change"], "modified")
        with self.assertRaisesRegex(editor.TemplateRuleError, "同一合同类型"):
            self.service.compare_versions(
                initial["id"], self.service.current_published("TransferApplicationContract")["id"]
            )

    def test_batch_simulation_runs_multiple_isolated_scenarios(self):
        self._import_v1()
        draft = self.service.create_draft("TransferApplicationContract", actor_user_id=1)
        lock = self.service.acquire_lock(
            "TransferApplicationContract", draft["id"], owner_user_id=1
        )
        scenarios = [
            {
                "source_type": "system", "student_snapshot": {
                    "student_name": f"Student {index}", "student_id": f"T-BATCH-{index}",
                    "date_of_birth": "2011-02-03", "current_year_level": "Year 7",
                    "applying_year_level": "Year 8", "preferred_intake_year": "2027",
                    "company_authorised_person": "Sonny Zhang",
                    "father_passport_first_name": "Father",
                    "father_passport_last_name": str(index), "father_mobile": "0400000000",
                    "father_email": f"father{index}@example.test",
                }, "signer_source": "father", "signer_overrides": {},
                "selected_services": list(contract_domain.default_service_ids("TransferApplicationContract")),
            }
            for index in (1, 2)
        ]
        results = self.service.simulate_batch(
            draft["id"], scenarios=scenarios, output_root=self.root / "previews", actor_user_id=1,
            expected_revision=draft["revision"], lock_token=lock["token"],
        )
        self.assertEqual(len(results), 2)
        self.assertTrue(all(result["status"] == "Passed" for result in results))

    def test_commercial_and_service_autosave_enforces_contract_type_catalog(self):
        self._import_v1()
        draft = self.service.create_draft("LocalStudentApplicationContract", actor_user_id=1)
        lock = self.service.acquire_lock("LocalStudentApplicationContract", draft["id"], owner_user_id=1)
        changed = self.service.autosave_settings(
            draft["id"], commercial={**self._valid_commercial(), "service_fee": "1800.00"},
            services=[
                {"service_id": "study_planning_consulting", "enabled": True, "default_selected": True},
                {"service_id": "student_visa_lodgement", "enabled": True, "default_selected": True},
            ], expected_revision=draft["revision"], lock_token=lock["token"], actor_user_id=1,
        )

        self.assertEqual(changed["revision"], draft["revision"] + 1)
        snapshot = json.loads(changed["snapshot_json"])
        self.assertEqual(snapshot["commercial"]["service_fee"], "1800.00")
        self.assertNotIn("student_visa_lodgement", {item["service_id"] for item in snapshot["services"]})

    def test_authority_v2_commercial_autosave_preserves_policy_and_representative(self):
        self._import_v1()
        draft = next(
            item
            for item in self.service.create_authority_v2_drafts(actor_user_id=1)
            if item["template_type"] == "NewInternationalStudentContract"
        )
        original = json.loads(draft["snapshot_json"])["commercial"]
        lock = self.service.acquire_lock(
            "NewInternationalStudentContract", draft["id"], owner_user_id=1
        )
        edited = {
            key: original[key]
            for key in (
                "currency", "service_fee", "company_name", "company_abn",
                "company_address", "company_contact",
            )
        }
        edited.update({
            "payment_terms": "CRM_TEST CONFIG BINDING PAYMENT TERMS",
            "refund_policy": "CRM_TEST CONFIG BINDING REFUND WORDING",
            "company_authorised_person": "Frozen Representative",
        })
        changed = self.service.autosave_settings(
            draft["id"], commercial=edited,
            services=[
                {
                    "service_id": service_id,
                    "enabled": True,
                    "default_selected": True,
                }
                for service_id in contract_domain.default_service_ids(
                    "NewInternationalStudentContract"
                )
            ],
            expected_revision=draft["revision"], lock_token=lock["token"],
            actor_user_id=1,
        )

        commercial = json.loads(changed["snapshot_json"])["commercial"]
        self.assertEqual(commercial.get("commercial_schema"), "authority_v2")
        self.assertEqual(commercial.get("standard_service_fee_ex_gst"), "2400.00")
        self.assertEqual(commercial.get("gst_rate"), "0.10")
        self.assertEqual(commercial.get("visa_refusal_retained_ex_gst"), "900.00")
        self.assertEqual(
            commercial["payment_terms"],
            "CRM_TEST CONFIG BINDING PAYMENT TERMS",
        )
        self.assertEqual(
            commercial["refund_policy"],
            "CRM_TEST CONFIG BINDING REFUND WORDING",
        )
        self.assertEqual(
            commercial["company_authorised_person"], "Frozen Representative"
        )

    def test_editor_view_is_three_template_responsive_single_logic(self):
        self._import_v1()
        html = editor.render_template_editor_index(
            self.service.list_templates(), csrf="token", can_edit=True, can_publish=True, is_admin=True
        )

        self.assertIn("新留学生合同", html)
        self.assertIn("转学申请合同", html)
        self.assertIn("本地生申请合同", html)
        self.assertIn('data-contract-template-editor="v1"', html)
        self.assertIn("当前发布版本", html)
        self.assertIn("历史版本", html)
        self.assertIn("合同模板编辑器 V1", html)
        self.assertNotIn("block_0", html)
        self.assertIn("合同模板权限", html)

        workspace = self.service.workspace("NewInternationalStudentContract")
        version_id = workspace["active_version"]["id"]
        detail = editor.render_template_editor_detail(
            workspace, csrf="token", can_edit=True, can_publish=True,
            comparison={"left_id": version_id, "right_id": version_id, "blocks": [], "services_changed": False, "commercial_changed": False},
        )
        self.assertIn('data-template-action="preview"', detail)
        self.assertIn("比较版本", detail)
        self.assertIn("重新发布为新版本", detail)

    def test_template_management_panel_uses_v1_service_after_migration_and_import(self):
        import app as app_module
        from database import Database

        self._import_v1()
        original_db = app_module.db
        app_module.db = Database(self.root / "editor.sqlite3")
        self.addCleanup(setattr, app_module, "db", original_db)
        handler = app_module.CRMHandler.__new__(app_module.CRMHandler)
        handler.user = {"id": 1, "role_name": "Admin"}
        handler.csrf = "csrf-token"

        panel = handler.contract_template_editor_v1_panel()

        self.assertIn("新留学生合同", panel)
        self.assertIn("转学申请合同", panel)
        self.assertIn("本地生申请合同", panel)
        self.assertIn("/templates/v1/NewInternationalStudentContract", panel)

    def test_template_management_create_and_edit_capabilities_render_simple_edit_action(self):
        """Catch module create/edit authority being ignored by the V1 permission adapter."""
        import app as app_module

        self._import_v1()
        handler = app_module.CRMHandler.__new__(app_module.CRMHandler)
        handler.user = {"id": 9, "role_name": "Manager + Student Cases"}

        with mock.patch.object(
            app_module,
            "can_module_action",
            side_effect=lambda _user, module, action: (
                module == "template_management" and action in {"create", "edit"}
            ),
        ):
            permission, is_admin = handler.contract_template_editor_v1_permissions(
                self.service
            )

        rendered = editor.render_template_editor_index(
            self.service.list_templates(),
            csrf="tong-yu-csrf",
            can_edit=permission["can_edit"],
            can_publish=permission["can_publish"],
            is_admin=is_admin,
        )
        self.assertFalse(is_admin)
        self.assertTrue(permission["can_edit"])
        self.assertFalse(permission["can_publish"])
        self.assertEqual(rendered.count(">编辑模板</button>"), 3)
        self.assertNotIn(">新建草稿</button>", rendered)
        self.assertNotIn(">新建锁定主版式草稿</button>", rendered)
        self.assertIn("可编辑，应用需授权", rendered)

    def test_template_manager_ui_exposes_edit_save_apply_without_advanced_version_actions(self):
        """Ordinary template managers should not need version-management vocabulary."""
        self._import_v1()
        rendered = editor.render_template_editor_index(
            self.service.list_templates(),
            csrf="tong-yu-csrf",
            can_edit=True,
            can_publish=True,
            is_admin=False,
        )
        self.assertEqual(rendered.count(">编辑模板</button>"), 3)
        self.assertNotIn(">新建草稿</button>", rendered)
        self.assertNotIn(">新建锁定主版式草稿</button>", rendered)

        draft = self.service.create_draft(
            "TransferApplicationContract", actor_user_id=2
        )
        workspace = self.service.workspace(
            "TransferApplicationContract", draft["id"]
        )
        detail = editor.render_template_editor_detail(
            workspace,
            csrf="tong-yu-csrf",
            can_edit=True,
            can_publish=True,
            is_admin=False,
        )
        self.assertIn('data-template-action="apply"', detail)
        self.assertIn(">应用模板</button>", detail)
        self.assertIn("自动保存：", detail)
        self.assertIn("已保存", detail)
        self.assertIn("应用状态：", detail)
        self.assertIn('data-template-action-message', detail)
        self.assertNotIn('data-template-action="validate"', detail)
        self.assertNotIn('data-template-action="preview"', detail)
        self.assertNotIn('data-template-action="publish"', detail)
        self.assertNotIn("版本差异与历史版本", detail)

    def test_admin_ui_retains_locked_master_and_advanced_publication_controls(self):
        self._import_v1()
        rendered = editor.render_template_editor_index(
            self.service.list_templates(),
            csrf="admin-csrf",
            can_edit=True,
            can_publish=True,
            is_admin=True,
        )
        self.assertEqual(rendered.count(">新建锁定主版式草稿</button>"), 3)
        workspace = self.service.workspace("TransferApplicationContract")
        detail = editor.render_template_editor_detail(
            workspace,
            csrf="admin-csrf",
            can_edit=True,
            can_publish=True,
            is_admin=True,
        )
        self.assertIn('data-template-action="validate"', detail)
        self.assertIn('data-template-action="preview"', detail)
        self.assertIn('data-template-action="publish"', detail)
        self.assertIn("版本差异与历史版本", detail)

    def test_template_manager_apply_runs_all_gates_and_updates_current_pointer(self):
        self._import_v1()
        self.service.set_user_permission(
            2,
            can_edit=False,
            can_publish=True,
            actor_user_id=1,
            actor_is_admin=True,
        )
        previous = self.service.current_published("TransferApplicationContract")
        draft = self.service.create_draft(
            "TransferApplicationContract", actor_user_id=2
        )
        lock = self.service.acquire_lock(
            "TransferApplicationContract", draft["id"], owner_user_id=2
        )
        draft = self.service.autosave_settings(
            draft["id"],
            commercial=self._valid_commercial(),
            services=[
                {
                    "service_id": service_id,
                    "enabled": True,
                    "default_selected": True,
                }
                for service_id in contract_domain.default_service_ids(
                    "TransferApplicationContract"
                )
            ],
            expected_revision=draft["revision"],
            lock_token=lock["token"],
            actor_user_id=2,
        )

        applied = self.service.apply_template(
            draft["id"],
            actor_user_id=2,
            lock_token=lock["token"],
            expected_revision=draft["revision"],
            output_root=self.root / "apply-previews",
            actor_is_admin=False,
        )

        self.assertEqual(applied["status"], "Published")
        self.assertEqual(applied["id"], draft["id"])
        self.assertNotEqual(applied["id"], previous["id"])
        self.assertEqual(
            self.service.current_published("TransferApplicationContract")["id"],
            draft["id"],
        )
        self.assertTrue(self.service.permission_for(2, actor_is_admin=False)["can_publish"])
        actions = [
            row[0]
            for row in self.conn.execute(
                "SELECT action FROM contract_template_audit_events "
                "WHERE version_id=? ORDER BY id",
                (draft["id"],),
            )
        ]
        expected = [
            "template_simulation_passed",
            "template_validation_passed",
            "template_preview_confirmed",
            "template_published",
        ]
        positions = [actions.index(action) for action in expected]
        self.assertEqual(positions, sorted(positions))

    def test_template_manager_apply_reclaims_own_expired_lock_after_final_autosave(self):
        """Catch Apply rejecting a clean draft solely because the page was suspended."""
        self._import_v1()
        self.service.set_user_permission(
            2, can_edit=False, can_publish=True,
            actor_user_id=1, actor_is_admin=True,
        )
        draft = self.service.create_draft(
            "TransferApplicationContract", actor_user_id=2
        )
        lock = self.service.acquire_lock(
            "TransferApplicationContract", draft["id"], owner_user_id=2
        )
        draft = self.service.autosave_settings(
            draft["id"], commercial=self._valid_commercial(),
            services=[
                {
                    "service_id": service_id,
                    "enabled": True,
                    "default_selected": True,
                }
                for service_id in contract_domain.default_service_ids(
                    "TransferApplicationContract"
                )
            ],
            expected_revision=draft["revision"], lock_token=lock["token"],
            actor_user_id=2,
        )
        self.conn.execute(
            "UPDATE contract_template_edit_locks SET expires_at=? WHERE version_id=?",
            ("2000-01-01T00:00:00+00:00", draft["id"]),
        )
        self.conn.commit()

        applied = self.service.apply_template(
            draft["id"], actor_user_id=2, lock_token=lock["token"],
            expected_revision=draft["revision"],
            output_root=self.root / "expired-apply-previews",
            actor_is_admin=False,
        )

        self.assertEqual(applied["status"], "Published")
        self.assertEqual(
            self.service.current_published("TransferApplicationContract")["id"],
            draft["id"],
        )

    def test_apply_route_requires_edit_and_publish_and_never_requires_admin(self):
        import app as app_module

        class RecordingService:
            def __init__(self):
                self.calls = []

            def apply_template(self, *args, **kwargs):
                self.calls.append((args, kwargs))
                return {"id": 41, "status": "Published"}

        def invoke(permission):
            service = RecordingService()
            handler = app_module.CRMHandler.__new__(app_module.CRMHandler)
            handler.user = {"id": 2, "role_name": "Manager + Student Cases"}
            handler.form = {
                "version_id": "41",
                "revision": "7",
                "lock_token": "manager-lock",
            }
            handler.headers = {"X-Requested-With": "fetch"}
            handler.contract_template_editor_v1_service = lambda: service
            handler.contract_template_editor_v1_permissions = lambda _service: (
                permission,
                False,
            )
            captured = {}
            handler.json_response = lambda payload, status=200: captured.update(
                payload=payload, status=int(status)
            )
            with mock.patch.object(app_module, "can_access_module", return_value=True):
                handler.contract_template_editor_v1_action(
                    "POST", "TransferApplicationContract", "apply"
                )
            return service, captured

        service, captured = invoke(
            {"can_view": True, "can_edit": True, "can_publish": True}
        )
        self.assertEqual(captured["status"], 200)
        self.assertEqual(len(service.calls), 1)
        self.assertFalse(service.calls[0][1]["actor_is_admin"])

        for permission in (
            {"can_view": True, "can_edit": False, "can_publish": True},
            {"can_view": True, "can_edit": True, "can_publish": False},
            {"can_view": True, "can_edit": False, "can_publish": False},
        ):
            with self.subTest(permission=permission):
                service, captured = invoke(permission)
                self.assertEqual(captured["status"], 403)
                self.assertEqual(service.calls, [])

    def test_lock_and_revision_conflicts_have_distinct_http_error_codes(self):
        import app as app_module

        class ConflictService:
            def __init__(self, error):
                self.error = error

            def acquire_lock(self, *_args, **_kwargs):
                raise self.error

            def autosave_block(self, *_args, **_kwargs):
                raise self.error

        def invoke(action, error):
            handler = app_module.CRMHandler.__new__(app_module.CRMHandler)
            handler.user = {"id": 2, "role_name": "Manager + Student Cases"}
            handler.form = {
                "version_id": "41", "revision": "7", "lock_token": "old-token",
                "block_key": "profile_3", "title": "Title", "body": "Body",
                "enabled": "1",
            }
            handler.headers = {"X-Requested-With": "fetch"}
            handler.contract_template_editor_v1_service = lambda: ConflictService(error)
            handler.contract_template_editor_v1_permissions = lambda _service: (
                {"can_view": True, "can_edit": True, "can_publish": True}, False,
            )
            captured = {}
            handler.json_response = lambda payload, status=200: captured.update(
                payload=payload, status=int(status)
            )
            with mock.patch.object(app_module, "can_access_module", return_value=True):
                handler.contract_template_editor_v1_action(
                    "POST", "TransferApplicationContract", action
                )
            return captured

        lock_conflict = invoke(
            "lock",
            editor.EditLockConflict(
                "此模板正在被另一位用户编辑，你当前修改尚未应用。"
            ),
        )
        self.assertEqual(lock_conflict["status"], 409)
        self.assertEqual(
            lock_conflict["payload"]["error_code"], "edit_lock_conflict"
        )
        self.assertIn("另一位用户", lock_conflict["payload"]["error"])

        stale = invoke(
            "autosave-block",
            editor.StaleRevision("模板版本已变化，请刷新后重试。"),
        )
        self.assertEqual(stale["status"], 409)
        self.assertEqual(stale["payload"]["error_code"], "stale_revision")

    def test_template_management_create_and_edit_capabilities_authorize_both_draft_routes(self):
        """Catch UI-authorized draft actions returning 403 at the HTTP permission gate."""
        import app as app_module

        self._import_v1()
        cases = (
            ("draft", "TransferApplicationContract"),
            ("master-draft", "LocalStudentApplicationContract"),
        )
        for action, template_type in cases:
            with self.subTest(action=action, template_type=template_type):
                handler = app_module.CRMHandler.__new__(app_module.CRMHandler)
                handler.user = {"id": 9, "role_name": "Manager + Student Cases"}
                handler.form = {}
                handler.headers = {"X-Requested-With": "fetch"}
                handler.contract_template_editor_v1_service = lambda: self.service
                captured = {}
                handler.json_response = lambda payload, status=200: captured.update(
                    payload=payload, status=int(status)
                )

                with mock.patch.object(
                    app_module, "can_access_module", return_value=True
                ), mock.patch.object(
                    app_module,
                    "can_module_action",
                    side_effect=lambda _user, module, module_action: (
                        module == "template_management"
                        and module_action in {"create", "edit"}
                    ),
                ):
                    handler.contract_template_editor_v1_action(
                        "POST", template_type, action
                    )

                self.assertEqual(captured["status"], 200)
                self.assertTrue(captured["payload"]["ok"])
                created = self.service.workspace(template_type)["active_version"]
                self.assertEqual(created["status"], "Draft")

    def test_template_management_without_create_or_edit_capability_stays_read_only(self):
        """Catch module access alone accidentally expanding staff template authority."""
        import app as app_module

        self._import_v1()
        handler = app_module.CRMHandler.__new__(app_module.CRMHandler)
        handler.user = {"id": 10, "role_name": "View Only"}
        with mock.patch.object(app_module, "can_module_action", return_value=False):
            permission, is_admin = handler.contract_template_editor_v1_permissions(
                self.service
            )

        rendered = editor.render_template_editor_index(
            self.service.list_templates(),
            csrf="read-only-csrf",
            can_edit=permission["can_edit"],
            can_publish=permission["can_publish"],
            is_admin=is_admin,
        )
        self.assertFalse(permission["can_edit"])
        self.assertFalse(permission["can_publish"])
        self.assertNotIn(">新建草稿</button>", rendered)
        self.assertNotIn(">新建锁定主版式草稿</button>", rendered)
        self.assertIn("只读", rendered)

    def test_validate_route_rejects_without_calling_service_when_actor_cannot_publish(self):
        import app as app_module

        class RecordingService:
            def __init__(self):
                self.validation_calls = []

            def validate_for_publish(self, *args, **kwargs):
                self.validation_calls.append((args, kwargs))
                return {"status": "Passed"}

        service = RecordingService()
        handler = app_module.CRMHandler.__new__(app_module.CRMHandler)
        handler.user = {"id": 9, "role_name": "Staff"}
        handler.form = {"version_id": "41", "lock_token": "forged-token"}
        handler.headers = {"X-Requested-With": "fetch"}
        handler.contract_template_editor_v1_service = lambda: service
        handler.contract_template_editor_v1_permissions = lambda _service: (
            {"can_view": True, "can_edit": False, "can_publish": False},
            False,
        )
        captured = {}
        handler.json_response = lambda payload, status=200: captured.update(
            payload=payload, status=int(status)
        )

        with mock.patch.object(app_module, "can_access_module", return_value=True):
            handler.contract_template_editor_v1_action(
                "POST", "TransferApplicationContract", "validate"
            )

        self.assertEqual(captured["status"], 403)
        self.assertFalse(captured["payload"]["ok"])
        self.assertEqual(service.validation_calls, [])

    def test_permission_route_rejects_forged_admin_from_live_role_without_writes(self):
        import app as app_module

        self.conn.executescript(
            """
            CREATE TABLE roles(id INTEGER PRIMARY KEY,name TEXT NOT NULL);
            CREATE TABLE users(id INTEGER PRIMARY KEY,role_id INTEGER NOT NULL);
            INSERT INTO roles(id,name) VALUES(2,'Staff');
            INSERT INTO users(id,role_id) VALUES(2,2);
            """
        )
        before_audits = len(self.service.audit_events())
        handler = app_module.CRMHandler.__new__(app_module.CRMHandler)
        handler.user = {"id": 2, "role_name": "Admin"}
        handler.form = {"user_id": "9", "can_edit": "1", "can_publish": "1"}
        handler.headers = {"X-Requested-With": "fetch"}
        handler.contract_template_editor_v1_service = lambda: self.service
        handler.contract_template_editor_v1_permissions = lambda _service: (
            {"can_view": True, "can_edit": True, "can_publish": True},
            True,
        )
        captured = {}
        handler.json_response = lambda payload, status=200: captured.update(
            payload=payload, status=int(status)
        )

        with mock.patch.object(app_module, "can_access_module", return_value=True):
            handler.contract_template_editor_v1_action(
                "POST", "TransferApplicationContract", "permission"
            )

        self.assertEqual(captured["status"], 403)
        self.assertFalse(captured["payload"]["ok"])
        self.assertIsNone(self.conn.execute(
            "SELECT * FROM contract_template_permissions WHERE user_id=9"
        ).fetchone())
        self.assertEqual(len(self.service.audit_events()), before_audits)

    def test_permission_route_allows_live_admin_despite_stale_route_flag(self):
        import app as app_module

        self.conn.executescript(
            """
            CREATE TABLE roles(id INTEGER PRIMARY KEY,name TEXT NOT NULL);
            CREATE TABLE users(id INTEGER PRIMARY KEY,role_id INTEGER NOT NULL);
            INSERT INTO roles(id,name) VALUES(1,'Admin');
            INSERT INTO users(id,role_id) VALUES(1,1);
            """
        )
        handler = app_module.CRMHandler.__new__(app_module.CRMHandler)
        handler.user = {"id": 1, "role_name": "Staff"}
        handler.form = {"user_id": "9", "can_edit": "1", "can_publish": "1"}
        handler.headers = {"X-Requested-With": "fetch"}
        handler.contract_template_editor_v1_service = lambda: self.service
        handler.contract_template_editor_v1_permissions = lambda _service: (
            {"can_view": True, "can_edit": False, "can_publish": False},
            False,
        )
        captured = {}
        handler.json_response = lambda payload, status=200: captured.update(
            payload=payload, status=int(status)
        )

        with mock.patch.object(app_module, "can_access_module", return_value=True):
            handler.contract_template_editor_v1_action(
                "POST", "TransferApplicationContract", "permission"
            )

        self.assertEqual(captured["status"], 200)
        self.assertTrue(captured["payload"]["ok"])
        granted = self.conn.execute(
            "SELECT * FROM contract_template_permissions WHERE user_id=9"
        ).fetchone()
        self.assertEqual(granted["granted_by"], 1)
        self.assertTrue(granted["can_edit"])
        self.assertTrue(granted["can_publish"])


if __name__ == "__main__":
    unittest.main()
