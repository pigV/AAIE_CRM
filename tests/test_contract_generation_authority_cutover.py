import copy
import hashlib
import io
import json
import sqlite3
import tempfile
import unittest
from pathlib import Path
from urllib.parse import parse_qs, urlparse

from docx import Document
from pypdf import PdfReader

import contract_domain
import contract_template_editor as editor
from contract_generation_cutover import (
    ContractGenerationPersistenceError,
    ContractGenerationRequestConflict,
    create_contract_atomic,
    replace_draft_atomic,
    set_generation_authority_atomic,
)
from database import Database

try:
    import contract_generation_authority as authority
except ModuleNotFoundError:
    authority = None


ROOT = Path(__file__).resolve().parents[1]
MIGRATION = ROOT / "migrations" / "20260818_025_contract_template_editor_v1.sql"


LEGACY_V1_INTERNATIONAL_PROFILE = {
    "三、服务费用与第三方费用": (
        "本合同服务费对应澳洲留学咨询、学校申请、材料整理、流程协调及签证相关咨询协助。\n"
        "学校、保险机构、住宿机构、移民局或其他第三方收取的申请费、学费、OSHC、住宿费、体检费、签证费及其他费用由乙方另行承担。\n"
        "第三方是否受理、录取、签发或退费，以该第三方的政策、标准及最终决定为准。"
    ),
    "四、甲方权利义务": (
        "甲方按照本合同服务范围提供澳洲教育制度、学校申请、材料准备、签证制度及相关流程的咨询与协调，并对乙方资料承担保密义务。"
    ),
    "五、乙方权利义务": (
        "乙方应保证资料真实、完整、有效，及时配合学校申请、保险、签证、付款、面试及补充材料，并及时通知联系方式或重要情况的变化。"
    ),
    "六、退费及违约责任": (
        "乙方确认，甲方已实际提供的留学规划、学校申请、材料整理、签证指导及流程协调服务部分不予退还。\n"
        "学校、保险机构、住宿机构、移民局或其他第三方费用按相应第三方政策处理。\n"
        "如乙方提供虚假材料、隐瞒重要事实或因自身原因终止服务，应承担由此产生的后果。"
    ),
    "七、不可抗力、结果说明及免责": (
        "学校录取、教育部门审批、保险安排、住宿安排、签证结果及其他第三方事项由相关第三方独立决定，甲方不保证录取、签证或其他批准结果。\n"
        "因不可抗力、政策变化、系统故障或第三方审理延误导致服务变化的，双方应友好协商。"
    ),
    "九、合同补充、变更和生效": (
        "本合同的补充、变更或修改应采用书面形式，经双方确认后与本合同具有同等法律效力。\n"
        "未满十八周岁的学生应由法定代理人或监护人签署或确认。\n"
        "本合同可采用纸质、扫描件或电子签名方式签署，自双方签字或盖章之日起生效。"
    ),
}


def frozen_legacy_v1_snapshot(source):
    """Recreate the exact pre-master V1 snapshot used by historical hashes."""
    blocks = [
        {
            "key": "fixed_parties",
            "title": "合同双方",
            "body": "学生：{{student_name}}\n签署人：{{signer_name}}\n电话：{{signer_phone}}\n邮箱：{{signer_email}}",
            "enabled": True,
            "fixed_layout": True,
            "required": True,
            "legal_kind": "identity",
            "required_variables": [
                "{{student_name}}", "{{signer_name}}", "{{signer_phone}}", "{{signer_email}}",
            ],
        },
        {
            "key": "fixed_service_scope",
            "title": "二、服务项目及服务范围",
            "body": "{{service_scope}}",
            "enabled": True,
            "fixed_layout": True,
            "required": True,
            "legal_kind": "service_scope",
            "required_variables": ["{{service_scope}}"],
        },
    ]
    for title, body in LEGACY_V1_INTERNATIONAL_PROFILE.items():
        blocks.append({
            "key": editor.PROFILE_KEYS[title],
            "title": title,
            "body": body,
            "enabled": True,
            "fixed_layout": False,
            "required": True,
            "legal_kind": {
                "六、退费及违约责任": "refund",
                "九、合同补充、变更和生效": "termination",
            }.get(title, ""),
            "required_variables": [],
        })
    blocks.extend([
        {
            "key": "legal_privacy",
            "title": "十、隐私与资料保护",
            "body": "双方应依法保护在服务过程中取得的个人信息和文件。",
            "enabled": True,
            "fixed_layout": False,
            "required": True,
            "legal_kind": "privacy",
            "required_variables": [],
        },
        {
            "key": "legal_dispute",
            "title": "八、适用法律及争议解决",
            "body": "争议应先友好协商；协商不成时按适用法律处理。",
            "enabled": True,
            "fixed_layout": False,
            "required": True,
            "legal_kind": "dispute",
            "required_variables": [],
        },
    ])
    return {
        "schema_version": 1,
        "template_type": source["template_type"],
        "display_name": source["display_name"],
        "contract_title": source["contract_title"],
        "clause_profile": source["clause_profile"],
        "blocks": blocks,
        "services": copy.deepcopy(source["services"]),
        "commercial": copy.deepcopy(source["commercial"]),
    }


class ContractGenerationAuthorityCutoverTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.conn = sqlite3.connect(self.root / "authority.sqlite3")
        self.conn.row_factory = sqlite3.Row
        self.addCleanup(self.conn.close)
        self.conn.execute("PRAGMA foreign_keys=ON")
        self.conn.executescript(MIGRATION.read_text(encoding="utf-8"))
        self.conn.execute("CREATE TABLE system_settings(key TEXT PRIMARY KEY,value TEXT NOT NULL)")
        self.conn.executemany(
            "INSERT INTO system_settings(key,value) VALUES(?,?)",
            [
                ("company_legal_name", "AAIE Test Pty Ltd"),
                ("company_abn_acn", "12 345 678 901"),
                ("company_address", "Melbourne VIC"),
                ("company_phone", "+61 3 0000 0000"),
                ("company_email", "contracts@example.test"),
                ("company_authorised_person", "Sonny Zhang"),
            ],
        )
        self.service = editor.ContractTemplateEditorService(self.conn)
        self.service.import_domain_v1(actor_user_id=1)

    @staticmethod
    def _student():
        return {
            "student_name": "Frozen Student",
            "student_id": "T-AUTH-001",
            "date_of_birth": "2010-01-01",
            "current_year_level": "Year 7",
            "applying_year_level": "Year 8",
            "preferred_intake_year": "2027",
            "preferred_intake_term": "02",
            "father_passport_first_name": "Father",
            "father_passport_last_name": "Only",
            "father_mobile": "0411111111",
            "father_email": "father@example.test",
            "mother_passport_first_name": "Mother",
            "mother_passport_last_name": "Other",
            "mother_mobile": "0422222222",
            "mother_email": "mother@example.test",
        }

    @staticmethod
    def _payload(contract_type="NewInternationalStudentContract"):
        config = contract_domain.contract_type_config(contract_type)
        services = list(config.default_services)
        return {
            "contract_type": contract_type,
            "contract_uid": "CONTRACT-AUTH-001",
            "contract_date": "2026-08-20",
            "generation_request_id": "req-authority-001",
            "student_full_name": "Frozen Student",
            "student_id_or_passport_number": "T-AUTH-001",
            "student_date_of_birth": "2010-01-01",
            "student_snapshot": {
                "student_full_name": "Frozen Student",
                "student_id": "T-AUTH-001",
                "date_of_birth": "2010-01-01",
                "current_year_level": "Year 7",
                "applying_year_level": "Year 8",
                "preferred_intake_year": "2027",
                "preferred_intake_term": "02",
            },
            "company_authorised_person": "Sonny Zhang",
            "service_scope_ids": services,
            "service_scope_json": json.dumps(services),
            "payment_method": "国内银行",
            "signer": {
                "source": "father",
                "name": "Father Only",
                "relationship": "父亲",
                "phone": "0411111111",
                "email": "father@example.test",
            },
            # These hostile legacy fields must never become fee authority.
            "consulting_and_visa_service_fee": "9999.00",
            "discount_amount": "8888.00",
            "final_fee_amount": "1111.00",
            "service_fee_currency": "USD",
        }

    def _drafts(self):
        return self.service.create_authority_v2_drafts(actor_user_id=1)

    def _published_v2(self, contract_type="NewInternationalStudentContract"):
        current = self.service.current_published(contract_type)
        if (
            json.loads(current["snapshot_json"]).get("commercial", {}).get(
                "commercial_schema"
            )
            == "authority_v2"
        ):
            return current
        draft = next(item for item in self._drafts() if item["template_type"] == contract_type)
        lock = self.service.acquire_lock(contract_type, draft["id"], owner_user_id=1)
        draft = self.service.configure_authority_v2_draft(
            draft["id"],
            gst_treatment="EXCLUSIVE",
            expected_revision=draft["revision"],
            lock_token=lock["token"],
            actor_user_id=1,
        )
        selected = list(contract_domain.contract_type_config(contract_type).default_services)
        self.service.simulate(
            draft["id"],
            source_type="system",
            student_snapshot=self._student(),
            signer_source="father",
            signer_overrides={},
            selected_services=selected,
            output_root=self.root / "simulation",
            actor_user_id=1,
            expected_revision=draft["revision"],
            lock_token=lock["token"],
        )
        return self.service.publish(
            draft["id"], actor_user_id=1, lock_token=lock["token"],
            expected_revision=draft["revision"], actor_is_admin=True,
        )

    def test_v2_conversion_creates_three_idempotent_drafts_without_mutating_v1(self):
        before = {
            row["id"]: (row["content_hash"], row["snapshot_json"], row["status"], row["immutable"])
            for row in self.conn.execute("SELECT * FROM contract_template_versions")
        }
        first = self._drafts()
        second = self._drafts()

        self.assertEqual(len(first), 3)
        self.assertEqual([row["id"] for row in first], [row["id"] for row in second])
        self.assertEqual(
            before,
            {
                row["id"]: (row["content_hash"], row["snapshot_json"], row["status"], row["immutable"])
                for row in self.conn.execute(
                    "SELECT * FROM contract_template_versions WHERE version_number=1"
                )
            },
        )
        for row in first:
            snapshot = json.loads(row["snapshot_json"])
            commercial = snapshot["commercial"]
            self.assertEqual(row["status"], "Draft")
            self.assertEqual(row["version_number"], 2)
            self.assertEqual(commercial["commercial_schema"], "authority_v2")
            self.assertEqual(commercial["service_fee"], "2400.00")
            self.assertEqual(commercial["standard_service_fee_ex_gst"], "2400.00")
            self.assertEqual(commercial["currency"], "AUD")
            self.assertEqual(commercial["gst_treatment"], "EXCLUSIVE")
            self.assertEqual(commercial["gst_rate"], "0.10")
            self.assertEqual(commercial["default_gst_amount"], "240.00")
            self.assertEqual(commercial["default_final_total"], "2640.00")
            self.assertEqual(commercial["base_fee_override_allowed"], "false")

    def test_missing_or_unknown_gst_blocks_v2_publication_validation(self):
        draft = self._drafts()[0]
        snapshot = json.loads(draft["snapshot_json"])
        for value, expected in (("", "gst_treatment_required"), ("UNKNOWN", "gst_treatment_invalid")):
            snapshot["commercial"]["gst_treatment"] = value
            issues = editor.authority_v2_commercial_issues(snapshot["commercial"])
            self.assertIn(expected, {item["code"] for item in issues})

    def test_v1_placeholder_and_zero_fee_are_not_valid_generation_authorities(self):
        self.assertIsNotNone(authority, "Published generation authority module is missing")
        with self.assertRaises(authority.PublishedTemplateIntegrityError):
            authority.load_published_template(self.conn, "NewInternationalStudentContract")

    def test_published_v2_owns_standard_fee_and_rejects_post_fee_authority(self):
        self._published_v2()
        package = authority.load_published_template(self.conn, "NewInternationalStudentContract")
        bound = authority.bind_new_contract(
            package, self._payload(),
            actor={"id": 2, "role_name": "Manager"},
            generated_at="2026-08-20T06:00:00+00:00",
        )

        self.assertEqual(bound["commercial"]["standard_service_fee"], "2400.00")
        self.assertEqual(bound["commercial"]["net_service_fee_ex_gst"], "2400.00")
        self.assertEqual(bound["commercial"]["discount_amount"], "0.00")
        self.assertEqual(bound["commercial"]["gst_treatment"], "NOT_APPLICABLE")
        self.assertEqual(bound["commercial"]["gst_amount"], "0.00")
        self.assertEqual(bound["commercial"]["final_fee_amount"], "2400.00")
        self.assertEqual(bound["commercial"]["final_total"], "2400.00")
        self.assertEqual(bound["commercial"]["payment_method"], "国内银行")
        self.assertEqual(bound["commercial"]["payment_plan"]["mode"], "single_payment")

    def test_payment_method_and_discount_define_exact_gst_totals(self):
        self._published_v2()
        package = authority.load_published_template(self.conn, "NewInternationalStudentContract")
        cases = (
            ("微信", None, "2400.00", "0.00", "2400.00", "NOT_APPLICABLE"),
            ("支付宝", None, "2400.00", "0.00", "2400.00", "NOT_APPLICABLE"),
            ("国内银行", None, "2400.00", "0.00", "2400.00", "NOT_APPLICABLE"),
            ("澳洲私户", None, "2400.00", "0.00", "2400.00", "NOT_APPLICABLE"),
            ("澳洲公户", None, "2400.00", "240.00", "2640.00", "EXCLUSIVE"),
            ("微信", "100.00", "2300.00", "0.00", "2300.00", "NOT_APPLICABLE"),
            ("澳洲公户", "100.00", "2300.00", "230.00", "2530.00", "EXCLUSIVE"),
        )
        for payment_method, discount_amount, net, gst, total, treatment in cases:
            with self.subTest(payment_method=payment_method, discount=discount_amount):
                payload = self._payload()
                payload["payment_method"] = payment_method
                approval = None
                if discount_amount:
                    approval = {
                        "discount_amount": discount_amount,
                        "reason": "Approved exact discount",
                        "approved_by": 1,
                        "operator_id": 1,
                        "approved_at": "2026-08-20T06:10:00+00:00",
                    }
                bound = authority.bind_new_contract(
                    package, payload, actor={"id": 1, "role_name": "Admin"},
                    discount_approval=approval,
                )
                self.assertEqual(bound["commercial"]["net_service_fee_ex_gst"], net)
                self.assertEqual(bound["commercial"]["gst_amount"], gst)
                self.assertEqual(bound["commercial"]["final_total"], total)
                self.assertEqual(bound["commercial"]["gst_treatment"], treatment)

    def test_visa_refusal_refund_uses_authoritative_final_total_for_four_payment_cases(self):
        self._published_v2()
        package = authority.load_published_template(
            self.conn, "NewInternationalStudentContract"
        )
        cases = (
            ("微信", None, "2400.00", "900.00", "1500.00"),
            ("微信", "100.00", "2300.00", "900.00", "1400.00"),
            ("澳洲公户", None, "2640.00", "990.00", "1650.00"),
            ("澳洲公户", "100.00", "2530.00", "990.00", "1540.00"),
        )
        for payment_method, discount_amount, paid_total, retained, refund in cases:
            with self.subTest(
                payment_method=payment_method, discount_amount=discount_amount
            ):
                payload = self._payload()
                payload["payment_method"] = payment_method
                approval = None
                if discount_amount:
                    approval = {
                        "discount_amount": discount_amount,
                        "reason": "Approved exact discount",
                        "approved_by": 1,
                        "operator_id": 1,
                        "approved_at": "2026-08-22T07:00:00+00:00",
                    }
                bound = authority.bind_new_contract(
                    package,
                    payload,
                    actor={"id": 1, "role_name": "Admin"},
                    discount_approval=approval,
                )
                result = bound["commercial"]["visa_refusal_refund"]
                self.assertEqual(result["actual_paid_total"], paid_total)
                self.assertEqual(result["retained_total"], retained)
                self.assertEqual(result["actual_refund_amount"], refund)
                self.assertEqual(
                    bound["template_snapshot"]["commercial"]["visa_refusal_refund"],
                    result,
                )
                rendered = authority.render_bound_contract(bound)
                self.assertIn(
                    f"签证拒签退款（按实际已支付额）：AUD {refund}",
                    rendered.authority_content,
                )

    def test_published_commercial_configuration_is_frozen_into_all_renderers(self):
        self._published_v2()
        source = authority.load_published_template(
            self.conn, "NewInternationalStudentContract"
        )
        snapshot = copy.deepcopy(source.snapshot)
        snapshot["commercial"].update({
            "payment_terms": "CRM_TEST CONFIG BINDING PAYMENT TERMS",
            "refund_policy": "CRM_TEST CONFIG BINDING REFUND WORDING",
            "company_name": "Frozen Contract Company Pty Ltd",
            "company_abn": "99 888 777 666",
            "company_address": "Frozen Contract Address",
            "company_contact": "+61 3 9999 8888 / frozen@example.test",
            "company_authorised_person": "Frozen Representative",
        })
        package = authority.PublishedTemplatePackage(
            identity_id=source.identity_id,
            version_id=source.version_id,
            version_number=source.version_number,
            contract_type=source.contract_type,
            content_hash=authority._hash(snapshot),
            snapshot_hash=authority._hash(snapshot),
            snapshot=snapshot,
        )
        payload = self._payload()
        payload["company_authorised_person"] = "Changed Live Setting"
        bound = authority.bind_new_contract(
            package, payload, actor={"id": 1, "role_name": "Admin"}
        )
        self.assertEqual(
            bound["contract_instance"]["company_authorised_person"],
            "Frozen Representative",
        )
        self.assertEqual(
            bound["commercial"]["payment_plan_display"],
            "CRM_TEST CONFIG BINDING PAYMENT TERMS",
        )
        rendered = authority.render_bound_contract(bound)
        for value in (
            "CRM_TEST CONFIG BINDING PAYMENT TERMS",
            "CRM_TEST CONFIG BINDING REFUND WORDING",
            "Frozen Contract Company Pty Ltd",
            "99 888 777 666",
            "Frozen Contract Address",
            "+61 3 9999 8888",
            "frozen@example.test",
            "Frozen Representative",
        ):
            self.assertIn(value, rendered.authority_content)
        document = Document(io.BytesIO(rendered.docx_bytes))
        docx_text = "\n".join(
            [paragraph.text for paragraph in document.paragraphs]
            + [
                cell.text
                for table in document.tables
                for row in table.rows
                for cell in row.cells
            ]
        )
        pdf_text = "\n".join(
            page.extract_text() or ""
            for page in PdfReader(io.BytesIO(rendered.pdf_bytes)).pages
        )
        signing_text = rendered.signing_preview_bytes.decode("utf-8")
        for surface, text in (
            ("docx", docx_text), ("pdf", pdf_text), ("signing", signing_text)
        ):
            with self.subTest(surface=surface):
                self.assertIn("CRM_TEST CONFIG BINDING PAYMENT TERMS", text)
                self.assertIn("CRM_TEST CONFIG BINDING REFUND WORDING", text)
                self.assertIn("Frozen Representative", text)

    def test_both_student_year_fields_are_required_for_new_contracts(self):
        self._published_v2()
        package = authority.load_published_template(self.conn, "NewInternationalStudentContract")
        for missing_field, message in (
            ("current_year_level", "缺少当前年级"),
            ("applying_year_level", "缺少申请年级"),
        ):
            with self.subTest(missing_field=missing_field):
                payload = self._payload()
                payload["student_snapshot"][missing_field] = ""
                with self.assertRaisesRegex(
                    authority.PublishedTemplateIntegrityError,
                    message,
                ):
                    authority.bind_new_contract(
                        package, payload, actor={"id": 1, "role_name": "Admin"}
                    )

    def test_mandatory_student_display_bindings_are_frozen_and_rendered_on_every_surface(self):
        self._published_v2()
        package = authority.load_published_template(
            self.conn, "NewInternationalStudentContract"
        )
        bound = authority.bind_new_contract(
            package, self._payload(), actor={"id": 1, "role_name": "Admin"}
        )

        bindings = bound["display_bindings"]
        self.assertEqual(bindings["schema_version"], 1)
        self.assertEqual(bindings["contract_type"], "NewInternationalStudentContract")
        expected = {
            "current_grade": ("当前年级", "Year 7", "students.current_year_level"),
            "application_grade": ("申请年级", "Year 8", "students.applying_year_level"),
            "planned_entry_year": ("计划入读年份", "2027", "students.preferred_intake_year"),
            "discounted_fee": ("优惠后服务费（不含 GST）", "2400.00", "contract.commercial.net_service_fee_ex_gst"),
        }
        for key, (label, value, source) in expected.items():
            with self.subTest(key=key):
                self.assertEqual(bindings["fields"][key]["label"], label)
                self.assertEqual(bindings["fields"][key]["value"], value)
                self.assertEqual(bindings["fields"][key]["source"], source)
                self.assertIn(key, bindings["required_fields"])

        rendered = authority.render_bound_contract(bound)
        document = Document(io.BytesIO(rendered.docx_bytes))
        docx_text = "\n".join(
            [paragraph.text for paragraph in document.paragraphs]
            + [
                cell.text
                for table in document.tables
                for row in table.rows
                for cell in row.cells
            ]
        )
        pdf_text = "\n".join(
            page.extract_text() or ""
            for page in PdfReader(io.BytesIO(rendered.pdf_bytes)).pages
        )
        signing_text = rendered.signing_preview_bytes.decode("utf-8")
        for surface, text in (
            ("docx", docx_text), ("pdf", pdf_text), ("signing", signing_text)
        ):
            with self.subTest(surface=surface):
                for label, value, _source in expected.values():
                    self.assertIn(label, text)
                    self.assertIn(value, text)

    def test_master_protected_display_uses_frozen_bindings_not_renderer_arguments(self):
        self._published_v2()
        package = authority.load_published_template(
            self.conn, "NewInternationalStudentContract"
        )
        bound = authority.bind_new_contract(
            package, self._payload(), actor={"id": 1, "role_name": "Admin"}
        )
        hostile_snapshot = copy.deepcopy(bound["template_snapshot"])
        hostile_snapshot["commercial"]["standard_service_fee_ex_gst"] = "9999.00"
        hostile_snapshot["commercial"]["net_service_fee_ex_gst"] = "9999.00"
        hostile_snapshot["commercial"]["final_total"] = "9999.00"
        hostile_instance = copy.deepcopy(bound["contract_instance"])
        hostile_instance["student_applying_year_level"] = "HOSTILE GRADE"
        hostile_signer = copy.deepcopy(bound["signer"])
        hostile_signer["name"] = "HOSTILE SIGNER"

        artifacts = editor.render_published_contract_artifacts(
            hostile_snapshot,
            student_name="HOSTILE STUDENT",
            signer=hostile_signer,
            selected_services=list(bound["selected_service_ids"]),
            contract_instance=hostile_instance,
            display_bindings=bound["display_bindings"],
        )
        visible = artifacts["authority_content"]

        self.assertIn("Year 8", visible)
        self.assertIn("Father Only", visible)
        self.assertIn("AUD 2400.00", visible)
        self.assertNotIn("HOSTILE GRADE", visible)
        self.assertNotIn("HOSTILE SIGNER", visible)
        self.assertNotIn("9999.00", visible)

    def test_missing_application_grade_and_planned_entry_year_have_exact_errors(self):
        self._published_v2()
        package = authority.load_published_template(
            self.conn, "NewInternationalStudentContract"
        )
        for missing_field, message in (
            ("applying_year_level", "缺少申请年级"),
            ("preferred_intake_year", "缺少计划入读年份"),
        ):
            with self.subTest(missing_field=missing_field):
                payload = self._payload()
                payload["student_snapshot"][missing_field] = ""
                with self.assertRaisesRegex(
                    authority.PublishedTemplateIntegrityError, message
                ):
                    authority.bind_new_contract(
                        package, payload, actor={"id": 1, "role_name": "Admin"}
                    )

    def test_only_admin_can_freeze_discount_approval(self):
        self._published_v2()
        package = authority.load_published_template(self.conn, "NewInternationalStudentContract")
        approval = {
            "discount_amount": "200.00",
            "reason": "Approved family support discount",
            "approved_by": 1,
            "operator_id": 1,
            "approved_at": "2026-08-20T06:10:00+00:00",
        }
        with self.assertRaises(authority.PublishedTemplateIntegrityError):
            authority.bind_new_contract(
                package, self._payload(), actor={"id": 2, "role_name": "Manager"},
                discount_approval=approval,
            )
        discounted_payload = self._payload()
        discounted_payload["payment_method"] = "澳洲公户"
        bound = authority.bind_new_contract(
            package, discounted_payload, actor={"id": 1, "role_name": "Admin"},
            discount_approval=approval,
        )
        self.assertEqual(bound["commercial"]["discount_amount_ex_gst"], "200.00")
        self.assertEqual(bound["commercial"]["net_service_fee_ex_gst"], "2200.00")
        self.assertEqual(bound["commercial"]["gst_amount"], "220.00")
        self.assertEqual(bound["commercial"]["final_fee_amount"], "2420.00")
        self.assertEqual(bound["commercial"]["discount_approval"], approval)
        zero_fee = dict(approval)
        zero_fee["discount_amount"] = "2400.00"
        with self.assertRaises(authority.PublishedTemplateIntegrityError):
            authority.bind_new_contract(
                package, self._payload(), actor={"id": 1, "role_name": "Admin"},
                discount_approval=zero_fee,
            )

    def test_discount_uses_decimal_round_half_up_and_browser_totals_are_ignored(self):
        self._published_v2()
        package = authority.load_published_template(self.conn, "NewInternationalStudentContract")
        approval = {
            "discount_amount": "123.45", "reason": "Approved exact discount",
            "approved_by": 1, "operator_id": 1,
            "approved_at": "2026-08-20T06:10:00+00:00",
        }
        hostile = self._payload()
        hostile["payment_method"] = "澳洲公户"
        hostile.update({
            "gst_treatment": "GST_FREE", "gst_amount": "0", "final_total": "1",
            "actual_paid_total": "9999", "refund_policy": "hostile refund",
            "discount_approval": {"reason": "browser"},
        })
        bound = authority.bind_new_contract(
            package, hostile, actor={"id": 1, "role_name": "Admin"},
            discount_approval=approval,
        )
        self.assertEqual(bound["commercial"]["net_service_fee_ex_gst"], "2276.55")
        self.assertEqual(bound["commercial"]["gst_amount"], "227.66")
        self.assertEqual(bound["commercial"]["final_total"], "2504.21")
        self.assertEqual(bound["commercial"]["gst_treatment"], "EXCLUSIVE")
        for key in (
            "gst_treatment", "gst_amount", "final_total", "actual_paid_total",
            "refund_policy", "discount_approval",
        ):
            self.assertNotIn(key, bound["contract_instance"])

    def test_gst_free_requires_authorized_complete_approval_and_is_frozen(self):
        self._published_v2("LocalStudentApplicationContract")
        package = authority.load_published_template(self.conn, "LocalStudentApplicationContract")
        approval = {
            "reason": "Exported service is GST-free", "approved_by": 7,
            "operator_id": 7, "approved_at": "2026-08-20T06:15:00+00:00",
        }
        public_payload = self._payload("LocalStudentApplicationContract")
        public_payload["payment_method"] = "澳洲公户"
        with self.assertRaises(authority.PublishedTemplateIntegrityError):
            authority.bind_new_contract(
                package, public_payload,
                actor={"id": 7, "role_name": "Manager"}, gst_free_approval=approval,
            )
        actor = {
            "id": 7, "role_name": "Manager",
            "actor_permissions": ["contracts.gst_free.approve"],
        }
        bound = authority.bind_new_contract(
            package, public_payload,
            actor=actor, gst_free_approval=approval,
        )
        self.assertEqual(bound["commercial"]["gst_treatment"], "GST_FREE")
        self.assertEqual(bound["commercial"]["gst_amount"], "0.00")
        self.assertEqual(bound["commercial"]["final_total"], "2400.00")
        self.assertEqual(bound["commercial"]["gst_free_approval"], approval)
        incomplete = dict(approval)
        incomplete["reason"] = ""
        with self.assertRaises(authority.PublishedTemplateIntegrityError):
            authority.bind_new_contract(
                package, public_payload,
                actor=actor, gst_free_approval=incomplete,
            )
        with self.assertRaises(authority.PublishedTemplateIntegrityError):
            authority.bind_new_contract(
                package, self._payload("LocalStudentApplicationContract"),
                actor=actor, gst_free_approval=approval,
            )

    def test_installment_plan_requires_admin_and_is_frozen(self):
        self._published_v2("TransferApplicationContract")
        package = authority.load_published_template(self.conn, "TransferApplicationContract")
        plan = {
            "mode": "installments",
            "reason": "Approved two-stage plan",
            "approved_by": 1,
            "operator_id": 1,
            "approved_at": "2026-08-20T06:20:00+00:00",
            "installments": [
                {"amount": "1320.00", "due_date": "2026-08-20"},
                {"amount": "1320.00", "due_date": "2026-09-20"},
            ],
        }
        with self.assertRaises(authority.PublishedTemplateIntegrityError):
            authority.bind_new_contract(
                package, self._payload("TransferApplicationContract"),
                actor={"id": 2, "role_name": "Manager"}, payment_plan_approval=plan,
                trusted_transfer_source=self._trusted_transfer_source(self._payload("TransferApplicationContract")),
            )
        installment_payload = self._payload("TransferApplicationContract")
        installment_payload["payment_method"] = "澳洲公户"
        bound = authority.bind_new_contract(
            package, installment_payload,
            actor={"id": 1, "role_name": "Admin"}, payment_plan_approval=plan,
            trusted_transfer_source=self._trusted_transfer_source(installment_payload),
        )
        self.assertEqual(bound["commercial"]["payment_plan"], plan)

        invalid = copy.deepcopy(plan)
        invalid["installments"][0].pop("due_date")
        with self.assertRaises(authority.PublishedTemplateIntegrityError):
            authority.bind_new_contract(
                package, self._payload("TransferApplicationContract"),
                actor={"id": 1, "role_name": "Admin"}, payment_plan_approval=invalid,
                trusted_transfer_source=self._trusted_transfer_source(self._payload("TransferApplicationContract")),
            )

    def test_visa_refusal_refund_is_type_scoped_and_uses_frozen_paid_total(self):
        self._published_v2("NewInternationalStudentContract")
        package = authority.load_published_template(self.conn, "NewInternationalStudentContract")
        public_payload = self._payload()
        public_payload["payment_method"] = "澳洲公户"
        exclusive = authority.bind_new_contract(
            package, public_payload, actor={"id": 1, "role_name": "Admin"},
            actual_paid_total="2640.00",
            trusted_transfer_source=self._trusted_transfer_source(public_payload),
        )
        refund = exclusive["commercial"]["visa_refusal_refund"]
        self.assertEqual(refund["retained_ex_gst"], "900.00")
        self.assertEqual(refund["gst_adjustment"], "90.00")
        self.assertEqual(refund["retained_total"], "990.00")
        self.assertEqual(refund["actual_refund_amount"], "1650.00")
        self.assertEqual(refund["template_version_id"], package.version_id)
        self.assertEqual(refund["template_content_hash"], package.content_hash)

        approval = {
            "reason": "Exported service is GST-free", "approved_by": 1,
            "operator_id": 1, "approved_at": "2026-08-20T06:25:00+00:00",
        }
        gst_free = authority.bind_new_contract(
            package, public_payload, actor={"id": 1, "role_name": "Admin"},
            gst_free_approval=approval, actual_paid_total="2400.00",
            trusted_transfer_source=self._trusted_transfer_source(public_payload),
        )
        self.assertEqual(gst_free["commercial"]["visa_refusal_refund"]["retained_total"], "900.00")
        self.assertEqual(gst_free["commercial"]["visa_refusal_refund"]["actual_refund_amount"], "1500.00")

        discount = {
            "discount_amount": "200.00", "reason": "Approved discount",
            "approved_by": 1, "operator_id": 1,
            "approved_at": "2026-08-20T06:26:00+00:00",
        }
        discounted = authority.bind_new_contract(
            package, public_payload, actor={"id": 1, "role_name": "Admin"},
            discount_approval=discount, actual_paid_total="2420.00",
            trusted_transfer_source=self._trusted_transfer_source(public_payload),
        )
        self.assertEqual(discounted["commercial"]["visa_refusal_refund"]["actual_refund_amount"], "1430.00")

        for contract_type in ("TransferApplicationContract", "LocalStudentApplicationContract"):
            self._published_v2(contract_type)
            other = authority.bind_new_contract(
                authority.load_published_template(self.conn, contract_type),
                self._payload(contract_type), actor={"id": 1, "role_name": "Admin"},
                actual_paid_total="2640.00",
                trusted_transfer_source=self._trusted_transfer_source(self._payload(contract_type)),
            )
            self.assertNotIn("visa_refusal_refund", other["commercial"])
            self.assertNotIn("拒签", str(other["commercial"].get("refund_policy") or ""))

    def test_rebind_uses_original_renderer_student_signer_and_commercial_facts(self):
        self._published_v2()
        package = authority.load_published_template(self.conn, "NewInternationalStudentContract")
        bound = authority.bind_new_contract(
            package, self._payload(), actor={"id": 2, "role_name": "Manager"},
            generated_at="2026-08-20T06:00:00+00:00",
        )
        hostile = self._payload()
        hostile["student_full_name"] = "Changed Master Student"
        hostile["student_snapshot"]["student_full_name"] = "Changed Master Student"
        hostile["signer"] = {
            "source": "mother", "name": "Changed Mother", "relationship": "母亲",
            "phone": "0499999999", "email": "changed@example.test",
        }
        hostile["service_scope_ids"] = hostile["service_scope_ids"][:-1]
        rebound = authority.rebind_existing_draft(
            bound, hostile, request_id="rebind-request-002",
            generated_at="2026-08-20T07:00:00+00:00",
        )
        self.assertEqual(rebound["student_snapshot"], bound["student_snapshot"])
        self.assertEqual(rebound["signer"], bound["signer"])
        self.assertEqual(rebound["commercial"], bound["commercial"])
        self.assertEqual(rebound["renderer_schema_hash"], bound["renderer_schema_hash"])
        self.assertEqual(rebound["selected_service_ids"], hostile["service_scope_ids"])

    def test_legacy_v1_renderer_snapshot_remains_renderable_and_rebindable(self):
        self._published_v2()
        package = authority.load_published_template(self.conn, "NewInternationalStudentContract")
        bound = authority.bind_new_contract(
            package, self._payload(), actor={"id": 2, "role_name": "Manager"},
            actual_paid_total="0.00",
        )
        self.assertEqual(bound["renderer_schema"]["version"], 4)

        legacy = copy.deepcopy(bound)
        frozen_published = frozen_legacy_v1_snapshot(
            legacy["published_template_snapshot"]
        )
        frozen_instance = frozen_legacy_v1_snapshot(legacy["template_snapshot"])
        legacy["published_template_snapshot"] = frozen_published
        legacy["template_snapshot"] = frozen_instance
        legacy["clause_blocks"] = copy.deepcopy(frozen_instance["blocks"])
        legacy["service_items"] = copy.deepcopy(frozen_instance["services"])
        legacy["template_content_hash"] = authority._hash(frozen_published)
        legacy["template_snapshot_hash"] = authority._hash(frozen_published)
        legacy["renderer_schema"] = copy.deepcopy(authority.LEGACY_RENDERER_SCHEMA)
        legacy["renderer_schema_hash"] = authority.LEGACY_RENDERER_SCHEMA_HASH
        for key in (
            "student_current_year_level",
            "student_applying_year_level",
            "company_authorised_person",
        ):
            legacy["contract_instance"].pop(key, None)
        legacy["company_configuration"].pop("company_authorised_person", None)
        legacy.pop("bound_snapshot_hash", None)
        legacy["bound_snapshot_hash"] = authority._hash(legacy)

        rendered = authority.render_bound_contract(legacy)
        self.assertNotIn("甲方授权代表", rendered.authority_content)
        self.assertNotIn("当前年级", rendered.authority_content)
        self.assertEqual(
            hashlib.sha256(rendered.authority_content.encode()).hexdigest(),
            "540488b591537fa7e8a5e1bad776f304ba17601b8b130a798e0c028bbf0799b4",
        )
        self.assertEqual(
            rendered.semantic_hashes["docx"],
            "6761c4aad1b33f895fa7d701ba5c8fcc5e209afc670ece81c80e4ff2fbd27fde",
        )
        self.assertEqual(
            hashlib.sha256(rendered.docx_bytes).hexdigest(),
            "37aeb78f4f729aa0ae8fd85d184b0ce623622e3266f5a2493af81a49aab3ccdb",
        )
        self.assertEqual(
            hashlib.sha256(rendered.pdf_bytes).hexdigest(),
            "ff461867f763712c6b6f0970ed570317719372a334e9aa34506059c940e786b5",
        )
        self.assertEqual(
            hashlib.sha256(rendered.signing_preview_bytes).hexdigest(),
            "feb0404ae0c9d4b169ba49c407d9f8cf70c25240fda833680c4e96f44d954b2f",
        )
        rebound = authority.rebind_existing_draft(
            legacy,
            {"service_scope_ids": legacy["selected_service_ids"][:-1]},
            request_id="legacy-v1-rebind-002",
        )
        self.assertEqual(rebound["renderer_schema"], authority.LEGACY_RENDERER_SCHEMA)
        rebound_rendered = authority.render_bound_contract(rebound)
        self.assertEqual(
            rebound_rendered.semantic_hashes["docx"],
            rebound_rendered.semantic_hashes["signing_preview"],
        )

    def test_frozen_master_v3_contract_remains_renderable_without_new_bindings(self):
        self._published_v2()
        package = authority.load_published_template(
            self.conn, "NewInternationalStudentContract"
        )
        bound = authority.bind_new_contract(
            package, self._payload(), actor={"id": 2, "role_name": "Manager"}
        )
        historical = copy.deepcopy(bound)
        historical["renderer_schema"] = copy.deepcopy(
            authority.LEGACY_RENDERER_SCHEMA_V3
        )
        historical["renderer_schema_hash"] = authority.LEGACY_RENDERER_SCHEMA_V3_HASH
        historical.pop("display_bindings", None)
        for key in (
            "student_planned_entry_year",
            "student_planned_entry_term",
            "intake_year",
            "intake_term",
        ):
            historical["contract_instance"].pop(key, None)
        historical.pop("bound_snapshot_hash", None)
        historical["bound_snapshot_hash"] = authority._hash(historical)

        rendered = authority.render_bound_contract(historical)

        self.assertIn("当前年级", rendered.authority_content)
        self.assertIn("申请年级", rendered.authority_content)
        self.assertNotIn("计划入读年份", rendered.authority_content)

    def test_manager_apply_only_affects_transfer_contracts_created_after_application(self):
        contract_type = "TransferApplicationContract"
        self._published_v2(contract_type)
        original_package = authority.load_published_template(
            self.conn, contract_type
        )
        original_payload = self._payload(contract_type)
        original = authority.bind_new_contract(
            original_package, original_payload,
            actor={"id": 2, "role_name": "Manager + Student Cases"},
            generated_at="2026-08-20T06:00:00+00:00",
            trusted_transfer_source=self._trusted_transfer_source(original_payload),
        )
        original_render = authority.render_bound_contract(original)
        original_bound_hash = original["bound_snapshot_hash"]
        self.service.set_user_permission(
            2, can_edit=False, can_publish=True,
            actor_user_id=1, actor_is_admin=True,
        )
        draft = self.service.create_draft(
            contract_type, actor_user_id=2
        )
        lock = self.service.acquire_lock(
            contract_type, draft["id"], owner_user_id=2
        )
        snapshot = json.loads(draft["snapshot_json"])
        block = next(item for item in snapshot["blocks"] if item["key"] == "profile_3")
        marker = "TRANSFER TEMPLATE MANAGER APPLY MARKER"
        draft = self.service.autosave_block(
            draft["id"], block["key"], title=block["title"],
            body=f'{block["body"]}\n{marker}', enabled=True,
            expected_revision=draft["revision"], lock_token=lock["token"],
            actor_user_id=2,
        )
        published = self.service.apply_template(
            draft["id"], actor_user_id=2, lock_token=lock["token"],
            expected_revision=draft["revision"],
            output_root=self.root / "manager-apply-simulation",
            actor_is_admin=False,
        )
        new_package = authority.load_published_template(
            self.conn, contract_type
        )
        new_payload = self._payload(contract_type)
        new_payload["generation_request_id"] = "req-authority-new-version"
        new_contract = authority.bind_new_contract(
            new_package, new_payload,
            actor={"id": 2, "role_name": "Manager + Student Cases"},
            trusted_transfer_source=self._trusted_transfer_source(new_payload),
        )
        new_render = authority.render_bound_contract(new_contract)

        self.assertNotEqual(original_package.version_id, published["id"])
        self.assertEqual(new_contract["template_version_id"], published["id"])
        self.assertEqual(original["template_version_id"], original_package.version_id)
        self.assertIn(marker, new_render.authority_content)
        self.assertNotIn(marker, original_render.authority_content)
        self.assertEqual(original["bound_snapshot_hash"], original_bound_hash)
        self.assertEqual(
            authority.render_bound_contract(original).docx_bytes,
            original_render.docx_bytes,
        )
        self.assertEqual(
            authority.render_bound_contract(original).pdf_bytes,
            original_render.pdf_bytes,
        )
        self.assertEqual(
            authority.render_bound_contract(original).signing_preview_bytes,
            original_render.signing_preview_bytes,
        )

    def test_renderer_schema_tamper_fails_closed(self):
        self._published_v2()
        package = authority.load_published_template(self.conn, "NewInternationalStudentContract")
        bound = authority.bind_new_contract(
            package, self._payload(), actor={"id": 2, "role_name": "Manager"}
        )
        bound["renderer_schema"]["version"] = 99
        bound["bound_snapshot_hash"] = authority._hash(
            {key: value for key, value in bound.items() if key != "bound_snapshot_hash"}
        )
        with self.assertRaises(authority.PublishedTemplateIntegrityError):
            authority.render_bound_contract(bound)

    def test_real_docx_pdf_and_signing_artifacts_share_visible_semantics(self):
        self._published_v2("LocalStudentApplicationContract")
        package = authority.load_published_template(self.conn, "LocalStudentApplicationContract")
        bound = authority.bind_new_contract(
            package, self._payload("LocalStudentApplicationContract"),
            actor={"id": 2, "role_name": "Manager"},
        )
        rendered = authority.render_bound_contract(bound)
        docx = self.root / "contract.docx"
        pdf = self.root / "contract.pdf"
        signing = self.root / "contract-signing.html"
        docx.write_bytes(rendered.docx_bytes)
        pdf.write_bytes(rendered.pdf_bytes)
        signing.write_bytes(rendered.signing_preview_bytes)

        expected = editor._semantic_trace_from_model(rendered.visible_model)
        self.assertEqual(editor._extract_docx_semantic_trace(docx), expected)
        self.assertEqual(
            editor._extract_verified_pdf_semantic_trace(
                pdf, editor._build_pdf_layout_plan(rendered.visible_model)
            ),
            expected,
        )
        self.assertEqual(editor._extract_signing_semantic_trace(signing), expected)
        self.assertTrue(rendered.docx_bytes.startswith(b"PK"))
        self.assertTrue(rendered.pdf_bytes.startswith(b"%PDF"))
        import app as app_module
        self.assertTrue(app_module.validate_docx_contains_contract_body(docx))
        for required in (
            "甲方授权代表", "Sonny Zhang", "当前年级", "Year 7", "申请年级", "Year 8",
        ):
            self.assertIn(required, rendered.authority_content)

    def test_all_three_types_use_published_service_allowlists(self):
        for contract_type, forbidden in {
            "NewInternationalStudentContract": (),
            "TransferApplicationContract": ("student_visa_lodgement",),
            "LocalStudentApplicationContract": (
                "student_visa_lodgement", "offer_coe_oshc_process_assistance",
            ),
        }.items():
            with self.subTest(contract_type=contract_type):
                self._published_v2(contract_type)
                package = authority.load_published_template(self.conn, contract_type)
                enabled = {
                    item["service_id"] for item in package.snapshot["services"]
                    if item.get("enabled", True)
                }
                self.assertFalse(enabled.intersection(forbidden))

    def test_v2_conversion_preserves_v1_clauses_services_and_domain_profiles(self):
        v1 = {
            row["template_type"]: json.loads(row["snapshot_json"])
            for row in self.conn.execute(
                """SELECT i.template_type,v.snapshot_json
                   FROM contract_template_identities i
                   JOIN contract_template_versions v ON v.identity_id=i.id
                   WHERE v.version_number=1"""
            )
        }
        drafts = {row["template_type"]: json.loads(row["snapshot_json"]) for row in self._drafts()}
        for contract_type, source in v1.items():
            with self.subTest(contract_type=contract_type):
                converted = drafts[contract_type]
                self.assertEqual(converted["blocks"], source["blocks"])
                self.assertEqual(converted["services"], source["services"])
                self.assertEqual(converted["clause_profile"], source["clause_profile"])
                self.assertEqual(
                    [item["service_id"] for item in converted["services"]],
                    list(contract_domain.contract_type_config(contract_type).allowed_services),
                )
                profile_text = "\n".join(
                    str(item["body"]) for item in converted["blocks"]
                )
                if contract_type == "NewInternationalStudentContract":
                    self.assertIn("签证", profile_text)
                else:
                    self.assertNotIn("移民局签证代理", profile_text)

    def test_real_four_surface_semantics_cover_all_types_590_partial_and_both_parents(self):
        cases = [
            ("NewInternationalStudentContract", "father", False),
            ("NewInternationalStudentContract", "mother", True),
            ("TransferApplicationContract", "father", False),
            ("LocalStudentApplicationContract", "mother", False),
        ]
        for contract_type, signer_source, include_590 in cases:
            with self.subTest(
                contract_type=contract_type,
                signer_source=signer_source,
                include_590=include_590,
            ):
                self._published_v2(contract_type)
                package = authority.load_published_template(self.conn, contract_type)
                payload = self._payload(contract_type)
                enabled = [
                    item["service_id"] for item in package.snapshot["services"]
                    if item.get("enabled", True)
                ]
                selected = enabled[::2] or enabled[:1]
                if include_590:
                    selected.extend(
                        item for item in (
                            "subclass_590_guardian_visa_lodgement",
                            "subclass_590_guardian_visa_consulting",
                        )
                        if item in enabled and item not in selected
                    )
                payload["service_scope_ids"] = selected
                payload["payment_method"] = "澳洲公户"
                if signer_source == "mother":
                    payload["signer"] = {
                        "source": "mother", "name": "Mother Other",
                        "relationship": "母亲", "phone": "0422222222",
                        "email": "mother@example.test",
                    }
                bound = authority.bind_new_contract(
                    package, payload, actor={"id": 1, "role_name": "Admin"}
                , trusted_transfer_source=self._trusted_transfer_source(payload))
                rendered = authority.render_bound_contract(bound)
                expected = editor._semantic_trace_from_model(rendered.visible_model)
                output = self.root / f"four-surface-{contract_type}-{signer_source}"
                output.mkdir()
                docx = output / "contract.docx"
                pdf = output / "contract.pdf"
                signing = output / "signing.html"
                docx.write_bytes(rendered.docx_bytes)
                pdf.write_bytes(rendered.pdf_bytes)
                signing.write_bytes(rendered.signing_preview_bytes)
                self.assertEqual(editor._extract_docx_semantic_trace(docx), expected)
                self.assertEqual(
                    editor._extract_verified_pdf_semantic_trace(
                        pdf, editor._build_pdf_layout_plan(rendered.visible_model)
                    ),
                    expected,
                )
                self.assertEqual(editor._extract_signing_semantic_trace(signing), expected)
                visible = rendered.authority_content
                for required in (
                    "当前年级", "Year 7",
                    "申请年级", "Year 8",
                    "计划入读年份", "2027",
                ):
                    self.assertIn(required, visible)
                self.assertIn("2400.00", visible)
                self.assertIn("240.00", visible)
                self.assertIn("2640.00", visible)
                for service_id in selected:
                    self.assertIn(contract_domain.service_label(service_id), visible)
                for service_id in set(enabled) - set(selected):
                    self.assertNotIn(contract_domain.service_label(service_id), visible)

    def test_real_four_surface_pipeline_covers_discount_gst_free_installment_and_refund(self):
        self._published_v2("NewInternationalStudentContract")
        new_package = authority.load_published_template(
            self.conn, "NewInternationalStudentContract"
        )
        discount = {
            "discount_amount": "123.45", "reason": "Approved exact discount",
            "approved_by": 1, "operator_id": 1,
            "approved_at": "2026-08-20T06:30:00+00:00",
        }
        gst_free = {
            "reason": "Exported service is GST-free", "approved_by": 1,
            "operator_id": 1, "approved_at": "2026-08-20T06:31:00+00:00",
        }
        self._published_v2("TransferApplicationContract")
        transfer_package = authority.load_published_template(
            self.conn, "TransferApplicationContract"
        )
        installments = {
            "mode": "installments", "reason": "Approved dated plan",
            "approved_by": 1, "operator_id": 1,
            "approved_at": "2026-08-20T06:32:00+00:00",
            "installments": [
                {"amount": "1320.00", "due_date": "2026-09-01"},
                {"amount": "1320.00", "due_date": "2026-10-01"},
            ],
        }
        cases = [
            (
                "discount-refund", new_package, self._payload(),
                {"discount_approval": discount, "actual_paid_total": "2504.21"},
                ("2276.55", "227.66", "2504.21", "1514.21"),
            ),
            (
                "gst-free-refund", new_package, self._payload(),
                {"gst_free_approval": gst_free, "actual_paid_total": "2400.00"},
                ("2400.00", "0.00", "2400.00", "1500.00"),
            ),
            (
                "installments", transfer_package,
                self._payload("TransferApplicationContract"),
                {"payment_plan_approval": installments},
                ("2400.00", "240.00", "2640.00", None),
            ),
        ]
        for name, package, payload, approvals, expected_values in cases:
            with self.subTest(name=name):
                payload["generation_request_id"] = f"four-surface-{name}"
                payload["payment_method"] = "澳洲公户"
                bound = authority.bind_new_contract(
                    package, payload, actor={"id": 1, "role_name": "Admin"},
                    **approvals,
                    trusted_transfer_source=self._trusted_transfer_source(payload),
                )
                rendered = authority.render_bound_contract(bound)
                expected = editor._semantic_trace_from_model(rendered.visible_model)
                output = self.root / name
                output.mkdir()
                docx = output / "contract.docx"
                pdf = output / "contract.pdf"
                signing = output / "signing.html"
                docx.write_bytes(rendered.docx_bytes)
                pdf.write_bytes(rendered.pdf_bytes)
                signing.write_bytes(rendered.signing_preview_bytes)
                self.assertEqual(editor._extract_docx_semantic_trace(docx), expected)
                self.assertEqual(
                    editor._extract_verified_pdf_semantic_trace(
                        pdf, editor._build_pdf_layout_plan(rendered.visible_model)
                    ),
                    expected,
                )
                self.assertEqual(editor._extract_signing_semantic_trace(signing), expected)
                for value in expected_values:
                    if value is not None:
                        self.assertIn(value, rendered.authority_content)

    def test_authority_signer_missing_contact_cannot_borrow_other_parent(self):
        import app as app_module

        database = self._persistence_database()
        student_data = self._student()
        student_data["father_mobile"] = ""
        student_data["father_email"] = ""
        database.execute(
            "UPDATE students SET form_data_json=? WHERE student_id=?",
            (json.dumps(student_data), "T-AUTH-001"),
        )
        student = database.one("SELECT * FROM students WHERE student_id='T-AUTH-001'")
        old_db = app_module.db
        try:
            app_module.db = database
            borrowed = {
                "parent_signer_source": "father",
                "parent_signer_name": "Father Only",
                "parent_signer_relationship": "父亲",
                "parent_phone": "0422222222",
                "parent_email": "mother@example.test",
            }
            with self.assertRaises(authority.PublishedTemplateAuthorityError):
                app_module.authority_contract_signer(student, borrowed)
            completed = dict(borrowed)
            completed["parent_phone"] = "0433333333"
            completed["parent_email"] = "father-fill@example.test"
            signer = app_module.authority_contract_signer(student, completed)
            self.assertEqual(
                signer,
                {
                    "source": "father", "name": "Father Only", "relationship": "父亲",
                    "phone": "0433333333", "email": "father-fill@example.test",
                },
            )
            self.assertEqual(
                json.loads(database.one(
                    "SELECT form_data_json FROM students WHERE student_id='T-AUTH-001'"
                )["form_data_json"])["father_mobile"],
                "",
            )
        finally:
            app_module.db = old_db

    def _persistence_database(self, *, all_types=False):
        database = Database(self.root / "persistence.sqlite3")
        database.initialize()
        with database.transaction() as connection:
            connection.executescript(MIGRATION.read_text(encoding="utf-8"))
            connection.executemany(
                "INSERT INTO system_settings(key,value,updated_at) VALUES(?,?,?) "
                "ON CONFLICT(key) DO UPDATE SET value=excluded.value",
                [
                    ("company_legal_name", "AAIE Test Pty Ltd", "2026-08-20"),
                    ("company_abn_acn", "12 345 678 901", "2026-08-20"),
                    ("company_address", "Melbourne VIC", "2026-08-20"),
                    ("company_phone", "+61 3 0000 0000", "2026-08-20"),
                    ("company_email", "contracts@example.test", "2026-08-20"),
                    ("company_authorised_person", "Sonny Zhang", "2026-08-20"),
                ],
            )
            connection.execute(
                """INSERT INTO students(
                   student_id,passport_last_name,passport_first_name,current_status,
                   date_of_birth,current_year_level,applying_year_level,preferred_intake_year,
                   preferred_intake_term,created_date,updated_at,nas_status,form_data_json
                ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                (
                    "T-AUTH-001", "Student", "Frozen", "Active", "2010-01-01",
                    "Year 7", "Year 8",
                    "2027", "02", "2026-08-20", "2026-08-20", "ready",
                    json.dumps(self._student()),
                ),
            )
        service = editor.ContractTemplateEditorService(database.connect())
        try:
            service.import_domain_v1(actor_user_id=1)
            drafts = service.create_authority_v2_drafts(actor_user_id=1)
            if not all_types:
                drafts = [
                    item for item in drafts
                    if item["template_type"] == "NewInternationalStudentContract"
                ]
            for original in drafts:
                contract_type = original["template_type"]
                lock = service.acquire_lock(
                    contract_type, original["id"], owner_user_id=1
                )
                draft = service.configure_authority_v2_draft(
                    original["id"], gst_treatment="EXCLUSIVE",
                    expected_revision=original["revision"], lock_token=lock["token"],
                    actor_user_id=1,
                )
                service.simulate(
                    draft["id"], source_type="system", student_snapshot=self._student(),
                    signer_source="father", signer_overrides={},
                    selected_services=list(
                        contract_domain.contract_type_config(contract_type).default_services
                    ),
                    output_root=self.root / f"persistence-simulation-{contract_type}",
                    actor_user_id=1, expected_revision=draft["revision"],
                    lock_token=lock["token"],
                )
                service.publish(
                    draft["id"], actor_user_id=1, lock_token=lock["token"],
                    expected_revision=draft["revision"], actor_is_admin=True,
                )
        finally:
            service.conn.close()
        return database

    @staticmethod
    def _trusted_transfer_source(payload):
        if payload.get("contract_type") != "TransferApplicationContract":
            return None
        return {"authority_id": 1, "transfer_contract_id": 1,
                "student_id": payload["student_snapshot"]["student_id"],
                "source_classification": "new_transfer", "rule_version": "transfer-source-v1"}

    def _bound_and_rendered(
        self, database, request_id="persist-request-001", *, actual_paid_total=None,
        discount_approval=None, gst_free_approval=None, payment_plan_approval=None,
        payment_method="国内银行",
    ):
        connection = database.connect()
        try:
            package = authority.load_published_template(
                connection, "NewInternationalStudentContract"
            )
        finally:
            connection.close()
        payload = self._payload()
        payload["generation_request_id"] = request_id
        payload["payment_method"] = payment_method
        bound = authority.bind_new_contract(
            package, payload, actor={"id": 1, "role_name": "Admin"},
            actual_paid_total=actual_paid_total,
            discount_approval=discount_approval,
            gst_free_approval=gst_free_approval,
            payment_plan_approval=payment_plan_approval,
            generated_at="2026-08-20T08:00:00+00:00",
        )
        return bound, authority.render_bound_contract(bound)

    def test_atomic_persistence_replays_same_request_without_duplicate_rows_or_files(self):
        database = self._persistence_database()
        bound, rendered = self._bound_and_rendered(database)
        target = self.root / "forms" / "PublishedContract_V1.docx"
        user = {"id": 1, "username": "admin"}
        first = create_contract_atomic(
            database, bound=bound, rendered=rendered, student_data=self._student(),
            student_id="T-AUTH-001", contract_uid="AUTH-001", target=target,
            file_version=1, user=user, occurred_at="2026-08-20T08:00:00+00:00",
        )
        replay = create_contract_atomic(
            database, bound=bound, rendered=rendered, student_data=self._student(),
            student_id="T-AUTH-001", contract_uid="AUTH-IGNORED", target=target,
            file_version=1, user=user, occurred_at="2026-08-20T08:01:00+00:00",
        )

        self.assertTrue(first.created)
        self.assertFalse(replay.created)
        self.assertEqual(first.contract_id, replay.contract_id)
        self.assertEqual(database.one("SELECT COUNT(*) n FROM contracts")["n"], 1)
        self.assertEqual(database.one("SELECT COUNT(*) n FROM student_documents")["n"], 1)
        self.assertEqual(len([p for p in target.parent.glob("*.docx") if not p.name.startswith("._")]), 1)
        stored = database.one("SELECT * FROM contracts WHERE id=?", (first.contract_id,))
        self.assertEqual(stored["consulting_and_visa_service_fee"], "2400.00")
        self.assertEqual(stored["service_fee_currency"], "国内银行")
        self.assertEqual(stored["final_fee_amount"], "2400.00")
        stored_bound = contract_domain.snapshot_payload(stored["standalone_data_json"])
        self.assertEqual(
            authority.render_bound_contract(stored_bound).semantic_hashes,
            rendered.semantic_hashes,
        )

    def test_persistence_freezes_refund_columns_and_immutable_commercial_audit(self):
        database = self._persistence_database()
        bound, rendered = self._bound_and_rendered(
            database, request_id="persist-commercial-audit-001",
            actual_paid_total="2640.00",
            payment_method="澳洲公户",
        )
        result = create_contract_atomic(
            database, bound=bound, rendered=rendered, student_data=self._student(),
            student_id="T-AUTH-001", contract_uid="AUTH-COMMERCIAL-001",
            target=self.root / "forms" / "PublishedCommercialAudit.docx",
            file_version=1, user={"id": 1, "username": "admin"},
            occurred_at="2026-08-20T08:00:00+00:00",
        )
        stored = database.one("SELECT * FROM contracts WHERE id=?", (result.contract_id,))
        self.assertEqual(stored["refusal_deduction_amount"], "990.00")
        self.assertEqual(stored["refusal_refund_amount"], "1650.00")
        audit = database.one(
            "SELECT new_value FROM audit_logs WHERE contract_id=? "
            "AND action_type='published_v2_contract_generated'",
            (str(result.contract_id),),
        )
        evidence = json.loads(audit["new_value"])
        self.assertEqual(
            evidence["visa_refusal_refund"],
            bound["commercial"]["visa_refusal_refund"],
        )
        self.assertEqual(len(evidence["commercial_snapshot_hash"]), 64)

    def test_same_request_id_with_different_fingerprint_fails_closed(self):
        database = self._persistence_database()
        bound, rendered = self._bound_and_rendered(database)
        target = self.root / "forms" / "PublishedContract_V1.docx"
        user = {"id": 1, "username": "admin"}
        create_contract_atomic(
            database, bound=bound, rendered=rendered, student_data=self._student(),
            student_id="T-AUTH-001", contract_uid="AUTH-001", target=target,
            file_version=1, user=user, occurred_at="2026-08-20T08:00:00+00:00",
        )
        changed = copy.deepcopy(bound)
        changed["request_fingerprint"] = "different"
        changed.pop("bound_snapshot_hash")
        changed["bound_snapshot_hash"] = authority._hash(changed)
        with self.assertRaises(ContractGenerationRequestConflict):
            create_contract_atomic(
                database, bound=changed, rendered=rendered, student_data=self._student(),
                student_id="T-AUTH-001", contract_uid="AUTH-002",
                target=self.root / "forms" / "PublishedContract_V2.docx",
                file_version=2, user=user, occurred_at="2026-08-20T08:01:00+00:00",
            )
        self.assertEqual(database.one("SELECT COUNT(*) n FROM contracts")["n"], 1)

    def test_request_fingerprint_covers_frozen_instance_terms(self):
        database = self._persistence_database()
        connection = database.connect()
        try:
            package = authority.load_published_template(
                connection, "NewInternationalStudentContract"
            )
        finally:
            connection.close()
        first_payload = self._payload()
        first = authority.bind_new_contract(
            package, first_payload, actor={"id": 1, "role_name": "Admin"},
            generated_at="2026-08-20T08:00:00+00:00",
        )
        changed_payload = copy.deepcopy(first_payload)
        changed_payload["special_terms"] = "a changed frozen term"
        changed = authority.bind_new_contract(
            package, changed_payload, actor={"id": 1, "role_name": "Admin"},
            generated_at="2026-08-20T08:00:00+00:00",
        )
        self.assertNotEqual(first["request_fingerprint"], changed["request_fingerprint"])

    def test_generation_authority_switch_is_admin_only_and_atomically_audited(self):
        database = self._persistence_database()
        with self.assertRaises(ContractGenerationPersistenceError):
            set_generation_authority_atomic(
                database, mode="published_v2",
                actor={"id": 2, "username": "staff", "role_name": "Staff"},
                occurred_at="2026-08-20T08:00:00+00:00",
            )
        self.assertIsNone(database.one(
            "SELECT value FROM system_settings WHERE key='contract_generation_authority'"
        ))
        before, after = set_generation_authority_atomic(
            database, mode="published_v2",
            actor={"id": 1, "username": "admin", "role_name": "Admin"},
            occurred_at="2026-08-20T08:01:00+00:00",
        )
        self.assertEqual((before, after), ("legacy_v6", "published_v2"))
        self.assertEqual(database.setting("contract_generation_authority"), "published_v2")
        audit = database.one(
            "SELECT * FROM audit_logs WHERE action_type='contract_generation_authority_switched'"
        )
        self.assertEqual(
            (audit["old_value"], audit["new_value"]),
            ("legacy_v6", "published_v2"),
        )
        self.assertEqual(audit["status"], "success")
        set_generation_authority_atomic(
            database, mode="published_v2",
            actor={"id": 1, "username": "admin", "role_name": "Admin"},
            occurred_at="2026-08-20T08:02:00+00:00",
        )
        self.assertEqual(database.one(
            "SELECT COUNT(*) n FROM audit_logs "
            "WHERE action_type='contract_generation_authority_switched'"
        )["n"], 1)

    def test_persistence_failure_rolls_back_rows_and_promoted_file(self):
        database = self._persistence_database()
        bound, rendered = self._bound_and_rendered(database)
        target = self.root / "forms" / "PublishedContract_V1.docx"
        with self.assertRaises(sqlite3.IntegrityError):
            create_contract_atomic(
                database, bound=bound, rendered=rendered, student_data=self._student(),
                student_id="T-AUTH-001", contract_uid="AUTH-001", target=target,
                file_version=1, user={"id": 999, "username": "missing"},
                occurred_at="2026-08-20T08:00:00+00:00",
            )
        self.assertFalse(target.exists())
        self.assertEqual(database.one("SELECT COUNT(*) n FROM contracts")["n"], 0)
        self.assertEqual(database.one("SELECT COUNT(*) n FROM student_documents")["n"], 0)

    def test_draft_rebuild_is_idempotent_and_keeps_frozen_master_facts(self):
        database = self._persistence_database()
        bound, rendered = self._bound_and_rendered(database)
        user = {"id": 1, "username": "admin"}
        first_target = self.root / "forms" / "PublishedContract_V1.docx"
        created = create_contract_atomic(
            database, bound=bound, rendered=rendered, student_data=self._student(),
            student_id="T-AUTH-001", contract_uid="AUTH-001", target=first_target,
            file_version=1, user=user, occurred_at="2026-08-20T08:00:00+00:00",
        )
        updated = authority.rebind_existing_draft(
            bound, {"service_scope_ids": bound["selected_service_ids"][:-1]},
            request_id="persist-rebuild-002", generated_at="2026-08-20T09:00:00+00:00",
        )
        updated_rendered = authority.render_bound_contract(updated)
        second_target = self.root / "forms" / "PublishedContract_V2.docx"
        rebuilt = replace_draft_atomic(
            database, contract_id=created.contract_id, student_id="T-AUTH-001",
            bound=updated, rendered=updated_rendered, student_data={"changed": True},
            target=second_target, file_version=2, user=user,
            occurred_at="2026-08-20T09:00:00+00:00",
        )
        replay = replace_draft_atomic(
            database, contract_id=created.contract_id, student_id="T-AUTH-001",
            bound=updated, rendered=updated_rendered, student_data={"changed": True},
            target=second_target, file_version=2, user=user,
            occurred_at="2026-08-20T09:01:00+00:00",
        )
        stored = contract_domain.snapshot_payload(
            database.one("SELECT standalone_data_json FROM contracts WHERE id=?", (created.contract_id,))[
                "standalone_data_json"
            ]
        )
        self.assertTrue(rebuilt.created)
        self.assertFalse(replay.created)
        self.assertEqual(stored["student_snapshot"], bound["student_snapshot"])
        self.assertEqual(stored["signer"], bound["signer"])
        self.assertEqual(database.one("SELECT COUNT(*) n FROM contracts")["n"], 1)
        self.assertEqual(database.one("SELECT COUNT(*) n FROM student_documents")["n"], 2)

    def test_application_route_uses_published_v2_and_replays_event_without_duplicate_contract(self):
        import app as app_module

        class FakeNas:
            def student_folder(self, student):
                return Path(student["folder_path"]).resolve()

            def require_category_access(self, student, category, *, write=False):
                target = self.student_folder(student) / category
                target.mkdir(parents=True, exist_ok=True)
                return target

        database = self._persistence_database()
        database.set_setting(
            app_module.CONTRACT_GENERATION_AUTHORITY_SETTING,
            app_module.CONTRACT_GENERATION_AUTHORITY_PUBLISHED_V2,
            1,
        )
        student_folder = self.root / "route-student"
        student_folder.mkdir()
        database.execute(
            "UPDATE students SET folder_path=?,form_data_json=? WHERE student_id=?",
            (str(student_folder), json.dumps(self._student()), "T-AUTH-001"),
        )
        old_db, old_nas = app_module.db, app_module.nas
        old_emit = app_module.emit_contract_business_event
        old_validate = app_module.validate_pdf_contains_contract_body
        events = []
        event_attempts = {"count": 0}
        try:
            app_module.db = database
            app_module.nas = FakeNas()
            app_module.validate_pdf_contains_contract_body = lambda _path: None
            def emit_with_one_transient_failure(contract, **kwargs):
                event_attempts["count"] += 1
                events.append((contract["id"], kwargs["event_type"]))
                if event_attempts["count"] == 1:
                    raise RuntimeError("simulated transient fact failure")

            app_module.emit_contract_business_event = emit_with_one_transient_failure
            handler = object.__new__(app_module.CRMHandler)
            handler.form = {
                "contract_type": "NewInternationalStudentContract",
                "contract_date": "2026-08-20",
                "generation_request_id": "route-authority-001",
                "parent_signer_source": "father",
                "parent_signer_name": "Father Only",
                "parent_signer_relationship": "父亲",
                "parent_phone": "0411111111",
                "parent_email": "father@example.test",
                "service_scope": list(
                    contract_domain.contract_type_config(
                        "NewInternationalStudentContract"
                    ).default_services
                ),
                "service_fee_currency": "国内银行",
                "consulting_and_visa_service_fee": "9999.00",
                "final_fee_amount": "1.00",
            }
            handler.user = {
                "id": 1, "username": "admin", "role_name": "Admin",
                "display_name": "Admin",
                "permissions_json": json.dumps(["*"]),
            }
            handler.client_address = ("127.0.0.1", 12345)
            handler.headers = {}
            handler.sent_headers = []
            handler.send_response = lambda status: setattr(handler, "response_status", status)
            handler.send_header = lambda key, value: handler.sent_headers.append((key, value))
            handler.end_headers = lambda: None
            signing_pdf = self.root / "route-signing-source.pdf"
            signing_pdf.write_bytes(b"%PDF-1.4\n%%EOF\n")
            handler.contract_signing_source_pdf = lambda _contract: signing_pdf

            with self.assertRaisesRegex(RuntimeError, "transient fact failure"):
                handler.contract_generate("POST", "T-AUTH-001")
            handler.contract_generate("POST", "T-AUTH-001")
            handler.contract_generate("POST", "T-AUTH-001")

            self.assertEqual(handler.response_status, 303)
            self.assertEqual(database.one("SELECT COUNT(*) n FROM contracts")["n"], 1)
            self.assertEqual(database.one("SELECT COUNT(*) n FROM student_documents")["n"], 1)
            contract = database.one("SELECT * FROM contracts")
            frozen = contract_domain.snapshot_payload(contract["standalone_data_json"])
            self.assertEqual(frozen["generation_authority"], authority.AUTHORITY_ID)
            self.assertEqual(frozen["commercial"]["standard_service_fee"], "2400.00")
            self.assertEqual(frozen["commercial"]["gst_amount"], "0.00")
            self.assertEqual(frozen["commercial"]["final_fee_amount"], "2400.00")
            self.assertEqual(events, [(contract["id"], "contract.created")] * 3)
            self.assertEqual(database.one("SELECT COUNT(*) n FROM contract_signing_links")["n"], 1)
            signing_link = database.one("SELECT * FROM contract_signing_links")
            self.assertEqual(signing_link["contract_id"], contract["id"])
            self.assertEqual(signing_link["status"], "active")
            self.assertEqual(
                database.one("SELECT contract_status FROM contracts WHERE id=?", (contract["id"],))["contract_status"],
                "Sent",
            )
            location = dict(handler.sent_headers)["Location"]
            redirect = urlparse(location)
            query = parse_qs(redirect.query)
            self.assertEqual(query["toast_copy"], [signing_link["signing_url"]])
            self.assertEqual(
                query["toast_message"],
                ["✓ 新合同已生成\n✓ 新签署链接已生成"],
            )
            self.assertEqual(
                query["toast_copy_success"],
                ["✓ 新合同已生成\n✓ 新签署链接已生成\n✓ 链接已复制"],
            )
            self.assertNotIn("toast_copy_failure", query)
        finally:
            app_module.db, app_module.nas = old_db, old_nas
            app_module.emit_contract_business_event = old_emit
            app_module.validate_pdf_contains_contract_body = old_validate

    def test_application_route_preserves_contract_when_automatic_signing_link_fails(self):
        import app as app_module

        class FakeNas:
            def student_folder(self, student):
                return Path(student["folder_path"]).resolve()

            def require_category_access(self, student, category, *, write=False):
                target = self.student_folder(student) / category
                target.mkdir(parents=True, exist_ok=True)
                return target

        database = self._persistence_database()
        database.set_setting(
            app_module.CONTRACT_GENERATION_AUTHORITY_SETTING,
            app_module.CONTRACT_GENERATION_AUTHORITY_PUBLISHED_V2,
            1,
        )
        student_folder = self.root / "route-link-failure-student"
        student_folder.mkdir()
        database.execute(
            "UPDATE students SET folder_path=?,form_data_json=? WHERE student_id=?",
            (str(student_folder), json.dumps(self._student()), "T-AUTH-001"),
        )
        old_db, old_nas = app_module.db, app_module.nas
        old_emit = app_module.emit_contract_business_event
        old_log = app_module.append_service_log
        logged = []
        try:
            app_module.db = database
            app_module.nas = FakeNas()
            app_module.emit_contract_business_event = lambda *_args, **_kwargs: None
            def failing_log(*args, **kwargs):
                logged.append((args, kwargs))
                raise OSError("simulated log directory permission failure")

            app_module.append_service_log = failing_log
            handler = object.__new__(app_module.CRMHandler)
            handler.form = {
                "contract_type": "NewInternationalStudentContract",
                "contract_date": "2026-08-20",
                "generation_request_id": "route-link-failure-001",
                "parent_signer_source": "father",
                "parent_signer_name": "Father Only",
                "parent_signer_relationship": "父亲",
                "parent_phone": "0411111111",
                "parent_email": "father@example.test",
                "service_scope": list(
                    contract_domain.contract_type_config(
                        "NewInternationalStudentContract"
                    ).default_services
                ),
                "service_fee_currency": "国内银行",
            }
            handler.user = {
                "id": 1, "username": "admin", "role_name": "Admin",
                "permissions_json": json.dumps(["*"]),
            }
            handler.client_address = ("127.0.0.1", 12345)
            handler.headers = {}
            handler.sent_headers = []
            handler.send_response = lambda status: setattr(handler, "response_status", status)
            handler.send_header = lambda key, value: handler.sent_headers.append((key, value))
            handler.end_headers = lambda: None
            handler.contract_signing_create_link = lambda *_args, **_kwargs: (_ for _ in ()).throw(
                RuntimeError("simulated signing PDF conversion failure")
            )

            handler.contract_generate("POST", "T-AUTH-001")

            self.assertEqual(handler.response_status, 303)
            self.assertEqual(database.one("SELECT COUNT(*) n FROM contracts")["n"], 1)
            self.assertEqual(database.one("SELECT COUNT(*) n FROM student_documents")["n"], 1)
            self.assertEqual(database.one("SELECT COUNT(*) n FROM contract_signing_links")["n"], 0)
            contract = database.one("SELECT id,contract_status FROM contracts")
            self.assertEqual(contract["contract_status"], "Draft")
            location = dict(handler.sent_headers)["Location"]
            query = parse_qs(urlparse(location).query)
            self.assertEqual(query["toast_kind"], ["warn"])
            self.assertEqual(query["toast_title"], ["合同已生成，但签署链接生成失败"])
            self.assertIn("合同已保留", query["toast_message"][0])
            self.assertNotIn("toast_copy", query)
            self.assertEqual(logged[0][1]["contract_id"], contract["id"])
        finally:
            app_module.db, app_module.nas = old_db, old_nas
            app_module.emit_contract_business_event = old_emit
            app_module.append_service_log = old_log

    def test_authority_binding_uses_frozen_representative_when_live_setting_is_blank(self):
        import app as app_module

        database = self._persistence_database()
        database.set_setting("company_authorised_person", "", 1)
        student = database.one("SELECT * FROM students WHERE student_id='T-AUTH-001'")
        connection = database.connect()
        try:
            package = authority.load_published_template(
                connection, "NewInternationalStudentContract"
            )
        finally:
            connection.close()
        form = {
            "contract_type": "NewInternationalStudentContract",
            "contract_date": "2026-08-20",
            "generation_request_id": "blank-company-representative",
            "parent_signer_source": "father",
            "parent_signer_name": "Father Only",
            "parent_signer_relationship": "父亲",
            "parent_phone": "0411111111",
            "parent_email": "father@example.test",
            "service_scope": list(
                contract_domain.contract_type_config(
                    "NewInternationalStudentContract"
                ).default_services
            ),
            "service_fee_currency": "国内银行",
        }
        user = {
            "id": 1, "username": "admin", "role_name": "Admin",
            "permissions_json": json.dumps(["*"]),
        }
        old_db = app_module.db
        try:
            app_module.db = database
            bound, _rendered = app_module.authority_contract_bindings(
                student, form, user, package, "AUTH-BLANK-REP"
            )
            self.assertEqual(
                bound["contract_instance"]["company_authorised_person"],
                "Sonny Zhang",
            )
            self.assertEqual(database.one("SELECT COUNT(*) n FROM contracts")["n"], 0)
        finally:
            app_module.db = old_db

    def test_application_route_fails_closed_before_persistence_when_published_is_incomplete(self):
        import app as app_module

        database = self._persistence_database()
        database.set_setting(
            app_module.CONTRACT_GENERATION_AUTHORITY_SETTING,
            app_module.CONTRACT_GENERATION_AUTHORITY_PUBLISHED_V2,
            1,
        )
        with database.transaction() as connection:
            identity = connection.execute(
                "SELECT * FROM contract_template_identities "
                "WHERE template_type='NewInternationalStudentContract'"
            ).fetchone()
            current = connection.execute(
                "SELECT * FROM contract_template_versions WHERE id=?",
                (identity["current_published_version_id"],),
            ).fetchone()
            invalid = json.loads(current["snapshot_json"])
            invalid["commercial"]["standard_service_fee"] = "0.00"
            invalid["commercial"]["service_fee"] = "0.00"
            invalid_id = connection.execute(
                """INSERT INTO contract_template_versions(
                   identity_id,version_number,status,source_version_id,revision,snapshot_json,
                   content_hash,immutable,created_by,created_at,published_by,published_at
                ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?)""",
                (
                    identity["id"], 99, "Draft", current["id"], 1,
                    editor._canonical_json(invalid), editor._hash(invalid), 0, 1,
                    "2026-08-20", None, None,
                ),
            ).lastrowid
            connection.execute(
                """UPDATE contract_template_versions
                   SET status='Published',immutable=1,published_by=1,published_at='2026-08-20'
                   WHERE id=?""",
                (invalid_id,),
            )
            connection.execute(
                "UPDATE contract_template_identities SET current_published_version_id=? WHERE id=?",
                (invalid_id, identity["id"]),
            )
        old_db = app_module.db
        try:
            app_module.db = database
            handler = object.__new__(app_module.CRMHandler)
            handler.form = {
                "contract_type": "NewInternationalStudentContract",
                "generation_request_id": "route-authority-missing",
            }
            handler.user = {
                "id": 1, "username": "admin", "role_name": "Admin",
                "permissions_json": json.dumps(["*"]),
            }
            handler.client_address = ("127.0.0.1", 12345)
            handler.headers = {}
            handler.sent_headers = []
            handler.send_response = lambda status: setattr(handler, "response_status", status)
            handler.send_header = lambda key, value: handler.sent_headers.append((key, value))
            handler.end_headers = lambda: None
            handler.contract_signing_create_link = lambda *_args, **_kwargs: self.fail(
                "Signing Link must not be attempted when contract creation fails"
            )
            handler.contract_generate("POST", "T-AUTH-001")
            self.assertEqual(handler.response_status, 303)
            self.assertEqual(database.one("SELECT COUNT(*) n FROM contracts")["n"], 0)
            self.assertEqual(database.one("SELECT COUNT(*) n FROM student_documents")["n"], 0)
        finally:
            app_module.db = old_db

    def test_shared_desktop_mobile_form_uses_published_catalog_and_admin_controls(self):
        import app as app_module

        database = self._persistence_database(all_types=True)
        database.set_setting(
            app_module.CONTRACT_GENERATION_AUTHORITY_SETTING,
            app_module.CONTRACT_GENERATION_AUTHORITY_PUBLISHED_V2,
            1,
        )
        student = database.one("SELECT * FROM students WHERE student_id='T-AUTH-001'")
        old_db = app_module.db
        try:
            app_module.db = database
            handler = object.__new__(app_module.CRMHandler)
            handler.user = {
                "id": 1, "username": "admin", "role_name": "Admin",
                "permissions_json": json.dumps(["*"]),
            }
            handler.csrf = "csrf-test"
            handler._student_lifecycle_record_contexts = {}
            html = handler.contract_panel(student, [], [])

            self.assertIn('name="generation_request_id"', html)
            self.assertIn("Published V2唯一权威", html)
            self.assertIn("AUD 2400.00", html)
            self.assertIn("GST金额", html)
            self.assertIn("AUD 0.00", html)
            self.assertIn("最终合同总额", html)
            self.assertIn("Admin折扣", html)
            self.assertIn("特殊分期", html)
            self.assertIn('data-contract-service-labels=', html)
            self.assertIn('data-contract-published-commercials=', html)
            self.assertIn('data-contract-published-payment-method', html)
            self.assertIn('data-contract-published-standard-fee', html)
            self.assertIn('data-contract-published-net-fee', html)
            self.assertIn('data-contract-published-gst-amount', html)
            self.assertIn('data-contract-published-final-total', html)
            self.assertIn('data-contract-published-refund', html)
            for package in app_module.published_contract_packages().values():
                commercial = package.snapshot["commercial"]
                self.assertIn(str(commercial["refund_policy"]), html)
            self.assertIn("NewInternationalStudentContract", html)
            self.assertIn("TransferApplicationContract", html)
            self.assertIn("LocalStudentApplicationContract", html)
            self.assertIn('name="parent_signer_name"', html)
            self.assertIn('name="parent_signer_name" placeholder="请输入监护人姓名" required value="Mother Other" readonly', html)
            self.assertNotIn('name="refund_policy_type"', html)
            self.assertNotIn('name="consulting_and_visa_service_fee"', html)
            self.assertNotIn("legacy_v6", html)
            self.assertNotIn("@media (min-width", html)
        finally:
            app_module.db = old_db


if __name__ == "__main__":
    unittest.main()
