"""Versioned three-type contract-template editor domain service.

The module is deliberately independent from HTTP, Production storage, signing,
stamping and Workflow.  Production generation continues to use
``contract_domain.py`` until a separately authorised cutover.
"""

from __future__ import annotations

import hashlib
import html
import json
import re
import sqlite3
import uuid
from contextlib import contextmanager
from copy import deepcopy
from datetime import datetime, timedelta, timezone
from html.parser import HTMLParser
from pathlib import Path
from typing import Iterable, Mapping

import contract_domain
import contract_display_bindings
from contract_master_layout import (
    EDITABLE_MASTER_SLOTS,
    HEADINGS as MASTER_HEADINGS,
    MASTER_LAYOUT_POLICY,
    MASTER_SLOT_ORDER,
    MasterLayoutError,
    build_master_render_model,
    is_master_snapshot,
    master_block_parent,
    render_master_docx_bytes,
    render_master_pdf_bytes,
    render_master_signing_html,
    validate_master_layout_snapshot,
)
from student_doc import (
    DEFAULT_PARTY_A_ADDRESS,
    DEFAULT_PARTY_A_EMAIL,
    DEFAULT_PARTY_A_NAME,
    DEFAULT_PARTY_A_PHONE,
)


class ContractTemplateEditorError(ValueError):
    pass


class TemplateRuleError(ContractTemplateEditorError):
    pass


class EditLockConflict(ContractTemplateEditorError):
    pass


class StaleRevision(ContractTemplateEditorError):
    pass


class ImmutableVersionError(ContractTemplateEditorError):
    pass


class PublishBlocked(ContractTemplateEditorError):
    pass


class PermissionDenied(ContractTemplateEditorError):
    pass


REQUIRED_LEGAL_KINDS = {"privacy", "refund", "termination", "dispute"}
PROFILE_KEYS = {
    "三、服务费用与第三方费用": "profile_3",
    "四、甲方权利义务": "profile_4",
    "五、乙方权利义务": "profile_5",
    "六、退费及违约责任": "profile_6",
    "七、不可抗力、结果说明及免责": "profile_7",
    "九、合同补充、变更和生效": "profile_9",
}
VARIABLE_PATTERN = re.compile(r"\{\{[a-zA-Z0-9_]+\}\}")


def _now() -> str:
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat()


def _canonical_json(value: object) -> str:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"))


def _hash(value: object) -> str:
    return hashlib.sha256(_canonical_json(value).encode("utf-8")).hexdigest()


def _file_hash(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _visible_paragraphs(value: object) -> list[dict[str, list[str]]]:
    """Normalize visible text without losing paragraph or line boundaries."""
    text = str(value or "").replace("\r\n", "\n").replace("\r", "\n")
    chunks = re.split(r"\n[ \t]*\n", text)
    return [{"lines": chunk.split("\n")} for chunk in chunks]


def _visible_cell(value: object) -> dict[str, list[str]]:
    return {"lines": str(value or "").replace("\r\n", "\n").replace("\r", "\n").split("\n")}


def build_visible_render_model(
    snapshot: Mapping[str, object], replacements: Mapping[str, str],
    signer: Mapping[str, object],
) -> dict[str, object]:
    """Build the one normalized visible model consumed by every renderer."""
    content: list[dict[str, object]] = []
    commercial = dict(snapshot.get("commercial") or {})
    commercial_rows = [
        ("公司名称", commercial.get("company_name", "")),
        ("ABN / ACN", commercial.get("company_abn", "")),
        ("公司地址", commercial.get("company_address", "")),
        ("公司联系方式", commercial.get("company_contact", "")),
    ]
    if commercial.get("commercial_schema") == "authority_v2":
        commercial_rows.extend([
            ("标准服务费（不含GST）", f"{commercial.get('currency', '')} {commercial.get('standard_service_fee_ex_gst', commercial.get('standard_service_fee', ''))}".strip()),
            ("批准折扣（不含GST）", f"{commercial.get('currency', '')} {commercial.get('discount_amount_ex_gst', commercial.get('discount_amount', '0.00'))}".strip()),
            ("折后金额（不含GST）", f"{commercial.get('currency', '')} {commercial.get('net_service_fee_ex_gst', '')}".strip()),
            ("GST处理", commercial.get("gst_display", commercial.get("gst_treatment", ""))),
            ("GST税率", commercial.get("gst_rate", "")),
            ("GST金额", f"{commercial.get('currency', '')} {commercial.get('gst_amount', '')}".strip()),
            ("合同总额", f"{commercial.get('currency', '')} {commercial.get('final_total', commercial.get('final_fee_amount', ''))}".strip()),
            ("付款方式", commercial.get("payment_method", "")),
            ("付款计划", commercial.get("payment_plan_display", commercial.get("payment_terms", ""))),
            ("退款配置", commercial.get("refund_policy", "")),
        ])
        refusal = commercial.get("visa_refusal_refund")
        if isinstance(refusal, Mapping):
            commercial_rows.extend([
                ("拒签保留金额（不含GST）", f"{commercial.get('currency', '')} {refusal.get('retained_ex_gst', '')}".strip()),
                ("拒签GST调整", f"{commercial.get('currency', '')} {refusal.get('gst_adjustment', '')}".strip()),
                ("拒签实际保留总额", f"{commercial.get('currency', '')} {refusal.get('retained_total', '')}".strip()),
                ("按已支付额计算的拒签退款", f"{commercial.get('currency', '')} {refusal.get('actual_refund_amount', '')}".strip()),
            ])
    else:
        commercial_rows.extend([
            ("服务费用", f"{commercial.get('currency', '')} {commercial.get('service_fee', '')}".strip()),
            ("付款条款", commercial.get("payment_terms", "")),
            ("退款配置", commercial.get("refund_policy", "")),
        ])
    commercial_node = {
        "kind": "table",
        "key": "commercial",
        "title": "合同主体、费用与付款配置",
        "rows": [
            {"cells": [_visible_cell(label), _visible_cell(value)]}
            for label, value in commercial_rows
        ],
    }
    commercial_inserted = False
    for block in snapshot.get("blocks") or []:
        if not block.get("enabled", True):
            continue
        body = str(block.get("body") or "")
        for variable, value in replacements.items():
            body = body.replace(variable, value)
        content.append({
            "kind": "section",
            "key": str(block.get("key") or ""),
            "title": str(block.get("title") or ""),
            "paragraphs": _visible_paragraphs(body),
        })
        if not commercial_inserted:
            content.append(commercial_node)
            commercial_inserted = True
    if not commercial_inserted:
        content.append(commercial_node)

    footer_lines: list[str] = []
    for footer_value in (
        commercial.get("company_name"), commercial.get("company_contact"),
    ):
        footer_lines.extend(
            str(footer_value or "").strip().replace("\r\n", "\n").replace("\r", "\n").split("\n")
        )
    header_lines = (
        str(snapshot.get("contract_title") or "").strip()
        .replace("\r\n", "\n").replace("\r", "\n").split("\n")
    )
    return {
        "schema_version": 1,
        "header": [{"lines": header_lines}],
        "content": content,
        "footer": [{"lines": footer_lines}],
        "signature_region": {
            "title": "签署确认",
            "paragraphs": [{
                "lines": [
                    f"学生：{replacements.get('{{student_name}}', '')}",
                    f"签署人：{signer.get('name', '')}",
                    f"关系：{signer.get('relationship', '')}",
                    f"电话：{signer.get('phone', '')}",
                    f"邮箱：{signer.get('email', '')}",
                ]
            }],
        },
    }


def _derive_render_expectation(
    snapshot: Mapping[str, object], input_payload: Mapping[str, object],
) -> dict[str, object]:
    """Rebuild the canonical signer, services and visible result from bound input."""
    required_keys = {
        "source_type", "student_snapshot", "signer_source",
        "signer_overrides", "selected_services",
    }
    if not isinstance(input_payload, dict) or set(input_payload) != required_keys:
        raise TemplateRuleError("模拟输入结构无效。")
    source_type = input_payload.get("source_type")
    student_snapshot = input_payload.get("student_snapshot")
    signer_source = input_payload.get("signer_source")
    signer_overrides = input_payload.get("signer_overrides")
    selected_services = input_payload.get("selected_services")
    if source_type not in {"system", "student_readonly"}:
        raise TemplateRuleError("模拟资料来源无效。")
    if not isinstance(student_snapshot, dict) or not isinstance(signer_overrides, dict):
        raise TemplateRuleError("模拟学生或签署人资料无效。")
    if not isinstance(signer_source, str) or not isinstance(selected_services, list):
        raise TemplateRuleError("模拟签署人或服务范围无效。")
    if any(not isinstance(value, str) for value in selected_services):
        raise TemplateRuleError("模拟服务范围无效。")
    selected = list(dict.fromkeys(selected_services))
    if selected != selected_services:
        raise TemplateRuleError("模拟服务范围包含重复项。")
    allowed = {
        str(item["service_id"])
        for item in snapshot.get("services") or []
        if item.get("enabled")
    }
    if any(value not in allowed for value in selected):
        raise TemplateRuleError("服务范围与合同类型不匹配。")
    try:
        signer = contract_domain.resolve_signer(
            signer_source, student_snapshot, (), overrides=signer_overrides
        )
    except (contract_domain.ContractDomainError, TypeError, ValueError) as exc:
        raise TemplateRuleError("模拟签署人资料无效。") from exc
    replacements = {
        "{{student_name}}": str(student_snapshot.get("student_name") or "").strip(),
        "{{signer_name}}": signer["name"],
        "{{signer_phone}}": signer["phone"],
        "{{signer_email}}": signer["email"],
        "{{service_scope}}": contract_domain.service_scope_text(selected),
    }
    if is_master_snapshot(snapshot):
        student_name = str(student_snapshot.get("student_name") or student_snapshot.get("student_full_name") or "").strip()
        preview_instance = {
            "contract_type": snapshot.get("template_type", ""),
            "contract_uid": "PREVIEW",
            "contract_date": "PREVIEW",
            "student_full_name": student_name,
            "student_date_of_birth": student_snapshot.get("date_of_birth", ""),
            "student_id_or_passport_number": student_snapshot.get("student_id", ""),
            "student_current_year_level": student_snapshot.get("current_year_level", ""),
            "student_applying_year_level": student_snapshot.get("applying_year_level", ""),
            "student_planned_entry_year": student_snapshot.get("preferred_intake_year", ""),
            "student_planned_entry_term": student_snapshot.get("preferred_intake_term", ""),
            "intake_year": student_snapshot.get("preferred_intake_year", ""),
            "intake_term": student_snapshot.get("preferred_intake_term", ""),
            "parent_address": student_snapshot.get("home_address", ""),
            "company_authorised_person": (
                (snapshot.get("commercial") or {}).get("company_authorised_person")
                or student_snapshot.get("company_authorised_person", "")
            ),
        }
        preview_commercial = dict(snapshot.get("commercial") or {})
        preview_commercial.setdefault("payment_method", "PREVIEW")
        preview_commercial.setdefault("discount_amount_ex_gst", "0.00")
        preview_commercial.setdefault(
            "net_service_fee_ex_gst",
            preview_commercial.get("standard_service_fee_ex_gst")
            or preview_commercial.get("standard_service_fee")
            or preview_commercial.get("service_fee")
            or "0.00",
        )
        preview_commercial.setdefault("gst_amount", preview_commercial.get("default_gst_amount", "0.00"))
        preview_commercial.setdefault("final_total", preview_commercial.get("default_final_total", "0.00"))
        display_bindings = contract_display_bindings.build_frozen_display_bindings(
            snapshot.get("template_type"),
            student_snapshot=student_snapshot,
            signer=signer,
            commercial=preview_commercial,
            contract_instance=preview_instance,
        )
        blocking_codes = {
            "mandatory_display_binding_current_grade_missing",
            "mandatory_display_binding_application_grade_missing",
            "mandatory_display_binding_planned_entry_year_missing",
        }
        for issue in contract_display_bindings.required_value_issues(display_bindings):
            if issue["code"] in blocking_codes:
                raise TemplateRuleError(issue["message"])
        visible_model = build_master_render_model(
            snapshot,
            student_name=student_name,
            signer=signer,
            selected_services=selected,
            contract_instance=preview_instance,
            display_bindings=display_bindings,
        )
    else:
        visible_model = build_visible_render_model(snapshot, replacements, signer)
    authority_content = _authority_content_from_model(visible_model)
    semantic_trace = _semantic_trace_from_model(visible_model)
    semantic_json = _canonical_json(semantic_trace)
    return {
        "source_type": source_type,
        "signer": signer,
        "selected_services": selected,
        "visible_model": visible_model,
        "authority_content": authority_content,
        "authority_content_hash": hashlib.sha256(authority_content.encode("utf-8")).hexdigest(),
        "semantic_trace": semantic_trace,
        "semantic_json": semantic_json,
        "semantic_hash": hashlib.sha256(semantic_json.encode("utf-8")).hexdigest(),
    }


def _authority_content_from_model(model: Mapping[str, object]) -> str:
    sections: list[str] = []
    for paragraph in model["header"]:
        sections.append("\n".join(paragraph["lines"]))
    for node in model["content"]:
        lines = [str(node["title"])]
        if node["kind"] == "section":
            for paragraph in node["paragraphs"]:
                lines.append("\n".join(paragraph["lines"]))
        else:
            for row in node["rows"]:
                cells = ["\n".join(cell["lines"]) for cell in row["cells"]]
                lines.append("：".join(cells))
        sections.append("\n".join(lines))
    signature = model["signature_region"]
    signature_lines = [str(signature["title"])]
    signature_lines.extend("\n".join(paragraph["lines"]) for paragraph in signature["paragraphs"])
    sections.append("\n".join(signature_lines))
    sections.extend("\n".join(paragraph["lines"]) for paragraph in model["footer"])
    return "\n\n".join(sections)


def _normalize_semantic_trace(trace: Mapping[str, object]) -> dict[str, object]:
    """Keep logical regions/boundaries while removing renderer-only whitespace."""
    normalized = deepcopy(trace)
    for lines in _trace_line_lists(normalized):
        lines[:] = [re.sub(r"\s+", "", str(line)) for line in lines]
    return normalized


def _semantic_trace_from_model(model: Mapping[str, object]) -> dict[str, object]:
    """Project the visible model into renderer-neutral visible structures."""
    content: list[dict[str, object]] = []
    for node in model["content"]:
        content.append({"kind": "heading", "lines": [str(node["title"])]})
        if node["kind"] == "section":
            content.extend({"kind": "paragraph", "lines": list(paragraph["lines"])}
                           for paragraph in node["paragraphs"])
        else:
            content.append({
                "kind": "table",
                "rows": [
                    {"cells": [{"lines": list(cell["lines"])} for cell in row["cells"]]}
                    for row in node["rows"]
                ],
            })
    signature = model["signature_region"]
    content.append({"kind": "heading", "lines": [str(signature["title"])]})
    content.extend({"kind": "paragraph", "lines": list(paragraph["lines"])}
                   for paragraph in signature["paragraphs"])
    return _normalize_semantic_trace({
        "schema_version": 1,
        "header": [{"lines": list(paragraph["lines"])} for paragraph in model["header"]],
        "content": content,
        "footer": [{"lines": list(paragraph["lines"])} for paragraph in model["footer"]],
    })


DOCX_FONT_NAME = "Noto Sans SC"
DOCX_FONT_PART = "word/fonts/NotoSansSC-Regular.odttf"
DOCX_FONT_RELATIONSHIP_ID = "rIdAAIENotoSansSCRegular"
DOCX_FONT_RELATIONSHIP_TYPE = (
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/font"
)


def _obfuscated_docx_font(font_bytes: bytes) -> tuple[bytes, str]:
    font_key = uuid.UUID(hashlib.sha256(font_bytes).hexdigest()[:32])
    key = bytes.fromhex(font_key.hex)[::-1]
    obfuscated = bytearray(font_bytes)
    for index in range(min(32, len(obfuscated))):
        obfuscated[index] ^= key[index % len(key)]
    return bytes(obfuscated), "{" + str(font_key).upper() + "}"


def _embed_docx_font(package_bytes: bytes) -> bytes:
    import io
    import zipfile
    from lxml import etree

    font_bytes = PDF_FONT_FILE.read_bytes()
    embedded_font, font_key = _obfuscated_docx_font(font_bytes)
    with zipfile.ZipFile(io.BytesIO(package_bytes), "r") as source:
        names = source.namelist()
        if len(names) != len(set(names)):
            raise TemplateRuleError("DOCX成品包包含重复部件。")
        parts = {name: source.read(name) for name in names}

    word_namespace = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    office_rel_namespace = (
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    )
    package_rel_namespace = "http://schemas.openxmlformats.org/package/2006/relationships"
    content_type_namespace = "http://schemas.openxmlformats.org/package/2006/content-types"

    def qualified(namespace, name):
        return f"{{{namespace}}}{name}"

    font_table = etree.fromstring(parts["word/fontTable.xml"])
    for existing in font_table.findall(qualified(word_namespace, "font")):
        if existing.get(qualified(word_namespace, "name")) == DOCX_FONT_NAME:
            font_table.remove(existing)
    font = etree.SubElement(
        font_table, qualified(word_namespace, "font"),
        {qualified(word_namespace, "name"): DOCX_FONT_NAME},
    )
    etree.SubElement(
        font, qualified(word_namespace, "charset"),
        {qualified(word_namespace, "val"): "86"},
    )
    etree.SubElement(
        font, qualified(word_namespace, "family"),
        {qualified(word_namespace, "val"): "swiss"},
    )
    etree.SubElement(
        font, qualified(word_namespace, "pitch"),
        {qualified(word_namespace, "val"): "variable"},
    )
    etree.SubElement(
        font, qualified(word_namespace, "embedRegular"),
        {
            qualified(office_rel_namespace, "id"): DOCX_FONT_RELATIONSHIP_ID,
            qualified(word_namespace, "fontKey"): font_key,
        },
    )
    parts["word/fontTable.xml"] = etree.tostring(
        font_table, xml_declaration=True, encoding="UTF-8", standalone=True
    )

    relationship_path = "word/_rels/fontTable.xml.rels"
    if relationship_path in parts:
        relationships = etree.fromstring(parts[relationship_path])
    else:
        relationships = etree.Element(
            qualified(package_rel_namespace, "Relationships"),
            nsmap={None: package_rel_namespace},
        )
    for relationship in list(relationships):
        if (
            relationship.get("Id") == DOCX_FONT_RELATIONSHIP_ID
            or relationship.get("Target") == "fonts/NotoSansSC-Regular.odttf"
        ):
            relationships.remove(relationship)
    etree.SubElement(
        relationships,
        qualified(package_rel_namespace, "Relationship"),
        {
            "Id": DOCX_FONT_RELATIONSHIP_ID,
            "Type": DOCX_FONT_RELATIONSHIP_TYPE,
            "Target": "fonts/NotoSansSC-Regular.odttf",
        },
    )
    parts[relationship_path] = etree.tostring(
        relationships, xml_declaration=True, encoding="UTF-8", standalone=True
    )

    settings = etree.fromstring(parts["word/settings.xml"])
    for name in ("embedTrueTypeFonts", "saveSubsetFonts"):
        tag = qualified(word_namespace, name)
        if settings.find(tag) is None:
            settings.append(etree.Element(tag))
    parts["word/settings.xml"] = etree.tostring(
        settings, xml_declaration=True, encoding="UTF-8", standalone=True
    )

    content_types = etree.fromstring(parts["[Content_Types].xml"])
    if not any(
        item.get("Extension") == "odttf"
        for item in content_types.findall(qualified(content_type_namespace, "Default"))
    ):
        etree.SubElement(
            content_types,
            qualified(content_type_namespace, "Default"),
            {
                "Extension": "odttf",
                "ContentType": (
                    "application/vnd.openxmlformats-officedocument.obfuscatedFont"
                ),
            },
        )
    parts["[Content_Types].xml"] = etree.tostring(
        content_types, xml_declaration=True, encoding="UTF-8", standalone=True
    )
    parts[DOCX_FONT_PART] = embedded_font

    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as target:
        for name in sorted(parts):
            information = zipfile.ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
            information.compress_type = (
                zipfile.ZIP_STORED if name == DOCX_FONT_PART else zipfile.ZIP_DEFLATED
            )
            information.external_attr = 0o600 << 16
            target.writestr(information, parts[name])
    return output.getvalue()


def _build_docx_artifact_bytes(model: Mapping[str, object]) -> bytes:
    import io
    from docx import Document
    from docx.oxml.ns import qn

    def write_lines(paragraph, lines):
        paragraph.clear()
        for index, line in enumerate(lines):
            if index:
                paragraph.add_run().add_break()
            paragraph.add_run(str(line))

    document = Document()
    for style_name in ("Normal", "Heading 2", "Header", "Footer"):
        style = document.styles[style_name]
        style.font.name = DOCX_FONT_NAME
        for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
            style.element.rPr.rFonts.set(qn(f"w:{attribute}"), DOCX_FONT_NAME)
    header = document.sections[0].header
    write_lines(header.paragraphs[0], model["header"][0]["lines"])
    footer = document.sections[0].footer
    write_lines(footer.paragraphs[0], model["footer"][0]["lines"])
    for node in model["content"]:
        document.add_heading(str(node["title"]), level=2)
        if node["kind"] == "section":
            for paragraph_model in node["paragraphs"]:
                write_lines(document.add_paragraph(), paragraph_model["lines"])
        else:
            rows = node["rows"]
            table = document.add_table(rows=len(rows), cols=2)
            table.style = "Table Grid"
            for row_index, row_model in enumerate(rows):
                for cell_index, cell_model in enumerate(row_model["cells"]):
                    write_lines(
                        table.cell(row_index, cell_index).paragraphs[0],
                        cell_model["lines"],
                    )
    signature = model["signature_region"]
    document.add_heading(str(signature["title"]), level=2)
    for paragraph_model in signature["paragraphs"]:
        write_lines(document.add_paragraph(), paragraph_model["lines"])
    output = io.BytesIO()
    document.save(output)
    return _embed_docx_font(output.getvalue())


def _docx_layout_fingerprint(source: Path | bytes) -> str:
    import io
    import zipfile
    from lxml import etree

    raw = source.read_bytes() if isinstance(source, Path) else bytes(source)
    canonical_parts: dict[str, dict[str, object]] = {}
    with zipfile.ZipFile(io.BytesIO(raw), "r") as package:
        names = package.namelist()
        if len(names) != len(set(names)):
            raise TemplateRuleError("DOCX成品包包含重复部件。")
        for name in sorted(names):
            if name.startswith("docProps/"):
                continue
            value = package.read(name)
            if name.endswith((".xml", ".rels")) or name == "[Content_Types].xml":
                try:
                    root = etree.fromstring(value)
                    value = etree.tostring(
                        root, method="c14n", exclusive=False, with_comments=False
                    )
                except etree.XMLSyntaxError as exc:
                    raise TemplateRuleError("DOCX成品XML结构无效。") from exc
            canonical_parts[name] = {
                "bytes": len(value),
                "sha256": hashlib.sha256(value).hexdigest(),
            }
    return _hash(canonical_parts)


def _extract_docx_semantic_trace(path: Path) -> dict[str, object]:
    import base64
    import gzip
    from docx import Document
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    import zipfile
    document = Document(path)
    with zipfile.ZipFile(path, "r") as package:
        packed_xml = package.read("customXml/aaieMasterModel.xml") if "customXml/aaieMasterModel.xml" in package.namelist() else b""
    if packed_xml:
        try:
            from lxml import etree
            packed = str(etree.fromstring(packed_xml).text or "")
            model = json.loads(gzip.decompress(base64.b64decode(packed)).decode("utf-8"))
        except Exception as exc:
            raise TemplateRuleError("保存后DOCX主版式语义绑定无效。") from exc
        defaults = document.styles.element.find(qn("w:docDefaults"))
        reject_targets = [defaults, document._element.body]
        reject_targets.extend(
            document.styles[style_name].element.rPr
            for style_name in ("Normal", "Heading 2", "Header", "Footer")
        )
        for properties in reject_targets:
            if properties is not None:
                for node in properties.iter():
                    if node.tag in {qn("w:vanish"), qn("w:webHidden"), qn("w:specVanish")}:
                        raise TemplateRuleError("保存后DOCX成品包含隐藏文本。")
                    if node.tag in {qn("w:sz"), qn("w:szCs")} and int(node.get(qn("w:val")) or 0) < 12:
                        raise TemplateRuleError("保存后DOCX成品字号不可见。")
                    if node.tag == qn("w:color") and str(node.get(qn("w:val")) or "").upper() == "FFFFFF":
                        raise TemplateRuleError("保存后DOCX成品颜色不可见。")
                    if node.tag == qn("w:ind"):
                        for attr in ("left", "right", "start", "end"):
                            raw = node.get(qn(f"w:{attr}"))
                            if raw is not None and abs(int(raw)) > 12000:
                                raise TemplateRuleError("保存后DOCX成品段落位于页面外。")
        visible = "\n".join(paragraph.text for paragraph in document.paragraphs)
        visible += "\n" + "\n".join(
            cell.text for table in document.tables for row in table.rows for cell in row.cells
        )
        if len(document.inline_shapes) < 1 or not all(
            value in visible for value in ("一、合同双方", "十一、签署", "AAIE / 甲方（签字 / 盖章）", "委托方（乙方签字）")
        ):
            raise TemplateRuleError("保存后DOCX主版式物理锚点无效。")
        return _semantic_trace_from_model(model)

    def reject_unsafe_properties(properties, label):
        if properties is None:
            return
        for node in properties.iter():
            if node.tag in {
                qn("w:vanish"), qn("w:webHidden"), qn("w:specVanish"),
            }:
                raise TemplateRuleError(f"保存后DOCX成品{label}包含隐藏文本。")
            if node.tag in {qn("w:sz"), qn("w:szCs")}:
                try:
                    if int(node.get(qn("w:val"))) < 12:
                        raise TemplateRuleError(f"保存后DOCX成品{label}字号不可见。")
                except (TypeError, ValueError) as exc:
                    raise TemplateRuleError(f"保存后DOCX成品{label}字号无效。") from exc
            if node.tag == qn("w:color") and (
                str(node.get(qn("w:val")) or "").upper() == "FFFFFF"
                or str(node.get(qn("w:themeColor")) or "").lower().startswith("background")
            ):
                raise TemplateRuleError(f"保存后DOCX成品{label}颜色不可见。")
            if node.tag == qn("w:w"):
                try:
                    if int(node.get(qn("w:val"))) < 50:
                        raise TemplateRuleError(f"保存后DOCX成品{label}文本缩放不可见。")
                except (TypeError, ValueError) as exc:
                    raise TemplateRuleError(f"保存后DOCX成品{label}文本缩放无效。") from exc

    defaults = document.styles.element.find(qn("w:docDefaults"))
    reject_unsafe_properties(defaults, "默认样式")

    for style_name in ("Normal", "Heading 2", "Header", "Footer"):
        style = document.styles[style_name]
        properties = style.element.rPr
        reject_unsafe_properties(properties, "样式")
        if style.font.hidden is True or getattr(style.font, "spec_vanish", None) is True:
            raise TemplateRuleError("保存后DOCX成品样式包含隐藏文本。")
        if style.font.size is not None and style.font.size.pt < 6:
            raise TemplateRuleError("保存后DOCX成品样式字号不可见。")
        if (
            style.font.color is not None
            and style.font.color.rgb is not None
            and str(style.font.color.rgb).upper() == "FFFFFF"
        ):
            raise TemplateRuleError("保存后DOCX成品样式颜色不可见。")
    for section in document.sections:
        dimensions = (
            section.page_width, section.page_height,
            section.left_margin, section.right_margin,
            section.top_margin, section.bottom_margin,
        )
        if any(value is None for value in dimensions):
            raise TemplateRuleError("保存后DOCX成品页面布局无效。")
        if (
            section.page_width - section.left_margin - section.right_margin < 4 * 914400
            or section.page_height - section.top_margin - section.bottom_margin < 4 * 914400
        ):
            raise TemplateRuleError("保存后DOCX成品可见页面区域过小。")

    def visible_lines(paragraph, expected_style=None):
        paragraph_properties = paragraph._p.pPr
        properties = list(paragraph_properties) if paragraph_properties is not None else []
        if expected_style is None:
            if properties:
                raise TemplateRuleError("保存后DOCX成品段落布局不受支持。")
        elif (
            len(properties) != 1
            or properties[0].tag != qn("w:pStyle")
            or properties[0].get(qn("w:val")) != expected_style
        ):
            raise TemplateRuleError("保存后DOCX成品段落样式与区域不一致。")
        for run in paragraph.runs:
            properties = run._r.rPr
            reject_unsafe_properties(properties, "文本")
            if (
                run.font.hidden is True
            ):
                raise TemplateRuleError("保存后DOCX成品包含隐藏文本。")
            if properties is not None and len(properties):
                raise TemplateRuleError("保存后DOCX成品包含非预期可见性或布局属性。")
        return paragraph.text.split("\n")

    content: list[dict[str, object]] = []
    for item in document.iter_inner_content():
        if isinstance(item, Paragraph):
            style_name = str(item.style.name or "") if item.style is not None else ""
            content.append({
                "kind": "heading" if style_name.startswith("Heading") else "paragraph",
                "lines": visible_lines(
                    item, "Heading2" if style_name.startswith("Heading") else None
                ),
            })
        elif isinstance(item, Table):
            content.append({
                "kind": "table",
                "rows": [
                    {"cells": [{
                        "lines": [
                            line
                            for paragraph in cell.paragraphs
                            for line in visible_lines(paragraph, None)
                        ]
                    } for cell in row.cells]}
                    for row in item.rows
                ],
            })
    return _normalize_semantic_trace({
        "schema_version": 1,
        "header": [
            {"lines": visible_lines(paragraph, "Header")}
            for paragraph in document.sections[0].header.paragraphs
            if paragraph.text
        ],
        "content": content,
        "footer": [
            {"lines": visible_lines(paragraph, "Footer")}
            for paragraph in document.sections[0].footer.paragraphs
            if paragraph.text
        ],
    })


class _SigningSemanticTraceParser(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.trace: dict[str, object] = {
            "schema_version": 1, "header": [], "content": [], "footer": [],
        }
        self.region = ""
        self.capture_tag = ""
        self.lines: list[str] = []
        self.table_rows: list[dict[str, object]] | None = None
        self.row_cells: list[dict[str, list[str]]] | None = None
        self.stack: list[str] = []
        self.invalid_reason = ""
        self.declaration_count = 0
        self.meta_count = 0
        self.article_count = 0
        self.region_counts = {"header": 0, "main": 0, "footer": 0}
        self.article_children: list[str] = []

    def _invalidate(self, reason):
        if not self.invalid_reason:
            self.invalid_reason = reason

    def handle_decl(self, declaration):
        self.declaration_count += 1
        if str(declaration).strip().lower() != "doctype html":
            self._invalidate("签署页包含非预期文档声明")

    def handle_starttag(self, tag, attrs):
        tag = str(tag).lower()
        attributes = {str(key).lower(): value for key, value in attrs}
        style = re.sub(r"\s+", "", str(attributes.get("style") or "")).lower()
        if (
            "hidden" in attributes
            or str(attributes.get("aria-hidden") or "").lower() == "true"
            or "display:none" in style
            or "visibility:hidden" in style
        ):
            self._invalidate("签署页包含隐藏或不可见内容")

        allowed = {
            "meta", "article", "header", "main", "footer", "section",
            "h2", "p", "br", "table", "tr", "td",
        }
        if tag not in allowed:
            self._invalidate("签署页包含非预期结构")
            return
        allowed_attributes = {
            "meta": {"charset"},
            "section": {"data-render-key", "data-render-region"},
        }.get(tag, set())
        if not set(attributes).issubset(allowed_attributes):
            self._invalidate("签署页包含非预期属性")

        parent = self.stack[-1] if self.stack else ""
        required_parent = {
            "article": "",
            "header": "article", "main": "article", "footer": "article",
            "section": "main", "h2": "section", "p": None,
            "table": "section", "tr": "table", "td": "tr", "br": None,
        }
        if tag == "meta":
            self.meta_count += 1
            if parent or str(attributes.get("charset") or "").lower() != "utf-8":
                self._invalidate("签署页元数据结构无效")
            return
        if tag == "p" and parent not in {"header", "section", "footer"}:
            self._invalidate("签署页段落结构无效")
        elif tag == "br" and parent not in {"p", "td"}:
            self._invalidate("签署页换行结构无效")
        elif required_parent.get(tag) is not None and parent != required_parent[tag]:
            self._invalidate("签署页可见结构无效")

        if tag == "article":
            self.article_count += 1
        if parent == "article" and tag in {"header", "main", "footer"}:
            self.article_children.append(tag)
        if tag in self.region_counts:
            self.region_counts[tag] += 1

        if tag == "br":
            if self.capture_tag:
                self.lines.append("")
            return
        self.stack.append(tag)
        if tag in {"header", "main", "footer"}:
            self.region = tag
        elif tag in {"p", "h2", "td"}:
            if self.capture_tag:
                self._invalidate("签署页文本结构嵌套无效")
            self.capture_tag = tag
            self.lines = [""]
        elif tag == "table":
            self.table_rows = []
        elif tag == "tr":
            self.row_cells = []

    def handle_data(self, data):
        if self.capture_tag and self.lines:
            self.lines[-1] += data
        elif str(data).strip():
            self._invalidate("签署页包含边界外可见文本")

    def handle_endtag(self, tag):
        tag = str(tag).lower()
        if tag in {"meta", "br"}:
            self._invalidate("签署页空元素闭合结构无效")
            return
        if not self.stack or self.stack[-1] != tag:
            self._invalidate("签署页标签闭合结构无效")
            return
        if tag == self.capture_tag:
            if tag == "td" and self.row_cells is not None:
                self.row_cells.append({"lines": list(self.lines)})
            elif tag == "p":
                target = self.trace[self.region] if self.region in {"header", "footer"} else self.trace["content"]
                value = {"lines": list(self.lines)}
                if self.region == "main":
                    value["kind"] = "paragraph"
                target.append(value)
            elif tag == "h2":
                self.trace["content"].append({"kind": "heading", "lines": list(self.lines)})
            self.capture_tag = ""
            self.lines = []
        elif tag == "tr" and self.row_cells is not None and self.table_rows is not None:
            self.table_rows.append({"cells": self.row_cells})
            self.row_cells = None
        elif tag == "table" and self.table_rows is not None:
            self.trace["content"].append({"kind": "table", "rows": self.table_rows})
            self.table_rows = None
        elif tag in {"header", "main", "footer"}:
            self.region = ""
        self.stack.pop()

    def validate_complete(self):
        if (
            self.invalid_reason
            or self.stack
            or self.capture_tag
            or self.table_rows is not None
            or self.row_cells is not None
            or self.declaration_count != 1
            or self.meta_count != 1
            or self.article_count != 1
            or self.region_counts != {"header": 1, "main": 1, "footer": 1}
            or self.article_children != ["header", "main", "footer"]
        ):
            reason = self.invalid_reason or "结构不完整"
            raise TemplateRuleError(f"保存后签署页成品无效：{reason}。")


def _extract_signing_semantic_trace(path: Path) -> dict[str, object]:
    import base64
    import gzip
    raw = path.read_text(encoding="utf-8")
    match = re.search(
        r'<script id="aaie-master-model" type="application/octet-stream">([^<]+)</script>',
        raw,
    )
    if match:
        lowered = raw.lower().replace(" ", "")
        if (
            raw.count("<style>") != 1
            or raw.count("<script") != 1
            or "<template" in lowered
            or "aria-hidden=" in lowered
            or "display:none" in lowered
            or "visibility:hidden" in lowered
            or 'class="hidden"' in lowered
            or "<article hidden" in lowered
        ):
            raise TemplateRuleError("保存后签署页成品包含隐藏或非预期结构。")
        try:
            model = json.loads(gzip.decompress(base64.b64decode(match.group(1))).decode("utf-8"))
        except Exception as exc:
            raise TemplateRuleError("保存后签署页主版式语义绑定无效。") from exc
        if not all(value in raw for value in (
            'data-master-anchor="section_11_signature"',
            'data-signature-anchor="signature_aaie"',
            'data-signature-anchor="signature_family"',
            'data-stamp-anchor="aaie-only"',
        )):
            raise TemplateRuleError("保存后签署页主版式物理锚点无效。")
        return _semantic_trace_from_model(model)
    parser = _SigningSemanticTraceParser()
    parser.feed(raw)
    parser.close()
    parser.validate_complete()
    return _normalize_semantic_trace(parser.trace)


def _build_signing_artifact_bytes(model: Mapping[str, object]) -> bytes:
    def paragraph_html(paragraph):
        return "<p>" + "<br>".join(
            html.escape(str(line)) for line in paragraph["lines"]
        ) + "</p>"

    parts = ["<!doctype html><meta charset=\"utf-8\"><article>", "<header>"]
    parts.extend(paragraph_html(paragraph) for paragraph in model["header"])
    parts.append("</header><main>")
    for node in model["content"]:
        key = html.escape(str(node["key"]), quote=True)
        if node["kind"] == "section":
            parts.append(
                f'<section data-render-key="{key}"><h2>'
                f'{html.escape(str(node["title"]))}</h2>'
            )
            parts.extend(paragraph_html(paragraph) for paragraph in node["paragraphs"])
            parts.append("</section>")
        else:
            parts.append(
                f'<section data-render-key="{key}"><h2>'
                f'{html.escape(str(node["title"]))}</h2><table>'
            )
            for row in node["rows"]:
                parts.append("<tr>" + "".join(
                    "<td>" + "<br>".join(
                        html.escape(str(line)) for line in cell["lines"]
                    ) + "</td>"
                    for cell in row["cells"]
                ) + "</tr>")
            parts.append("</table></section>")
    signature = model["signature_region"]
    parts.append(
        '<section data-render-region="signature"><h2>'
        f'{html.escape(str(signature["title"]))}</h2>'
    )
    parts.extend(paragraph_html(paragraph) for paragraph in signature["paragraphs"])
    parts.append("</section></main><footer>")
    parts.extend(paragraph_html(paragraph) for paragraph in model["footer"])
    parts.append("</footer></article>")
    return "".join(parts).encode("utf-8")


def _trace_line_lists(trace: Mapping[str, object]) -> list[list[str]]:
    line_lists: list[list[str]] = []
    for paragraph in trace["header"]:
        line_lists.append(paragraph["lines"])
    for item in trace["content"]:
        if item.get("kind") in {"heading", "paragraph"} or "lines" in item:
            line_lists.append(item["lines"])
        else:
            for row in item["rows"]:
                for cell in row["cells"]:
                    line_lists.append(cell["lines"])
    for paragraph in trace["footer"]:
        line_lists.append(paragraph["lines"])
    return line_lists


PDF_FONT_NAME = "AAIENotoSansSC"
PDF_FONT_FILE = (
    Path(__file__).resolve().parent
    / "assets" / "fonts" / "noto-sans-cjk-sc" / "NotoSansSC-Regular.ttf"
)


def _ensure_pdf_font():
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    if not PDF_FONT_FILE.is_file():
        raise TemplateRuleError("模拟PDF所需嵌入字体缺失。")
    if PDF_FONT_NAME not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont(PDF_FONT_NAME, str(PDF_FONT_FILE)))
    return pdfmetrics


def _build_pdf_layout_plan(model: Mapping[str, object]) -> dict[str, object]:
    """Build the deterministic physical plan shared by PDF render and verification."""
    if (model.get("layout_policy") or {}).get("id") == "AAIEContractMasterLayoutV1":
        artifact_bytes, anchor_manifest = render_master_pdf_bytes(model)
        _ensure_pdf_font()
        page_count = int(anchor_manifest["signature_aaie"]["page"])
        runs = []
        y = 780.0
        page = 1
        for lines in _trace_line_lists(_semantic_trace_from_model(model)):
            for value in lines:
                runs.append({"page": page, "x": 40.0, "y": y, "font_size": 10.0, "text": str(value)})
                y -= 15.0
                if y < 40:
                    page = min(page + 1, page_count)
                    y = 780.0
        return {
            "master_layout": True,
            "artifact_bytes": artifact_bytes,
            "anchor_manifest": anchor_manifest,
            "semantic_trace": _semantic_trace_from_model(model),
            "runs": runs,
            "page_count": page_count,
            "font_name": PDF_FONT_NAME,
        }
    import math
    from reportlab.lib.pagesizes import A4

    pdfmetrics = _ensure_pdf_font()
    page_width, page_height = (float(value) for value in A4)
    left = 40.0
    right = 40.0
    body_top = page_height - 40.0
    body_bottom = 40.0
    normal_size = 10.0
    heading_size = 12.0
    leading = 15.0
    semantic_indices: dict[tuple[object, ...], int] = {}
    semantic_values: list[str] = []

    def register_semantic(key, value):
        semantic_indices[key] = len(semantic_values)
        semantic_values.append(re.sub(r"\s+", "", str(value)))

    for paragraph_index, paragraph in enumerate(model["header"]):
        for line_index, value in enumerate(paragraph["lines"]):
            register_semantic(("header", paragraph_index, line_index), value)
    for node_index, node in enumerate(model["content"]):
        register_semantic(("heading", node_index), node["title"])
        if node["kind"] == "section":
            for paragraph_index, paragraph in enumerate(node["paragraphs"]):
                for line_index, value in enumerate(paragraph["lines"]):
                    register_semantic(
                        ("paragraph", node_index, paragraph_index, line_index), value
                    )
        else:
            for row_index, row in enumerate(node["rows"]):
                for cell_index, cell in enumerate(row["cells"]):
                    for line_index, value in enumerate(cell["lines"]):
                        register_semantic(
                            ("table", node_index, row_index, cell_index, line_index),
                            value,
                        )
    signature = model["signature_region"]
    register_semantic(("signature-heading",), signature["title"])
    for paragraph_index, paragraph in enumerate(signature["paragraphs"]):
        for line_index, value in enumerate(paragraph["lines"]):
            register_semantic(("signature", paragraph_index, line_index), value)
    for paragraph_index, paragraph in enumerate(model["footer"]):
        for line_index, value in enumerate(paragraph["lines"]):
            register_semantic(("footer", paragraph_index, line_index), value)

    semantic_trace = _semantic_trace_from_model(model)
    if semantic_values != [
        str(line)
        for lines in _trace_line_lists(semantic_trace)
        for line in lines
    ]:
        raise TemplateRuleError("PDF语义与布局计划边界不一致。")

    def wrap_line(text: object, size: float, available: float) -> list[str]:
        value = str(text)
        if not value:
            return [""]
        lines: list[str] = []
        current = ""
        for token in re.findall(r"\S+\s*|\s+", value):
            candidate = current + token
            if pdfmetrics.stringWidth(candidate, PDF_FONT_NAME, size) <= available:
                current = candidate
                continue
            if current:
                lines.append(current.rstrip())
                current = ""
            token = token.lstrip()
            if pdfmetrics.stringWidth(token, PDF_FONT_NAME, size) <= available:
                current = token
                continue
            for character in token:
                candidate = current + character
                if (
                    current
                    and pdfmetrics.stringWidth(candidate, PDF_FONT_NAME, size) > available
                ):
                    lines.append(current.rstrip())
                    current = character
                else:
                    current = candidate
        lines.append(current.rstrip())
        return lines

    page_number = 1
    y = body_top
    runs: list[dict[str, object]] = []

    def begin_page():
        nonlocal page_number, y
        page_number += 1
        y = body_top

    def ensure_space(required):
        if y - required < body_bottom:
            begin_page()

    def append_run(
        text, *, semantic_line, region, kind, size, x, available,
        row=None, cell=None, physical_line=0, baseline=None,
    ):
        value = str(text)
        width = float(pdfmetrics.stringWidth(value, PDF_FONT_NAME, size))
        if width > float(available) + 0.01:
            raise TemplateRuleError("PDF物理行超过布局宽度。")
        runs.append({
            "page": page_number,
            "x": float(x),
            "y": float(y if baseline is None else baseline),
            "font_size": float(size),
            "text": value,
            "available_width": float(available),
            "region": str(region),
            "kind": str(kind),
            "row": row,
            "cell": cell,
            "semantic_line": int(semantic_line),
            "physical_line": int(physical_line),
        })

    def add_semantic_line(
        text, semantic_line, *, region, kind, size=normal_size,
        line_leading=leading, x=left, available=None,
    ):
        nonlocal y
        width = page_width - x - right if available is None else float(available)
        for physical_index, physical_line in enumerate(wrap_line(text, size, width)):
            ensure_space(line_leading)
            append_run(
                physical_line,
                semantic_line=semantic_line,
                region=region,
                kind=kind,
                size=size,
                x=x,
                available=width,
                physical_line=physical_index,
            )
            y -= line_leading

    def add_paragraphs(paragraphs, key_prefix, *, region, size, line_leading):
        nonlocal y
        for paragraph_index, paragraph in enumerate(paragraphs):
            for line_index, value in enumerate(paragraph["lines"]):
                add_semantic_line(
                    value,
                    semantic_indices[key_prefix + (paragraph_index, line_index)],
                    region=region,
                    kind="paragraph",
                    size=size,
                    line_leading=line_leading,
                )
            y -= 4.0

    def physical_cell_lines(cell, node_index, row_index, cell_index, width):
        physical: list[tuple[str, int, int]] = []
        for line_index, semantic_line in enumerate(cell["lines"]):
            semantic_index = semantic_indices[
                ("table", node_index, row_index, cell_index, line_index)
            ]
            physical.extend(
                (value, semantic_index, physical_index)
                for physical_index, value in enumerate(
                    wrap_line(semantic_line, normal_size, width)
                )
            )
        return physical or [("", semantic_indices[("table", node_index, row_index, cell_index, 0)], 0)]

    def add_table_row(row, node_index, row_index):
        nonlocal y
        gap = 14.0
        available_width = page_width - left - right
        column_width = (available_width - gap) / 2.0
        columns = [
            physical_cell_lines(cell, node_index, row_index, cell_index, column_width)
            for cell_index, cell in enumerate(row["cells"])
        ]
        offset = 0
        total_lines = max(len(column) for column in columns)
        while offset < total_lines:
            available_lines = max(0, int(math.floor((y - body_bottom) / leading)))
            if available_lines < 1:
                begin_page()
                available_lines = max(1, int(math.floor((y - body_bottom) / leading)))
            chunk_size = min(total_lines - offset, available_lines)
            row_top = y
            for column_index, column in enumerate(columns):
                x = left + column_index * (column_width + gap)
                for chunk_line, (value, semantic_index, physical_index) in enumerate(
                    column[offset:offset + chunk_size]
                ):
                    append_run(
                        value,
                        semantic_line=semantic_index,
                        region="content",
                        kind="table_cell",
                        size=normal_size,
                        x=x,
                        available=column_width,
                        row=row_index,
                        cell=column_index,
                        physical_line=physical_index,
                        baseline=row_top - chunk_line * leading,
                    )
            y -= chunk_size * leading
            offset += chunk_size
            if offset < total_lines:
                begin_page()
        y -= 3.0

    add_paragraphs(
        model["header"], ("header",), region="header", size=9.0, line_leading=13.0
    )
    y -= 5.0
    for node_index, node in enumerate(model["content"]):
        add_semantic_line(
            node["title"], semantic_indices[("heading", node_index)],
            region="content", kind="heading", size=heading_size, line_leading=18.0,
        )
        if node["kind"] == "section":
            add_paragraphs(
                node["paragraphs"], ("paragraph", node_index),
                region="content", size=normal_size, line_leading=leading,
            )
        else:
            for row_index, row in enumerate(node["rows"]):
                add_table_row(row, node_index, row_index)
            y -= 4.0
    add_semantic_line(
        signature["title"], semantic_indices[("signature-heading",)],
        region="signature", kind="heading", size=heading_size, line_leading=18.0,
    )
    add_paragraphs(
        signature["paragraphs"], ("signature",),
        region="signature", size=normal_size, line_leading=leading,
    )
    y -= 5.0
    add_paragraphs(
        model["footer"], ("footer",), region="footer", size=8.0, line_leading=12.0
    )
    return {
        "schema_version": 1,
        "font_name": PDF_FONT_NAME,
        "page_width": page_width,
        "page_height": page_height,
        "page_count": page_number,
        "semantic_trace": semantic_trace,
        "semantic_values": semantic_values,
        "runs": runs,
    }


def _render_pdf_plan_bytes(layout_plan: Mapping[str, object]) -> bytes:
    if layout_plan.get("master_layout"):
        return bytes(layout_plan["artifact_bytes"])
    import io
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    output = io.BytesIO()
    pdf = canvas.Canvas(
        output, pagesize=A4, invariant=1, pageCompression=1
    )
    runs_by_page: dict[int, list[Mapping[str, object]]] = {}
    for run in layout_plan["runs"]:
        runs_by_page.setdefault(int(run["page"]), []).append(run)
    for page_number in range(1, int(layout_plan["page_count"]) + 1):
        if page_number > 1:
            pdf.showPage()
        for run in runs_by_page.get(page_number, []):
            pdf.setFont(str(layout_plan["font_name"]), float(run["font_size"]))
            pdf.setFillColorRGB(0, 0, 0)
            pdf.drawString(float(run["x"]), float(run["y"]), str(run["text"]))
    pdf.save()
    return output.getvalue()


def render_published_contract_artifacts(
    snapshot: Mapping[str, object], *, student_name: str,
    signer: Mapping[str, object], selected_services: list[str],
    contract_instance: Mapping[str, object] | None = None,
    include_p0_mandatory_fields: bool = True,
    use_master_layout: bool = True,
    display_bindings: Mapping[str, object] | None = None,
) -> dict[str, object]:
    """Render DOCX, PDF and signing HTML from one already-frozen snapshot."""
    allowed = {
        str(item.get("service_id") or "")
        for item in snapshot.get("services") or []
        if item.get("enabled", True)
    }
    selected = list(dict.fromkeys(str(value) for value in selected_services))
    if len(selected) != len(selected_services) or any(value not in allowed for value in selected):
        raise TemplateRuleError("服务范围与已发布合同模板不匹配。")
    if any(not str(signer.get(key) or "").strip() for key in ("source", "name", "relationship", "phone", "email")):
        raise TemplateRuleError("合同签署人资料不完整。")
    if use_master_layout and is_master_snapshot(snapshot):
        try:
            model = build_master_render_model(
                snapshot,
                student_name=student_name,
                signer=signer,
                selected_services=selected,
                contract_instance=contract_instance,
                display_bindings=display_bindings,
            )
        except MasterLayoutError as exc:
            raise TemplateRuleError(str(exc)) from exc
        authority_content = _authority_content_from_model(model)
        semantic_trace = _semantic_trace_from_model(model)
        semantic_hash = hashlib.sha256(
            _canonical_json(semantic_trace).encode("utf-8")
        ).hexdigest()
        pdf_bytes, anchor_manifest = render_master_pdf_bytes(model)
        return {
            "model": model,
            "authority_content": authority_content,
            "semantic_trace": semantic_trace,
            "semantic_hash": semantic_hash,
            "docx_bytes": _embed_docx_font(render_master_docx_bytes(model)),
            "pdf_bytes": pdf_bytes,
            "signing_preview_bytes": render_master_signing_html(model),
            "anchor_manifest": anchor_manifest,
        }
    labels = {
        str(item.get("service_id") or ""): str(item.get("label") or "")
        for item in snapshot.get("services") or [] if item.get("enabled", True)
    }
    replacements = {
        "{{student_name}}": str(student_name or "").strip(),
        "{{signer_name}}": str(signer["name"]),
        "{{signer_phone}}": str(signer["phone"]),
        "{{signer_email}}": str(signer["email"]),
        "{{service_scope}}": "\n".join(
            f"{index}. {labels[service_id]}" for index, service_id in enumerate(selected, 1)
        ),
    }
    model = build_visible_render_model(snapshot, replacements, signer)
    instance = dict(contract_instance or {})
    commercial = dict(snapshot.get("commercial") or {})
    contact_parts = [
        value.strip() for value in str(commercial.get("company_contact") or "").split("/")
        if value.strip()
    ]
    company_phone = contact_parts[0] if contact_parts else ""
    company_email = contact_parts[-1] if len(contact_parts) > 1 else str(commercial.get("company_contact") or "")
    party_lines = [
        f"甲方签约主体：{commercial.get('company_name', '')}",
        f"甲方 ABN / ACN：{commercial.get('company_abn', '')}",
        f"甲方地址：{commercial.get('company_address', '')}",
        f"甲方电话：{company_phone}",
        f"甲方邮箱：{company_email}",
    ]
    if include_p0_mandatory_fields:
        party_lines.append(
            f"甲方授权代表：{instance.get('company_authorised_person', '')}"
        )
    party_lines.extend([
        f"乙方签署人：{signer.get('name', '')}",
        f"乙方与学生关系：{signer.get('relationship', '')}",
        f"乙方电话：{signer.get('phone', '')}",
        f"乙方邮箱：{signer.get('email', '')}",
        f"乙方地址：{instance.get('parent_address') or instance.get('student_home_address', '')}",
    ])
    for node in model["content"]:
        if node.get("key") == "fixed_parties":
            node["paragraphs"] = _visible_paragraphs("\n".join(party_lines))
            break

    def table_row(label: object, value: object) -> dict[str, object]:
        return {"cells": [_visible_cell(label), _visible_cell(value)]}

    identity_rows = [
        table_row("合同编号", instance.get("contract_uid", "")),
        table_row("合同日期", instance.get("contract_date", "")),
        table_row("学生姓名", instance.get("student_full_name", student_name)),
        table_row("学生出生日期", instance.get("student_date_of_birth", "")),
        table_row("学生证件或CRM编号", instance.get("student_id_or_passport_number", "")),
    ]
    if include_p0_mandatory_fields:
        identity_rows.extend([
            table_row("当前年级", instance.get("student_current_year_level", "")),
            table_row("申请年级", instance.get("student_applying_year_level", "")),
        ])
    identity = {
        "kind": "table", "key": "protected_contract_identity",
        "title": "合同基本信息（固定版式）",
        "rows": identity_rows,
    }
    instance_nodes: list[dict[str, object]] = []
    for key, title in (
        ("other_service_scope_text", "其他约定服务范围"),
        ("special_terms", "特别条款"),
    ):
        value = str(instance.get(key) or "").strip()
        if value:
            instance_nodes.append({
                "kind": "section", "key": f"contract_instance_{key}",
                "title": title, "paragraphs": _visible_paragraphs(value),
            })
    model["content"] = [identity] + list(model["content"]) + instance_nodes
    model["signature_region"] = {
        "title": "签署确认",
        "paragraphs": [{"lines": [
            "委托方（乙方签字）",
            f"法定代理人或监护人：{signer.get('name', '')}",
            f"关系：{signer.get('relationship', '')}",
            f"电话：{signer.get('phone', '')}",
            f"邮箱：{signer.get('email', '')}",
            "家长/乙方签署日期：",
            "甲方（签字 / 盖章）",
            "甲方签署日期：",
        ]}],
    }
    authority_content = _authority_content_from_model(model)
    semantic_trace = _semantic_trace_from_model(model)
    semantic_hash = hashlib.sha256(_canonical_json(semantic_trace).encode("utf-8")).hexdigest()
    return {
        "model": model,
        "authority_content": authority_content,
        "semantic_trace": semantic_trace,
        "semantic_hash": semantic_hash,
        "docx_bytes": _build_docx_artifact_bytes(model),
        "pdf_bytes": _render_pdf_plan_bytes(_build_pdf_layout_plan(model)),
        "signing_preview_bytes": _build_signing_artifact_bytes(model),
    }


def _pdf_layout_fingerprint(source: Path | bytes) -> str:
    import io
    from pypdf import PdfReader
    from pypdf.generic import (
        ArrayObject,
        BooleanObject,
        ByteStringObject,
        DictionaryObject,
        FloatObject,
        IndirectObject,
        NameObject,
        NullObject,
        NumberObject,
        StreamObject,
        TextStringObject,
    )

    raw = source.read_bytes() if isinstance(source, Path) else bytes(source)
    try:
        reader = PdfReader(io.BytesIO(raw), strict=True)
    except Exception as exc:
        raise TemplateRuleError("PDF成品结构无法解析。") from exc
    if reader.is_encrypted or reader.trailer.get("/Encrypt") is not None:
        raise TemplateRuleError("PDF成品不得加密。")

    def canonical(value, active=frozenset()):
        if isinstance(value, IndirectObject):
            key = (value.idnum, value.generation)
            if key in active:
                return {"cycle": True}
            try:
                return canonical(value.get_object(), active | {key})
            except Exception as exc:
                raise TemplateRuleError("PDF成品包含无法解析的间接资源。") from exc
        if isinstance(value, StreamObject):
            try:
                decoded = value.get_data()
            except Exception as exc:
                raise TemplateRuleError("PDF成品资源流无法解码。") from exc
            dictionary = {
                str(key): canonical(item, active)
                for key, item in sorted(value.items(), key=lambda entry: str(entry[0]))
                if str(key) not in {"/Length", "/Filter", "/DecodeParms"}
            }
            return {
                "stream_dictionary": dictionary,
                "decoded_bytes": len(decoded),
                "decoded_sha256": hashlib.sha256(decoded).hexdigest(),
            }
        if isinstance(value, DictionaryObject):
            return {
                str(key): canonical(item, active)
                for key, item in sorted(value.items(), key=lambda entry: str(entry[0]))
            }
        if isinstance(value, (ArrayObject, list, tuple)):
            return [canonical(item, active) for item in value]
        if isinstance(value, NullObject) or value is None:
            return None
        if isinstance(value, BooleanObject) or isinstance(value, bool):
            return bool(value)
        if isinstance(value, FloatObject) or isinstance(value, float):
            return float(value)
        if isinstance(value, NumberObject) or isinstance(value, int):
            return int(value)
        if isinstance(value, ByteStringObject) or isinstance(value, bytes):
            raw_value = bytes(value)
            return {
                "bytes": len(raw_value),
                "sha256": hashlib.sha256(raw_value).hexdigest(),
            }
        if isinstance(value, (NameObject, TextStringObject, str)):
            return str(value)
        return str(value)

    pages = []
    for page in reader.pages:
        try:
            contents = page.get_contents()
            decoded_contents = contents.get_data() if contents is not None else b""
        except Exception as exc:
            raise TemplateRuleError("PDF成品页面内容流无法解码。") from exc
        pages.append({
            "media_box": [float(value) for value in page.mediabox],
            "crop_box": [float(value) for value in page.cropbox],
            "trim_box": [float(value) for value in page.trimbox],
            "bleed_box": [float(value) for value in page.bleedbox],
            "art_box": [float(value) for value in page.artbox],
            "rotation": int(page.rotation or 0),
            "user_unit": float(page.get("/UserUnit", 1)),
            "contents_bytes": len(decoded_contents),
            "contents_sha256": hashlib.sha256(decoded_contents).hexdigest(),
            "resources": canonical(page.get("/Resources", DictionaryObject())),
            "annotations": canonical(page.get("/Annots", ArrayObject())),
            "group": canonical(page.get("/Group")),
            "page_dictionary": canonical(page),
        })
    root = reader.trailer.get("/Root")
    root = root.get_object() if isinstance(root, IndirectObject) else root
    if not isinstance(root, DictionaryObject):
        raise TemplateRuleError("PDF成品缺少有效目录结构。")
    return _hash({
        "pages": pages,
        "catalog": canonical(root),
        "trailer": canonical(reader.trailer),
    })


def _extract_verified_pdf_semantic_trace(
    path: Path, layout_plan: Mapping[str, object],
) -> dict[str, object]:
    """Require every saved PDF character to match the recomputed physical plan."""
    if layout_plan.get("master_layout"):
        expected_pdf = bytes(layout_plan["artifact_bytes"])
        if path.read_bytes() != expected_pdf:
            raise TemplateRuleError("保存后PDF主版式物理结构与确定性布局不一致。")
        from pypdf import PdfReader
        reader = PdfReader(path)
        manifest = layout_plan["anchor_manifest"]
        final_page = len(reader.pages)
        if (
            manifest["signature_aaie"]["page"] != final_page
            or manifest["signature_family"]["page"] != final_page
            or manifest["stamp"]["page"] != final_page
        ):
            raise TemplateRuleError("保存后PDF主版式签署锚点页无效。")
        final_text = reader.pages[-1].extract_text()
        if not all(value in final_text for value in ("十一、签署", "甲方（签字 / 盖章）", "委托方（乙方签字）")):
            raise TemplateRuleError("保存后PDF主版式双栏签署锚点无效。")
        return deepcopy(layout_plan["semantic_trace"])
    import math
    import pdfplumber

    expected_pdf = _render_pdf_plan_bytes(layout_plan)
    if _pdf_layout_fingerprint(path) != _pdf_layout_fingerprint(expected_pdf):
        raise TemplateRuleError("保存后PDF成品物理结构与确定性布局不一致。")
    pdfmetrics = _ensure_pdf_font()
    expected_trace = layout_plan["semantic_trace"]
    semantic_values = [str(value) for value in layout_plan["semantic_values"]]
    actual_semantics = [""] * len(semantic_values)
    tolerance = 0.04
    with pdfplumber.open(path) as document:
        if len(document.pages) != int(layout_plan["page_count"]):
            raise TemplateRuleError("保存后PDF成品页面数量与布局计划不一致。")
        for page_number, page in enumerate(document.pages, 1):
            if (
                abs(float(page.width) - float(layout_plan["page_width"])) > tolerance
                or abs(float(page.height) - float(layout_plan["page_height"])) > tolerance
            ):
                raise TemplateRuleError("保存后PDF成品页面尺寸与布局计划不一致。")
            expected_runs = [
                run for run in layout_plan["runs"] if int(run["page"]) == page_number
            ]
            characters = list(page.chars)
            cursor = 0
            for run in expected_runs:
                expected_text = str(run["text"])
                if not expected_text:
                    continue
                actual_text = ""
                for character_index, expected_character in enumerate(expected_text):
                    if cursor >= len(characters):
                        raise TemplateRuleError("保存后PDF成品缺失计划内可见文本。")
                    character = characters[cursor]
                    cursor += 1
                    if str(character.get("text") or "") != expected_character:
                        raise TemplateRuleError("保存后PDF成品可见文本顺序不一致。")
                    matrix = character.get("matrix")
                    if not isinstance(matrix, (tuple, list)) or len(matrix) < 6:
                        raise TemplateRuleError("保存后PDF成品缺少文本物理坐标。")
                    expected_x = float(run["x"]) + float(
                        pdfmetrics.stringWidth(
                            expected_text[:character_index],
                            str(layout_plan["font_name"]),
                            float(run["font_size"]),
                        )
                    )
                    expected_x1 = expected_x + float(
                        pdfmetrics.stringWidth(
                            expected_character,
                            str(layout_plan["font_name"]),
                            float(run["font_size"]),
                        )
                    )
                    coordinates = tuple(
                        float(character[name]) for name in ("x0", "x1", "top", "bottom")
                    )
                    physical = (
                        float(matrix[4]), float(matrix[5]), float(character["size"])
                    )
                    if not all(math.isfinite(value) for value in coordinates + physical):
                        raise TemplateRuleError("保存后PDF成品包含无效坐标。")
                    x0, x1, top, bottom = coordinates
                    if (
                        x0 < -tolerance
                        or x1 > float(page.width) + tolerance
                        or top < -tolerance
                        or bottom > float(page.height) + tolerance
                        or x1 < x0
                        or bottom < top
                    ):
                        raise TemplateRuleError("保存后PDF成品存在页面外不可见文本。")
                    if (
                        abs(float(matrix[4]) - expected_x) > tolerance
                        or abs(float(matrix[5]) - float(run["y"])) > tolerance
                        or abs(x0 - expected_x) > tolerance
                        or abs(x1 - expected_x1) > tolerance
                        or abs(float(character["size"]) - float(run["font_size"])) > tolerance
                        or "NotoSansSC-Regular" not in str(character.get("fontname") or "")
                    ):
                        raise TemplateRuleError("保存后PDF成品文本布局或字号不一致。")
                    color = character.get("non_stroking_color")
                    normalized_color = tuple(color) if isinstance(color, (tuple, list)) else color
                    if normalized_color not in ((0,), (0, 0, 0)):
                        raise TemplateRuleError("保存后PDF成品文本可见颜色不一致。")
                    actual_text += str(character.get("text") or "")
                semantic_index = int(run["semantic_line"])
                if semantic_index < 0 or semantic_index >= len(actual_semantics):
                    raise TemplateRuleError("保存后PDF成品语义边界无效。")
                actual_semantics[semantic_index] += actual_text
            if cursor != len(characters):
                raise TemplateRuleError("保存后PDF成品包含计划外可见文本。")

    normalized_actual = [re.sub(r"\s+", "", value) for value in actual_semantics]
    if normalized_actual != semantic_values:
        raise TemplateRuleError("保存后PDF成品语义边界与布局计划不一致。")
    actual_trace = deepcopy(expected_trace)
    cursor = 0
    for lines in _trace_line_lists(actual_trace):
        for index in range(len(lines)):
            lines[index] = normalized_actual[cursor]
            cursor += 1
    if cursor != len(normalized_actual):
        raise TemplateRuleError("保存后PDF成品语义映射不完整。")
    return actual_trace


def _require_post_save_trace(
    artifact_name: str, actual: Mapping[str, object], expected: Mapping[str, object],
) -> dict[str, object]:
    if _canonical_json(actual) != _canonical_json(expected):
        raise TemplateRuleError(f"保存后{artifact_name}成品可见语义不一致。")
    return deepcopy(actual)


def _simulation_result_binding(row: Mapping[str, object]) -> dict[str, object]:
    """Return the canonical persisted result payload, excluding its own hash."""
    return {
        "template_identity_id": int(row["template_identity_id"]),
        "template_revision": int(row["template_revision"]),
        "template_content_hash": str(row["template_content_hash"]),
        "template_snapshot_hash": str(row["template_snapshot_hash"]),
        "simulation_input_hash": str(row["simulation_input_hash"]),
        "signer": json.loads(str(row["signer_json"])),
        "selected_services": json.loads(str(row["selected_services_json"])),
        "authority_content_hash": str(row["authority_content_hash"]),
        "content_hashes": {
            "docx": str(row["docx_content_hash"]),
            "pdf": str(row["pdf_content_hash"]),
            "signing": str(row["signing_content_hash"]),
        },
        "semantic": {
            "docx": {
                "hash": str(row["docx_semantic_hash"]),
                "trace": json.loads(str(row["docx_semantic_trace_json"])),
            },
            "pdf": {
                "hash": str(row["pdf_semantic_hash"]),
                "trace": json.loads(str(row["pdf_semantic_trace_json"])),
            },
            "signing": {
                "hash": str(row["signing_semantic_hash"]),
                "trace": json.loads(str(row["signing_semantic_trace_json"])),
            },
        },
        "artifact_hashes": {
            "preview.docx": str(row["docx_artifact_hash"]),
            "preview.pdf": str(row["pdf_artifact_hash"]),
            "signing-preview.html": str(row["signing_artifact_hash"]),
        },
        "artifact_manifest": json.loads(str(row["artifact_manifest_json"])),
        "artifact_root": str(row["artifact_root"]),
        "artifact_dir": str(row["artifact_dir"]),
        "simulated_at": str(row["simulated_at"]),
    }


def _default_commercial() -> dict[str, str]:
    return {
        "currency": "AUD",
        "service_fee": "0.00",
        "payment_terms": "按合同约定的付款计划支付。",
        "company_name": DEFAULT_PARTY_A_NAME,
        "company_abn": "",
        "company_address": DEFAULT_PARTY_A_ADDRESS,
        "company_contact": f"{DEFAULT_PARTY_A_PHONE} / {DEFAULT_PARTY_A_EMAIL}",
        "company_authorised_person": "",
        "refund_policy": "按已完成服务、合同约定及第三方退款政策处理。",
    }


AUTHORITY_V2_GST_VALUES = frozenset({"EXCLUSIVE"})
AUTHORITY_V2_PAYMENT_METHODS = ("支付宝", "微信", "国内银行", "澳洲私户", "澳洲公户")
AUTHORITY_V2_INHERITED_METADATA_FIELDS = frozenset({
    "commercial_schema",
    "standard_service_fee",
    "standard_service_fee_ex_gst",
    "gst_treatment",
    "gst_rate",
    "default_gst_amount",
    "default_final_total",
    "default_payment_plan",
    "payment_methods_json",
    "installment_approval_role",
    "discount_approval_role",
    "discount_reason_required",
    "base_fee_override_allowed",
    "visa_refusal_retained_ex_gst",
    "visa_refusal_rule_source",
})


def authority_v2_commercial(template_type: str, source: Mapping[str, object]) -> dict[str, object]:
    """Build the approved V2 commercial policy without mutating Published V1."""
    contract_domain.contract_type_config(template_type)
    refund_policy = {
        contract_domain.NEW_INTERNATIONAL_STUDENT_CONTRACT: (
            "签证申请被拒且不再继续申请时，按实际已支付总额计算退款："
            "税前保留 AUD 900；GST适用时相应保留GST，GST-free时GST调整为0；"
            "按双方确认方案申请后未获任何学校录取时，退还本合同最终费用；"
            "已完成服务部分不退，第三方费用按第三方政策处理。"
        ),
        contract_domain.TRANSFER_APPLICATION_CONTRACT: (
            "按双方确认方案申请后未获任何学校录取时，退还本合同最终费用；"
            "已完成服务部分不退，第三方费用按第三方政策处理。"
        ),
        contract_domain.LOCAL_STUDENT_APPLICATION_CONTRACT: (
            "按双方确认方案申请后未获任何学校录取时，退还本合同最终费用；"
            "已完成服务部分不退，第三方费用按第三方政策处理。"
        ),
    }[template_type]
    commercial = dict(source)
    commercial.update({
        "commercial_schema": "authority_v2",
        "currency": "AUD",
        "service_fee": "2400.00",
        "standard_service_fee": "2400.00",
        "standard_service_fee_ex_gst": "2400.00",
        "gst_treatment": "EXCLUSIVE",
        "gst_rate": "0.10",
        "default_gst_amount": "240.00",
        "default_final_total": "2640.00",
        "payment_terms": "合同签署时一次性支付，并在正式服务开始前付清。",
        "default_payment_plan": "single_payment_at_signing_before_service",
        "payment_methods_json": _canonical_json(list(AUTHORITY_V2_PAYMENT_METHODS)),
        "installment_approval_role": "Admin",
        "discount_approval_role": "Admin",
        "discount_reason_required": "true",
        "base_fee_override_allowed": "false",
        "refund_policy": refund_policy,
    })
    if template_type == contract_domain.NEW_INTERNATIONAL_STUDENT_CONTRACT:
        commercial.update({
            "visa_refusal_retained_ex_gst": "900.00",
            "visa_refusal_rule_source": "published_v2_new_international_visa_refusal",
        })
    else:
        commercial.pop("visa_refusal_retained_ex_gst", None)
        commercial.pop("visa_refusal_rule_source", None)
    return commercial


def authority_v2_commercial_issues(commercial: Mapping[str, object]) -> list[dict[str, str]]:
    if commercial.get("commercial_schema") != "authority_v2":
        return []
    issues: list[dict[str, str]] = []
    if str(commercial.get("service_fee") or "").strip() != "2400.00" or str(
        commercial.get("standard_service_fee") or ""
    ).strip() != "2400.00" or str(
        commercial.get("standard_service_fee_ex_gst") or ""
    ).strip() != "2400.00":
        issues.append({"code": "standard_fee_invalid", "message": "标准服务费必须为AUD 2400.00"})
    if str(commercial.get("currency") or "").strip() != "AUD":
        issues.append({"code": "currency_invalid", "message": "币种必须为AUD"})
    gst = str(commercial.get("gst_treatment") or "").strip()
    if not gst:
        issues.append({"code": "gst_treatment_required", "message": "GST处理尚未确认"})
    elif gst not in AUTHORITY_V2_GST_VALUES:
        issues.append({"code": "gst_treatment_invalid", "message": "GST处理值无效"})
    if str(commercial.get("gst_rate") or "").strip() != "0.10":
        issues.append({"code": "gst_rate_invalid", "message": "默认GST税率必须为10%"})
    if str(commercial.get("default_gst_amount") or "").strip() != "240.00":
        issues.append({"code": "default_gst_amount_invalid", "message": "默认GST金额必须为AUD 240.00"})
    if str(commercial.get("default_final_total") or "").strip() != "2640.00":
        issues.append({"code": "default_total_invalid", "message": "默认合同总额必须为AUD 2640.00"})
    if commercial.get("default_payment_plan") != "single_payment_at_signing_before_service":
        issues.append({"code": "payment_plan_invalid", "message": "默认付款节点无效"})
    try:
        methods = json.loads(str(commercial.get("payment_methods_json") or "[]"))
    except json.JSONDecodeError:
        methods = []
    if methods != list(AUTHORITY_V2_PAYMENT_METHODS):
        issues.append({"code": "payment_methods_invalid", "message": "付款方式目录无效"})
    for key, expected in {
        "installment_approval_role": "Admin",
        "discount_approval_role": "Admin",
        "discount_reason_required": "true",
        "base_fee_override_allowed": "false",
    }.items():
        if str(commercial.get(key) or "") != expected:
            issues.append({"code": f"{key}_invalid", "message": f"{key}配置无效"})
    return issues


def build_domain_v1_snapshot(
    template_type: str, commercial: Mapping[str, object] | None = None
) -> dict[str, object]:
    """Deterministically convert the current contract domain into editor V1."""
    config = contract_domain.contract_type_config(template_type)
    commercial_data = _default_commercial()
    commercial_data.update({key: str(value or "").strip() for key, value in (commercial or {}).items()})
    sections = contract_domain.compose_contract_sections(template_type)
    blocks: list[dict[str, object]] = [
        {
            "key": "fixed_parties",
            "title": "一、合同双方",
            "body": "学生：{{student_name}}\n签署人：{{signer_name}}\n电话：{{signer_phone}}\n邮箱：{{signer_email}}",
            "enabled": True,
            "fixed_layout": True,
            "required": True,
            "legal_kind": "identity",
            "required_variables": ["{{student_name}}", "{{signer_name}}", "{{signer_phone}}", "{{signer_email}}"],
        },
        {
            "key": "fixed_service_scope",
            "title": "二、服务项目及服务范围",
            "body": "{{service_scope}}",
            "enabled": True,
            "fixed_layout": True,
            "required": True,
            "legal_kind": "service_scope",
            "required_variables": ["{{service_scope}}"],
        },
    ]
    for title, body in sections.items():
        legal_kind = {
            "四、甲方权利义务": "privacy",
            "六、退费及违约责任": "refund",
            "九、合同补充、变更和生效": "termination",
        }.get(title, "")
        blocks.append({
            "key": PROFILE_KEYS[title], "title": title, "body": body, "enabled": True,
            "fixed_layout": False, "required": True, "legal_kind": legal_kind,
            "required_variables": [],
        })
    dispute_block = {
            "key": "legal_dispute", "title": "八、适用法律及争议解决",
            "body": "本合同适用澳大利亚维多利亚州法律。双方应先友好协商；协商不成的，按适用法律及相关管辖规则处理。",
            "enabled": True, "fixed_layout": False, "required": True,
            "legal_kind": "dispute", "required_variables": [],
        }
    special_block = {
            "key": "special_terms", "title": "十、特别条款",
            "body": "无。", "enabled": True, "fixed_layout": False,
            "required": True, "legal_kind": "special_terms", "required_variables": [],
        }
    by_key = {block["key"]: block for block in blocks}
    by_key["legal_dispute"] = dispute_block
    by_key["special_terms"] = special_block
    blocks = [by_key[key] for key in MASTER_SLOT_ORDER]
    services = [
        {
            "service_id": service_id,
            "label": contract_domain.service_label(service_id),
            "enabled": True,
            "default_selected": service_id in config.default_services,
        }
        for service_id in config.allowed_services
    ]
    return {
        "schema_version": 1,
        "template_type": config.type_id,
        "display_name": config.label,
        "contract_title": config.title,
        "clause_profile": config.clause_profile,
        "layout_policy": deepcopy(MASTER_LAYOUT_POLICY),
        "mandatory_display_bindings": contract_display_bindings.mandatory_profile(
            config.type_id
        ),
        "blocks": blocks,
        "services": services,
        "commercial": commercial_data,
    }


def render_template_editor_index(
    identities: Iterable[Mapping[str, object]], *, csrf: str,
    can_edit: bool, can_publish: bool, is_admin: bool = False,
) -> str:
    """Render one responsive V1 index used by Desktop and Mobile."""
    cards = []
    for identity in identities:
        template_type = html.escape(str(identity["template_type"]), quote=True)
        display_name = html.escape(str(identity["display_name"]), quote=True)
        version_id = html.escape(str(identity.get("current_published_version_id") or ""), quote=True)
        version_number = html.escape(
            str(identity.get("current_published_version_number") or ""), quote=True
        )
        published_version = f"V{version_number} · #{version_id}" if version_number else "未发布"
        actions = []
        if can_edit:
            actions.append(
                f'<form method="post" action="/templates/v1/{template_type}/actions/draft">'
                f'<input type="hidden" name="csrf" value="{html.escape(csrf, quote=True)}">'
                '<button class="primary">编辑模板</button></form>'
            )
        else:
            actions.append(
                f'<a class="button-link secondary" href="/templates/v1/{template_type}">查看当前模板</a>'
            )
        if is_admin:
            actions.append(
                f'<a class="button-link secondary" href="/templates/v1/{template_type}">查看与预览</a>'
            )
            actions.append(
                f'<form method="post" action="/templates/v1/{template_type}/actions/master-draft">'
                f'<input type="hidden" name="csrf" value="{html.escape(csrf, quote=True)}">'
                '<button class="secondary">新建锁定主版式草稿</button></form>'
            )
        if is_admin:
            template_facts = (
                f'<dl><div><dt>当前发布版本</dt><dd>{published_version}</dd></div>'
                '<div><dt>草稿版本</dt><dd>进入模板查看</dd></div>'
                '<div><dt>历史版本</dt><dd>永久保留并可重新发布</dd></div></dl>'
            )
        else:
            template_facts = (
                f'<dl><div><dt>当前模板</dt><dd>{published_version}</dd></div>'
                '<div><dt>使用范围</dt><dd>应用后供新建合同使用</dd></div></dl>'
            )
        cards.append(
            '<article class="contract-template-v1-card">'
            f'<header><h3>{display_name}</h3>{f"<code>{template_type}</code>" if is_admin else ""}</header>'
            f'{template_facts}'
            f'<div class="row-actions">{"".join(actions)}</div></article>'
        )
    if is_admin:
        publish_note = "可编辑及发布" if can_publish else ("可编辑，发布需授权" if can_edit else "只读")
    else:
        publish_note = "可编辑并可应用" if can_publish else ("可编辑，应用需授权" if can_edit else "只读")
    permission_panel = ""
    if is_admin:
        permission_panel = (
            '<details class="contract-template-permission-panel"><summary>合同模板权限</summary>'
            '<form method="post" action="/templates/v1/NewInternationalStudentContract/actions/permission">'
            f'<input type="hidden" name="csrf" value="{html.escape(csrf, quote=True)}">'
            '<label>员工User ID<input name="user_id" type="number" min="1" required></label>'
            '<label class="choice"><input name="can_edit" value="1" type="checkbox">允许编辑</label>'
            '<label class="choice"><input name="can_publish" value="1" type="checkbox">允许发布</label>'
            '<button class="primary">保存权限</button></form></details>'
        )
    return (
        '<section class="contract-template-editor-v1" data-contract-template-editor="v1">'
        f'<div class="panel-head"><div><h2>{"合同模板编辑器 V1" if is_admin else "合同模板管理"}</h2>'
        f'<p>{"三类合同使用独立草稿、发布版本和历史版本；旧V5/V6编辑器保持兼容。" if is_admin else "编辑并保存模板，应用后新建合同将自动使用当前模板。"}</p>'
        f'</div><span class="status-badge">{publish_note}</span></div>'
        f'<div class="contract-template-v1-grid">{"".join(cards)}</div>{permission_panel}</section>'
    )


def render_template_editor_detail(
    workspace: Mapping[str, object], *, csrf: str, can_edit: bool,
    can_publish: bool, lock: Mapping[str, object] | None = None,
    comparison: Mapping[str, object] | None = None,
    is_admin: bool = True,
) -> str:
    identity = workspace["identity"]
    active = workspace["active_version"]
    snapshot = workspace["snapshot"]
    editable = (
        can_edit
        and not bool(active["immutable"])
        and active["status"] in {"Draft", "Simulation", "Validation", "Preview"}
    )
    token = html.escape(str((lock or {}).get("token") or ""), quote=True)
    revision = html.escape(str(active["revision"]), quote=True)
    version_id = html.escape(str(active["id"]), quote=True)
    template_type = html.escape(str(identity["template_type"]), quote=True)
    master = is_master_snapshot(snapshot)
    block_cards = []
    for block in snapshot["blocks"]:
        key = html.escape(str(block["key"]), quote=True)
        can_reorder = editable and not bool(block.get("required"))
        block_cards.append(
            f'<article class="contract-template-block" draggable="{str(can_reorder).lower()}" data-block-key="{key}">'
            f'<label>条款标题<input name="title" value="{html.escape(str(block["title"]), quote=True)}" '
            f'{"" if editable and not block.get("fixed_layout") else "disabled"}></label>'
            f'<label>条款正文<textarea name="body" rows="7" {"" if editable else "disabled"}>'
            f'{html.escape(str(block["body"]))}</textarea></label>'
            f'<label class="choice"><input type="checkbox" name="enabled" {"checked" if block.get("enabled") else ""} '
            f'{"" if editable and not block.get("fixed_layout") else "disabled"}>启用</label>'
            f'<small>{"固定版式 · 变量受保护" if block.get("fixed_layout") else ("必要章节 · 标题须保留序号" if block.get("required") else "可增删、可在所属章节内排序")}</small>'
            f'{"" if block.get("required") or not editable else "<button type=\"button\" class=\"danger compact\" data-template-delete-block>删除区块</button>"}'
            '</article>'
        )
    service_rows = "".join(
        f'<label class="choice"><input type="checkbox" name="service_id" data-template-service value="{html.escape(str(item["service_id"]), quote=True)}" '
        f'{"checked" if item.get("default_selected") else ""} {"" if editable else "disabled"}>'
        f'{html.escape(str(item["label"]))}</label>'
        for item in snapshot["services"]
    )
    history_rows = []
    for row in workspace["versions"]:
        republish = ""
        if is_admin and can_edit and row["status"] in {"Published", "Historical"}:
            republish = (
                f' <button type="button" class="secondary compact" data-template-republish-version="{row["id"]}">'
                '重新发布为新版本</button>'
            )
        history_rows.append(
            f'<tr><td>V{row["version_number"]}</td><td>{html.escape(str(row["status"]))}</td>'
            f'<td>{html.escape(str(row["published_at"] or row["created_at"]))}</td>'
            f'<td><a href="/templates/v1/{template_type}?version_id={row["id"]}">查看</a>{republish}</td></tr>'
        )
    history = "".join(history_rows)
    compare_options = "".join(
        f'<option value="{row["id"]}">V{row["version_number"]} · {html.escape(str(row["status"]))}</option>'
        for row in workspace["versions"]
    )
    comparison_html = ""
    if comparison is not None:
        comparison_html = (
            '<div class="contract-template-comparison" data-template-comparison>'
            f'<p>区块变化：{len(comparison.get("blocks", []))}；服务范围变化：'
            f'{"是" if comparison.get("services_changed") else "否"}；费用配置变化：'
            f'{"是" if comparison.get("commercial_changed") else "否"}</p>'
            f'<pre>{html.escape(json.dumps(dict(comparison), ensure_ascii=False, indent=2))}</pre></div>'
        )
    commercial_labels = {
        "service_fee": "服务费用", "currency": "币种", "payment_terms": "付款条款",
        "company_name": "公司名称", "company_abn": "ABN", "company_address": "公司地址",
        "company_contact": "公司联系方式", "company_authorised_person": "授权代表",
        "refund_policy": "退款配置",
    }
    commercial_fields = "".join(
        f'<label>{label}<textarea rows="2" data-template-commercial="{key}" '
        f'{"" if editable else "disabled"}>{html.escape(str(snapshot["commercial"].get(key) or ""))}</textarea></label>'
        for key, label in commercial_labels.items()
    )
    add_block_slot = ""
    if master:
        add_block_slot = (
            '<label>新增条款所属章节<select data-template-add-block-slot '
            f'{"" if editable else "disabled"}>'
            + "".join(
                f'<option value="{html.escape(key, quote=True)}">{html.escape(str(next(block["title"] for block in snapshot["blocks"] if block["key"] == key)))}</option>'
                for key in EDITABLE_MASTER_SLOTS
            )
            + '</select></label>'
        )
    editor_mode = "自动保存已启用" if editable else "当前模板只读"
    if is_admin:
        editor_heading = (
            f"V{active['version_number']} · {html.escape(str(active['status']))} · {editor_mode}"
        )
        publication_controls = f"""
          <section><h3>模拟与预览</h3>
            <p>支持系统资料及只读真实学生资料，输出仅写入独立预览目录。</p>
            <label>真实Student ID（留空使用系统资料）<input data-template-student-id {"" if editable else "disabled"}></label>
            <label>签署人<select data-template-signer {"" if editable else "disabled"}><option value="father">父亲</option><option value="mother">母亲</option></select></label>
            <label>补填电话<input data-template-signer-phone {"" if editable else "disabled"}></label>
            <label>补填邮箱<input data-template-signer-email type="email" {"" if editable else "disabled"}></label>
            <button type="button" data-template-action="simulate" {"" if editable else "disabled"}>运行模拟</button>
            <button type="button" data-template-action="validate" {"" if can_publish and editable else "disabled"}>发布验证</button>
            <button type="button" data-template-action="preview" {"" if can_publish and editable else "disabled"}>确认发布前预览</button>
            <button type="button" data-template-action="publish" {"" if can_publish and editable else "disabled"}>发布新版本</button>
            <p class="contract-template-action-state" aria-live="polite">
              <strong>操作状态：</strong><span data-template-action-message>尚未执行</span>
            </p>
          </section>"""
        history_section = f"""
      <section><h3>版本差异与历史版本</h3>
        <form method="get" class="contract-template-compare-form">
          <label>版本一<select name="compare_left">{compare_options}</select></label>
          <label>版本二<select name="compare_right">{compare_options}</select></label>
          <button class="secondary">比较版本</button>
        </form>{comparison_html}<div class="table-wrap"><table>
        <thead><tr><th>版本</th><th>状态</th><th>时间</th><th>操作</th></tr></thead><tbody>{history}</tbody>
      </table></div></section>"""
    else:
        editor_heading = f"正在编辑 · {editor_mode}" if editable else editor_mode
        publication_controls = f"""
          <section class="contract-template-apply"><h3>应用模板</h3>
            <p>应用后，之后新建的该类型合同将自动使用此模板。</p>
            <button type="button" class="primary" data-template-action="apply"
              {"" if can_publish and editable else "disabled"}>应用模板</button>
            <p class="contract-template-action-state" aria-live="polite">
              <strong>应用状态：</strong><span data-template-action-message>尚未应用</span>
            </p>
          </section>"""
        history_section = ""
    return f"""
    <section class="contract-template-editor-v1 contract-template-editor-detail"
      data-contract-template-editor="v1" data-version-id="{version_id}"
      data-revision="{revision}" data-lock-token="{token}" data-template-type="{template_type}"
      data-editable="{str(editable).lower()}">
      <div class="panel-head"><div><h2>{html.escape(str(identity['display_name']))}</h2>
      <p>{editor_heading}</p></div>
      <a class="secondary button-link" href="/templates">返回模板管理</a></div>
      <div class="contract-template-editor-layout">
        <main><h3>条款区块</h3><div class="contract-template-block-list">{''.join(block_cards)}</div>{add_block_slot}
        <button type="button" data-template-add-block {"" if editable else "disabled"}>新增条款区块</button></main>
        <aside>
          <section><h3>服务范围</h3>{service_rows}</section>
          <section class="contract-template-commercial"><h3>费用、公司及退款配置</h3>{commercial_fields}</section>
          {publication_controls}
        </aside>
      </div>
      {history_section}
      <input type="hidden" data-template-csrf value="{html.escape(csrf, quote=True)}">
      <p class="contract-template-save-state" aria-live="polite">
        <strong class="contract-template-save-state-label">自动保存：</strong>
        <span data-template-save-message>已保存</span>
      </p>
    </section>"""


class ContractTemplateEditorService:
    def __init__(self, connection: sqlite3.Connection):
        self.conn = connection
        self.conn.row_factory = sqlite3.Row

    @staticmethod
    def _dict(row: sqlite3.Row | None) -> dict[str, object] | None:
        return dict(row) if row is not None else None

    def _audit(self, action: str, actor: int | None, identity_id=None, version_id=None, before=None, after=None):
        self.conn.execute(
            """INSERT INTO contract_template_audit_events(
               identity_id,version_id,actor_user_id,action,before_json,after_json,created_at
               ) VALUES(?,?,?,?,?,?,?)""",
            (identity_id, version_id, actor, action,
             _canonical_json(before) if before is not None else None,
             _canonical_json(after) if after is not None else None, _now()),
        )

    def _identity(self, template_type: str) -> sqlite3.Row:
        row = self.conn.execute(
            "SELECT * FROM contract_template_identities WHERE template_type=?", (template_type,)
        ).fetchone()
        if not row:
            raise ContractTemplateEditorError("合同模板不存在。")
        return row

    def _version(self, version_id: int) -> sqlite3.Row:
        row = self.conn.execute("SELECT * FROM contract_template_versions WHERE id=?", (version_id,)).fetchone()
        if not row:
            raise ContractTemplateEditorError("模板版本不存在。")
        return row

    def _snapshot(self, version_id: int) -> dict[str, object]:
        return json.loads(self._version(version_id)["snapshot_json"])

    def _insert_snapshot(self, identity_id: int, version_number: int, status: str,
                         snapshot: Mapping[str, object], actor: int, source_version_id=None,
                         immutable=False) -> int:
        stamp = _now()
        initial_status = "Draft" if immutable else status
        cursor = self.conn.execute(
            """INSERT INTO contract_template_versions(
               identity_id,version_number,status,source_version_id,revision,snapshot_json,
               content_hash,immutable,created_by,created_at,published_by,published_at
               ) VALUES(?,?,?,?,1,?,?,?,?,?,?,?)""",
            (identity_id, version_number, initial_status, source_version_id,
             _canonical_json(snapshot), _hash(snapshot), 0, actor, stamp,
             None, None),
        )
        version_id = int(cursor.lastrowid)
        for order, block in enumerate(snapshot["blocks"], 1):
            self.conn.execute(
                """INSERT INTO contract_template_clause_blocks(
                   version_id,block_key,title,body,sort_order,enabled,fixed_layout,required,
                   legal_kind,required_variables_json) VALUES(?,?,?,?,?,?,?,?,?,?)""",
                (version_id, block["key"], block["title"], block["body"], order,
                 int(bool(block.get("enabled", True))), int(bool(block.get("fixed_layout"))),
                 int(bool(block.get("required"))), block.get("legal_kind") or None,
                 _canonical_json(block.get("required_variables", []))),
            )
        for order, item in enumerate(snapshot["services"], 1):
            self.conn.execute(
                """INSERT INTO contract_template_service_items(
                   version_id,service_id,label,sort_order,enabled,default_selected
                   ) VALUES(?,?,?,?,?,?)""",
                (version_id, item["service_id"], item["label"], order,
                 int(bool(item.get("enabled", True))), int(bool(item.get("default_selected")))),
            )
        if immutable:
            self.conn.execute(
                """UPDATE contract_template_versions
                   SET status=?,immutable=1,published_by=?,published_at=? WHERE id=?""",
                (status, actor, stamp, version_id),
            )
        return version_id

    def _existing_company_commercial(self) -> dict[str, str]:
        try:
            rows = self.conn.execute(
                "SELECT key,value FROM system_settings WHERE key IN (?,?,?,?,?,?)",
                (
                    "company_legal_name", "company_abn_acn", "company_address",
                    "company_phone", "company_email", "company_authorised_person",
                ),
            ).fetchall()
        except sqlite3.OperationalError:
            return {}
        settings = {str(row["key"]): str(row["value"] or "").strip() for row in rows}
        contact = " / ".join(filter(None, (settings.get("company_phone"), settings.get("company_email"))))
        return {
            key: value
            for key, value in {
                "company_name": settings.get("company_legal_name", ""),
                "company_abn": settings.get("company_abn_acn", ""),
                "company_address": settings.get("company_address", ""),
                "company_contact": contact,
                "company_authorised_person": settings.get("company_authorised_person", ""),
            }.items()
            if value
        }

    def import_domain_v1(self, actor_user_id: int) -> list[dict[str, object]]:
        imported = []
        commercial = self._existing_company_commercial()
        for template_type, display_name in contract_domain.contract_type_options():
            existing = self.conn.execute(
                "SELECT id FROM contract_template_identities WHERE template_type=?", (template_type,)
            ).fetchone()
            if existing:
                imported.append(self.current_published(template_type))
                continue
            stamp = _now()
            cursor = self.conn.execute(
                "INSERT INTO contract_template_identities(template_type,display_name,created_at,updated_at) VALUES(?,?,?,?)",
                (template_type, display_name, stamp, stamp),
            )
            identity_id = int(cursor.lastrowid)
            snapshot = build_domain_v1_snapshot(template_type, commercial)
            version_id = self._insert_snapshot(identity_id, 1, "Published", snapshot, actor_user_id, immutable=True)
            self.conn.execute(
                "UPDATE contract_template_identities SET current_published_version_id=? WHERE id=?",
                (version_id, identity_id),
            )
            self._audit("template_v1_imported", actor_user_id, identity_id, version_id, after=snapshot)
            imported.append(self._dict(self._version(version_id)))
        self.conn.commit()
        return imported

    def create_authority_v2_drafts(self, actor_user_id: int) -> list[dict[str, object]]:
        """Create one idempotent, non-publishable Authority Draft V2 per identity."""
        drafts: list[dict[str, object]] = []
        for template_type, _display_name in contract_domain.contract_type_options():
            identity = self._identity(template_type)
            source_id = int(identity["current_published_version_id"])
            source = self._version(source_id)
            source_snapshot = json.loads(source["snapshot_json"])
            if (
                source["status"] == "Published"
                and (source_snapshot.get("commercial") or {}).get("commercial_schema")
                == "authority_v2"
            ):
                # Conversion is a one-time V1 -> Draft V2 operation.  A later
                # Published V2 must never cause an implicit V3 draft.
                continue
            existing = None
            for row in self.conn.execute(
                "SELECT * FROM contract_template_versions WHERE identity_id=? AND source_version_id=? "
                "AND status IN ('Draft','Simulation','Validation','Preview') ORDER BY id",
                (identity["id"], source_id),
            ):
                snapshot = json.loads(row["snapshot_json"])
                if (snapshot.get("commercial") or {}).get("commercial_schema") == "authority_v2":
                    existing = dict(row)
                    break
            if existing is None:
                snapshot = deepcopy(source_snapshot)
                snapshot["mandatory_display_bindings"] = (
                    contract_display_bindings.mandatory_profile(template_type)
                )
                snapshot["commercial"] = authority_v2_commercial(
                    template_type, snapshot.get("commercial") or {}
                )
                next_number = int(self.conn.execute(
                    "SELECT COALESCE(MAX(version_number),0)+1 FROM contract_template_versions WHERE identity_id=?",
                    (identity["id"],),
                ).fetchone()[0])
                version_id = self._insert_snapshot(
                    int(identity["id"]), next_number, "Draft", snapshot,
                    actor_user_id, source_version_id=source_id, immutable=False,
                )
                self._audit(
                    "template_authority_v2_draft_created", actor_user_id,
                    int(identity["id"]), version_id,
                    after={"source_version_id": source_id, "gst_treatment": "EXCLUSIVE"},
                )
                existing = dict(self._version(version_id))
            existing["template_type"] = template_type
            drafts.append(existing)
        self.conn.commit()
        return drafts

    def configure_authority_v2_draft(
        self, version_id: int, *, gst_treatment: str, expected_revision: int,
        lock_token: str, actor_user_id: int,
    ) -> dict[str, object]:
        version = self._require_edit(version_id, expected_revision, lock_token, actor_user_id)
        snapshot = self._snapshot(version_id)
        commercial = dict(snapshot.get("commercial") or {})
        if commercial.get("commercial_schema") != "authority_v2":
            raise TemplateRuleError("模板不是Authority V2草稿。")
        gst = str(gst_treatment or "").strip()
        if gst not in AUTHORITY_V2_GST_VALUES:
            raise TemplateRuleError("GST处理值无效。")
        commercial["gst_treatment"] = gst
        snapshot["commercial"] = commercial
        next_revision = int(version["revision"]) + 1
        self.conn.execute(
            "UPDATE contract_template_versions SET revision=?,snapshot_json=?,content_hash=? WHERE id=?",
            (next_revision, _canonical_json(snapshot), _hash(snapshot), version_id),
        )
        self._audit(
            "template_authority_v2_gst_configured", actor_user_id,
            int(version["identity_id"]), version_id,
            after={"gst_treatment": gst},
        )
        self.conn.commit()
        return dict(self._version(version_id))

    def list_templates(self) -> list[dict[str, object]]:
        order = {value: index for index, (value, _label) in enumerate(contract_domain.contract_type_options())}
        rows = [dict(row) for row in self.conn.execute(
            """SELECT identities.*,
                      versions.version_number AS current_published_version_number
               FROM contract_template_identities identities
               LEFT JOIN contract_template_versions versions
                 ON versions.id=identities.current_published_version_id"""
        )]
        return sorted(rows, key=lambda row: order.get(str(row["template_type"]), 99))

    def current_published(self, template_type: str) -> dict[str, object]:
        identity = self._identity(template_type)
        row = self._version(int(identity["current_published_version_id"]))
        return dict(row)

    def version_history(self, template_type: str) -> list[dict[str, object]]:
        identity = self._identity(template_type)
        return [dict(row) for row in self.conn.execute(
            "SELECT * FROM contract_template_versions WHERE identity_id=? ORDER BY version_number DESC",
            (identity["id"],),
        )]

    def workspace(self, template_type: str, version_id: int | None = None) -> dict[str, object]:
        identity = dict(self._identity(template_type))
        versions = self.version_history(template_type)
        if version_id is None:
            draft = next((row for row in versions if row["status"] in {"Draft", "Simulation", "Validation", "Preview"}), None)
            active = draft or self.current_published(template_type)
        else:
            active = dict(self._version(version_id))
            if int(active["identity_id"]) != int(identity["id"]):
                raise TemplateRuleError("模板版本不属于所选合同类型。")
        return {
            "identity": identity, "active_version": active,
            "snapshot": json.loads(active["snapshot_json"]), "versions": versions,
        }

    def _active_draft(self, identity_id: int) -> sqlite3.Row | None:
        return self.conn.execute(
            """SELECT * FROM contract_template_versions
               WHERE identity_id=?
                 AND status IN ('Draft','Simulation','Validation','Preview')
               ORDER BY version_number DESC,id DESC LIMIT 1""",
            (identity_id,),
        ).fetchone()

    def _repair_missing_inherited_metadata_on_reuse(
        self, version: sqlite3.Row, actor_user_id: int
    ) -> sqlite3.Row:
        """Materialize missing authoritative metadata for a legacy draft only.

        Older active drafts can predate the version-snapshot profile or can
        have lost Authority V2 metadata in an old master-layout clone.  Reuse
        is the safe healing point: mandatory bindings come from the canonical
        code profile, while commercial metadata comes only from the draft's
        source version and only for absent allowlisted keys.  Present values
        are never overwritten.
        """
        self._require_mutable(int(version["id"]))
        snapshot = json.loads(str(version["snapshot_json"]))
        expected = contract_display_bindings.mandatory_profile(
            snapshot.get("template_type")
        )
        actual = snapshot.get("mandatory_display_bindings")
        repaired_profile = actual is None
        if actual is not None:
            if actual != expected:
                raise TemplateRuleError(
                    "mandatory display binding profile 已存在但与当前权威规则不匹配，"
                    "不允许自动覆盖。"
                )
        else:
            snapshot["mandatory_display_bindings"] = expected

        commercial = snapshot.get("commercial")
        if not isinstance(commercial, dict):
            raise TemplateRuleError("模板 commercial metadata 结构无效。")
        restored_commercial: dict[str, object] = {}
        source_version_id = version["source_version_id"]
        if source_version_id is not None:
            source = self._version(int(source_version_id))
            if int(source["identity_id"]) != int(version["identity_id"]):
                raise TemplateRuleError("Draft 来源版本与模板身份不匹配。")
            source_snapshot = json.loads(str(source["snapshot_json"]))
            source_commercial = source_snapshot.get("commercial")
            if (
                isinstance(source_commercial, dict)
                and source_commercial.get("commercial_schema") == "authority_v2"
            ):
                for key in sorted(AUTHORITY_V2_INHERITED_METADATA_FIELDS):
                    if key not in commercial and key in source_commercial:
                        commercial[key] = deepcopy(source_commercial[key])
                        restored_commercial[key] = deepcopy(source_commercial[key])
        snapshot["commercial"] = commercial

        if not repaired_profile and not restored_commercial:
            return version

        before_hash = str(version["content_hash"])
        next_revision = int(version["revision"]) + 1
        after_hash = _hash(snapshot)
        cursor = self.conn.execute(
            """UPDATE contract_template_versions
               SET revision=?,snapshot_json=?,content_hash=?
               WHERE id=? AND revision=?""",
            (
                next_revision,
                _canonical_json(snapshot),
                after_hash,
                int(version["id"]),
                int(version["revision"]),
            ),
        )
        if cursor.rowcount != 1:
            raise StaleRevision("模板版本已变化，未执行 mandatory profile 治理。")
        audit_common_before = {
            "content_hash": before_hash,
            "revision": int(version["revision"]),
        }
        audit_common_after = {
            "content_hash": after_hash,
            "revision": next_revision,
        }
        if repaired_profile:
            self._audit(
                "template_mandatory_display_binding_profile_repaired",
                actor_user_id,
                int(version["identity_id"]),
                int(version["id"]),
                before={
                    **audit_common_before,
                    "mandatory_display_bindings": None,
                },
                after={
                    **audit_common_after,
                    "mandatory_display_bindings": expected,
                },
            )
        if restored_commercial:
            self._audit(
                "template_authority_v2_commercial_metadata_repaired",
                actor_user_id,
                int(version["identity_id"]),
                int(version["id"]),
                before={
                    **audit_common_before,
                    "missing_fields": sorted(restored_commercial),
                    "source_version_id": int(source_version_id),
                },
                after={
                    **audit_common_after,
                    "restored_fields": sorted(restored_commercial),
                    "source_version_id": int(source_version_id),
                },
            )
        return self._version(int(version["id"]))

    def create_draft(self, template_type: str, actor_user_id: int) -> dict[str, object]:
        with self._immediate_transaction():
            identity = self._identity(template_type)
            existing = self._active_draft(int(identity["id"]))
            if existing:
                return dict(
                    self._repair_missing_inherited_metadata_on_reuse(
                        existing, actor_user_id
                    )
                )
            published = self.current_published(template_type)
            return self.create_draft_from_version(
                int(published["id"]), actor_user_id, commit=False
            )

    def create_master_layout_draft(self, template_type: str, actor_user_id: int) -> dict[str, object]:
        """Create an explicit canonical master draft without rewriting Published history."""
        with self._immediate_transaction():
            identity = self._identity(template_type)
            existing = self._active_draft(int(identity["id"]))
            if existing:
                return dict(
                    self._repair_missing_inherited_metadata_on_reuse(
                        existing, actor_user_id
                    )
                )
            published = self.current_published(template_type)
            source_snapshot = json.loads(str(published["snapshot_json"]))
            commercial = dict(source_snapshot.get("commercial") or self._existing_company_commercial())
            snapshot = build_domain_v1_snapshot(template_type, commercial)
            identity_id = int(published["identity_id"])
            next_number = int(self.conn.execute(
                "SELECT COALESCE(MAX(version_number),0)+1 FROM contract_template_versions WHERE identity_id=?",
                (identity_id,),
            ).fetchone()[0])
            draft_id = self._insert_snapshot(
                identity_id, next_number, "Draft", snapshot, actor_user_id,
                source_version_id=int(published["id"]), immutable=False,
            )
            self._audit(
                "template_master_layout_draft_created", actor_user_id,
                identity_id, draft_id, after={"source_version_id": int(published["id"]), "layout_policy": MASTER_LAYOUT_POLICY},
            )
            return dict(self._version(draft_id))

    def create_draft_from_version(
        self, version_id: int, actor_user_id: int, *, commit: bool = True
    ) -> dict[str, object]:
        if commit:
            with self._immediate_transaction():
                return self.create_draft_from_version(
                    version_id, actor_user_id, commit=False
                )
        source = self._version(version_id)
        identity_id = int(source["identity_id"])
        existing = self._active_draft(identity_id)
        if existing:
            return dict(
                self._repair_missing_inherited_metadata_on_reuse(
                    existing, actor_user_id
                )
            )
        next_number = int(self.conn.execute(
            "SELECT COALESCE(MAX(version_number),0)+1 FROM contract_template_versions WHERE identity_id=?",
            (identity_id,),
        ).fetchone()[0])
        snapshot = json.loads(source["snapshot_json"])
        snapshot["mandatory_display_bindings"] = contract_display_bindings.mandatory_profile(
            snapshot.get("template_type")
        )
        draft_id = self._insert_snapshot(
            identity_id, next_number, "Draft", snapshot,
            actor_user_id, source_version_id=version_id, immutable=False,
        )
        self._audit("template_draft_created", actor_user_id, identity_id, draft_id,
                    after={"source_version_id": version_id})
        return dict(self._version(draft_id))

    def acquire_lock(self, template_type: str, version_id: int, owner_user_id: int,
                     ttl_minutes: int = 15) -> dict[str, object]:
        now = datetime.now(timezone.utc)
        expires = (now + timedelta(minutes=ttl_minutes)).replace(microsecond=0).isoformat()
        with self._immediate_transaction():
            identity = self._identity(template_type)
            version = self._version(version_id)
            if int(version["identity_id"]) != int(identity["id"]):
                raise TemplateRuleError("模板与版本不匹配。")
            self._require_mutable(version_id)
            existing = self.conn.execute(
                "SELECT * FROM contract_template_edit_locks WHERE identity_id=?",
                (identity["id"],),
            ).fetchone()
            existing_is_live = False
            if existing:
                try:
                    existing_is_live = (
                        datetime.fromisoformat(str(existing["expires_at"])) > now
                    )
                except (TypeError, ValueError):
                    existing_is_live = False
            if (
                existing
                and existing_is_live
                and int(existing["owner_user_id"]) != int(owner_user_id)
            ):
                raise EditLockConflict(
                    "此模板正在被另一位用户编辑，你当前修改尚未应用。"
                )
            same_owner_version = bool(
                existing
                and int(existing["owner_user_id"]) == int(owner_user_id)
                and int(existing["version_id"]) == int(version_id)
            )
            token = str(existing["token"]) if same_owner_version else uuid.uuid4().hex
            acquired_at = (
                str(existing["acquired_at"])
                if same_owner_version
                else now.replace(microsecond=0).isoformat()
            )
            self.conn.execute(
                """INSERT INTO contract_template_edit_locks(
                   identity_id,version_id,owner_user_id,token,acquired_at,expires_at
                   ) VALUES(?,?,?,?,?,?) ON CONFLICT(identity_id) DO UPDATE SET
                   version_id=excluded.version_id,owner_user_id=excluded.owner_user_id,
                   token=excluded.token,acquired_at=excluded.acquired_at,
                   expires_at=excluded.expires_at""",
                (
                    identity["id"], version_id, owner_user_id, token,
                    acquired_at, expires,
                ),
            )
            self._audit(
                "template_lock_renewed" if same_owner_version else "template_lock_acquired",
                owner_user_id, identity["id"], version_id,
            )
        return dict(self.conn.execute(
            "SELECT * FROM contract_template_edit_locks WHERE identity_id=?", (identity["id"],)
        ).fetchone())

    def live_lock(
        self, version_id: int, owner_user_id: int
    ) -> dict[str, object] | None:
        """Return only a currently usable lock for editor page hydration."""
        lock = self.conn.execute(
            "SELECT * FROM contract_template_edit_locks "
            "WHERE version_id=? AND owner_user_id=?",
            (version_id, owner_user_id),
        ).fetchone()
        if not lock:
            return None
        try:
            if (
                datetime.fromisoformat(str(lock["expires_at"]))
                <= datetime.now(timezone.utc)
            ):
                return None
        except (TypeError, ValueError):
            return None
        return dict(lock)

    def _require_mutable(self, version_id: int) -> sqlite3.Row:
        version = self._version(version_id)
        if version["immutable"] or version["status"] not in {"Draft", "Simulation", "Validation", "Preview"}:
            raise ImmutableVersionError("已发布或历史模板版本不可修改。")
        return version

    def _require_live_lock(
        self, version_id: int, lock_token: str, actor_user_id: int
    ) -> sqlite3.Row:
        lock = self.conn.execute(
            "SELECT * FROM contract_template_edit_locks WHERE version_id=? AND token=?",
            (version_id, lock_token),
        ).fetchone()
        if not lock or int(lock["owner_user_id"]) != int(actor_user_id):
            current = self.conn.execute(
                "SELECT * FROM contract_template_edit_locks WHERE version_id=?",
                (version_id,),
            ).fetchone()
            if current:
                try:
                    current_is_live = (
                        datetime.fromisoformat(str(current["expires_at"]))
                        > datetime.now(timezone.utc)
                    )
                except (TypeError, ValueError):
                    current_is_live = False
                if (
                    current_is_live
                    and int(current["owner_user_id"]) != int(actor_user_id)
                ):
                    raise EditLockConflict(
                        "此模板正在被另一位用户编辑，你当前修改尚未应用。"
                    )
            raise EditLockConflict(
                "编辑会话已更新，当前修改尚未保存，系统将安全重试。"
            )
        try:
            expired = (
                datetime.fromisoformat(str(lock["expires_at"]))
                <= datetime.now(timezone.utc)
            )
        except (TypeError, ValueError):
            expired = True
        if expired:
            was_in_transaction = self.conn.in_transaction
            expires = (
                datetime.now(timezone.utc) + timedelta(minutes=15)
            ).replace(microsecond=0).isoformat()
            cursor = self.conn.execute(
                """UPDATE contract_template_edit_locks SET expires_at=?
                   WHERE version_id=? AND token=? AND owner_user_id=?
                     AND expires_at=?""",
                (
                    expires, version_id, lock_token, actor_user_id,
                    str(lock["expires_at"]),
                ),
            )
            if cursor.rowcount != 1:
                if not was_in_transaction:
                    self.conn.rollback()
                raise EditLockConflict(
                    "编辑会话已更新，当前修改尚未保存，系统将安全重试。"
                )
            version = self._version(version_id)
            self._audit(
                "template_lock_reclaimed", actor_user_id,
                version["identity_id"], version_id,
            )
            if not was_in_transaction:
                self.conn.commit()
            lock = self.conn.execute(
                "SELECT * FROM contract_template_edit_locks "
                "WHERE version_id=? AND token=? AND owner_user_id=?",
                (version_id, lock_token, actor_user_id),
            ).fetchone()
        return lock

    @contextmanager
    def _immediate_transaction(self):
        if self.conn.in_transaction:
            raise TemplateRuleError("当前数据库连接已有未完成事务，无法安全发布。")
        self.conn.execute("BEGIN IMMEDIATE")
        try:
            yield
        except Exception:
            self.conn.rollback()
            raise
        else:
            self.conn.commit()

    def _require_edit(self, version_id: int, expected_revision: int | None, lock_token: str, actor_user_id: int):
        version = self._require_mutable(version_id)
        if expected_revision is not None and int(version["revision"]) != int(expected_revision):
            raise StaleRevision("模板版本已变化，请刷新后重试。")
        self._require_live_lock(version_id, lock_token, actor_user_id)
        return version

    def _refresh_snapshot(self, version_id: int, revision: int):
        snapshot = self._snapshot(version_id)
        snapshot["blocks"] = [
            {
                "key": row["block_key"], "title": row["title"], "body": row["body"],
                "enabled": bool(row["enabled"]), "fixed_layout": bool(row["fixed_layout"]),
                "required": bool(row["required"]), "legal_kind": row["legal_kind"] or "",
                "required_variables": json.loads(row["required_variables_json"]),
            }
            for row in self.conn.execute(
                "SELECT * FROM contract_template_clause_blocks WHERE version_id=? ORDER BY sort_order", (version_id,)
            )
        ]
        snapshot["services"] = [
            {
                "service_id": row["service_id"], "label": row["label"],
                "enabled": bool(row["enabled"]), "default_selected": bool(row["default_selected"]),
            }
            for row in self.conn.execute(
                "SELECT * FROM contract_template_service_items WHERE version_id=? ORDER BY sort_order", (version_id,)
            )
        ]
        self.conn.execute(
            "UPDATE contract_template_versions SET revision=?,snapshot_json=?,content_hash=? WHERE id=?",
            (revision, _canonical_json(snapshot), _hash(snapshot), version_id),
        )

    def autosave_block(self, version_id: int, block_key: str, *, title: str, body: str,
                       enabled: bool, expected_revision: int, lock_token: str,
                       actor_user_id: int) -> dict[str, object]:
        version = self._require_edit(version_id, expected_revision, lock_token, actor_user_id)
        block = self.conn.execute(
            "SELECT * FROM contract_template_clause_blocks WHERE version_id=? AND block_key=?",
            (version_id, block_key),
        ).fetchone()
        if not block:
            raise TemplateRuleError("条款区块不存在。")
        required_variables = json.loads(block["required_variables_json"])
        missing = [value for value in required_variables if value not in body]
        if missing:
            raise TemplateRuleError("必要变量不得删除或改变。")
        snapshot = self._snapshot(version_id)
        if is_master_snapshot(snapshot):
            parent = master_block_parent(block_key)
            if parent is None:
                raise TemplateRuleError("条款不属于受保护的主版式章节。")
            if block_key in MASTER_SLOT_ORDER:
                if not enabled:
                    raise TemplateRuleError("主版式必要章节不得停用。")
                if block_key in EDITABLE_MASTER_SLOTS:
                    ordinal = MASTER_HEADINGS[block_key].split("、", 1)[0] + "、"
                    if not str(title).strip().startswith(ordinal):
                        raise TemplateRuleError("可编辑章节标题必须保留主版式序号。")
                elif str(title).strip() != str(block["title"]):
                    raise TemplateRuleError("主版式固定标题位置与锚点不得改变。")
            elif enabled and (not str(title).strip() or not str(body).strip()):
                raise TemplateRuleError("已启用的新增条款标题和正文不得为空。")
        if block["fixed_layout"] and str(title).strip() != str(block["title"]):
            raise TemplateRuleError("固定版式标题不得改变。")
        if block["fixed_layout"] and VARIABLE_PATTERN.findall(body) != required_variables:
            raise TemplateRuleError("固定版式中的必要变量位置和顺序不得改变。")
        if block["fixed_layout"] and not enabled:
            raise TemplateRuleError("固定版式区块不得停用。")
        before = dict(block)
        self.conn.execute(
            "UPDATE contract_template_clause_blocks SET title=?,body=?,enabled=? WHERE id=?",
            (str(title).strip(), str(body), int(bool(enabled)), block["id"]),
        )
        self._refresh_snapshot(version_id, int(version["revision"]) + 1)
        self._audit("template_block_autosaved", actor_user_id, version["identity_id"], version_id,
                    before=before, after={"block_key": block_key, "title": title, "body": body, "enabled": enabled})
        self.conn.commit()
        return dict(self._version(version_id))

    def autosave_settings(self, version_id: int, *, commercial: Mapping[str, object],
                          services: Iterable[Mapping[str, object]], expected_revision: int,
                          lock_token: str, actor_user_id: int) -> dict[str, object]:
        version = self._require_edit(version_id, expected_revision, lock_token, actor_user_id)
        snapshot = self._snapshot(version_id)
        template_type = str(snapshot["template_type"])
        allowed = contract_domain.allowed_service_ids(template_type)
        submitted = {str(item.get("service_id") or ""): item for item in services}
        normalized_services = []
        for service_id in allowed:
            item = submitted.get(service_id, {})
            normalized_services.append({
                "service_id": service_id,
                "label": contract_domain.service_label(service_id),
                "enabled": bool(item.get("enabled", False)),
                "default_selected": bool(item.get("default_selected", False)) and bool(item.get("enabled", False)),
            })
        current_commercial = dict(snapshot.get("commercial") or {})
        current_commercial.update({
            key: str(commercial.get(key) or "").strip()
            for key in _default_commercial()
        })
        snapshot["commercial"] = current_commercial
        snapshot["services"] = normalized_services
        self.conn.execute("DELETE FROM contract_template_service_items WHERE version_id=?", (version_id,))
        for order, item in enumerate(normalized_services, 1):
            self.conn.execute(
                """INSERT INTO contract_template_service_items(
                   version_id,service_id,label,sort_order,enabled,default_selected
                   ) VALUES(?,?,?,?,?,?)""",
                (version_id, item["service_id"], item["label"], order,
                 int(item["enabled"]), int(item["default_selected"])),
            )
        next_revision = int(version["revision"]) + 1
        self.conn.execute(
            "UPDATE contract_template_versions SET revision=?,snapshot_json=?,content_hash=? WHERE id=?",
            (next_revision, _canonical_json(snapshot), _hash(snapshot), version_id),
        )
        self._audit("template_settings_autosaved", actor_user_id, version["identity_id"], version_id,
                    after={"commercial": snapshot["commercial"], "services": normalized_services})
        self.conn.commit()
        return dict(self._version(version_id))

    def reorder_blocks(self, version_id: int, ordered_keys: Iterable[str], *, expected_revision: int,
                       lock_token: str, actor_user_id: int) -> dict[str, object]:
        version = self._require_edit(version_id, expected_revision, lock_token, actor_user_id)
        master = is_master_snapshot(self._snapshot(version_id))
        all_rows = list(self.conn.execute(
            "SELECT block_key FROM contract_template_clause_blocks WHERE version_id=? ORDER BY sort_order", (version_id,)
        ))
        existing = [row["block_key"] for row in all_rows]
        requested = list(dict.fromkeys(str(key) for key in ordered_keys))
        if any(key not in existing for key in requested):
            raise TemplateRuleError("条款排序包含未知区块。")
        if master:
            if len(requested) != len(existing) or set(requested) != set(existing):
                raise TemplateRuleError("主版式排序必须包含全部章节与条款。")
            if [key for key in requested if key in MASTER_SLOT_ORDER] != list(MASTER_SLOT_ORDER):
                raise TemplateRuleError("主版式章节锚点顺序不得改变。")
            active_parent = None
            for key in requested:
                if key in MASTER_SLOT_ORDER:
                    active_parent = key
                elif master_block_parent(key) != active_parent:
                    raise TemplateRuleError("可编辑条款只能在所属主版式章节内排序。")
            final = requested
        else:
            final = requested + [key for key in existing if key not in requested]
        for index, key in enumerate(final, 1):
            self.conn.execute(
                "UPDATE contract_template_clause_blocks SET sort_order=? WHERE version_id=? AND block_key=?",
                (-index, version_id, key),
            )
        for index, key in enumerate(final, 1):
            self.conn.execute(
                "UPDATE contract_template_clause_blocks SET sort_order=? WHERE version_id=? AND block_key=?",
                (index, version_id, key),
            )
        self._refresh_snapshot(version_id, int(version["revision"]) + 1)
        self._audit("template_blocks_reordered", actor_user_id, version["identity_id"], version_id, after=final)
        self.conn.commit()
        return dict(self._version(version_id))

    def create_block(self, version_id: int, block_key: str, title: str, body: str,
                     expected_revision: int, lock_token: str, actor_user_id: int) -> dict[str, object]:
        version = self._require_edit(version_id, expected_revision, lock_token, actor_user_id)
        master = is_master_snapshot(self._snapshot(version_id))
        parent = master_block_parent(block_key) if master else None
        if master and (block_key in MASTER_SLOT_ORDER or parent not in EDITABLE_MASTER_SLOTS):
            raise TemplateRuleError("主版式只允许在正文章节内新增内容条款。")
        if master and (not str(title).strip() or not str(body).strip()):
            raise TemplateRuleError("新增内容条款标题和正文不得为空。")
        if master:
            rows = list(self.conn.execute(
                "SELECT block_key,sort_order FROM contract_template_clause_blocks WHERE version_id=? ORDER BY sort_order",
                (version_id,),
            ))
            owned_orders = [
                int(row["sort_order"])
                for row in rows
                if row["block_key"] == parent or master_block_parent(row["block_key"]) == parent
            ]
            next_order = max(owned_orders) + 1
            displaced = [row for row in rows if int(row["sort_order"]) >= next_order]
            for row in displaced:
                self.conn.execute(
                    "UPDATE contract_template_clause_blocks SET sort_order=? WHERE version_id=? AND block_key=?",
                    (-int(row["sort_order"]), version_id, row["block_key"]),
                )
            for row in displaced:
                self.conn.execute(
                    "UPDATE contract_template_clause_blocks SET sort_order=? WHERE version_id=? AND block_key=?",
                    (int(row["sort_order"]) + 1, version_id, row["block_key"]),
                )
        else:
            next_order = int(self.conn.execute(
                "SELECT COALESCE(MAX(sort_order),0)+1 FROM contract_template_clause_blocks WHERE version_id=?",
                (version_id,),
            ).fetchone()[0])
        self.conn.execute(
            """INSERT INTO contract_template_clause_blocks(
               version_id,block_key,title,body,sort_order,enabled,fixed_layout,required,required_variables_json
               ) VALUES(?,?,?,?,?,1,0,0,'[]')""",
            (version_id, block_key, title, body, next_order),
        )
        self._refresh_snapshot(version_id, int(version["revision"]) + 1)
        self._audit("template_block_created", actor_user_id, version["identity_id"], version_id, after={"block_key": block_key})
        self.conn.commit()
        return dict(self._version(version_id))

    def delete_block(self, version_id: int, block_key: str, expected_revision: int,
                     lock_token: str, actor_user_id: int) -> dict[str, object]:
        version = self._require_edit(version_id, expected_revision, lock_token, actor_user_id)
        master = is_master_snapshot(self._snapshot(version_id))
        if master and block_key in MASTER_SLOT_ORDER:
            raise TemplateRuleError("主版式必要章节锚点不得删除。")
        if master and master_block_parent(block_key) not in EDITABLE_MASTER_SLOTS:
            raise TemplateRuleError("条款不属于可编辑正文章节。")
        block = self.conn.execute(
            "SELECT * FROM contract_template_clause_blocks WHERE version_id=? AND block_key=?", (version_id, block_key)
        ).fetchone()
        if not block:
            raise TemplateRuleError("条款区块不存在。")
        if block["fixed_layout"] or block["required"]:
            raise TemplateRuleError("固定版式或必要条款不得删除。")
        self.conn.execute("DELETE FROM contract_template_clause_blocks WHERE id=?", (block["id"],))
        remaining = list(self.conn.execute(
            "SELECT block_key,sort_order FROM contract_template_clause_blocks WHERE version_id=? ORDER BY sort_order",
            (version_id,),
        ))
        for row in remaining:
            self.conn.execute(
                "UPDATE contract_template_clause_blocks SET sort_order=? WHERE version_id=? AND block_key=?",
                (-int(row["sort_order"]), version_id, row["block_key"]),
            )
        for index, row in enumerate(remaining, 1):
            self.conn.execute(
                "UPDATE contract_template_clause_blocks SET sort_order=? WHERE version_id=? AND block_key=?",
                (index, version_id, row["block_key"]),
            )
        self._refresh_snapshot(version_id, int(version["revision"]) + 1)
        self._audit("template_block_deleted", actor_user_id, version["identity_id"], version_id, before=dict(block))
        self.conn.commit()
        return dict(self._version(version_id))

    @staticmethod
    def _safe_output_root(path: Path) -> Path:
        resolved = path.expanduser().resolve()
        if str(resolved).startswith("/Volumes/Server/Students"):
            raise TemplateRuleError("模拟文件不得写入正式学生NAS目录。")
        resolved.mkdir(parents=True, exist_ok=True)
        return resolved

    @staticmethod
    def _render_docx_artifact(model: Mapping[str, object], path: Path) -> dict[str, object]:
        master = (model.get("layout_policy") or {}).get("id") == "AAIEContractMasterLayoutV1"
        artifact_bytes = (
            _embed_docx_font(render_master_docx_bytes(model))
            if master else _build_docx_artifact_bytes(model)
        )
        path.write_bytes(artifact_bytes)
        if _docx_layout_fingerprint(path) != _docx_layout_fingerprint(artifact_bytes):
            raise TemplateRuleError("保存后DOCX成品布局包与确定性生成结果不一致。")
        expected_trace = _semantic_trace_from_model(model)
        if master:
            from docx import Document
            document = Document(path)
            visible = "\n".join(paragraph.text for paragraph in document.paragraphs)
            visible += "\n" + "\n".join(
                cell.text for table in document.tables for row in table.rows for cell in row.cells
            )
            if len(document.inline_shapes) < 1 or not all(
                marker in visible for marker in ("一、合同双方", "十、特别条款", "十一、签署", "AAIE / 甲方（签字 / 盖章）", "委托方（乙方签字）")
            ):
                raise TemplateRuleError("保存后DOCX成品缺少主版式或签署锚点。")
            semantic_trace = expected_trace
        else:
            semantic_trace = _require_post_save_trace(
                "DOCX", _extract_docx_semantic_trace(path), expected_trace
            )
        return {
            "semantic_trace": semantic_trace,
            "artifact_hash": _file_hash(path),
            "bytes": path.stat().st_size,
        }

    @staticmethod
    def _render_pdf_artifact(model: Mapping[str, object], path: Path) -> dict[str, object]:
        if (model.get("layout_policy") or {}).get("id") == "AAIEContractMasterLayoutV1":
            artifact_bytes, manifest = render_master_pdf_bytes(model)
            path.write_bytes(artifact_bytes)
            from pypdf import PdfReader
            reader = PdfReader(path)
            final_text = reader.pages[-1].extract_text()
            if not all(value in final_text for value in ("十一、签署", "甲方（签字 / 盖章）", "委托方（乙方签字）")):
                raise TemplateRuleError("保存后PDF成品缺少同页双栏签署锚点。")
            if manifest["signature_aaie"]["page"] != len(reader.pages) or manifest["signature_family"]["page"] != len(reader.pages):
                raise TemplateRuleError("保存后PDF签署锚点未保留在最终页。")
            return {
                "semantic_trace": _semantic_trace_from_model(model),
                "artifact_hash": _file_hash(path),
                "bytes": path.stat().st_size,
            }
        layout_plan = _build_pdf_layout_plan(model)
        path.write_bytes(_render_pdf_plan_bytes(layout_plan))
        semantic_trace = _extract_verified_pdf_semantic_trace(path, layout_plan)
        return {
            "semantic_trace": semantic_trace,
            "artifact_hash": _file_hash(path),
            "bytes": path.stat().st_size,
        }

    @staticmethod
    def _write_signing_artifact(model: Mapping[str, object], path: Path):
        if (model.get("layout_policy") or {}).get("id") == "AAIEContractMasterLayoutV1":
            path.write_bytes(render_master_signing_html(model))
        else:
            path.write_bytes(_build_signing_artifact_bytes(model))

    def _render_signing_artifact(self, model: Mapping[str, object], path: Path) -> dict[str, object]:
        self._write_signing_artifact(model, path)
        expected_trace = _semantic_trace_from_model(model)
        if (model.get("layout_policy") or {}).get("id") == "AAIEContractMasterLayoutV1":
            raw = path.read_text(encoding="utf-8")
            if not all(value in raw for value in (
                'data-master-anchor="section_11_signature"',
                'data-signature-anchor="signature_aaie"',
                'data-signature-anchor="signature_family"',
                'data-stamp-anchor="aaie-only"',
            )):
                raise TemplateRuleError("保存后签署页成品缺少受保护锚点。")
            semantic_trace = expected_trace
        else:
            semantic_trace = _require_post_save_trace(
                "签署页", _extract_signing_semantic_trace(path), expected_trace
            )
        return {
            "semantic_trace": semantic_trace,
            "artifact_hash": _file_hash(path),
            "bytes": path.stat().st_size,
        }

    def simulate(self, version_id: int, *, source_type: str,
                 student_snapshot: Mapping[str, object], signer_source: str,
                 signer_overrides: Mapping[str, object], selected_services: Iterable[str],
                 output_root: Path, actor_user_id: int, expected_revision: int,
                 lock_token: str) -> dict[str, object]:
        version = self._require_edit(
            version_id, expected_revision, lock_token, actor_user_id
        )
        snapshot = self._snapshot(version_id)
        snapshot_hash = _hash(snapshot)
        if str(version["content_hash"]) != snapshot_hash:
            raise TemplateRuleError("模板内容指纹与快照不一致。")
        selected = list(dict.fromkeys(str(value) for value in selected_services))
        input_payload = {
            "source_type": source_type,
            "student_snapshot": dict(student_snapshot),
            "signer_source": signer_source,
            "signer_overrides": dict(signer_overrides),
            "selected_services": selected,
        }
        input_json = _canonical_json(input_payload)
        input_hash = _hash(input_payload)
        expectation = _derive_render_expectation(snapshot, input_payload)
        signer = expectation["signer"]
        visible_model = expectation["visible_model"]
        authority_content = expectation["authority_content"]
        content_hash = expectation["authority_content_hash"]
        root = self._safe_output_root(Path(output_root))
        artifact_dir = root / f"template-{version_id}-{uuid.uuid4().hex[:10]}"
        artifact_dir.mkdir(parents=True, exist_ok=False)
        docx_path = artifact_dir / "preview.docx"
        pdf_path = artifact_dir / "preview.pdf"
        signing_path = artifact_dir / "signing-preview.html"
        rendered = {
            "docx": self._render_docx_artifact(visible_model, docx_path),
            "pdf": self._render_pdf_artifact(visible_model, pdf_path),
            "signing": self._render_signing_artifact(visible_model, signing_path),
        }
        semantic_json = {
            key: _canonical_json(value["semantic_trace"])
            for key, value in rendered.items()
        }
        if len(set(semantic_json.values())) != 1:
            raise TemplateRuleError("模拟输出语义不一致。")
        semantic_hashes = {
            key: hashlib.sha256(value.encode("utf-8")).hexdigest()
            for key, value in semantic_json.items()
        }
        manifest = {
            "preview.docx": {
                "sha256": rendered["docx"]["artifact_hash"],
                "bytes": rendered["docx"]["bytes"],
            },
            "preview.pdf": {
                "sha256": rendered["pdf"]["artifact_hash"],
                "bytes": rendered["pdf"]["bytes"],
            },
            "signing-preview.html": {
                "sha256": rendered["signing"]["artifact_hash"],
                "bytes": rendered["signing"]["bytes"],
            },
        }
        stamp = _now()
        simulation_values = {
            "template_identity_id": int(version["identity_id"]),
            "template_revision": int(version["revision"]),
            "template_content_hash": str(version["content_hash"]),
            "template_snapshot_hash": snapshot_hash,
            "simulation_input_hash": input_hash,
            "signer_json": _canonical_json(signer),
            "selected_services_json": _canonical_json(selected),
            "authority_content_hash": content_hash,
            "docx_content_hash": semantic_hashes["docx"],
            "pdf_content_hash": semantic_hashes["pdf"],
            "signing_content_hash": semantic_hashes["signing"],
            "docx_semantic_hash": semantic_hashes["docx"],
            "pdf_semantic_hash": semantic_hashes["pdf"],
            "signing_semantic_hash": semantic_hashes["signing"],
            "docx_semantic_trace_json": semantic_json["docx"],
            "pdf_semantic_trace_json": semantic_json["pdf"],
            "signing_semantic_trace_json": semantic_json["signing"],
            "docx_artifact_hash": rendered["docx"]["artifact_hash"],
            "pdf_artifact_hash": rendered["pdf"]["artifact_hash"],
            "signing_artifact_hash": rendered["signing"]["artifact_hash"],
            "artifact_manifest_json": _canonical_json(manifest),
            "artifact_root": str(root),
            "artifact_dir": str(artifact_dir),
            "simulated_at": stamp,
        }
        simulation_result_hash = _hash(_simulation_result_binding(simulation_values))
        with self._immediate_transaction():
            current = self._require_edit(
                version_id, expected_revision, lock_token, actor_user_id
            )
            current_snapshot = self._snapshot(version_id)
            if (
                int(current["identity_id"]) != int(version["identity_id"])
                or int(current["revision"]) != int(version["revision"])
                or str(current["content_hash"]) != str(version["content_hash"])
                or _hash(current_snapshot) != snapshot_hash
            ):
                raise StaleRevision("模板版本已变化，请刷新后重新模拟。")
            self.conn.execute(
                """INSERT INTO contract_template_simulations(
                   version_id,template_identity_id,template_revision,template_content_hash,
                   template_snapshot_hash,source_type,input_json,simulation_input_hash,signer_json,
                   selected_services_json,authority_content_hash,docx_content_hash,pdf_content_hash,
                   signing_content_hash,docx_semantic_hash,pdf_semantic_hash,signing_semantic_hash,
                   docx_semantic_trace_json,pdf_semantic_trace_json,signing_semantic_trace_json,
                   docx_artifact_hash,pdf_artifact_hash,signing_artifact_hash,artifact_manifest_json,
                   simulation_result_hash,artifact_root,artifact_dir,status,created_by,created_at,simulated_at
                   ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,'Passed',?,?,?)""",
                (
                    version_id, simulation_values["template_identity_id"],
                    simulation_values["template_revision"], simulation_values["template_content_hash"],
                    simulation_values["template_snapshot_hash"], source_type, input_json,
                    simulation_values["simulation_input_hash"], simulation_values["signer_json"],
                    simulation_values["selected_services_json"], simulation_values["authority_content_hash"],
                    simulation_values["docx_content_hash"], simulation_values["pdf_content_hash"],
                    simulation_values["signing_content_hash"], simulation_values["docx_semantic_hash"],
                    simulation_values["pdf_semantic_hash"], simulation_values["signing_semantic_hash"],
                    simulation_values["docx_semantic_trace_json"], simulation_values["pdf_semantic_trace_json"],
                    simulation_values["signing_semantic_trace_json"], simulation_values["docx_artifact_hash"],
                    simulation_values["pdf_artifact_hash"], simulation_values["signing_artifact_hash"],
                    simulation_values["artifact_manifest_json"], simulation_result_hash,
                    simulation_values["artifact_root"], simulation_values["artifact_dir"],
                    actor_user_id, stamp, stamp,
                ),
            )
            self.conn.execute(
                "UPDATE contract_template_versions SET status='Simulation' WHERE id=?",
                (version_id,),
            )
            self._audit(
                "template_simulation_passed", actor_user_id,
                current["identity_id"], version_id,
                after={
                    "content_hash": content_hash, "source_type": source_type,
                    "simulation_result_hash": simulation_result_hash,
                },
            )
        return {
            "version_id": version_id, "signer": signer, "authority_content": authority_content,
            "authority_content_hash": content_hash,
            "docx_content_hash": semantic_hashes["docx"],
            "pdf_content_hash": semantic_hashes["pdf"],
            "signing_content_hash": semantic_hashes["signing"],
            "docx_semantic_hash": semantic_hashes["docx"],
            "pdf_semantic_hash": semantic_hashes["pdf"],
            "signing_semantic_hash": semantic_hashes["signing"],
            "docx_artifact_hash": rendered["docx"]["artifact_hash"],
            "pdf_artifact_hash": rendered["pdf"]["artifact_hash"],
            "signing_artifact_hash": rendered["signing"]["artifact_hash"],
            "artifact_manifest": manifest,
            "simulation_input_hash": input_hash,
            "simulation_result_hash": simulation_result_hash,
            "visible_render_model": visible_model,
            "artifact_root": str(root),
            "artifact_dir": str(artifact_dir), "status": "Passed",
        }

    def simulate_batch(self, version_id: int, *, scenarios: Iterable[Mapping[str, object]],
                       output_root: Path, actor_user_id: int, expected_revision: int,
                       lock_token: str) -> list[dict[str, object]]:
        results = []
        for scenario in scenarios:
            results.append(self.simulate(
                version_id,
                source_type=str(scenario.get("source_type") or "system"),
                student_snapshot=dict(scenario.get("student_snapshot") or {}),
                signer_source=str(scenario.get("signer_source") or "mother"),
                signer_overrides=dict(scenario.get("signer_overrides") or {}),
                selected_services=list(scenario.get("selected_services") or []),
                output_root=output_root,
                actor_user_id=actor_user_id,
                expected_revision=expected_revision,
                lock_token=lock_token,
            ))
        if not results:
            raise TemplateRuleError("批量模拟至少需要一组资料。")
        return results

    @staticmethod
    def _simulation_binding_issues(
        simulation: sqlite3.Row, version: sqlite3.Row,
        snapshot: Mapping[str, object],
    ) -> list[dict[str, str]]:
        issues: list[dict[str, str]] = []
        seen: set[str] = set()

        def add(code: str, message: str):
            if code not in seen:
                issues.append({"code": code, "message": message})
                seen.add(code)

        simulation_identity = simulation["template_identity_id"]
        simulation_revision = simulation["template_revision"]
        if type(simulation_identity) is not int or type(simulation_revision) is not int:
            add("simulation_binding_invalid", "模拟归属或修订格式无效")
        else:
            if simulation_identity != int(version["identity_id"]):
                add("simulation_stale_identity", "模拟所属模板已变化")
            if simulation_revision != int(version["revision"]):
                add("simulation_stale_revision", "模板修订已变化，必须重新模拟")
        if str(simulation["template_content_hash"]) != str(version["content_hash"]):
            add("simulation_stale_content_hash", "模板内容指纹已变化，必须重新模拟")
        if str(simulation["template_snapshot_hash"]) != _hash(snapshot):
            add("simulation_stale_snapshot_hash", "模板快照已变化，必须重新模拟")

        expectation = None
        try:
            input_payload = json.loads(str(simulation["input_json"]))
            if not isinstance(input_payload, dict):
                raise ValueError("input must be an object")
            if _hash(input_payload) != str(simulation["simulation_input_hash"]):
                raise ValueError("input hash mismatch")
            if input_payload.get("source_type") != simulation["source_type"]:
                raise ValueError("source type mismatch")
            expectation = _derive_render_expectation(snapshot, input_payload)
        except (KeyError, TypeError, ValueError, json.JSONDecodeError, TemplateRuleError):
            input_payload = None
            add("simulation_input_mismatch", "模拟输入绑定已失效")

        stored_traces: dict[str, dict[str, object]] = {}
        semantic_valid = True
        for renderer in ("docx", "pdf", "signing"):
            try:
                trace = json.loads(str(simulation[f"{renderer}_semantic_trace_json"]))
                if not isinstance(trace, dict):
                    raise ValueError("trace must be an object")
                trace_json = _canonical_json(trace)
                trace_hash = hashlib.sha256(trace_json.encode("utf-8")).hexdigest()
            except (TypeError, ValueError, json.JSONDecodeError):
                semantic_valid = False
                continue
            stored_hash = str(simulation[f"{renderer}_semantic_hash"])
            content_hash = str(simulation[f"{renderer}_content_hash"])
            if trace_hash != stored_hash or content_hash != stored_hash:
                semantic_valid = False
            stored_traces[renderer] = trace
            if expectation is not None and (
                trace_json != expectation["semantic_json"]
                or stored_hash != expectation["semantic_hash"]
            ):
                semantic_valid = False
        if not semantic_valid or len(stored_traces) != 3:
            add("simulation_semantic_mismatch", "DOCX、PDF与签署页可见语义不一致")

        try:
            stored_signer = json.loads(str(simulation["signer_json"]))
            stored_services = json.loads(str(simulation["selected_services_json"]))
        except (TypeError, ValueError, json.JSONDecodeError):
            stored_signer = stored_services = None
        if expectation is not None and (
            stored_signer != expectation["signer"]
            or stored_services != expectation["selected_services"]
            or str(simulation["authority_content_hash"])
            != expectation["authority_content_hash"]
        ):
            add("simulation_result_mismatch", "模拟结果与当前模板不一致")

        try:
            stored_binding = _simulation_result_binding(simulation)
            recomputed_result_hash = _hash(stored_binding)
        except (KeyError, TypeError, ValueError, json.JSONDecodeError):
            stored_binding = None
            recomputed_result_hash = None
        if recomputed_result_hash != str(simulation["simulation_result_hash"]):
            add("simulation_result_mismatch", "模拟结果绑定已失效")

        stored_artifact_hashes = {
            "preview.docx": str(simulation["docx_artifact_hash"]),
            "preview.pdf": str(simulation["pdf_artifact_hash"]),
            "signing-preview.html": str(simulation["signing_artifact_hash"]),
        }
        try:
            manifest = json.loads(str(simulation["artifact_manifest_json"]))
        except (TypeError, ValueError, json.JSONDecodeError):
            manifest = None
        manifest_valid = isinstance(manifest, dict) and set(manifest) == set(stored_artifact_hashes)
        if manifest_valid:
            for filename, expected_hash in stored_artifact_hashes.items():
                entry = manifest.get(filename)
                if (
                    not isinstance(entry, dict)
                    or entry.get("sha256") != expected_hash
                    or not isinstance(entry.get("bytes"), int)
                    or entry["bytes"] < 1
                ):
                    manifest_valid = False
                    break
        if not manifest_valid:
            add("simulation_artifact_manifest_invalid", "模拟成品清单无效")

        artifact_paths: dict[str, Path] = {}
        artifact_root = artifact_dir = None
        try:
            raw_root = Path(str(simulation["artifact_root"])).expanduser()
            raw_dir = Path(str(simulation["artifact_dir"])).expanduser()
            if not raw_root.is_absolute() or not raw_dir.is_absolute():
                raise ValueError("artifact paths must be absolute")
            forbidden_root = Path("/Volumes/Server/Students")
            if (
                raw_root == forbidden_root
                or forbidden_root in raw_root.parents
                or raw_dir == forbidden_root
                or forbidden_root in raw_dir.parents
            ):
                raise ValueError("formal student storage is forbidden")
            if raw_root.is_symlink() or raw_dir.is_symlink():
                raise ValueError("artifact path cannot be a symlink")
            artifact_root = raw_root.resolve()
            artifact_dir = raw_dir.resolve()
            if (
                artifact_root == forbidden_root
                or forbidden_root in artifact_root.parents
                or artifact_dir == forbidden_root
                or forbidden_root in artifact_dir.parents
            ):
                raise ValueError("formal student storage is forbidden")
            expected_name = f"template-{int(version['id'])}-"
            if (
                artifact_dir.parent != artifact_root
                or not artifact_dir.name.startswith(expected_name)
                or re.fullmatch(
                    rf"template-{int(version['id'])}-[0-9a-f]{{10}}", artifact_dir.name
                ) is None
            ):
                raise ValueError("artifact directory ownership mismatch")
        except (KeyError, TypeError, ValueError, OSError, RuntimeError):
            artifact_root = artifact_dir = None
            artifact_paths.clear()
            add("simulation_artifact_path_invalid", "模拟成品路径归属无效")

        if artifact_dir is not None:
            for filename in stored_artifact_hashes:
                try:
                    raw_path = artifact_dir / filename
                    if raw_path.is_symlink():
                        raise ValueError("artifact cannot be a symlink")
                    artifact_path = raw_path.resolve()
                    forbidden_root = Path("/Volumes/Server/Students")
                    if (
                        artifact_path.parent != artifact_dir
                        or artifact_path == forbidden_root
                        or forbidden_root in artifact_path.parents
                    ):
                        raise ValueError("artifact escaped its directory")
                    if not artifact_path.is_file():
                        add("simulation_artifact_missing", "模拟成品文件缺失")
                        continue
                    artifact_paths[filename] = artifact_path
                except (ValueError, OSError, RuntimeError):
                    artifact_root = artifact_dir = None
                    artifact_paths.clear()
                    add("simulation_artifact_path_invalid", "模拟成品文件路径无效")
                    break

        actual_manifest: dict[str, dict[str, object]] = {}
        if len(artifact_paths) == len(stored_artifact_hashes):
            if expectation is not None:
                try:
                    visible_model = expectation["visible_model"]
                    master = (visible_model.get("layout_policy") or {}).get("id") == "AAIEContractMasterLayoutV1"
                    if master:
                        expected_pdf, _anchor_manifest = render_master_pdf_bytes(visible_model)
                        expected_artifacts = {
                            "preview.docx": _embed_docx_font(render_master_docx_bytes(visible_model)),
                            "preview.pdf": expected_pdf,
                            "signing-preview.html": render_master_signing_html(visible_model),
                        }
                        pdf_plan = None
                    else:
                        pdf_plan = _build_pdf_layout_plan(visible_model)
                        expected_artifacts = {
                            "preview.docx": _build_docx_artifact_bytes(visible_model),
                            "preview.pdf": _render_pdf_plan_bytes(pdf_plan),
                            "signing-preview.html": _build_signing_artifact_bytes(
                                visible_model
                            ),
                        }
                    for filename, expected_bytes in expected_artifacts.items():
                        actual_bytes = artifact_paths[filename].read_bytes()
                        if (
                            len(actual_bytes) != len(expected_bytes)
                            or hashlib.sha256(actual_bytes).hexdigest()
                            != hashlib.sha256(expected_bytes).hexdigest()
                        ):
                            raise TemplateRuleError(
                                "保存后的模拟成品字节与当前确定性模型不一致。"
                            )
                    expected_docx = expected_artifacts["preview.docx"]
                    if _docx_layout_fingerprint(
                        artifact_paths["preview.docx"]
                    ) != _docx_layout_fingerprint(expected_docx):
                        raise TemplateRuleError(
                            "保存后DOCX成品布局包与当前确定性模型不一致。"
                        )
                    actual_traces = (
                        {renderer: expectation["semantic_trace"] for renderer in ("docx", "pdf", "signing")}
                        if master else {
                            "docx": _extract_docx_semantic_trace(artifact_paths["preview.docx"]),
                            "pdf": _extract_verified_pdf_semantic_trace(
                                artifact_paths["preview.pdf"],
                                pdf_plan,
                            ),
                            "signing": _extract_signing_semantic_trace(
                                artifact_paths["signing-preview.html"]
                            ),
                        }
                    )
                    if any(
                        _canonical_json(trace) != expectation["semantic_json"]
                        for trace in actual_traces.values()
                    ):
                        add("simulation_semantic_mismatch", "保存后的模拟成品语义不一致")
                except TemplateRuleError:
                    add("simulation_semantic_mismatch", "保存后的模拟成品语义不一致")
                except Exception:
                    add("simulation_artifact_unreadable", "模拟成品无法安全读取")
            try:
                for filename, artifact_path in artifact_paths.items():
                    actual_manifest[filename] = {
                        "sha256": _file_hash(artifact_path),
                        "bytes": artifact_path.stat().st_size,
                    }
            except Exception:
                actual_manifest = {}
                add("simulation_artifact_unreadable", "模拟成品无法安全读取")

        if actual_manifest:
            for filename, actual in actual_manifest.items():
                if (
                    actual["sha256"] != stored_artifact_hashes[filename]
                    or not manifest_valid
                    or actual != manifest.get(filename)
                ):
                    add("simulation_artifact_tampered", "模拟成品文件已被修改")

        if expectation is not None and actual_manifest and artifact_root is not None and artifact_dir is not None:
            expected_binding = {
                "template_identity_id": int(version["identity_id"]),
                "template_revision": int(version["revision"]),
                "template_content_hash": str(version["content_hash"]),
                "template_snapshot_hash": _hash(snapshot),
                "simulation_input_hash": _hash(input_payload),
                "signer": expectation["signer"],
                "selected_services": expectation["selected_services"],
                "authority_content_hash": expectation["authority_content_hash"],
                "content_hashes": {
                    renderer: expectation["semantic_hash"]
                    for renderer in ("docx", "pdf", "signing")
                },
                "semantic": {
                    renderer: {
                        "hash": expectation["semantic_hash"],
                        "trace": expectation["semantic_trace"],
                    }
                    for renderer in ("docx", "pdf", "signing")
                },
                "artifact_hashes": {
                    filename: str(details["sha256"])
                    for filename, details in actual_manifest.items()
                },
                "artifact_manifest": actual_manifest,
                "artifact_root": str(artifact_root),
                "artifact_dir": str(artifact_dir),
                "simulated_at": str(simulation["simulated_at"]),
            }
            if (
                stored_binding != expected_binding
                or str(simulation["simulation_result_hash"]) != _hash(expected_binding)
            ):
                add("simulation_result_mismatch", "模拟结果与当前模板或成品不一致")
        return issues

    def _publish_issues(
        self, version: sqlite3.Row, snapshot: Mapping[str, object]
    ) -> list[dict[str, str]]:
        version_id = int(version["id"])
        issues: list[dict[str, str]] = []
        if is_master_snapshot(snapshot):
            try:
                validate_master_layout_snapshot(snapshot)
            except MasterLayoutError as exc:
                issues.append({"code": "master_layout_invalid", "message": str(exc)})
            issues.extend(contract_display_bindings.mandatory_profile_issues(snapshot))
        commercial = snapshot.get("commercial", {})
        for field, code in {
            "service_fee": "fee_required", "currency": "currency_required",
            "payment_terms": "payment_terms_required", "company_name": "company_name_required",
            "company_abn": "company_abn_required", "company_address": "company_address_required",
            "company_contact": "company_contact_required", "refund_policy": "refund_policy_required",
        }.items():
            if not str(commercial.get(field) or "").strip():
                issues.append({"code": code, "message": f"{field}缺失"})
        issues.extend(authority_v2_commercial_issues(commercial))
        enabled_blocks = [block for block in snapshot["blocks"] if block.get("enabled", True)]
        present_legal = {block.get("legal_kind") for block in enabled_blocks if str(block.get("body") or "").strip()}
        for legal_kind in sorted(REQUIRED_LEGAL_KINDS - present_legal):
            issues.append({"code": f"{legal_kind}_clause_required", "message": f"{legal_kind}条款缺失"})
        for block in enabled_blocks:
            for variable in block.get("required_variables", []):
                if variable not in str(block.get("body") or ""):
                    issues.append({"code": "required_variable_changed", "message": f"必要变量{variable}缺失"})
        simulation = self.conn.execute(
            "SELECT * FROM contract_template_simulations WHERE version_id=? ORDER BY id DESC LIMIT 1", (version_id,)
        ).fetchone()
        if not simulation or simulation["status"] != "Passed":
            issues.append({"code": "simulation_required", "message": "必须先通过模拟"})
        else:
            try:
                issues.extend(self._simulation_binding_issues(simulation, version, snapshot))
            except Exception:
                issues.append({
                    "code": "simulation_binding_invalid",
                    "message": "模拟持久化来源无效，必须重新模拟",
                })
        return issues

    def _require_publish_permission(
        self, actor_user_id: int, actor_is_admin: bool
    ) -> None:
        permission = self.permission_for(
            actor_user_id, actor_is_admin=actor_is_admin
        )
        if not permission["can_publish"]:
            raise PermissionDenied("没有合同模板发布权限。")

    def _record_validation(
        self, version: sqlite3.Row, issues: list[dict[str, str]], actor_user_id: int
    ) -> dict[str, object]:
        version_id = int(version["id"])
        passed = not issues
        self.conn.execute(
            "INSERT INTO contract_template_validations(version_id,passed,issues_json,validated_by,validated_at) VALUES(?,?,?,?,?)",
            (version_id, int(passed), _canonical_json(issues), actor_user_id, _now()),
        )
        self.conn.execute("UPDATE contract_template_versions SET status='Validation' WHERE id=?", (version_id,))
        self._audit("template_validation_passed" if passed else "template_validation_failed",
                    actor_user_id, version["identity_id"], version_id, after=issues)
        return {"passed": passed, "issues": issues}

    def _publish_provenance(
        self, version: sqlite3.Row, snapshot: Mapping[str, object]
    ) -> dict[str, object]:
        simulation = self.conn.execute(
            "SELECT * FROM contract_template_simulations WHERE version_id=? "
            "ORDER BY id DESC LIMIT 1",
            (version["id"],),
        ).fetchone()
        return {
            "identity_id": int(version["identity_id"]),
            "revision": int(version["revision"]),
            "content_hash": str(version["content_hash"]),
            "snapshot_hash": _hash(snapshot),
            "simulation": dict(simulation) if simulation else None,
        }

    def validate_for_publish(self, version_id: int, *, actor_user_id: int,
                             lock_token: str, expected_revision: int | None = None,
                             actor_is_admin: bool = False) -> dict[str, object]:
        self._require_publish_permission(actor_user_id, actor_is_admin)
        version = self._require_edit(
            version_id, expected_revision, lock_token, actor_user_id
        )
        snapshot = self._snapshot(version_id)
        result = self._record_validation(
            version, self._publish_issues(version, snapshot), actor_user_id
        )
        self.conn.commit()
        return result

    def prepare_preview(self, version_id: int, *, actor_user_id: int, lock_token: str,
                        expected_revision: int | None = None,
                        actor_is_admin: bool = False) -> dict[str, object]:
        validation = self.validate_for_publish(
            version_id, actor_user_id=actor_user_id, lock_token=lock_token,
            expected_revision=expected_revision, actor_is_admin=actor_is_admin,
        )
        if not validation["passed"]:
            raise PublishBlocked("模板发布前预览门禁未通过。")
        self._require_publish_permission(actor_user_id, actor_is_admin)
        version = self._require_edit(
            version_id, expected_revision, lock_token, actor_user_id
        )
        self.conn.execute("UPDATE contract_template_versions SET status='Preview' WHERE id=?", (version_id,))
        self._audit("template_preview_confirmed", actor_user_id, version["identity_id"], version_id)
        self.conn.commit()
        return dict(self._version(version_id))

    def publish(self, version_id: int, *, actor_user_id: int, lock_token: str,
                expected_revision: int | None = None,
                actor_is_admin: bool = False) -> dict[str, object]:
        self._require_publish_permission(actor_user_id, actor_is_admin)
        with self._immediate_transaction():
            self._require_publish_permission(actor_user_id, actor_is_admin)
            version = self._require_edit(
                version_id, expected_revision, lock_token, actor_user_id
            )
            snapshot = self._snapshot(version_id)
            initial_provenance = self._publish_provenance(version, snapshot)
            issues = self._publish_issues(version, snapshot)
            validation = self._record_validation(version, issues, actor_user_id)
            if not validation["passed"]:
                raise PublishBlocked("模板发布门禁未通过。")

            self.conn.execute(
                "UPDATE contract_template_versions SET status='Preview' WHERE id=?",
                (version_id,),
            )
            self._audit(
                "template_preview_confirmed", actor_user_id,
                version["identity_id"], version_id,
            )

            current = self._require_edit(
                version_id, expected_revision, lock_token, actor_user_id
            )
            current_snapshot = self._snapshot(version_id)
            final_provenance = self._publish_provenance(current, current_snapshot)
            final_issues = self._publish_issues(current, current_snapshot)
            if final_provenance != initial_provenance or final_issues:
                raise PublishBlocked("模板或模拟来源已变化，必须重新模拟和验证。")

            identity_id = int(current["identity_id"])
            previous_id = self.conn.execute(
                "SELECT current_published_version_id FROM contract_template_identities WHERE id=?",
                (identity_id,),
            ).fetchone()[0]
            stamp = _now()
            self.conn.execute(
                "UPDATE contract_template_versions SET status='Published',immutable=1,"
                "published_by=?,published_at=? WHERE id=?",
                (actor_user_id, stamp, version_id),
            )
            self.conn.execute(
                "UPDATE contract_template_identities SET current_published_version_id=?,updated_at=? WHERE id=?",
                (version_id, stamp, identity_id),
            )
            if previous_id is not None and int(previous_id) != int(version_id):
                self.conn.execute(
                    "UPDATE contract_template_versions SET status='Historical' WHERE id=?",
                    (previous_id,),
                )
            self.conn.execute(
                "DELETE FROM contract_template_edit_locks WHERE identity_id=?",
                (identity_id,),
            )
            self._audit(
                "template_published", actor_user_id, identity_id, version_id,
                before={"published_version_id": previous_id},
                after={"published_version_id": version_id},
            )
        return dict(self._version(version_id))

    def apply_template(self, version_id: int, *, actor_user_id: int,
                       lock_token: str, expected_revision: int,
                       output_root: Path,
                       actor_is_admin: bool = False) -> dict[str, object]:
        """Run the existing publication gates behind one employee-facing action."""
        self._require_publish_permission(actor_user_id, actor_is_admin)
        version = self._require_edit(
            version_id, expected_revision, lock_token, actor_user_id
        )
        snapshot = self._snapshot(version_id)
        selected_services = [
            str(item["service_id"])
            for item in snapshot.get("services", [])
            if item.get("enabled", True) and item.get("default_selected")
        ]
        if not selected_services:
            selected_services = [
                str(item["service_id"])
                for item in snapshot.get("services", [])
                if item.get("enabled", True)
            ]
        self.simulate(
            version_id,
            source_type="system",
            student_snapshot={
                "student_name": "Template Apply Verification",
                "student_id": "TEMPLATE-APPLY-GATE",
                "date_of_birth": "2011-02-03",
                "current_year_level": "Year 7",
                "applying_year_level": "Year 8",
                "preferred_intake_year": "2027",
                "preferred_intake_term": "02",
                "father_passport_first_name": "Template",
                "father_passport_last_name": "Verifier",
                "father_mobile": "0400000000",
                "father_email": "template-apply@example.test",
                "mother_passport_first_name": "Template",
                "mother_passport_last_name": "Verifier",
                "mother_mobile": "0400000001",
                "mother_email": "template-apply@example.test",
            },
            signer_source="father",
            signer_overrides={},
            selected_services=selected_services,
            output_root=output_root,
            actor_user_id=actor_user_id,
            expected_revision=int(version["revision"]),
            lock_token=lock_token,
        )
        return self.publish(
            version_id,
            actor_user_id=actor_user_id,
            lock_token=lock_token,
            expected_revision=int(version["revision"]),
            actor_is_admin=actor_is_admin,
        )

    def compare_versions(self, left_id: int, right_id: int) -> dict[str, object]:
        left_version = self._version(left_id)
        right_version = self._version(right_id)
        if int(left_version["identity_id"]) != int(right_version["identity_id"]):
            raise TemplateRuleError("只能比较同一合同类型的模板版本。")
        left = self._snapshot(left_id)
        right = self._snapshot(right_id)
        left_blocks = {block["key"]: block for block in left["blocks"]}
        right_blocks = {block["key"]: block for block in right["blocks"]}
        changes = []
        for key in sorted(set(left_blocks) | set(right_blocks)):
            if key not in left_blocks:
                changes.append({"key": key, "change": "added"})
            elif key not in right_blocks:
                changes.append({"key": key, "change": "removed"})
            elif left_blocks[key] != right_blocks[key]:
                changes.append({"key": key, "change": "modified", "before": left_blocks[key], "after": right_blocks[key]})
        return {
            "left_id": left_id,
            "right_id": right_id,
            "blocks": changes,
            "services_changed": left["services"] != right["services"],
            "commercial_changed": left["commercial"] != right["commercial"],
        }

    def set_user_permission(self, user_id: int, *, can_edit: bool, can_publish: bool,
                            actor_user_id: int, actor_is_admin: bool):
        with self._immediate_transaction():
            database_admin = self._database_admin_state(actor_user_id)
            effective_admin = actor_is_admin if database_admin is None else database_admin
            if not effective_admin:
                raise PermissionDenied("只有Admin可以修改合同模板权限。")
            before = self.conn.execute(
                "SELECT * FROM contract_template_permissions WHERE user_id=?", (user_id,)
            ).fetchone()
            stamp = _now()
            self.conn.execute(
                """INSERT INTO contract_template_permissions(
                   user_id,can_edit,can_publish,granted_by,granted_at,updated_at
                   ) VALUES(?,?,?,?,?,?) ON CONFLICT(user_id) DO UPDATE SET
                   can_edit=excluded.can_edit,can_publish=excluded.can_publish,
                   granted_by=excluded.granted_by,updated_at=excluded.updated_at""",
                (user_id, int(can_edit), int(can_publish), actor_user_id, stamp, stamp),
            )
            self._audit(
                "template_permission_changed", actor_user_id,
                before=dict(before) if before else None,
                after={
                    "user_id": user_id,
                    "can_edit": can_edit,
                    "can_publish": can_publish,
                },
            )

    def _database_admin_state(self, user_id: int) -> bool | None:
        auth_tables = self.conn.execute(
            "SELECT COUNT(*) FROM sqlite_master WHERE type='table' AND name IN ('users','roles')"
        ).fetchone()[0]
        if int(auth_tables) != 2:
            return None
        try:
            role = self.conn.execute(
                "SELECT r.name FROM users u JOIN roles r ON r.id=u.role_id WHERE u.id=?",
                (user_id,),
            ).fetchone()
            return bool(role and str(role["name"]) == "Admin")
        except sqlite3.DatabaseError:
            return False

    def permission_for(self, user_id: int, *, actor_is_admin: bool) -> dict[str, bool]:
        database_admin = self._database_admin_state(user_id)
        effective_admin = actor_is_admin if database_admin is None else database_admin
        if effective_admin:
            return {"can_view": True, "can_edit": True, "can_publish": True}
        row = self.conn.execute(
            "SELECT can_edit,can_publish FROM contract_template_permissions WHERE user_id=?", (user_id,)
        ).fetchone()
        return {
            "can_view": True,
            "can_edit": bool(row["can_edit"]) if row else False,
            "can_publish": bool(row["can_publish"]) if row else False,
        }

    def audit_events(self) -> list[dict[str, object]]:
        return [dict(row) for row in self.conn.execute(
            "SELECT * FROM contract_template_audit_events ORDER BY id"
        )]
