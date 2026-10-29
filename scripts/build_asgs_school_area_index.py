#!/usr/bin/env python3
import json
import re
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
ASGS_DIR = ROOT / "data" / "asgs"
MB_FILE = ASGS_DIR / "MB_2021_AUST.xlsx"
SAL_FILE = ASGS_DIR / "SAL_2021_AUST.xlsx"
LGA_FILE = ASGS_DIR / "LGA_2021_AUST.xlsx"
OUTPUT_FILE = ASGS_DIR / "asgs_school_area_index_2021.json"

SOURCE_BASE_URL = (
    "https://www.abs.gov.au/statistics/standards/australian-statistical-geography-standard-asgs/"
    "edition-3-july-2021-june-2026/access-and-downloads/allocation-files/"
)


def norm(value):
    return re.sub(r"\s+", " ", str(value or "").strip()).lower()


def rows(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    iterator = ws.iter_rows(values_only=True)
    headers = [str(value or "").strip() for value in next(iterator)]
    for row in iterator:
        yield dict(zip(headers, row))


def as_float(value):
    try:
        return float(value or 0)
    except (TypeError, ValueError):
        return 0.0


def max_area_item(area_by_key):
    if not area_by_key:
        return ""
    return max(area_by_key.items(), key=lambda item: item[1])[0]


def main():
    missing = [str(path) for path in (MB_FILE, SAL_FILE, LGA_FILE) if not path.exists()]
    if missing:
        raise SystemExit("Missing ASGS files: " + ", ".join(missing))

    mb_info = {}
    for row in rows(MB_FILE):
        mb_code = str(row.get("MB_CODE_2021") or "").strip()
        if not mb_code:
            continue
        mb_info[mb_code] = {
            "sa1_code": str(row.get("SA1_CODE_2021") or "").strip(),
            "sa2_code": str(row.get("SA2_CODE_2021") or "").strip(),
            "sa2_name": str(row.get("SA2_NAME_2021") or "").strip(),
            "sa3_code": str(row.get("SA3_CODE_2021") or "").strip(),
            "sa3_name": str(row.get("SA3_NAME_2021") or "").strip(),
            "sa4_code": str(row.get("SA4_CODE_2021") or "").strip(),
            "sa4_name": str(row.get("SA4_NAME_2021") or "").strip(),
        }

    lga_by_mb = {}
    for row in rows(LGA_FILE):
        mb_code = str(row.get("MB_CODE_2021") or "").strip()
        if not mb_code:
            continue
        lga_by_mb[mb_code] = {
            "lga_code": str(row.get("LGA_CODE_2021") or "").strip(),
            "lga_name": str(row.get("LGA_NAME_2021") or "").strip(),
        }

    grouped = {}
    area_buckets = defaultdict(lambda: {
        "sa1": defaultdict(float),
        "sa2": defaultdict(float),
        "sa3": defaultdict(float),
        "sa4": defaultdict(float),
        "lga": defaultdict(float),
    })
    for row in rows(SAL_FILE):
        mb_code = str(row.get("MB_CODE_2021") or "").strip()
        info = mb_info.get(mb_code)
        if not info:
            continue
        state_code = str(row.get("STATE_CODE_2021") or "").strip()
        sal_name = str(row.get("SAL_NAME_2021") or "").strip()
        if not state_code or not sal_name:
            continue
        key = f"{state_code}|{norm(sal_name)}"
        area = as_float(row.get("AREA_ALBERS_SQKM"))
        grouped.setdefault(key, {
            "suburb": sal_name,
            "state_code": state_code,
            "state": str(row.get("STATE_NAME_2021") or "").strip(),
            "sal_code": str(row.get("SAL_CODE_2021") or "").strip(),
            "source": "ABS ASGS Edition 3 allocation files",
            "confidence": 92,
        })
        buckets = area_buckets[key]
        buckets["sa1"][info["sa1_code"]] += area
        buckets["sa2"][(info["sa2_code"], info["sa2_name"])] += area
        buckets["sa3"][(info["sa3_code"], info["sa3_name"])] += area
        buckets["sa4"][(info["sa4_code"], info["sa4_name"])] += area
        lga = lga_by_mb.get(mb_code)
        if lga:
            buckets["lga"][(lga["lga_code"], lga["lga_name"])] += area

    areas = {}
    for key, base in grouped.items():
        buckets = area_buckets[key]
        sa2_code, sa2_name = max_area_item(buckets["sa2"]) or ("", "")
        sa3_code, sa3_name = max_area_item(buckets["sa3"]) or ("", "")
        sa4_code, sa4_name = max_area_item(buckets["sa4"]) or ("", "")
        lga_code, lga_name = max_area_item(buckets["lga"]) or ("", "")
        areas[key] = {
            **base,
            "sa1_code": max_area_item(buckets["sa1"]),
            "sa2_code": sa2_code,
            "sa2_name": sa2_name,
            "sa3_code": sa3_code,
            "sa3_name": sa3_name,
            "sa4_code": sa4_code,
            "sa4_name": sa4_name,
            "lga_code": lga_code,
            "lga_name": lga_name,
        }

    payload = {
        "source": "ABS ASGS Edition 3 allocation files",
        "source_urls": [
            SOURCE_BASE_URL + "MB_2021_AUST.xlsx",
            SOURCE_BASE_URL + "SAL_2021_AUST.xlsx",
            SOURCE_BASE_URL + "LGA_2021_AUST.xlsx",
        ],
        "areas": dict(sorted(areas.items())),
    }
    OUTPUT_FILE.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"Wrote {OUTPUT_FILE} with {len(areas)} suburb/locality records.")


if __name__ == "__main__":
    main()
