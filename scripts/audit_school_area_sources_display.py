#!/usr/bin/env python3
import csv
import json
import math
import re
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

import app  # noqa: E402
from scripts.backfill_unresolved_school_area_by_boundary import (  # noqa: E402
    boundary_match,
    dbf_fields_and_rows,
    load_asgs_by_sa2,
    load_sa2_boundaries,
    point_in_ring,
    shp_records,
)


REPORT_ROOT = ROOT / "reports" / "school_area_profiles"
SA2_XLSX = ROOT / "data" / "asgs" / "SA2_2021_AUST.xlsx"
LGA_BOUNDARY_ZIP = ROOT / "data" / "asgs" / "LGA_2021_AUST_GDA2020_SHP.zip"
SAMPLE_SCHOOL_NAMES = {
    "Abbotsford Primary School",
    "Auburn High School",
    "Melbourne High School",
    "Australian International High School",
}
SAMPLE_SCHOOL_IDS = {1041}


def stamp_for_path():
    return datetime.now().strftime("%Y%m%dT%H%M%S")


def write_csv(path, rows):
    path.parent.mkdir(parents=True, exist_ok=True)
    if not rows:
        path.write_text("", encoding="utf-8")
        return
    fields = []
    for row in rows:
        for key in row.keys():
            if key not in fields:
                fields.append(key)
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)


def load_sa2_reference():
    wb = load_workbook(SA2_XLSX, read_only=True, data_only=True)
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    header = [str(value or "") for value in next(rows)]
    index = {name: pos for pos, name in enumerate(header)}
    by_code = {}
    for row in rows:
        code = str(row[index["SA2_CODE_2021"]] or "").strip()
        if not code:
            continue
        by_code[code] = {
            "sa2_name": str(row[index["SA2_NAME_2021"]] or "").strip(),
            "sa3_code": str(row[index["SA3_CODE_2021"]] or "").strip(),
            "sa3_name": str(row[index["SA3_NAME_2021"]] or "").strip(),
            "sa4_code": str(row[index["SA4_CODE_2021"]] or "").strip(),
            "sa4_name": str(row[index["SA4_NAME_2021"]] or "").strip(),
            "state_code": str(row[index["STATE_CODE_2021"]] or "").strip(),
            "state_name": str(row[index["STATE_NAME_2021"]] or "").strip(),
            "asgs_loci_uri": str(row[index["ASGS_LOCI_URI_2021"]] or "").strip(),
        }
    return by_code


def load_lga_boundaries():
    if not LGA_BOUNDARY_ZIP.exists():
        return []
    from zipfile import ZipFile

    with ZipFile(LGA_BOUNDARY_ZIP) as archive:
        dbf_name = next(name for name in archive.namelist() if name.lower().endswith(".dbf"))
        shp_name = next(name for name in archive.namelist() if name.lower().endswith(".shp"))
        rows = dbf_fields_and_rows(archive.read(dbf_name))
        shapes = list(shp_records(archive.read(shp_name)))
    features = []
    for row, (_record_number, bbox, rings) in zip(rows, shapes):
        if not bbox or not rings:
            continue
        features.append(
            {
                "lga_code": str(row.get("LGA_CODE21") or "").strip(),
                "lga_name": str(row.get("LGA_NAME21") or "").strip(),
                "state_code": str(row.get("STE_CODE21") or "").strip(),
                "bbox": bbox,
                "rings": rings,
            }
        )
    return features


def lga_boundary_match(features, latitude, longitude, state):
    try:
        lat = float(latitude)
        lon = float(longitude)
    except (TypeError, ValueError):
        return None
    if state == "VIC" and lat > 0:
        lat = -lat
    state_code = app.ABS_STATE_CODE_BY_STATE.get(str(state or "").upper(), "")
    for feature in features:
        if state_code and feature["state_code"] != state_code:
            continue
        xmin, ymin, xmax, ymax = feature["bbox"]
        if not (xmin <= lon <= xmax and ymin <= lat <= ymax):
            continue
        contains = False
        for ring_xmin, ring_ymin, ring_xmax, ring_ymax, ring in feature["rings"]:
            if ring_xmin <= lon <= ring_xmax and ring_ymin <= lat <= ring_ymax and point_in_ring(lon, lat, ring):
                contains = not contains
        if contains:
            return {
                "lga_code": feature["lga_code"],
                "lga_name": feature["lga_name"],
                "source": "ABS ASGS 2021 LGA boundary point-in-polygon",
            }
    return None


def latest_profiles():
    return {
        row["school_id"]: row
        for row in app.db.all(
            """SELECT sap.*
               FROM school_area_profiles sap
               JOIN (
                 SELECT school_id,MAX(id) id
                 FROM school_area_profiles
                 GROUP BY school_id
               ) x ON x.id=sap.id"""
        )
    }


def latest_metric_rows(profile_ids):
    if not profile_ids:
        return {}
    placeholders = ",".join("?" for _ in profile_ids)
    rows = app.db.all(
        f"""SELECT *
            FROM school_area_profile_metrics
            WHERE profile_id IN ({placeholders})
            ORDER BY profile_id,id""",
        tuple(profile_ids),
    )
    by_profile = {}
    for row in rows:
        by_profile.setdefault(row["profile_id"], {})[row["metric_key"]] = row
    return by_profile


def linked_suburb_rows(profile_ids):
    if not profile_ids:
        return {}
    placeholders = ",".join("?" for _ in profile_ids)
    rows = app.db.all(
        f"""SELECT *
            FROM school_suburb_data
            WHERE json_valid(other_indicators_json)=1
              AND json_extract(other_indicators_json,'$.school_area_profile_id') IN ({placeholders})""",
        tuple(profile_ids),
    )
    by_profile = {}
    for row in rows:
        try:
            meta = json.loads(row["other_indicators_json"] or "{}")
        except Exception:
            meta = {}
        by_profile[int(meta.get("school_area_profile_id") or 0)] = row
    return by_profile


def json_object(value):
    try:
        parsed = json.loads(value or "{}")
        return parsed if isinstance(parsed, dict) else {}
    except Exception:
        return {}


def numbers_equal(left, right):
    if left is None and right is None:
        return True
    if left in ("", None) and right in ("", None):
        return True
    try:
        return math.isclose(float(left), float(right), rel_tol=0, abs_tol=0.000001)
    except (TypeError, ValueError):
        return str(left or "") == str(right or "")


def metric_display(row):
    if not row or row["metric_value"] is None:
        return "未提供"
    try:
        number = float(row["metric_value"])
        text = f"{number:.1f}" if not number.is_integer() else str(int(number))
    except (TypeError, ValueError):
        text = str(row["metric_value"])
    unit = row["metric_unit"] or ""
    return f"{text}{' ' + unit if unit and unit != 'count' else ''}"


def expected_abs_g01_url(sa2_code):
    return app.build_abs_census_sa2_data_api_url("G01", sa2_code, "2021")


def audit():
    report_dir = REPORT_ROOT / f"source_display_audit_{stamp_for_path()}"
    report_dir.mkdir(parents=True, exist_ok=True)
    profiles = latest_profiles()
    profile_ids = [row["id"] for row in profiles.values()]
    metrics_by_profile = latest_metric_rows(profile_ids)
    linked_by_profile = linked_suburb_rows(profile_ids)
    sa2_ref = load_sa2_reference()
    asgs_by_sa2 = load_asgs_by_sa2()
    boundaries = load_sa2_boundaries()
    lga_boundaries = load_lga_boundaries()
    schools = app.db.all(
        """SELECT id,school_name,address,suburb,state,postcode,latitude,longitude,suburb_data_status,zone_data_json
           FROM schools
           WHERE deleted_at IS NULL
           ORDER BY id"""
    )

    summary = {
        "schools": len(schools),
        "latest_profiles": len(profiles),
        "approved_profiles": 0,
        "approved_school_status": 0,
        "abs_source_urls": 0,
        "expected_abs_g01_urls": 0,
        "data_year_2021": 0,
        "blank_error_message": 0,
        "sa2_boundary_matches": 0,
        "sa2_boundary_mismatches": 0,
        "sa2_reference_matches": 0,
        "sa2_reference_mismatches": 0,
        "asgs_lga_matches_or_blank": 0,
        "asgs_lga_mismatches": 0,
        "lga_boundary_matches_or_blank": 0,
        "lga_boundary_mismatches": 0,
        "lga_boundary_not_checked": 0,
        "metric_raw_matches": 0,
        "metric_raw_mismatches": 0,
        "linked_zone_rows": 0,
        "linked_zone_matches": 0,
        "display_samples_checked": 0,
        "display_sample_failures": 0,
    }
    issues = []
    lga_allocation_advisories = []
    sample_display = []

    for school in schools:
        profile = profiles.get(school["id"])
        if not profile:
            issues.append({"school_id": school["id"], "school_name": school["school_name"], "issue": "missing_latest_profile"})
            continue
        metrics = metrics_by_profile.get(profile["id"], {})
        linked = linked_by_profile.get(profile["id"])
        if profile["review_status"] == "Approved":
            summary["approved_profiles"] += 1
        if school["suburb_data_status"] == "Approved":
            summary["approved_school_status"] += 1
        if str(profile["source_url"] or "").startswith("https://data.api.abs.gov.au/"):
            summary["abs_source_urls"] += 1
        if profile["source_url"] == expected_abs_g01_url(profile["sa2_code"]):
            summary["expected_abs_g01_urls"] += 1
        if profile["data_year"] == "2021":
            summary["data_year_2021"] += 1
        if not profile["error_message"]:
            summary["blank_error_message"] += 1

        boundary = boundary_match(boundaries, school["latitude"], school["longitude"], str(school["state"] or "").upper())
        if boundary and str(boundary["sa2_code"]) == str(profile["sa2_code"]):
            summary["sa2_boundary_matches"] += 1
        else:
            summary["sa2_boundary_mismatches"] += 1
            issues.append(
                {
                    "school_id": school["id"],
                    "school_name": school["school_name"],
                    "issue": "sa2_boundary_mismatch",
                    "profile_sa2": profile["sa2_code"],
                    "boundary_sa2": boundary.get("sa2_code") if boundary else "",
                }
            )

        ref = sa2_ref.get(str(profile["sa2_code"] or ""))
        ref_ok = bool(ref) and all(
            str(profile[key] or "") == str(ref[ref_key] or "")
            for key, ref_key in (
                ("sa2_name", "sa2_name"),
                ("sa3_code", "sa3_code"),
                ("sa3_name", "sa3_name"),
                ("sa4_code", "sa4_code"),
                ("sa4_name", "sa4_name"),
            )
        )
        if ref_ok:
            summary["sa2_reference_matches"] += 1
        else:
            summary["sa2_reference_mismatches"] += 1
            issues.append({"school_id": school["id"], "school_name": school["school_name"], "issue": "sa2_reference_mismatch", "sa2_code": profile["sa2_code"]})

        asgs = asgs_by_sa2.get(str(profile["sa2_code"] or "")) or {}
        expected_lga = str(asgs.get("lga_name") or "")
        if not expected_lga or str(profile["lga_name"] or "") == expected_lga:
            summary["asgs_lga_matches_or_blank"] += 1
        else:
            summary["asgs_lga_mismatches"] += 1
            lga_allocation_advisories.append(
                {
                    "school_id": school["id"],
                    "school_name": school["school_name"],
                    "advisory": "sa2_allocation_lga_differs_from_coordinate_lga",
                    "profile_lga": profile["lga_name"],
                    "asgs_lga": expected_lga,
                }
            )

        lga_boundary = lga_boundary_match(lga_boundaries, school["latitude"], school["longitude"], str(school["state"] or "").upper())
        if not lga_boundaries or not lga_boundary:
            summary["lga_boundary_not_checked"] += 1
        elif not profile["lga_name"] or str(profile["lga_name"] or "") == str(lga_boundary["lga_name"] or ""):
            summary["lga_boundary_matches_or_blank"] += 1
        else:
            summary["lga_boundary_mismatches"] += 1
            issues.append(
                {
                    "school_id": school["id"],
                    "school_name": school["school_name"],
                    "issue": "lga_boundary_mismatch",
                    "profile_lga_code": profile["lga_code"],
                    "profile_lga_name": profile["lga_name"],
                    "boundary_lga_code": lga_boundary["lga_code"],
                    "boundary_lga_name": lga_boundary["lga_name"],
                }
            )

        raw = json_object(profile["raw_json"])
        flows = raw.get("flows") or raw.get("payloads") or {}
        normalized = app.normalize_school_area_profile_metrics(app.normalize_abs_census_sa2_profile(flows, profile["sa2_code"])) if flows else {}
        metric_mismatches = []
        for key, _label, _unit in app.SCHOOL_AREA_PROFILE_METRICS:
            row = metrics.get(key)
            stored = row["metric_value"] if row else None
            expected = normalized.get(key)
            if key in app.SCHOOL_AREA_PROFILE_NUMERIC_METRICS and not numbers_equal(stored, expected):
                metric_mismatches.append(f"{key}: stored={stored} expected={expected}")
        if metric_mismatches:
            summary["metric_raw_mismatches"] += 1
            issues.append(
                {
                    "school_id": school["id"],
                    "school_name": school["school_name"],
                    "issue": "metric_raw_mismatch",
                    "details": " | ".join(metric_mismatches[:8]),
                }
            )
        else:
            summary["metric_raw_matches"] += 1

        if linked:
            summary["linked_zone_rows"] += 1
            zone_json = json_object(school["zone_data_json"])
            linked_ok = (
                linked["geography_code"] == profile["sa2_code"]
                and linked["source_url"] == profile["source_url"]
                and linked["reviewed_status"] == "Approved"
                and str(zone_json.get("school_area_profile_id") or "") == str(profile["id"])
            )
            if linked_ok:
                summary["linked_zone_matches"] += 1
            else:
                issues.append(
                    {
                        "school_id": school["id"],
                        "school_name": school["school_name"],
                        "issue": "linked_zone_display_source_mismatch",
                        "profile_id": profile["id"],
                        "linked_id": linked["id"],
                    }
                )
        else:
            issues.append({"school_id": school["id"], "school_name": school["school_name"], "issue": "missing_linked_school_suburb_data"})

        if school["school_name"] in SAMPLE_SCHOOL_NAMES or school["id"] in SAMPLE_SCHOOL_IDS:
            summary["display_samples_checked"] += 1
            zone_json = json_object(school["zone_data_json"])
            display = {
                "school_id": school["id"],
                "school_name": school["school_name"],
                "profile_status_badge": profile["review_status"],
                "school_zone_status_badge": school["suburb_data_status"],
                "display_suburb": profile["suburb"] or school["suburb"],
                "display_postcode": profile["postcode"] or school["postcode"],
                "display_sa2_code": profile["sa2_code"],
                "display_sa2_name": profile["sa2_name"],
                "display_lga_name": profile["lga_name"],
                "display_source": profile["source"],
                "display_source_url": profile["source_url"],
                "zone_json_profile_id": zone_json.get("school_area_profile_id"),
                "zone_json_suburb_data_id": zone_json.get("suburb_data_id"),
                "ui_metric_total_population": metric_display(metrics.get("total_population")),
                "ui_metric_median_age": metric_display(metrics.get("median_age")),
                "ui_metric_median_household_income": metric_display(metrics.get("median_household_income")),
                "ui_metric_median_rent": metric_display(metrics.get("median_rent")),
            }
            display["display_check"] = "PASS" if (
                display["profile_status_badge"] == "Approved"
                and display["school_zone_status_badge"] == "Approved"
                and str(display["zone_json_profile_id"]) == str(profile["id"])
                and display["display_source_url"] == expected_abs_g01_url(profile["sa2_code"])
            ) else "FAIL"
            if display["display_check"] != "PASS":
                summary["display_sample_failures"] += 1
            sample_display.append(display)

    summary["issue_count"] = len(issues)
    summary["lga_allocation_advisory_count"] = len(lga_allocation_advisories)
    summary["generated_at"] = datetime.now().isoformat(timespec="seconds")
    write_csv(report_dir / "issues.csv", issues)
    write_csv(report_dir / "lga_allocation_advisories.csv", lga_allocation_advisories)
    write_csv(report_dir / "sample_display_checks.csv", sample_display)
    (report_dir / "summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    readme = [
        "# School Area Source And Display Audit",
        "",
        "## Sources Cross Checked",
        "",
        "- CRM latest `school_area_profiles`",
        "- CRM linked `school_suburb_data` and `schools.zone_data_json`",
        "- ABS Census 2021 Data API raw payload saved in each profile",
        "- ABS ASGS Edition 3 SA2 boundary shapefile point-in-polygon",
        "- ABS ASGS Edition 3 LGA boundary shapefile point-in-polygon",
        "- ABS ASGS Edition 3 SA2 official Excel table",
        "- Local ASGS allocation/index built from ABS MB/SAL/LGA allocation files",
        "",
        "## Summary",
        "",
        "```json",
        json.dumps(summary, ensure_ascii=False, indent=2),
        "```",
        "",
        "## Files",
        "",
        "- `summary.json`",
        "- `issues.csv`",
        "- `lga_allocation_advisories.csv`",
        "- `sample_display_checks.csv`",
    ]
    (report_dir / "README.md").write_text("\n".join(readme), encoding="utf-8")
    print(json.dumps({"report_dir": str(report_dir), "summary": summary}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    audit()
