#!/usr/bin/env python3
import argparse
import json
import re
import sys
import urllib.parse
import urllib.request
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from app import csa_crime_area_key, csa_crime_number, refresh_school_area_crime_profile, upsert_csa_crime_area_stat
from database import Database

DATAVIC_PACKAGE_SHOW = "https://discover.data.vic.gov.au/api/3/action/package_show"
DATA_DIR = ROOT / "school_data_sources" / "vic_crime"
DB_PATH = ROOT / "database" / "aaie_crm.db"
USER_AGENT = "AAIE-CRM-DataVic-CSA-Importer/1.0"

CSA_PACKAGES = {
    "recorded_offences": "data-tables-recorded-offences",
    "criminal_incidents": "criminal-incident",
    "family_incidents": "data-tables-family-incidents",
    "geographic_classifications": "crime-statistics-agency-data-tables-geographical-classifications",
}


def fetch_json(url):
    request = urllib.request.Request(url, headers={"User-Agent": USER_AGENT, "Accept": "application/json"})
    with urllib.request.urlopen(request, timeout=45) as response:
        return json.loads(response.read().decode("utf-8"))


def package_metadata(package_id):
    url = DATAVIC_PACKAGE_SHOW + "?" + urllib.parse.urlencode({"id": package_id})
    payload = fetch_json(url)
    if not payload.get("success"):
        raise RuntimeError(f"DataVic package lookup failed: {package_id}")
    return payload["result"]


def period_sort_key(text):
    months = {
        "jan": 1, "january": 1, "feb": 2, "february": 2, "mar": 3, "march": 3,
        "apr": 4, "april": 4, "may": 5, "jun": 6, "june": 6, "jul": 7, "july": 7,
        "aug": 8, "august": 8, "sep": 9, "sept": 9, "september": 9,
        "oct": 10, "october": 10, "nov": 11, "november": 11, "dec": 12, "december": 12,
    }
    lowered = str(text or "").lower()
    year_match = re.search(r"(20\d{2}|19\d{2})", lowered)
    year = int(year_match.group(1)) if year_match else 0
    month = 0
    for name, number in months.items():
        if re.search(rf"\b{name}\b", lowered):
            month = number
            break
    return year, month, lowered


def xlsx_resources(package):
    rows = []
    for resource in package.get("resources", []):
        url = resource.get("url") or ""
        fmt = (resource.get("format") or Path(url).suffix).lower()
        if not url or "xls" not in fmt:
            continue
        name = resource.get("name") or Path(url).name
        rows.append({
            "name": name,
            "url": url,
            "format": fmt,
            "period_key": period_sort_key(name),
        })
    return rows


def latest_resources(package):
    resources = xlsx_resources(package)
    if not resources:
        return []
    latest_period = max((item["period_key"][0], item["period_key"][1]) for item in resources)
    return [item for item in resources if (item["period_key"][0], item["period_key"][1]) == latest_period]


def safe_filename(name, url):
    suffix = Path(urllib.parse.urlparse(url).path).suffix or ".xlsx"
    stem = re.sub(r"[^A-Za-z0-9._-]+", "_", str(name or "csa_resource")).strip("._")
    return f"{stem}{suffix}" if not stem.lower().endswith(suffix.lower()) else stem


def download_resource(resource, target_dir):
    target_dir.mkdir(parents=True, exist_ok=True)
    target = target_dir / safe_filename(resource["name"], resource["url"])
    if target.exists() and target.stat().st_size > 0:
        return target
    request = urllib.request.Request(resource["url"], headers={"User-Agent": USER_AGENT})
    with urllib.request.urlopen(request, timeout=120) as response:
        target.write_bytes(response.read())
    return target


def clean_header(value):
    return re.sub(r"\s+", " ", str(value or "").strip())


def workbook_dict_rows(path, sheet_names=None):
    selected_sheets = set(sheet_names or [])
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        for sheet in workbook.worksheets:
            if selected_sheets and sheet.title not in selected_sheets:
                continue
            header_index = None
            headers = []
            for index, row in enumerate(sheet.iter_rows(values_only=True)):
                cleaned = [clean_header(cell) for cell in row]
                if header_index is None:
                    joined = " ".join(cleaned).lower()
                    if index < 40 and ("postcode" in joined or "local government" in joined or "lga" in joined or "suburb" in joined) and len([x for x in cleaned if x]) >= 2:
                        header_index = index
                        headers = cleaned
                    elif index >= 40:
                        break
                    continue
                item = {headers[i]: row[i] for i in range(min(len(headers), len(row))) if headers[i]}
                if len([value for value in item.values() if value not in {None, ""}]) >= 2:
                    yield sheet.title, item
    finally:
        workbook.close()


def workbook_dict_rows_legacy(path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        for sheet in workbook.worksheets:
            values = list(sheet.iter_rows(values_only=True))
            header_index = None
            headers = []
            for index, row in enumerate(values[:40]):
                cleaned = [clean_header(cell) for cell in row]
                joined = " ".join(cleaned).lower()
                if ("postcode" in joined or "local government" in joined or "lga" in joined or "suburb" in joined) and len([x for x in cleaned if x]) >= 2:
                    header_index = index
                    headers = cleaned
                    break
            if header_index is None:
                continue
            for row in values[header_index + 1:]:
                item = {headers[i]: row[i] for i in range(min(len(headers), len(row))) if headers[i]}
                if len([value for value in item.values() if value not in {None, ""}]) >= 2:
                    yield sheet.title, item
    finally:
        workbook.close()


def infer_data_level(resource_name, row):
    lowered_name = str(resource_name or "").lower()
    lowered_keys = " ".join(str(key).lower() for key in row)
    if "lga" in lowered_name or "local government" in lowered_keys:
        return "lga"
    if "postcode" in lowered_keys:
        return "postcode"
    if "suburb" in lowered_keys or "locality" in lowered_keys:
        return "suburb"
    return ""


def row_value(row, *names):
    lookup = {re.sub(r"[^a-z0-9]+", " ", str(key or "").lower()).strip(): value for key, value in row.items()}
    for name in names:
        key = re.sub(r"[^a-z0-9]+", " ", str(name or "").lower()).strip()
        if key in lookup and lookup[key] not in {None, ""}:
            return lookup[key]
    return ""


def year_ending(row):
    ending = str(row_value(row, "Year ending") or "").strip()
    year = str(row_value(row, "Year") or "").strip()
    if ending and re.search(r"(20\d{2}|19\d{2})", ending):
        return ending
    if ending and year:
        return f"{ending} {year}"
    return ending or year


def empty_stat(state, data_level, area_name, source_dataset, source_url, row):
    suburb = str(row_value(row, "Suburb/Town Name", "Suburb", "Locality") or "").strip() if data_level == "suburb" else ""
    postcode = re.sub(r"[^0-9]", "", str(row_value(row, "Postcode") or "")) if data_level == "postcode" else ""
    lga_name = str(row_value(row, "Local Government Area", "LGA") or "").strip()
    return {
        "state": state,
        "data_level": data_level,
        "area_name": area_name,
        "normalized_area_key": csa_crime_area_key(data_level, suburb=suburb, postcode=postcode, lga_name=lga_name, area_name=area_name),
        "suburb": suburb,
        "postcode": postcode,
        "lga_name": lga_name,
        "data_year_ending": year_ending(row),
        "source_dataset": source_dataset,
        "source": "Crime Statistics Agency / DataVic",
        "source_url": source_url,
        "raw_json": {},
    }


def offence_bucket(row):
    division = str(row_value(row, "Offence Division") or "").lower()
    subgroup = str(row_value(row, "Offence Subgroup", "Offence Subdivision") or "").lower()
    fields = []
    if "crimes against the person" in division:
        fields.append("crimes_against_person")
    elif "property and deception" in division:
        fields.append("property_and_deception_offences")
    elif "drug offences" in division:
        fields.append("drug_offences")
    elif "public order and security" in division:
        fields.append("public_order_and_security_offences")
    elif "justice procedures" in division:
        fields.append("justice_procedures_offences")
    elif "other offences" in division:
        fields.append("other_offences")
    if "residential" in subgroup and "burglary" in subgroup:
        fields.append("residential_burglary_count")
    if "steal from a motor vehicle" in subgroup or "theft from a motor vehicle" in subgroup:
        fields.append("theft_from_motor_vehicle_count")
    if "steal motor vehicle" in subgroup or "motor vehicle theft" in subgroup:
        fields.append("motor_vehicle_theft_count")
    if "robbery" in subgroup:
        fields.append("robbery_count")
    if "assault" in subgroup:
        fields.append("assault_count")
    if "property damage" in subgroup or "criminal damage" in subgroup:
        fields.append("property_damage_count")
    return fields


def add_stat(stats, stat, field_name, value, add=True):
    if not stat.get("area_name") or not stat.get("normalized_area_key") or not stat.get("data_year_ending"):
        return
    key = (stat["data_level"], stat["normalized_area_key"], stat["data_year_ending"], stat["source_dataset"])
    current = stats.setdefault(key, stat)
    if current.get("lga_name") in {None, ""} and stat.get("lga_name"):
        current["lga_name"] = stat["lga_name"]
    if current.get("suburb") in {None, ""} and stat.get("suburb"):
        current["suburb"] = stat["suburb"]
    if current.get("postcode") in {None, ""} and stat.get("postcode"):
        current["postcode"] = stat["postcode"]
    if field_name:
        number = csa_crime_number(value, integer=field_name.endswith("_total") or field_name.endswith("_count") or field_name in {"total_recorded_offences", "criminal_incidents_total", "family_incidents_total"})
        if number is not None:
            current[field_name] = (current.get(field_name) or 0) + number if add else number


def aggregate_recorded_offences(resource, path):
    stats = {}
    source_text = f"{resource.get('name', '')} {Path(path).name}".lower()
    sheet_names = {"Table 01"} if "lga" in source_text else {"Table 01", "Table 03"}
    for sheet, row in workbook_dict_rows(path, sheet_names=sheet_names):
        sheet_name = str(sheet)
        count = row_value(row, "Offence Count", "Total recorded offences")
        if sheet_name == "Table 01":
            lga = str(row_value(row, "Local Government Area") or "").strip()
            stat = empty_stat("VIC", "lga", lga, "recorded_offences", resource["url"], row)
            add_stat(stats, stat, "total_recorded_offences", count, add=False)
            add_stat(stats, stat, "offence_rate_per_100000", row_value(row, "Rate per 100,000 population"), add=False)
        elif sheet_name == "Table 03":
            postcode = re.sub(r"[^0-9]", "", str(row_value(row, "Postcode") or ""))
            suburb = str(row_value(row, "Suburb/Town Name") or "").strip()
            for level, area in (("postcode", postcode), ("suburb", suburb)):
                stat = empty_stat("VIC", level, area, "recorded_offences", resource["url"], row)
                add_stat(stats, stat, "total_recorded_offences", count)
                for bucket in offence_bucket(row):
                    add_stat(stats, stat, bucket, count)
    return list(stats.values())


def aggregate_lga_totals(resource, path, source_dataset, total_field, rate_field):
    stats = {}
    for sheet, row in workbook_dict_rows(path):
        if str(sheet) != "Table 01":
            continue
        lga = str(row_value(row, "Local Government Area") or "").strip()
        if not lga:
            continue
        stat = empty_stat("VIC", "lga", lga, source_dataset, resource["url"], row)
        add_stat(stats, stat, total_field, row_value(row, "Family Incidents", "Incidents Recorded", "Criminal Incidents"), add=False)
        add_stat(stats, stat, rate_field, row_value(row, "Rate per 100,000 population"), add=False)
    return list(stats.values())


def import_resource(db, source_dataset, resource, path, apply_changes):
    if source_dataset == "recorded_offences":
        stats = aggregate_recorded_offences(resource, path)
    elif source_dataset == "family_incidents":
        stats = aggregate_lga_totals(resource, path, source_dataset, "family_incidents_total", "family_incidents_rate_per_100000")
    elif source_dataset == "criminal_incidents":
        stats = aggregate_lga_totals(resource, path, source_dataset, "criminal_incidents_total", "criminal_incidents_rate_per_100000")
    else:
        stats = []
    imported = len(stats)
    skipped = 0
    if apply_changes:
        for stat in stats:
            upsert_csa_crime_area_stat(db, stat)
    return imported, skipped


def refresh_profiles(db, apply_changes):
    if not apply_changes:
        return 0
    refreshed = 0
    for school in db.all("SELECT id FROM schools WHERE state='VIC'"):
        if refresh_school_area_crime_profile(db, school["id"]):
            refreshed += 1
    return refreshed


def run(apply_changes=False, latest_only=True):
    db = Database(DB_PATH)
    db.initialize()
    report = {"apply": apply_changes, "resources": [], "profiles_refreshed": 0}
    for source_dataset, package_id in CSA_PACKAGES.items():
        package = package_metadata(package_id)
        resources = latest_resources(package) if latest_only else xlsx_resources(package)
        for resource in resources:
            target = download_resource(resource, DATA_DIR / source_dataset)
            imported, skipped = (0, 0)
            if source_dataset != "geographic_classifications":
                imported, skipped = import_resource(db, source_dataset, resource, target, apply_changes)
            report["resources"].append({
                "dataset": source_dataset,
                "package": package_id,
                "name": resource["name"],
                "url": resource["url"],
                "file": str(target.relative_to(ROOT)),
                "rows_imported": imported,
                "rows_skipped": skipped,
            })
    report["profiles_refreshed"] = refresh_profiles(db, apply_changes)
    return report


def main():
    parser = argparse.ArgumentParser(description="Import VIC CSA/DataVic crime statistics into AAIE CRM.")
    parser.add_argument("--apply", action="store_true", help="Write imported rows and refreshed school profiles to the CRM database.")
    parser.add_argument("--all-periods", action="store_true", help="Download and parse all DataVic resources instead of only the latest period.")
    args = parser.parse_args()
    report = run(apply_changes=args.apply, latest_only=not args.all_periods)
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
