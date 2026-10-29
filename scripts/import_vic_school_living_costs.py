#!/usr/bin/env python3
"""Import VIC school surrounding-area living cost data into AAIE CRM.

Default mode is dry-run. Use --apply to write to database/aaie_crm.db.
"""

import argparse
import csv
import json
import re
import sqlite3
import sys
import urllib.error
import urllib.parse
import urllib.request
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

APP_DIR = Path(__file__).resolve().parents[1]
if str(APP_DIR) not in sys.path:
    sys.path.insert(0, str(APP_DIR))

from app import refresh_school_living_cost_summary  # noqa: E402

DB_PATH = APP_DIR / "database" / "aaie_crm.db"
SOURCE_DIR = APP_DIR / "school_data_sources" / "vic_living_costs"

DATAVIC_API = "https://discover.data.vic.gov.au/api/3/action/package_show?id="
SCHOOL_LOCATIONS_PACKAGE = "school-locations-2025"
RENTAL_REPORT_PACKAGE = "rental-report-quarterly-data-tables"
PROPERTY_HOUSE_PACKAGE = "victorian-property-sales-report-median-house-by-suburb"
PROPERTY_UNIT_PACKAGE = "victorian-property-sales-report-median-unit-by-suburb"
ABS_ASGS_URL = "https://www.abs.gov.au/statistics/standards/australian-statistical-geography-standard-asgs-edition-3"
DOWNLOAD_HEADERS = {
    "User-Agent": "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 Chrome/126 Safari/537.36",
    "Referer": "https://www.land.vic.gov.au/valuations/resources-and-reports/property-sales-statistics",
}

OFFICIAL_SOURCES = {
    "school_locations": f"https://discover.data.vic.gov.au/dataset/{SCHOOL_LOCATIONS_PACKAGE}",
    "rental_report": f"https://discover.data.vic.gov.au/dataset/{RENTAL_REPORT_PACKAGE}",
    "property_house": f"https://discover.data.vic.gov.au/dataset/{PROPERTY_HOUSE_PACKAGE}",
    "property_unit": f"https://discover.data.vic.gov.au/dataset/{PROPERTY_UNIT_PACKAGE}",
    "abs_asgs": ABS_ASGS_URL,
}

RENTAL_SOURCE_NAME = "VIC Rental Report - Quarterly: Data Tables"
PROPERTY_SOURCE_NAME = "Victorian Property Sales Report / Valuer-General Victoria"


class DbAdapter:
    def __init__(self, conn):
        self.conn = conn

    def one(self, sql, params=()):
        return self.conn.execute(sql, params).fetchone()

    def all(self, sql, params=()):
        return self.conn.execute(sql, params).fetchall()

    def execute(self, sql, params=()):
        cur = self.conn.execute(sql, params)
        return cur.lastrowid


SCHOOL_ALIASES = {
    "name": ("school_name", "school name", "schoolname", "name"),
    "sector": ("sector", "education_sector", "school_sector"),
    "type": ("school_type", "school type", "type"),
    "address": ("address", "school_address", "street_address", "address_line_1"),
    "suburb": ("suburb", "town", "locality", "address_town", "postal_town"),
    "postcode": ("postcode", "post_code", "postal_code", "address_postcode", "postal_postcode"),
    "latitude": ("latitude", "lat", "y"),
    "longitude": ("longitude", "lon", "lng", "x"),
    "website": ("website", "website_url", "school_website"),
}


def stamp():
    return datetime.now(timezone.utc).astimezone().isoformat(timespec="seconds")


def dry_run_label(apply):
    return "apply" if apply else "dry_run"


def normalize_key(value):
    return re.sub(r"[^a-z0-9]+", "_", str(value or "").strip().lower()).strip("_")


def clean_number(value):
    text = re.sub(r"[^0-9.\-]", "", str(value or ""))
    if text.count("-") > 1 or ("-" in text and not text.startswith("-")):
        text = text.replace("-", "")
    if not text:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def ckan_package(package_name):
    url = DATAVIC_API + urllib.parse.quote(package_name)
    with urllib.request.urlopen(url, timeout=30) as response:
        payload = json.loads(response.read().decode("utf-8"))
    if not payload.get("success"):
        raise RuntimeError(f"DataVic package lookup failed: {package_name}")
    return payload["result"]


def download_resource(url, target_dir):
    target_dir.mkdir(parents=True, exist_ok=True)
    name = urllib.parse.unquote(Path(urllib.parse.urlparse(url).path).name) or "downloaded-resource"
    target = target_dir / name
    headers = dict(DOWNLOAD_HEADERS)
    try:
        request = urllib.request.Request(url, headers=headers)
        with urllib.request.urlopen(request, timeout=60) as response:
            target.write_bytes(response.read())
    except urllib.error.HTTPError as exc:
        if exc.code != 403:
            raise
        headers["Range"] = "bytes=0-"
        request = urllib.request.Request(url, headers=headers)
        with urllib.request.urlopen(request, timeout=60) as response:
            target.write_bytes(response.read())
    if not target.suffix:
        head = target.read_bytes()[:8]
        if head.startswith(b"PK"):
            target = target.rename(target.with_suffix(".xlsx"))
        elif head.startswith(b"\xd0\xcf\x11\xe0"):
            target = target.rename(target.with_suffix(".xls"))
    return target


def resource_sort_key(item):
    name = str(item.get("name") or item.get("url") or "").lower()
    year_match = re.findall(r"20\d{2}", name)
    year = int(year_match[-1]) if year_match else 0
    quarter_order = {
        "mar": 1, "march": 1, "q1": 1,
        "jun": 2, "june": 2, "q2": 2,
        "sep": 3, "sept": 3, "september": 3, "q3": 3,
        "dec": 4, "december": 4, "q4": 4,
    }
    quarter = 0
    for token, value in quarter_order.items():
        if re.search(rf"\b{token}\b", name):
            quarter = max(quarter, value)
    return year, quarter, name


def latest_resource_url(package_name, format_prefixes=("CSV", "XLS", "XLSX")):
    package = ckan_package(package_name)
    resources = package.get("resources") or []
    accepted = [
        item for item in resources
        if str(item.get("format") or "").upper() in format_prefixes and item.get("url")
    ]
    if not accepted:
        raise RuntimeError(f"No supported resource found for {package_name}")
    return sorted(accepted, key=resource_sort_key)[-1]["url"]


def row_value(row, aliases):
    normalized = {normalize_key(key): value for key, value in row.items()}
    for alias in aliases:
        key = normalize_key(alias)
        if key in normalized and normalized[key] not in (None, ""):
            return str(normalized[key]).strip()
    return ""


def import_school_locations(conn, path, apply=False):
    total = 0
    changed = 0
    skipped_new = 0
    now = stamp()
    with Path(path).open("r", encoding="utf-8-sig", newline="") as handle:
        for raw in csv.DictReader(handle):
            name = row_value(raw, SCHOOL_ALIASES["name"])
            if not name:
                continue
            total += 1
            sector_raw = row_value(raw, SCHOOL_ALIASES["sector"])
            sector_key = "private" if re.search(r"non|independent|catholic|private", sector_raw, re.I) else "public"
            school_type = row_value(raw, SCHOOL_ALIASES["type"]) or ("Private" if sector_key == "private" else "Public")
            values = {
                "school_name": name,
                "state": "VIC",
                "sector_key": sector_key,
                "school_type": school_type,
                "address": row_value(raw, SCHOOL_ALIASES["address"]),
                "suburb": row_value(raw, SCHOOL_ALIASES["suburb"]),
                "postcode": re.sub(r"[^0-9]", "", row_value(raw, SCHOOL_ALIASES["postcode"])),
                "latitude": clean_number(row_value(raw, SCHOOL_ALIASES["latitude"])),
                "longitude": clean_number(row_value(raw, SCHOOL_ALIASES["longitude"])),
                "website_url": row_value(raw, SCHOOL_ALIASES["website"]),
                "source_url": OFFICIAL_SOURCES["school_locations"],
                "data_quality_status": "Official",
                "updated_at": now,
                "created_at": now,
            }
            if not apply:
                continue
            existing = conn.execute(
                "SELECT id FROM schools WHERE state='VIC' AND school_name=?",
                (name,),
            ).fetchone()
            if existing:
                conn.execute(
                    """UPDATE schools SET sector_key=?, school_type=?, address=COALESCE(NULLIF(?,''),address),
                       suburb=COALESCE(NULLIF(?,''),suburb), postcode=COALESCE(NULLIF(?,''),postcode),
                       latitude=COALESCE(?,latitude), longitude=COALESCE(?,longitude),
                       website_url=COALESCE(NULLIF(?,''),website_url), source_url=?, data_quality_status='Official',
                       updated_at=?, school_data_updated_at=?
                       WHERE id=?""",
                    (
                        values["sector_key"], values["school_type"], values["address"], values["suburb"],
                        values["postcode"], values["latitude"], values["longitude"], values["website_url"],
                        values["source_url"], now, now, existing["id"],
                    ),
                )
            else:
                # This script enriches the CRM school library; it must not turn
                # the DataVic whole-state directory into new CRM school records.
                skipped_new += 1
                continue
            if values["suburb"] and values["latitude"] is not None and values["longitude"] is not None:
                conn.execute(
                    """INSERT INTO geo_areas(state,area_type,area_name,suburb,postcode,latitude,longitude)
                       VALUES('VIC','suburb',?,?,?,?,?)
                       ON CONFLICT(state,area_type,area_name,postcode) DO UPDATE SET
                         suburb=excluded.suburb, latitude=excluded.latitude, longitude=excluded.longitude""",
                    (values["suburb"], values["suburb"], values["postcode"], values["latitude"], values["longitude"]),
                )
            changed += 1
    return {"source": "School Locations 2025", "total": total, "changed": changed, "skipped_new_schools": skipped_new}


def worksheet_rows(path):
    workbook = load_workbook(path, data_only=True, read_only=True)
    for sheet in workbook.worksheets:
        rows = list(sheet.iter_rows(values_only=True))
        for index, row in enumerate(rows):
            headers = [normalize_key(cell) for cell in row]
            if "suburb" in headers or "locality" in headers or "town" in headers:
                for data_row in rows[index + 1:]:
                    yield dict(zip(headers, data_row))
                return


def property_period_from_path(path):
    name = Path(path).stem.lower()
    match = re.search(r"q([1-4])[-_ ]*(20\d{2})", name)
    if not match:
        return Path(path).stem
    quarter, year = match.groups()
    labels = {"1": "Jan-Mar", "2": "Apr-Jun", "3": "Jul-Sep", "4": "Oct-Dec"}
    return f"{labels[quarter]} {year}"


def printable_tokens(path):
    data = Path(path).read_bytes()
    text = "".join(chr(byte) if 32 <= byte <= 126 or byte in (10, 12, 13) else "\n" for byte in data)
    tokens = []
    for raw in re.split(r"[\r\n\f]+", text):
        value = raw.strip().strip("\x00").strip()
        if len(value) >= 2:
            tokens.append(value)
    return tokens


def legacy_property_xls_rows(path, period=""):
    tokens = printable_tokens(path)
    if "Sheet1" in tokens:
        sheet_index = tokens.index("Sheet1")
    else:
        sheet_index = 0
    start = next(
        (index + 1 for index, token in enumerate(tokens[sheet_index:], sheet_index) if token in {"Jan-Mar", "Apr-Jun", "Jul-Sep", "Oct-Dec"}),
        sheet_index,
    )
    end = next((index for index, token in enumerate(tokens[start:], start) if token.startswith("^ means")), len(tokens))
    chunk = tokens[start:end]
    numeric = re.compile(r"^-?\d+(?:\.\d+)?$")
    skip = {"Jan-Mar", "Apr-Jun", "Jul-Sep", "Oct-Dec", "Sheet1", "Locality", "No. of Sales", "Change (%)"}

    def is_locality(token):
        return (
            token not in skip
            and token.upper() == token
            and re.search(r"[A-Z]", token)
            and not numeric.match(token)
            and not token.startswith(("[", "_"))
            and "means" not in token
        )

    rows = []
    index = 0
    row_period = period or property_period_from_path(path)
    while index < len(chunk):
        token = chunk[index].strip()
        if not is_locality(token):
            index += 1
            continue
        locality = token
        values = []
        cursor = index + 1
        while cursor < len(chunk) and not is_locality(chunk[cursor].strip()):
            cleaned = chunk[cursor].replace(",", "").strip()
            if numeric.match(cleaned):
                values.append(cleaned)
            cursor += 1
        price_values = [market for market in (clean_number(value) for value in values) if market is not None and market >= 50000]
        if price_values:
            sales_count = None
            for value in values[1:]:
                number = clean_number(value)
                if number is not None and float(number).is_integer() and 0 < number < 10000:
                    sales_count = int(number)
                    break
            rows.append({
                "suburb": locality.title(),
                "median_sale_price": price_values[0],
                "period": row_period,
                "sales_count": sales_count,
            })
        index = cursor
    return rows


def rental_table12_rows(path):
    workbook = load_workbook(path, data_only=True, read_only=True)
    if "Table 12" not in workbook.sheetnames:
        return []
    sheet = workbook["Table 12"]
    rows = list(sheet.iter_rows(values_only=True))
    if len(rows) < 4:
        return []
    property_headers = list(rows[1])
    metric_headers = list(rows[2])
    groups = []
    current_property = ""
    for index, value in enumerate(property_headers):
        if value not in (None, ""):
            current_property = str(value).strip()
        metric = str(metric_headers[index] or "").strip().lower() if index < len(metric_headers) else ""
        if current_property and metric == "median":
            groups.append((current_property, index))
    parsed = []
    for data_row in rows[3:]:
        suburb = str(data_row[1] or "").strip() if len(data_row) > 1 else ""
        if not suburb:
            continue
        for property_type, median_index in groups:
            value = data_row[median_index] if median_index < len(data_row) else None
            rent = clean_number(value)
            if rent is None:
                continue
            bedroom_match = re.search(r"\b(\d+)\s*bed", property_type, flags=re.I)
            parsed.append({
                "suburb": suburb,
                "property_type": property_type,
                "bedrooms": bedroom_match.group(1) if bedroom_match else "",
                "median_weekly_rent": rent,
            })
    return parsed


def import_rental_market(conn, path, apply=False, quarter=""):
    total = 0
    changed = 0
    now = stamp()
    parsed_rows = rental_table12_rows(path)
    if not parsed_rows:
        parsed_rows = list(worksheet_rows(path))
    prepared = []
    for row in parsed_rows:
        suburb = row.get("suburb") or row.get("locality") or row.get("town")
        rent = clean_number(row.get("median_weekly_rent") or row.get("median_rent") or row.get("median"))
        if not suburb or rent is None:
            continue
        total += 1
        property_type = str(row.get("property_type") or row.get("dwelling_type") or "House").strip()
        bedrooms = str(row.get("bedrooms") or row.get("bedroom") or "").strip()
        row_quarter = quarter or str(row.get("quarter") or row.get("period") or Path(path).stem).strip()
        prepared.append((str(suburb).strip(), "", property_type, bedrooms, rent, row_quarter, RENTAL_SOURCE_NAME, now))
    if apply and prepared:
        for row_quarter in sorted({item[5] for item in prepared}):
            conn.execute(
                "DELETE FROM rental_market WHERE state='VIC' AND source=? AND quarter=?",
                (RENTAL_SOURCE_NAME, row_quarter),
            )
        for values in prepared:
            conn.execute(
                """INSERT INTO rental_market(state,suburb,postcode,property_type,bedrooms,median_weekly_rent,quarter,source,last_updated)
                   VALUES('VIC',?,?,?,?,?,?,?,?)""",
                values,
            )
            changed += 1
    return {"source": "Rental Report quarterly median rents", "total": total, "changed": changed}


def import_property_market(conn, path, property_type, apply=False, period=""):
    total = 0
    changed = 0
    now = stamp()
    prepared = []
    try:
        rows = list(worksheet_rows(path))
    except Exception:
        rows = legacy_property_xls_rows(path, period=period)
    if not rows and str(path).lower().endswith(".xls"):
        rows = legacy_property_xls_rows(path, period=period)
    for row in rows:
        suburb = row.get("suburb") or row.get("locality") or row.get("town")
        price = clean_number(row.get("median_sale_price") or row.get("median_price") or row.get("median"))
        if not suburb or price is None:
            continue
        total += 1
        row_period = period or str(row.get("period") or row.get("quarter") or Path(path).stem).strip()
        sales_count = clean_number(row.get("sales_count") or row.get("number_of_sales") or row.get("sales"))
        prepared.append((str(suburb).strip(), "", property_type, price, row_period, int(sales_count) if sales_count is not None else None, PROPERTY_SOURCE_NAME, now))
    if apply and prepared:
        for row_period in sorted({item[4] for item in prepared}):
            conn.execute(
                "DELETE FROM property_market WHERE state='VIC' AND source=? AND property_type=? AND period=?",
                (PROPERTY_SOURCE_NAME, property_type, row_period),
            )
        for values in prepared:
            conn.execute(
                """INSERT INTO property_market(state,suburb,postcode,property_type,median_sale_price,period,sales_count,source,last_updated)
                   VALUES('VIC',?,?,?,?,?,?,?,?)""",
                values,
            )
            changed += 1
    return {"source": f"VPSR median {property_type.lower()} by suburb", "total": total, "changed": changed}


def refresh_summaries(conn, apply=False, limit=0):
    rows = conn.execute(
        """SELECT id FROM schools
           WHERE UPPER(COALESCE(state,''))='VIC'
             AND latitude IS NOT NULL AND longitude IS NOT NULL
           ORDER BY id"""
    ).fetchall()
    if limit:
        rows = rows[:limit]
    if not apply:
        return {"source": "school_radius_market_summary", "total": len(rows), "changed": 0}
    adapter = DbAdapter(conn)
    for row in rows:
        refresh_school_living_cost_summary(adapter, row["id"])
    return {"source": "school_radius_market_summary", "total": len(rows), "changed": len(rows)}


def maybe_download_one(downloaded, errors, key, package_name, target, formats):
    try:
        downloaded[key] = download_resource(latest_resource_url(package_name, formats), target)
    except Exception as exc:
        errors[key] = str(exc)


def maybe_download(args):
    downloaded = {}
    download_errors = {}
    if not args.download:
        return downloaded, download_errors
    target = SOURCE_DIR / datetime.now().strftime("%Y%m%d_%H%M%S")
    if not args.school_locations:
        maybe_download_one(downloaded, download_errors, "school_locations", SCHOOL_LOCATIONS_PACKAGE, target, ("CSV",))
    if not args.rental_file:
        maybe_download_one(downloaded, download_errors, "rental_file", RENTAL_REPORT_PACKAGE, target, ("XLS", "XLSX"))
    if not args.property_unit_file:
        maybe_download_one(downloaded, download_errors, "property_unit_file", PROPERTY_UNIT_PACKAGE, target, ("XLS", "XLSX"))
    if not args.property_house_file:
        maybe_download_one(downloaded, download_errors, "property_house_file", PROPERTY_HOUSE_PACKAGE, target, ("XLS", "XLSX"))
    return downloaded, download_errors


def optional_path(value):
    return Path(value) if value else None


def main():
    parser = argparse.ArgumentParser(description="Import official VIC school surrounding living cost data.")
    parser.add_argument("--db", default=str(DB_PATH))
    parser.add_argument("--download", action="store_true", help="Download latest official resources from DataVic / land.vic.gov.au.")
    parser.add_argument("--school-locations")
    parser.add_argument("--rental-file")
    parser.add_argument("--property-house-file")
    parser.add_argument("--property-unit-file")
    parser.add_argument("--quarter", default="")
    parser.add_argument("--period", default="")
    parser.add_argument("--limit", type=int, default=0)
    parser.add_argument("--apply", action="store_true", help="Write imported data and summaries. Without this flag the script is dry_run.")
    args = parser.parse_args()
    downloaded, download_errors = maybe_download(args)
    school_locations = optional_path(args.school_locations or downloaded.get("school_locations"))
    rental_file = optional_path(args.rental_file or downloaded.get("rental_file"))
    house_file = optional_path(args.property_house_file or downloaded.get("property_house_file"))
    unit_file = optional_path(args.property_unit_file or downloaded.get("property_unit_file"))
    conn = sqlite3.connect(args.db)
    conn.row_factory = sqlite3.Row
    try:
        results = {"mode": dry_run_label(args.apply), "official_sources": OFFICIAL_SOURCES, "download_errors": download_errors, "import_errors": {}, "steps": []}
        import_jobs = [
            ("school_locations", school_locations, lambda: import_school_locations(conn, school_locations, apply=args.apply)),
            ("rental_file", rental_file, lambda: import_rental_market(conn, rental_file, apply=args.apply, quarter=args.quarter)),
            ("property_house_file", house_file, lambda: import_property_market(conn, house_file, "House", apply=args.apply, period=args.period)),
            ("property_unit_file", unit_file, lambda: import_property_market(conn, unit_file, "Unit", apply=args.apply, period=args.period)),
        ]
        for key, path, callback in import_jobs:
            if not path or not path.exists():
                continue
            try:
                results["steps"].append(callback())
            except Exception as exc:
                results["import_errors"][key] = str(exc)
        try:
            results["steps"].append(refresh_summaries(conn, apply=args.apply, limit=args.limit))
        except Exception as exc:
            results["import_errors"]["summary_refresh"] = str(exc)
        if args.apply:
            conn.commit()
        else:
            conn.rollback()
        print(json.dumps(results, ensure_ascii=False, indent=2))
    finally:
        conn.close()


if __name__ == "__main__":
    main()
