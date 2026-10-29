#!/usr/bin/env python3
import argparse
import csv
import io
import json
import re
import sys
import urllib.parse
import urllib.request
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from app import csa_crime_area_key, csa_crime_header_key, csa_crime_number, refresh_school_area_crime_profile, upsert_csa_crime_area_stat
from database import Database
from scripts import import_vic_crime_stats

DB_PATH = ROOT / "database" / "aaie_crm.db"
DATA_DIR = ROOT / "school_data_sources" / "au_crime"
USER_AGENT = "AAIE-CRM-OfficialCrimeDataImporter/1.0"

NSW_BOCSAR_SOURCE_PAGE = "https://bocsar.nsw.gov.au/content/dcj/bocsar/bocsar-home/statistics-dashboards/open-datasets/criminal-offences-data.html"
NSW_BOCSAR_LGA_URL = "https://bocsarblob.blob.core.windows.net/bocsar-open-data/RCI_offencebymonth.xlsm"
QLD_PACKAGE_SHOW = "https://www.data.qld.gov.au/api/3/action/package_show"
QLD_LGA_NUMBER_PACKAGE = "lga_reported_offences_number"
QLD_LGA_RATE_PACKAGE = "lga_reported_offences_rates"
SA_PACKAGE_SHOW = "https://data.sa.gov.au/data/api/3/action/package_show"
SA_CRIME_PACKAGE = "crime-statistics"
ACT_CRIME_URL = "https://www.data.act.gov.au/download/2egm-dieb/application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
ACT_CRIME_SOURCE_PAGE = "https://data.gov.au/data/dataset/2egm-dieb"
TAS_CRIME_SUPPLEMENT_URL = "https://www.police.tas.gov.au/uploads/DPFEM-Crime-Statistics-Supplement-2024-25.pdf"
WA_TIMESERIES_URL = "https://www.wa.gov.au/media/48429/download?inline"

MONTH_ORDER = {
    "jan": 1,
    "feb": 2,
    "mar": 3,
    "apr": 4,
    "may": 5,
    "jun": 6,
    "jul": 7,
    "aug": 8,
    "sep": 9,
    "sept": 9,
    "oct": 10,
    "nov": 11,
    "dec": 12,
}
QLD_BROAD_COLUMNS = {
    "Offences Against the Person": "crimes_against_person",
    "Offences Against Property": "property_and_deception_offences",
    "Drug Offences": "drug_offences",
    "Prostitution Offences": "public_order_and_security_offences",
    "Liquor (excl. Drunkenness)": "public_order_and_security_offences",
    "Gaming Racing & Betting Offences": "public_order_and_security_offences",
    "Trespassing and Vagrancy": "public_order_and_security_offences",
    "Weapons Act Offences": "public_order_and_security_offences",
    "Good Order Offences": "public_order_and_security_offences",
    "Traffic and Related Offences": "public_order_and_security_offences",
    "Breach Domestic Violence Protection Order": "justice_procedures_offences",
    "Miscellaneous Offences": "other_offences",
    "Other Offences": "other_offences",
}
QLD_DETAIL_COLUMNS = {
    "Assault": "assault_count",
    "Robbery": "robbery_count",
    "Unlawful Entry With Intent - Dwelling": "residential_burglary_count",
    "Vehicles (steal from/enter with intent)": "theft_from_motor_vehicle_count",
    "Unlawful Use of Motor Vehicle": "motor_vehicle_theft_count",
    "Other Property Damage": "property_damage_count",
}
ACT_EXCLUDED_CATEGORIES = {
    "tins speeding",
    "tins mobile use",
    "tins seatbelts",
    "tins other",
    "cins",
}


def fetch_json(url):
    request = urllib.request.Request(url, headers={"User-Agent": USER_AGENT, "Accept": "application/json"})
    with urllib.request.urlopen(request, timeout=60) as response:
        return json.loads(response.read().decode("utf-8"))


def package_metadata(base_url, package_id):
    url = base_url + "?" + urllib.parse.urlencode({"id": package_id})
    payload = fetch_json(url)
    if not payload.get("success"):
        raise RuntimeError(f"Official data package lookup failed: {package_id}")
    return payload["result"]


def downloadable_resource(package, allowed_formats=("csv", "xlsx", "xls", "xlsm")):
    resources = package.get("resources") or []
    candidates = []
    for resource in resources:
        url = resource.get("url") or ""
        fmt = (resource.get("format") or Path(urllib.parse.urlparse(url).path).suffix).lower()
        if url and any(fmt.strip(".") == item or item in fmt for item in allowed_formats):
            candidates.append(resource)
    if not candidates:
        return None
    return max(candidates, key=lambda item: period_sort_key(item.get("name") or item.get("url")))


def sa_crime_resource(package):
    resources = []
    for resource in package.get("resources") or []:
        name = resource.get("name") or ""
        url = resource.get("url") or ""
        fmt = (resource.get("format") or Path(urllib.parse.urlparse(url).path).suffix).lower()
        lowered = name.lower()
        if url and "csv" in fmt and "crime statistics" in lowered and "family" not in lowered and "domestic" not in lowered:
            resources.append(resource)
    if not resources:
        return downloadable_resource(package, ("csv",))
    return max(resources, key=lambda item: period_sort_key(item.get("name") or item.get("url")))


def period_sort_key(text):
    lowered = str(text or "").lower()
    year_match = re.search(r"(20\d{2}|19\d{2})", lowered)
    year = int(year_match.group(1)) if year_match else 0
    month = 0
    for name, number in MONTH_ORDER.items():
        if re.search(rf"\b{name}\w*\b", lowered):
            month = number
            break
    return year, month, lowered


def safe_filename(name, url, default_suffix=""):
    suffix = Path(urllib.parse.urlparse(url).path).suffix or default_suffix or ".dat"
    if default_suffix and suffix.lower() in {".sheet", ".vnd.openxmlformats-officedocument.spreadsheetml.sheet"}:
        suffix = default_suffix
    stem = re.sub(r"[^A-Za-z0-9._-]+", "_", str(name or Path(urllib.parse.urlparse(url).path).name or "official_crime_data")).strip("._")
    if not stem.lower().endswith(suffix.lower()):
        stem += suffix
    return stem


def download(url, target_dir, name="", default_suffix=""):
    target_dir.mkdir(parents=True, exist_ok=True)
    target = target_dir / safe_filename(name, url, default_suffix)
    if target.exists() and target.stat().st_size > 0:
        return target
    request = urllib.request.Request(url, headers={"User-Agent": USER_AGENT})
    with urllib.request.urlopen(request, timeout=180) as response:
        target.write_bytes(response.read())
    return target


def clean_text(value):
    return re.sub(r"\s+", " ", str(value or "").strip())


def number(value):
    if isinstance(value, (list, tuple)):
        return 0
    return csa_crime_number(value, integer=True) or 0


def month_label(year, month):
    names = {
        1: "January", 2: "February", 3: "March", 4: "April", 5: "May", 6: "June",
        7: "July", 8: "August", 9: "September", 10: "October", 11: "November", 12: "December",
    }
    return f"{names.get(month, '')} {year}".strip()


def rolling_months(months, count=12):
    ordered = sorted(set(months))
    if not ordered:
        return []
    return ordered[-count:]


def add_bucket(stat, offence_name, value):
    text = clean_text(offence_name).lower()
    amount = number(value)
    if amount == 0:
        return
    stat["total_recorded_offences"] = (stat.get("total_recorded_offences") or 0) + amount
    if any(part in text for part in ("homicide", "assault", "sexual", "robbery", "abduction", "intimidation", "stalking", "person")):
        stat["crimes_against_person"] = (stat.get("crimes_against_person") or 0) + amount
    elif any(part in text for part in ("break and enter", "burglary", "theft", "steal", "fraud", "property", "arson", "damage")):
        stat["property_and_deception_offences"] = (stat.get("property_and_deception_offences") or 0) + amount
    elif "drug" in text:
        stat["drug_offences"] = (stat.get("drug_offences") or 0) + amount
    elif any(part in text for part in ("public order", "weapons", "good order", "traffic")):
        stat["public_order_and_security_offences"] = (stat.get("public_order_and_security_offences") or 0) + amount
    elif any(part in text for part in ("justice", "breach", "court", "custody")):
        stat["justice_procedures_offences"] = (stat.get("justice_procedures_offences") or 0) + amount
    else:
        stat["other_offences"] = (stat.get("other_offences") or 0) + amount
    if "dwelling" in text or "residential" in text or "burglary" in text:
        if "break" in text or "burglary" in text:
            stat["residential_burglary_count"] = (stat.get("residential_burglary_count") or 0) + amount
    if "steal from motor vehicle" in text or "theft from motor vehicle" in text:
        stat["theft_from_motor_vehicle_count"] = (stat.get("theft_from_motor_vehicle_count") or 0) + amount
    if "motor vehicle theft" in text or "steal motor vehicle" in text or "theft of motor vehicle" in text:
        stat["motor_vehicle_theft_count"] = (stat.get("motor_vehicle_theft_count") or 0) + amount
    if "robbery" in text:
        stat["robbery_count"] = (stat.get("robbery_count") or 0) + amount
    if "assault" in text:
        stat["assault_count"] = (stat.get("assault_count") or 0) + amount
    if "property damage" in text or "malicious damage" in text or "damage" in text:
        stat["property_damage_count"] = (stat.get("property_damage_count") or 0) + amount


def add_stat_value(stat, field_name, value):
    amount = number(value)
    if amount:
        stat[field_name] = (stat.get(field_name) or 0) + amount
    return amount


def add_qld_row(stat, row):
    broad_found = any(key in row and number(row.get(key)) for key in QLD_BROAD_COLUMNS)
    if broad_found:
        for key, field_name in QLD_BROAD_COLUMNS.items():
            amount = add_stat_value(stat, field_name, row.get(key))
            stat["total_recorded_offences"] = (stat.get("total_recorded_offences") or 0) + amount
        for key, field_name in QLD_DETAIL_COLUMNS.items():
            add_stat_value(stat, field_name, row.get(key))
        return
    for key, value in row.items():
        if key and key not in {"LGA Name", "Month Year"}:
            add_bucket(stat, key, value)


def base_stat(state, data_level, area_name, data_year_ending, source_dataset, source, source_url, notes=""):
    area_name = clean_text(area_name)
    postcode = re.sub(r"[^0-9]", "", area_name) if data_level == "postcode" else ""
    suburb = area_name if data_level == "suburb" else ""
    lga_name = area_name if data_level == "lga" else ""
    return {
        "state": state,
        "data_level": data_level,
        "area_name": area_name,
        "normalized_area_key": csa_crime_area_key(data_level, suburb=suburb, postcode=postcode, lga_name=lga_name, area_name=area_name),
        "suburb": suburb,
        "postcode": postcode,
        "lga_name": lga_name,
        "data_year_ending": data_year_ending,
        "source_dataset": source_dataset,
        "source": source,
        "source_url": source_url,
        "notes": notes,
        "raw_json": {},
    }


def wa_area_name(value):
    text = clean_text(value).title()
    text = text.replace(" Wa ", " WA ")
    text = re.sub(r"\bWa\b", "WA", text)
    return text


def act_suburb_name(value):
    text = clean_text(value)
    if not text:
        return ""
    return " ".join(part.capitalize() for part in text.lower().split())


def tas_text_number(pattern, text, default=None):
    match = re.search(pattern, text, flags=re.I | re.S)
    if not match:
        return default
    return csa_crime_number(match.group(1), integer=True)


def parse_tas_crime_text(text, source_url):
    text = re.sub(r"\s+", " ", str(text or " "))
    year_match = re.search(r"Crime Statistics Supplement\s+(20\d{2})-(\d{2})", text, flags=re.I)
    ending_year = 2000 + int(year_match.group(2)) if year_match else None
    data_year_ending = f"June {ending_year}" if ending_year else ""
    total = tas_text_number(r"Total Offences\*\*.*?\s([0-9,]+)\s+[0-9,]+\s+[0-9]+%\s+[-0-9]+%", text)
    person = tas_text_number(r"A\.\s*Offences Against the Person.*?\s([0-9,]+)\s+[0-9,]+\s+[0-9]+%\s+[-0-9]+%", text)
    property_offences = tas_text_number(r"B\.\s*Offences Against Property.*?\s([0-9,]+)\s+[0-9,]+\s+[0-9]+%\s+[-0-9]+%", text)
    fraud = tas_text_number(r"C\.\s*Fraud and Similar Offences.*?\s([0-9,]+)\s+[0-9,]+\s+[0-9]+%\s+[-0-9]+%", text) or 0
    other = tas_text_number(r"D\.\s*Other \(Miscellaneous\) Offences\*\*.*?\s([0-9,]+)\s+[0-9,]+\s+[0-9]+%\s+[-0-9]+%", text)
    trend = tas_text_number(r"Total Offences\*\*.*?[0-9,]+\s+[0-9,]+\s+[0-9]+%\s+([-0-9]+)%", text)
    stat = base_stat(
        "TAS",
        "state",
        "Tasmania",
        data_year_ending,
        "tasmania_police_state_crime_statistics_supplement",
        "Tasmania Police",
        source_url,
        "Official Tasmania Police state-level Crime Statistics Supplement; not suburb, postcode, LGA, or school-zone data.",
    )
    stat["total_recorded_offences"] = total
    stat["crimes_against_person"] = person
    stat["property_and_deception_offences"] = (property_offences or 0) + fraud if property_offences is not None else None
    stat["other_offences"] = other
    stat["assault_count"] = tas_text_number(r"assaults decreased by\s+[0-9]+%\s+to\s+([0-9,]+)", text)
    stat["robbery_count"] = tas_text_number(r"robberies decreased by\s+[0-9]+%\s+to\s+([0-9,]+)", text)
    stat["motor_vehicle_theft_count"] = tas_text_number(r"stolen motor vehicles fell by\s+[0-9]+%\s+to\s+([0-9,]+)", text)
    stat["property_damage_count"] = tas_text_number(r"Injure/Destroy Property offences.*?to\s+([0-9,]+)", text)
    stat["trend_1_year_percent"] = trend
    return [stat] if total else []


def parse_tas_crime(path, source_url):
    import pdfplumber

    with pdfplumber.open(path) as pdf:
        text = "\n".join(page.extract_text() or "" for page in pdf.pages)
    return parse_tas_crime_text(text, source_url)


def parse_nsw_bocsar_lga(path, source_url):
    workbook = load_workbook(path, read_only=True, data_only=True, keep_vba=True)
    try:
        sheet = workbook["Data"]
        rows = sheet.iter_rows(values_only=True)
        headers = [clean_text(value) for value in next(rows)]
        month_columns = []
        for index, header in enumerate(headers):
            parsed = parse_excel_month(header)
            if parsed:
                month_columns.append((index, parsed))
        selected = rolling_months([item[1] for item in month_columns])
        selected_indexes = [index for index, parsed in month_columns if parsed in selected]
        data_year_ending = month_label(selected[-1][0], selected[-1][1]) if selected else ""
        stats = {}
        for row in rows:
            lga = clean_text(row[0] if len(row) > 0 else "")
            category = clean_text(row[1] if len(row) > 1 else "")
            subcategory = clean_text(row[2] if len(row) > 2 else "")
            if not lga or not selected_indexes:
                continue
            stat = stats.setdefault(
                lga,
                base_stat("NSW", "lga", lga, data_year_ending, "nsw_bocsar_lga_recorded_criminal_incidents", "NSW Bureau of Crime Statistics and Research", source_url),
            )
            total = sum(number(row[index] if index < len(row) else 0) for index in selected_indexes)
            add_bucket(stat, f"{category} {subcategory}", total)
        return list(stats.values())
    finally:
        workbook.close()


def parse_excel_month(value):
    if isinstance(value, datetime):
        return value.year, value.month
    text = clean_text(value).lower()
    if not text:
        return None
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d"):
        try:
            parsed_date = datetime.strptime(text, fmt)
            return parsed_date.year, parsed_date.month
        except ValueError:
            pass
    quarter_match = re.search(r"(20\d{2})\s*q([1-4])", text)
    if quarter_match:
        quarter_month = {"1": 3, "2": 6, "3": 9, "4": 12}[quarter_match.group(2)]
        return int(quarter_match.group(1)), quarter_month
    parsed = period_sort_key(text)
    if parsed[0] and parsed[1]:
        return parsed[0], parsed[1]
    match = re.match(r"([a-z]{3})(\d{2})$", text)
    if match and match.group(1) in MONTH_ORDER:
        year = int(match.group(2))
        return 2000 + year, MONTH_ORDER[match.group(1)]
    return None


def csv_rows(path):
    text = path.read_text(encoding="utf-8-sig", errors="replace")
    return list(csv.DictReader(io.StringIO(text)))


def parse_qld_month(value):
    match = re.match(r"([A-Za-z]{3})(\d{2})$", clean_text(value))
    if not match:
        return None
    month = MONTH_ORDER.get(match.group(1).lower())
    if not month:
        return None
    return 2000 + int(match.group(2)), month


def parse_qld_lga(number_path, rate_path, number_url, rate_url):
    rows = csv_rows(number_path)
    months = [parse_qld_month(row.get("Month Year")) for row in rows]
    selected = rolling_months([month for month in months if month])
    data_year_ending = month_label(selected[-1][0], selected[-1][1]) if selected else ""
    stats = {}
    for row in rows:
        parsed = parse_qld_month(row.get("Month Year"))
        if parsed not in selected:
            continue
        lga = clean_text(row.get("LGA Name"))
        if not lga:
            continue
        stat = stats.setdefault(
            lga,
            base_stat("QLD", "lga", lga, data_year_ending, "qld_police_lga_reported_offences", "Queensland Police Service Open Data", number_url),
        )
        add_qld_row(stat, row)
    rates = {}
    for row in csv_rows(rate_path):
        parsed = parse_qld_month(row.get("Month Year"))
        if parsed != (selected[-1] if selected else None):
            continue
        lga = clean_text(row.get("LGA Name"))
        rates[lga] = sum(csa_crime_number(row.get(key)) or 0 for key in QLD_BROAD_COLUMNS if key in row)
    for lga, rate in rates.items():
        if lga in stats and rate:
            stats[lga]["offence_rate_per_100000"] = rate
            stats[lga]["source_url"] = f"{number_url} | {rate_url}"
    return list(stats.values())


def parse_sa_date(value):
    text = clean_text(value)
    for fmt in ("%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            pass
    return None


def parse_sa_crime(path, source_url):
    rows = csv_rows(path)
    dates = [parse_sa_date(row.get("Reported Date")) for row in rows]
    dates = [item for item in dates if item]
    data_year_ending = month_label(max(dates).year, max(dates).month) if dates else ""
    stats = {}
    for row in rows:
        count = row.get("Offence count")
        offence = " ".join([
            clean_text(row.get("Offence Level 1 Description")),
            clean_text(row.get("Offence Level 2 Description")),
            clean_text(row.get("Offence Level 3 Description")),
        ])
        for data_level, area_name in (
            ("suburb", row.get("Suburb - Incident")),
            ("postcode", row.get("Postcode - Incident")),
        ):
            area_name = clean_text(area_name)
            if not area_name:
                continue
            if data_level == "postcode" and not re.sub(r"[^0-9]", "", area_name):
                continue
            key = (data_level, area_name)
            stat = stats.setdefault(
                key,
                base_stat("SA", data_level, area_name, data_year_ending, "sapol_recorded_offences", "South Australia Police / Data SA", source_url),
            )
            add_bucket(stat, offence, count)
    return list(stats.values())


def parse_act_crime(path, source_url):
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        stats = {}
        all_months = []
        sheet_rows = {}
        for sheet_name in workbook.sheetnames:
            rows = list(workbook[sheet_name].iter_rows(values_only=True))
            sheet_rows[sheet_name] = rows
            for row in rows:
                if row and row[0] in (None, ""):
                    all_months.extend(parsed for parsed in (parse_excel_month(value) for value in row[1:]) if parsed)
        selected = rolling_months(all_months, 4)
        data_year_ending = month_label(selected[-1][0], selected[-1][1]) if selected else ""
        for rows in sheet_rows.values():
            category = ""
            selected_indexes = []
            for row in rows:
                first = clean_text(row[0] if row else "")
                rest = row[1:] if row else []
                if first and all(value in (None, "") for value in rest):
                    lowered = first.lower()
                    if "offences and other activities" in lowered or lowered.startswith("promis"):
                        continue
                    category = first
                    selected_indexes = []
                    continue
                if row and row[0] in (None, ""):
                    indexes = [
                        index
                        for index, value in enumerate(row[1:], start=1)
                        if parse_excel_month(value) in selected
                    ]
                    if indexes:
                        selected_indexes = indexes
                    continue
                if not category or csa_crime_header_key(category) in ACT_EXCLUDED_CATEGORIES or not selected_indexes:
                    continue
                suburb = act_suburb_name(first)
                if not suburb or suburb.lower() == "total":
                    continue
                amount = sum(number(row[index] if index < len(row) else 0) for index in selected_indexes)
                if amount == 0:
                    continue
                stat = stats.setdefault(
                    suburb,
                    base_stat(
                        "ACT",
                        "suburb",
                        suburb,
                        data_year_ending,
                        "act_policing_suburb_offences",
                        "ACT Government Open Data / ACT Policing",
                        source_url,
                        "Official ACT suburb-level offences reported or becoming known to police; not school-zone data.",
                    ),
                )
                add_bucket(stat, category, amount)
        return list(stats.values())
    finally:
        workbook.close()


def parse_wa_crime(path, source_url):
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook["Data"]
        rows = sheet.iter_rows(values_only=True)
        headers = [clean_text(value) for value in next(rows)]
        row_payloads = []
        months = []
        for raw in rows:
            row = {header: (raw[index] if index < len(raw) else "") for index, header in enumerate(headers)}
            parsed = parse_excel_month(row.get("Period"))
            if parsed:
                months.append(parsed)
            row_payloads.append((row, parsed))
        selected = rolling_months(months)
        data_year_ending = month_label(selected[-1][0], selected[-1][1]) if selected else ""
        stats = {}
        for row, parsed in row_payloads:
            if parsed not in selected:
                continue
            area_name = wa_area_name(row.get("Website Region"))
            area_key = area_name.upper()
            if not area_name or area_key == "WESTERN AUSTRALIA":
                continue
            if area_key.endswith(" DISTRICT"):
                data_level = "police_district"
            elif area_key.endswith(" REGION"):
                data_level = "police_region"
            else:
                continue
            key = (data_level, area_name)
            stat = stats.setdefault(
                key,
                base_stat(
                    "WA",
                    data_level,
                    area_name,
                    data_year_ending,
                    "wa_police_reported_offences_timeseries",
                    "Western Australia Police Force",
                    source_url,
                    "Official WA Police time-series data by police region/district; not suburb, postcode, LGA, or school-zone data.",
                ),
            )
            add_bucket(stat, f"{row.get('WAPOL_Hierarchy_Lvl1')} {row.get('WAPOL_Hierarchy_Lvl2')}", row.get("Count"))
        return list(stats.values())
    finally:
        workbook.close()


def import_stats(db, stats, apply_changes):
    if apply_changes:
        for stat in stats:
            if stat.get("total_recorded_offences"):
                upsert_csa_crime_area_stat(db, stat)
    return len([stat for stat in stats if stat.get("total_recorded_offences")])


def refresh_school_area_crime_profiles(db, states, apply_changes):
    if not apply_changes:
        return 0
    placeholders = ",".join("?" for _ in states)
    refreshed = 0
    for school in db.all(f"SELECT id FROM schools WHERE UPPER(COALESCE(state,'')) IN ({placeholders})", tuple(states)):
        if refresh_school_area_crime_profile(db, school["id"]):
            refreshed += 1
    return refreshed


def run(apply_changes=False, states=None, latest_only=True):
    selected_states = {state.upper() for state in (states or ["VIC", "NSW", "QLD", "SA", "WA", "ACT", "TAS"])}
    db = Database(DB_PATH)
    db.initialize()
    report = {"apply": apply_changes, "states": {}, "profiles_refreshed": 0}

    if "VIC" in selected_states:
        report["states"]["VIC"] = import_vic_crime_stats.run(apply_changes=apply_changes, latest_only=latest_only)

    if "NSW" in selected_states:
        path = download(NSW_BOCSAR_LGA_URL, DATA_DIR / "nsw", "RCI_offencebymonth", ".xlsm")
        stats = parse_nsw_bocsar_lga(path, f"{NSW_BOCSAR_SOURCE_PAGE} | {NSW_BOCSAR_LGA_URL}")
        report["states"]["NSW"] = {"source": "NSW Bureau of Crime Statistics and Research", "file": str(path.relative_to(ROOT)), "rows_imported": import_stats(db, stats, apply_changes)}

    if "QLD" in selected_states:
        num_package = package_metadata(QLD_PACKAGE_SHOW, QLD_LGA_NUMBER_PACKAGE)
        rate_package = package_metadata(QLD_PACKAGE_SHOW, QLD_LGA_RATE_PACKAGE)
        num_resource = downloadable_resource(num_package, ("csv",))
        rate_resource = downloadable_resource(rate_package, ("csv",))
        num_path = download(num_resource["url"], DATA_DIR / "qld", num_resource.get("name"), ".csv")
        rate_path = download(rate_resource["url"], DATA_DIR / "qld", rate_resource.get("name"), ".csv")
        stats = parse_qld_lga(num_path, rate_path, num_resource["url"], rate_resource["url"])
        report["states"]["QLD"] = {"source": "Queensland Police Service Open Data", "file": str(num_path.relative_to(ROOT)), "rate_file": str(rate_path.relative_to(ROOT)), "rows_imported": import_stats(db, stats, apply_changes)}

    if "SA" in selected_states:
        package = package_metadata(SA_PACKAGE_SHOW, SA_CRIME_PACKAGE)
        resource = sa_crime_resource(package)
        path = download(resource["url"], DATA_DIR / "sa", resource.get("name"), ".csv")
        stats = parse_sa_crime(path, resource["url"])
        report["states"]["SA"] = {"source": "South Australia Police / Data SA", "file": str(path.relative_to(ROOT)), "rows_imported": import_stats(db, stats, apply_changes)}

    if "ACT" in selected_states:
        path = download(ACT_CRIME_URL, DATA_DIR / "act", "act-crime-statistics", ".xlsx")
        stats = parse_act_crime(path, f"{ACT_CRIME_SOURCE_PAGE} | {ACT_CRIME_URL}")
        report["states"]["ACT"] = {
            "source": "ACT Government Open Data / ACT Policing",
            "file": str(path.relative_to(ROOT)),
            "rows_imported": import_stats(db, stats, apply_changes),
        }

    if "WA" in selected_states:
        path = download(WA_TIMESERIES_URL, DATA_DIR / "wa", "wa-police-force-crime-timeseries", ".xlsx")
        stats = parse_wa_crime(path, WA_TIMESERIES_URL)
        report["states"]["WA"] = {
            "source": "Western Australia Police Force",
            "file": str(path.relative_to(ROOT)),
            "rows_imported": import_stats(db, stats, apply_changes),
            "notes": "Official workbook imported as WA Police region/district area-level data; not suburb, postcode, LGA, or school-zone data.",
        }

    if "TAS" in selected_states:
        path = download(TAS_CRIME_SUPPLEMENT_URL, DATA_DIR / "tas", "tasmania-police-crime-statistics-supplement-2024-25", ".pdf")
        stats = parse_tas_crime(path, TAS_CRIME_SUPPLEMENT_URL)
        report["states"]["TAS"] = {
            "source": "Tasmania Police",
            "file": str(path.relative_to(ROOT)),
            "rows_imported": import_stats(db, stats, apply_changes),
            "notes": "Official PDF supplement imported as Tasmania state-level area data; not suburb, postcode, LGA, or school-zone data.",
        }

    report["profiles_refreshed"] = refresh_school_area_crime_profiles(db, sorted(selected_states), apply_changes)
    return report


def main():
    parser = argparse.ArgumentParser(description="Import official Australian state crime statistics into AAIE CRM.")
    parser.add_argument("--apply", action="store_true", help="Write imported rows and refreshed school profiles to the CRM database.")
    parser.add_argument("--all-periods", action="store_true", help="Pass all-period import to the VIC DataVic importer.")
    parser.add_argument("--states", default="", help="Comma-separated states to import, e.g. NSW,QLD,SA. Defaults to VIC,NSW,QLD,SA,WA,ACT,TAS.")
    args = parser.parse_args()
    states = [item.strip().upper() for item in args.states.split(",") if item.strip()] or None
    report = run(apply_changes=args.apply, states=states, latest_only=not args.all_periods)
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
